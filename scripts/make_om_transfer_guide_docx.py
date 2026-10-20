# -*- coding: utf-8 -*-
"""將「跨OM組別調貨指引.md」轉為 Word (.docx)，供傳閱與簽核。

用法：
    python scripts/make_om_transfer_guide_docx.py

輸出：專案根目錄「跨OM組別調貨指引.docx」（覆寫）。
支援的 Markdown 結構（本指引實際使用的子集）：
    # / ## / ### 標題、> 引言、- 清單、- [ ] 核對清單、1. 編號清單、
    ``` 程式區塊（等寬）、| 表格 |、--- 分隔線、行內 **粗體** 與 `程式碼`。
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "跨OM組別調貨指引.md"
DST = ROOT / "跨OM組別調貨指引.docx"

BODY_FONT = "Microsoft JhengHei"   # 微軟正黑體（東亞字型）
LATIN_FONT = "Calibri"
CODE_FONT = "Consolas"
BODY_SIZE = 10.5
TABLE_SIZE = 9.5
HEAD_FILL = "E4EEF5"
QUOTE_FILL = "F2F7FB"
CODE_FILL = "F4F4F4"
ACCENT = RGBColor(0x1F, 0x49, 0x7D)


# ── 低階工具 ──────────────────────────────────────────────────────────────
def _east_asian(run, name: str = BODY_FONT) -> None:
    """設定 run 的東亞字型（Word 需另設 w:eastAsia 才不會回退成新細明體）。"""
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)


def _shade(element, fill: str) -> None:
    """在段落或表格儲存格的屬性上加底紋。"""
    pr = element.get_or_add_pPr() if hasattr(element, "get_or_add_pPr") else element
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def _cell_shade(cell, fill: str) -> None:
    _shade(cell._tc.get_or_add_tcPr(), fill)


def _add_runs(paragraph, text: str, *, size: float = BODY_SIZE, bold_all: bool = False,
              font: str = BODY_FONT) -> None:
    """把行內 **粗體** 與 `程式碼` 轉成對應的 run。"""
    for chunk in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not chunk:
            continue
        bold = bold_all
        code = False
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            chunk, bold = chunk[2:-2], True
        elif chunk.startswith("`") and chunk.endswith("`") and len(chunk) > 2:
            chunk, code = chunk[1:-1], True
        run = paragraph.add_run(chunk)
        run.bold = bold
        run.font.size = Pt(size - 0.5 if code else size)
        _east_asian(run, CODE_FONT if code else font)
        if code:
            run.font.name = CODE_FONT


def _add_page_number_footer(section, label: str = "跨 OM 組別調貨指引") -> None:
    """頁尾：文件名稱＋第 X 頁（Word 功能變數）。"""
    para = section.footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(f"{label}　·　第 ")
    run.font.size = Pt(8)
    _east_asian(run)
    _field(para, "PAGE")
    run = para.add_run(" 頁")
    run.font.size = Pt(8)
    _east_asian(run)


def _field(paragraph, code: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {code} "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run.font.size = Pt(8)
    run._element.append(begin)
    run._element.append(instr)
    run._element.append(end)


def _add_toc(doc: Document, levels: str = "1-2") -> None:
    """插入目錄功能變數（Word 開啟時自動更新；其他檢視器顯示提示文字）。"""
    label = doc.add_paragraph()
    run = label.add_run("目錄")
    run.bold = True
    run.font.size = Pt(12.5)
    run.font.color.rgb = ACCENT
    _east_asian(run)

    para = doc.add_paragraph()
    run = para.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f' TOC \\o "{levels}" \\h \\z \\u '
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    hint = OxmlElement("w:t")
    hint.text = "（Word 將於開啟時自動更新目錄；如未更新請按 F9）"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, sep, hint, end):
        run._element.append(node)


def _set_update_fields(doc: Document) -> None:
    """令 Word 開啟檔案時自動更新功能變數（目錄、頁碼）。"""
    settings = doc.settings.element
    node = settings.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields")
        settings.append(node)
    node.set(qn("w:val"), "true")


# ── 樣式 ──────────────────────────────────────────────────────────────────
def _setup_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(BODY_SIZE)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 15, ACCENT, 12, 6),
        ("Heading 2", 12.5, ACCENT, 10, 4),
        ("Heading 3", 11, RGBColor(0x33, 0x33, 0x33), 8, 3),
    ):
        style = doc.styles[name]
        style.font.name = LATIN_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)


# ── 區塊寫入 ──────────────────────────────────────────────────────────────
def _display_width(text: str) -> int:
    """顯示寬度估算（全形字元約 2、半形約 1），供欄寬比例分配用。"""
    return sum(2 if ord(ch) > 0x2E7F else 1 for ch in str(text))


def _add_table(doc: Document, rows: list[list[str]], *, total_cm: float = 17.0) -> None:
    header, body = rows[0], rows[1:]
    n_cols = len(header)
    table = doc.add_table(rows=1, cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # 依各欄最長內容按比例分配欄寬（避免 Word 自動排版把長文字欄壓得過窄）
    units = [
        max(6, min(max([_display_width(r[i]) for r in rows if i < len(r)] or [4]), 60))
        for i in range(n_cols)
    ]
    col_cm = [max(1.6, total_cm * u / sum(units)) for u in units]
    scale = total_cm / sum(col_cm)
    col_cm = [c * scale for c in col_cm]

    for idx, text in enumerate(header):
        cell = table.rows[0].cells[idx]
        cell.text = ""
        para = cell.paragraphs[0]
        para.paragraph_format.space_after = Pt(0)
        _add_runs(para, text, size=TABLE_SIZE, bold_all=True)
        _cell_shade(cell, HEAD_FILL)
    for record in body:
        cells = table.add_row().cells
        for idx, text in enumerate(record[:n_cols]):
            para = cells[idx].paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            _add_runs(para, text, size=TABLE_SIZE)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Cm(col_cm[idx])
    doc.add_paragraph()


def _add_quote(doc: Document, lines: list[str]) -> None:
    for line in lines:
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Cm(0.4)
        para.paragraph_format.space_before = Pt(2)
        para.paragraph_format.space_after = Pt(2)
        _add_runs(para, line, size=BODY_SIZE - 0.5)
        _shade(para._p.get_or_add_pPr(), QUOTE_FILL)


def _add_code(doc: Document, lines: list[str]) -> None:
    for line in lines:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.line_spacing = 1.1
        run = para.add_run(line if line else " ")
        run.font.name = CODE_FONT
        run.font.size = Pt(TABLE_SIZE)
        _east_asian(run, CODE_FONT)
        _shade(para._p.get_or_add_pPr(), CODE_FILL)
    doc.add_paragraph()


def _add_bullet(doc: Document, text: str, *, checklist: bool = False) -> None:
    para = doc.add_paragraph(style="List Bullet")
    para.paragraph_format.space_after = Pt(2)
    if checklist:
        text = "☐　" + text
    _add_runs(para, text)


def _add_number(doc: Document, number: str, text: str) -> None:
    """編號清單改以固定編號文字輸出，避免 Word 跨章節沿用同一組清單編號。"""
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Cm(0.85)
    para.paragraph_format.first_line_indent = Cm(-0.85)
    para.paragraph_format.space_after = Pt(2)
    run = para.add_run(f"{number}　")
    run.font.size = Pt(BODY_SIZE)
    _east_asian(run)
    _add_runs(para, text)


# ── 解析與主流程 ──────────────────────────────────────────────────────────
QUICK_HEADING = "### 七點速覽（簡化版）"
QUICK_DST = ROOT / "跨OM組別調貨指引_七點速覽.docx"


def _extract_section(lines: list[str], heading: str,
                     skip_prefixes: tuple[str, ...] = ("本節以下",)) -> list[str]:
    """取出指定 ### 標題到下一個 ### 之間的內容（供單頁速覽使用）。

    ``skip_prefixes`` 用於剔除只對主檔有意義的過場句（如「本節以下為逐條詳細判定…」）。
    """
    out: list[str] = []
    inside = False
    for line in lines:
        stripped = line.strip()
        if stripped == heading:
            inside = True
            continue
        if inside and stripped.startswith("### "):
            break
        if inside and stripped.lstrip(">").strip().startswith(skip_prefixes):
            continue
        if inside:
            out.append(line)
    return out


def _render_lines(doc: Document, lines: list[str]) -> None:
    """把 Markdown 行渲染進 doc（沿用兩份文件同一套規則）。"""
    idx, total = 0, len(lines)
    toc_added = False
    while idx < total:
        line = lines[idx].rstrip()
        stripped = line.strip()
        if not stripped:
            idx += 1
            continue
        if stripped == "---":
            idx += 1
            continue
        if stripped.startswith("```"):
            idx += 1
            block: list[str] = []
            while idx < total and not lines[idx].strip().startswith("```"):
                block.append(lines[idx].rstrip())
                idx += 1
            idx += 1
            _add_code(doc, block)
            continue
        if stripped.startswith("|"):
            block = []
            while idx < total and lines[idx].strip().startswith("|"):
                if not re.fullmatch(r"\|[\s:\-|]+\|", lines[idx].strip()):
                    block.append([c.strip() for c in lines[idx].strip().strip("|").split("|")])
                idx += 1
            _add_table(doc, block)
            continue
        if stripped.startswith(">"):
            block = []
            while idx < total and lines[idx].strip().startswith(">"):
                block.append(lines[idx].strip().lstrip(">").strip())
                idx += 1
            _add_quote(doc, block)
            continue
        if stripped.startswith("- [ ]"):
            _add_bullet(doc, stripped[5:].strip(), checklist=True)
            idx += 1
            continue
        if stripped.startswith("- "):
            _add_bullet(doc, stripped[2:].strip())
            idx += 1
            continue
        if re.match(r"^\d+\.\s", stripped):
            number = re.match(r"^(\d+)\.\s", stripped).group(1)
            _add_number(doc, f"{number}.", re.sub(r"^\d+\.\s", "", stripped))
            idx += 1
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=2)
            idx += 1
            continue
        if stripped.startswith("## "):
            if not toc_added:
                _add_toc(doc)
                toc_added = True
            doc.add_heading(stripped[3:].strip(), level=1)
            idx += 1
            continue
        if stripped.startswith("# "):
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(stripped[2:].strip())
            run.bold = True
            run.font.size = Pt(20)
            run.font.color.rgb = ACCENT
            _east_asian(run)
            idx += 1
            continue
        para = doc.add_paragraph()
        _add_runs(para, stripped)
        idx += 1


def _new_document(label: str = "跨 OM 組別調貨指引") -> Document:
    """建立 A4 文件（共同版面設定）。"""
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)   # A4
    for attr in ("left_margin", "right_margin"):
        setattr(section, attr, Cm(2.0))
    for attr in ("top_margin", "bottom_margin"):
        setattr(section, attr, Cm(1.8))
    _setup_styles(doc)
    _add_page_number_footer(section, label)
    return doc


def build(src: Path = SRC, dst: Path = DST) -> Path:
    """產生完整指引 Word。"""
    lines = src.read_text(encoding="utf-8").splitlines()
    doc = _new_document()
    _render_lines(doc, lines)
    _set_update_fields(doc)
    doc.save(dst)
    return dst


def build_quick(src: Path = SRC, dst: Path = QUICK_DST) -> Path:
    """由指引第 5 節抽出「七點速覽」，產生單頁派發版 Word（內容與主檔同源）。"""
    lines = src.read_text(encoding="utf-8").splitlines()
    quick = _extract_section(lines, QUICK_HEADING)
    doc = _new_document("跨 OM 組別調貨指引 — 七點速覽")
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run("跨 OM 組別調貨指引 — 七點速覽")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = ACCENT
    _east_asian(run)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_runs(
        sub,
        "《跨 OM 組別調貨指引》第 5 節「轉出優先順序」的簡化版；"
        "詳細判定、參數與核對清單見指引全文。系統邏輯現時未作任何修改。",
        size=BODY_SIZE - 0.5,
    )
    _render_lines(doc, quick)
    _set_update_fields(doc)
    doc.save(dst)
    return dst


def main() -> int:
    if not SRC.exists():
        print(f"找不到來源檔案：{SRC}", file=sys.stderr)
        return 1
    for out in (build(), build_quick()):
        print(f"已產生：{out}（{out.stat().st_size:,} bytes）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
