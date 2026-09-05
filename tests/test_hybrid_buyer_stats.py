"""Safety Stock x NetStock Review Buyer 統計（需調貨＋需調整 Safety）。

- Buyer 需調貨 SKU 統計與 Stock Checking 同口徑（同 SKU 首個非空 Buyer，已刪除重複）
- Buyer 需調整 Safety SKU 統計母體為 Safety 設定不足＋過高（有建議 Safety 設定值的列）
- Excel Review 摘要在既有 ①②③ 區塊後另加 ④⑤ Buyer 表；Streamlit 看板與 Review 摘要分頁同步顯示
"""
import warnings
warnings.filterwarnings("ignore")
import pandas as pd
from openpyxl import load_workbook
from transformer import (
    transform_hybrid,
    generate_excel_hybrid,
    transform_stock_checking,
    _hybrid_buyer_need_stats,
    _hybrid_buyer_safety_stats,
    _hybrid_stats_from_raw,
    _HYBRID_BUYER_NEED_TITLE,
    _HYBRID_BUYER_SAFETY_TITLE,
)

BASE = {
    "Buyer": "B1", "Brand": "A", "SKU": 1, "SKU Desc": "Test Product",
    "Item_Shop": "Prod-S01-ShopHK", "Site Code": "S01", "OM": "OM1",
    "Shop Nature": "Tourist", "Vendor MC": "V1", "Item Status": "A",
    "Adj. SST": "SST1",
    "Launch Date": "2024-01-01", "Item Class": "C1", "RP Type": "RF",
    "Shop Stock Status": "Normal", "Adj. Safety Stock": 3.0,
    "Sold Qty (P90D)": 90.0, "Excess Stock Qty": 2.0,
    "Shop Stock Qty (CM)": 5.0,
    "Shop Monthly Sold Qty (P90D)": 30.5,
    "Sold Qty (CM - MTD)": 8.0,
    "D001 Stock Qty (CM)": 200.0, "D001 Net Stock Qty (CM)": 99.0,
    "Turnday (P90D Qty)": 13.0,
}


def _row(**kw):
    r = dict(BASE)
    r.update(kw)
    return r


def _df(rows):
    return pd.DataFrame([_row(**r) for r in rows])


_SD10 = {"Tourist": 10, "Mix": 10, "Non Tourist": 10}

_TWO_BUYER = [
    {"SKU": 300, "Site Code": "T01", "Shop Nature": "Tourist", "Buyer": "Alice",
     "Sold Qty (P90D)": 360.0, "Shop Stock Qty (CM)": 47.0,
     "Shop Monthly Sold Qty (P90D)": 120.0, "Adj. Safety Stock": 30,
     "D001 Net Stock Qty (CM)": 10.0},
    {"SKU": 300, "Site Code": "N01", "Shop Nature": "Non Tourist", "Buyer": "Alice",
     "Sold Qty (P90D)": 90.0, "Shop Stock Qty (CM)": 30.0,
     "Shop Monthly Sold Qty (P90D)": 30.5, "Adj. Safety Stock": 30},
    {"SKU": 301, "Site Code": "T02", "Shop Nature": "Tourist", "Buyer": "Bob",
     "Sold Qty (P90D)": 360.0, "Shop Stock Qty (CM)": 47.0,
     "Shop Monthly Sold Qty (P90D)": 120.0, "Adj. Safety Stock": 5,
     "D001 Net Stock Qty (CM)": 10.0},
    {"SKU": 301, "Site Code": "N02", "Shop Nature": "Non Tourist", "Buyer": "Bob",
     "Sold Qty (P90D)": 90.0, "Shop Stock Qty (CM)": 30.0,
     "Shop Monthly Sold Qty (P90D)": 30.5, "Adj. Safety Stock": 30},
]


def test_buyer_need_matches_stock_checking_caliber():
    raw, _, _, _, _, stats, _ = transform_hybrid(_df(_TWO_BUYER), safety_days=_SD10)
    sc_raw, _, sc_stats, _ = transform_stock_checking(_df(_TWO_BUYER))
    assert stats["buyer_need_skus"] == sc_stats["buyer_need_skus"] == [("Alice", 1), ("Bob", 1)]
    assert _hybrid_buyer_need_stats(raw) == stats["buyer_need_skus"]
    assert sum(c for _, c in stats["buyer_need_skus"]) == stats["need_skus"]


def test_buyer_need_empty_buyer_grouped():
    rows = [
        {"SKU": 1, "Site Code": "S01", "Buyer": "", "Shop Nature": "Tourist",
         "Sold Qty (P90D)": 360.0, "Shop Stock Qty (CM)": 47.0,
         "Shop Monthly Sold Qty (P90D)": 120.0, "D001 Net Stock Qty (CM)": 10.0},
        {"SKU": 1, "Site Code": "N01", "Buyer": "", "Shop Nature": "Non Tourist",
         "Sold Qty (P90D)": 90.0, "Shop Stock Qty (CM)": 30.0,
         "Shop Monthly Sold Qty (P90D)": 30.5},
    ]
    _, _, _, _, _, stats, _ = transform_hybrid(_df(rows))
    assert stats["buyer_need_skus"] == [("(未填 Buyer)", 1)]


def test_buyer_safety_counts_and_amounts():
    # 10 日口徑：SKU 300（T01 +10、N01 −20 → 淨 −10）；SKU 301（T02 +35、N02 −20 → 淨 +15）
    raw, _, _, _, _, stats, _ = transform_hybrid(_df(_TWO_BUYER), safety_days=_SD10)
    assert stats["buyer_safety_fix"] == [("Alice", 1, -10), ("Bob", 1, 15)]
    assert _hybrid_buyer_safety_stats(raw) == stats["buyer_safety_fix"]
    # 合計對帳：SKU 數合計 = 需調整母體去重 SKU 數；調整量合計 = 不足＋過高淨額
    _fix_mask = raw["Safety Review"].isin(["Safety 設定不足", "Safety 設定過高"])
    assert sum(c for _, c, _ in stats["buyer_safety_fix"]) == int(raw.loc[_fix_mask, "SKU"].nunique())
    assert sum(v for _, _, v in stats["buyer_safety_fix"]) == stats["under_inc_total"] + stats["over_dec_total"]
    assert sum(c for _, c, _ in stats["buyer_safety_fix"]) == 2


def test_buyer_safety_empty_when_no_fix():
    df = _df([
        {"SKU": 1, "Site Code": "N01", "Shop Nature": "Non Tourist",
         "Sold Qty (P90D)": 90.0, "Shop Stock Qty (CM)": 30.0,
         "Shop Monthly Sold Qty (P90D)": 30.5},
    ])
    raw, _, _, _, _, stats, _ = transform_hybrid(df)
    assert stats["buyer_need_skus"] == []
    assert stats["buyer_safety_fix"] == []
    assert _hybrid_buyer_need_stats(raw) == []
    assert _hybrid_buyer_safety_stats(raw) == []
    assert _hybrid_buyer_need_stats(None) == []
    assert _hybrid_buyer_safety_stats(pd.DataFrame()) == []


def test_stats_fallback_matches_transform():
    raw, _, _, _, _, stats, _ = transform_hybrid(_df(_TWO_BUYER), safety_days=_SD10)
    fb = _hybrid_stats_from_raw(raw)
    assert fb["buyer_need_skus"] == stats["buyer_need_skus"]
    assert fb["buyer_safety_fix"] == stats["buyer_safety_fix"]


def test_excel_buyer_tables_after_safety_blocks():
    raw, sc, diag, diag_om, cross, stats, _ = transform_hybrid(_df(_TWO_BUYER), safety_days=_SD10)
    ws = load_workbook(generate_excel_hybrid(raw, sc, diag, diag_om, cross, stats))["Review 摘要"]
    titles = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value in (_HYBRID_BUYER_NEED_TITLE, _HYBRID_BUYER_SAFETY_TITLE):
                titles[c.value] = c.row
    assert set(titles) == {_HYBRID_BUYER_NEED_TITLE, _HYBRID_BUYER_SAFETY_TITLE}
    assert titles[_HYBRID_BUYER_NEED_TITLE] < titles[_HYBRID_BUYER_SAFETY_TITLE]
    # ④ 需調貨表：標題 → 表頭 → Alice → Bob → 總計
    r0 = titles[_HYBRID_BUYER_NEED_TITLE]
    assert ws.cell(row=r0 + 1, column=1).value == "Buyer"
    assert ws.cell(row=r0 + 1, column=2).value == "需調貨 SKU 數"
    assert ws.cell(row=r0 + 2, column=1).value == "Alice"
    assert ws.cell(row=r0 + 2, column=2).value == 1
    assert ws.cell(row=r0 + 3, column=1).value == "Bob"
    assert ws.cell(row=r0 + 4, column=1).value == "總計"
    assert ws.cell(row=r0 + 4, column=2).value == 2
    assert ws.cell(row=r0 + 2, column=3).number_format == "0.0%"
    # ⑤ Safety 表：標題 → 表頭 → Alice(−10) → Bob(+15) → 總計(淨 +5)
    r1 = titles[_HYBRID_BUYER_SAFETY_TITLE]
    assert ws.cell(row=r1 + 1, column=2).value == "需調整 Safety SKU 數"
    assert ws.cell(row=r1 + 1, column=3).value == "Safety 建議調整量合計"
    assert ws.cell(row=r1 + 2, column=1).value == "Alice"
    assert ws.cell(row=r1 + 2, column=3).value == -10
    assert ws.cell(row=r1 + 3, column=3).value == 15
    assert ws.cell(row=r1 + 4, column=1).value == "總計"
    assert ws.cell(row=r1 + 4, column=2).value == 2
    assert ws.cell(row=r1 + 4, column=3).value == 5


def test_excel_no_buyer_tables_when_empty():
    df = _df([
        {"SKU": 1, "Site Code": "N01", "Shop Nature": "Non Tourist",
         "Sold Qty (P90D)": 90.0, "Shop Stock Qty (CM)": 30.0,
         "Shop Monthly Sold Qty (P90D)": 30.5},
    ])
    raw, sc, diag, diag_om, cross, stats, _ = transform_hybrid(df)
    ws = load_workbook(generate_excel_hybrid(raw, sc, diag, diag_om, cross, stats))["Review 摘要"]
    texts = {c.value for row in ws.iter_rows() for c in row}
    assert _HYBRID_BUYER_NEED_TITLE not in texts
    assert _HYBRID_BUYER_SAFETY_TITLE not in texts
