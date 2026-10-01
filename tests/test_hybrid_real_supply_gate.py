"""Validation tests for Safety Stock x NetStock Review — 「有貨可調」改用實際可轉出量.

`real_supply_gate`（預設關閉）把「有貨可調」由 Nature 池加總 TurnDay 改為 SKU 級
**實際可轉出量 > 0**，消除兩類假陽性事件列：
- 池內沒有任何一列有 14 日超額（池看來健康但實際一貨可調也沒有）；
- 某 Nature 完全沒有列時其加總 TurnDay 為 inf 而令門檻空過。

同時新增「調貨補足狀態」三態欄（已補足／部分補足／無貨可調），令下游可分辨
「一件都調不到」與「部分補足」——兩者在現行「調入缺口」下無法區分。
"""
import warnings
warnings.filterwarnings("ignore")
import pandas as pd
from openpyxl import load_workbook
from transformer import (
    transform_hybrid,
    generate_excel_hybrid,
    _hybrid_dashboard_metrics,
    _hybrid_raw_display_columns,
    _hybrid_stats_from_raw,
    _HYBRID_DIAG_COLUMNS,
    _HYBRID_FULFILL_COLUMN,
    _HYBRID_FULFILL_FULL,
    _HYBRID_FULFILL_PARTIAL,
    _HYBRID_FULFILL_NONE,
)
from modes.issues import _build_hybrid_shop_issues


def _row(sku, nature, site, om, stock, sold, safety, rp="RF", d001=0.0):
    return {
        "Buyer": "B1", "Brand": "BR", "SKU": sku, "SKU Desc": f"D{sku}", "Adj. SST": "SST1",
        "Shop Nature": nature, "Site Code": site, "OM": om,
        "Shop Stock Qty (CM)": stock, "Sold Qty (P90D)": sold,
        "Adj. Safety Stock": safety,
        "Shop Monthly Sold Qty (P90D)": sold / 90 * 30.5,
        "RP Type": rp, "D001 Net Stock Qty (CM)": d001,
    }


def _df(rows):
    return pd.DataFrame(rows)


def _get(raw, site, col):
    return raw[raw["Site Code"] == site].iloc[0][col]


# ── 場景 ──────────────────────────────────────────────────────────────────
# SKU-N：NT 池每店剛好 14 日 → 池 TurnDay 14 過門檻，但逐列超額全為 0（假陽性來源）
# SKU-O：NT 池健康且有 26 件真實超額（對照組，不應受影響）
_ZERO_EXCESS = [
    _row("SKU-N", "Non Tourist", "N1", "OM1", 14, 90, 10),
    _row("SKU-N", "Tourist", "T1", "OM1", 0, 90, 7),
    _row("SKU-O", "Non Tourist", "N2", "OM1", 40, 90, 10),
    _row("SKU-O", "Tourist", "T2", "OM1", 0, 90, 7),
]

# NT 池超額僅 3 件；3 個 Tourist 需求列各需 5／3／2 件 → 供給競爭
_CONTENTION = [
    _row("SKU-S", "Non Tourist", "N1", "OM1", 17, 90, 10),
    _row("SKU-S", "Tourist", "T1", "OM1", 9, 90, 7),    # 需 5
    _row("SKU-S", "Tourist", "T2", "OM1", 11, 90, 7),   # 需 3
    _row("SKU-S", "Tourist", "T3", "OM1", 12, 90, 7),   # 需 2
]

# 空 Nature 池：只有 Tourist 與 NT，Mix 加總 TurnDay = inf → 門檻空過
_EMPTY_POOL = [
    _row("SKU-E", "Non Tourist", "N1", "OM1", 5, 90, 10),
    _row("SKU-E", "Tourist", "T1", "OM1", 0, 90, 7),
]


# ── 1. 關閉時維持現行契約 ─────────────────────────────────────────────────
def test_off_keeps_false_positive_event():
    raw, *_ = transform_hybrid(_df(_ZERO_EXCESS))
    assert _get(raw, "T1", "需要調貨") == "是"
    assert int(_get(raw, "T1", "實際調入量")) == 0
    assert int(_get(raw, "T1", "調入缺口")) == 14


def test_off_does_not_write_fulfill_column():
    raw, _, diag, _, _, _, _ = transform_hybrid(_df(_ZERO_EXCESS + _CONTENTION))
    assert _HYBRID_FULFILL_COLUMN not in raw.columns
    assert not any(_HYBRID_FULFILL_COLUMN in c for c in _hybrid_raw_display_columns(raw))
    assert _HYBRID_FULFILL_COLUMN not in diag.columns
    missing = [c for c in _HYBRID_DIAG_COLUMNS if c not in list(diag.columns)]
    assert not missing, f"主表契約欄缺失：{missing}"


def test_default_matches_explicit_false():
    a, *_ = transform_hybrid(_df(_ZERO_EXCESS + _CONTENTION))
    b, *_ = transform_hybrid(_df(_ZERO_EXCESS + _CONTENTION), real_supply_gate=False)
    assert a.columns.equals(b.columns) and a.equals(b)


# ── 2. 開啟後消除假陽性 ───────────────────────────────────────────────────
def test_on_removes_zero_excess_false_positive():
    raw, *_ = transform_hybrid(_df(_ZERO_EXCESS), real_supply_gate=True)
    assert _get(raw, "T1", "需要調貨") != "是"
    assert int(_get(raw, "T1", "調入需求量 (補到14日)")) == 0


def test_on_keeps_genuine_shortage():
    raw, *_ = transform_hybrid(_df(_ZERO_EXCESS), real_supply_gate=True)
    assert _get(raw, "T2", "需要調貨") == "是"
    assert int(_get(raw, "T2", "實際調入量")) == 14
    assert _get(raw, "T2", "調貨來源 (Site Code)") == "N2:14"


def test_on_removes_empty_pool_vacuous_pass():
    off, *_ = transform_hybrid(_df(_EMPTY_POOL))
    on, *_ = transform_hybrid(_df(_EMPTY_POOL), real_supply_gate=True)
    assert _get(off, "T1", "需要調貨") == "是"
    assert _get(on, "T1", "需要調貨") != "是"


def test_trigger_only_shrinks_and_total_actual_unchanged():
    rows = _ZERO_EXCESS + _CONTENTION + _EMPTY_POOL
    off, *_ = transform_hybrid(_df(rows))
    on, *_ = transform_hybrid(_df(rows), real_supply_gate=True)
    off_keys = set(map(tuple, off.loc[off["需要調貨"] == "是", ["SKU", "Site Code"]].values))
    on_keys = set(map(tuple, on.loc[on["需要調貨"] == "是", ["SKU", "Site Code"]].values))
    assert off_keys >= on_keys and on_keys, "觸發面必須只收縮且仍有事件"
    assert int(on["實際調入量"].sum()) == int(off["實際調入量"].sum())


def test_rp_dimension_unchanged_by_switch():
    """rp 維度本來就以 SKU 級實際量判定，兩種開關狀態應完全一致。"""
    rows = _ZERO_EXCESS + _CONTENTION
    off, *_ = transform_hybrid(_df(rows), dimension="rp")
    on, *_ = transform_hybrid(_df(rows), dimension="rp", real_supply_gate=True)
    assert off["需要調貨"].equals(on["需要調貨"])
    assert off["實際調入量"].equals(on["實際調入量"])


def test_bidirectional_removes_false_positive():
    """Mix 池整體缺貨（13.3 日 < 14，池級歸零）且 NT 池零超額 → 無真實供給。"""
    rows = [
        _row("SKU-B", "Mix", "M1", "OM1", 40, 180, 10),
        _row("SKU-B", "Mix", "M2", "OM1", 0, 90, 10),
        _row("SKU-B", "Non Tourist", "N1", "OM1", 9, 90, 10),
    ]
    off, *_ = transform_hybrid(_df(rows), bidirectional=True)
    on, *_ = transform_hybrid(_df(rows), bidirectional=True, real_supply_gate=True)
    assert _get(off, "M2", "需要調貨") == "是"
    assert int(_get(off, "M2", "實際調入量")) == 0
    assert _get(on, "M2", "需要調貨") != "是"


def test_bidirectional_keeps_cross_nature_supply():
    rows = [
        _row("SKU-B", "Mix", "M1", "OM1", 40, 180, 10),
        _row("SKU-B", "Mix", "M2", "OM1", 0, 90, 10),
        _row("SKU-B", "Non Tourist", "N1", "OM1", 40, 90, 10),
    ]
    on, *_ = transform_hybrid(_df(rows), bidirectional=True, real_supply_gate=True)
    assert _get(on, "M2", "需要調貨") == "是"
    assert int(_get(on, "M2", "實際調入量")) == 14
    assert _get(on, "M2", "調貨來源 (Site Code)") == "N1:14"


# ── 3. 調貨補足狀態三態欄 ────────────────────────────────────────────────
def test_fulfill_column_written_when_on():
    raw, *_ = transform_hybrid(_df(_ZERO_EXCESS), real_supply_gate=True)
    assert _HYBRID_FULFILL_COLUMN in raw.columns
    assert set(raw[_HYBRID_FULFILL_COLUMN].unique()) <= {"", _HYBRID_FULFILL_FULL}


def test_fulfill_states_match_actual_and_gap():
    raw, *_ = transform_hybrid(_df(_CONTENTION), real_supply_gate=True)
    for site in ("T1", "T2", "T3"):
        actual = int(_get(raw, site, "實際調入量"))
        gap = int(_get(raw, site, "調入缺口"))
        expected = (_HYBRID_FULFILL_FULL if gap == 0
                    else (_HYBRID_FULFILL_PARTIAL if actual > 0 else _HYBRID_FULFILL_NONE))
        assert _get(raw, site, _HYBRID_FULFILL_COLUMN) == expected
    states = set(raw[_HYBRID_FULFILL_COLUMN].unique())
    assert _HYBRID_FULFILL_PARTIAL in states and _HYBRID_FULFILL_NONE in states


def test_fulfill_blank_for_non_need_rows():
    raw, *_ = transform_hybrid(_df(_CONTENTION), real_supply_gate=True)
    non_need = raw[~raw["需要調貨"].eq("是")]
    assert (non_need[_HYBRID_FULFILL_COLUMN] == "").all()


def test_fulfill_raw_display_position():
    raw, *_ = transform_hybrid(_df(_CONTENTION), real_supply_gate=True)
    disp = _hybrid_raw_display_columns(raw)
    assert disp.index("需要調貨") < disp.index(_HYBRID_FULFILL_COLUMN)
    assert disp.index(_HYBRID_FULFILL_COLUMN) < disp.index("事件類型")


def test_fulfill_main_table_inserted_after_demand():
    _, _, diag, _, _, _, _ = transform_hybrid(
        _df(_CONTENTION), real_supply_gate=True)
    assert _HYBRID_FULFILL_COLUMN not in _HYBRID_DIAG_COLUMNS  # 不污染 29 欄契約
    assert _HYBRID_FULFILL_COLUMN in diag.columns
    cols = list(diag.columns)
    assert cols.index(_HYBRID_FULFILL_COLUMN) == cols.index("調入需求量 (補到14日)") + 1
    off_diag = transform_hybrid(_df(_CONTENTION))[2]
    assert _HYBRID_FULFILL_COLUMN not in off_diag.columns


def test_same_om_variant_without_same_om_source_is_still_fulfilled():
    """不限OM 可調入、同OM 變體無同 OM 來源 → 補足狀態以不限OM 實績為準。"""
    rows = [
        _row("SKU-O2", "Non Tourist", "N1", "OM1", 40, 90, 10),
        _row("SKU-O2", "Tourist", "T1", "OM2", 0, 90, 7),
        _row("SKU-O2", "Non Tourist", "N2", "OM9", 40, 90, 10),
    ]
    raw, *_ = transform_hybrid(_df(rows), real_supply_gate=True)
    assert int(_get(raw, "T1", "實際調入量")) == 14
    assert int(_get(raw, "T1", "實際調入量 (同OM)")) == 0
    assert _get(raw, "T1", _HYBRID_FULFILL_COLUMN) == _HYBRID_FULFILL_FULL


# ── 4. stats 與指標卡 ───────────────────────────────────────────────────
def test_stats_record_switch_and_fulfill_counts():
    raw, _, _, _, _, stats, _ = transform_hybrid(
        _df(_CONTENTION), real_supply_gate=True)
    need_n = int((raw["需要調貨"] == "是").sum())
    assert stats["real_supply_gate"] is True
    assert (stats["fulfill_full_rows"] + stats["fulfill_partial_rows"]
            + stats["no_supply_need_rows"]) == need_n
    assert stats["no_supply_need_rows"] == int(
        (raw[_HYBRID_FULFILL_COLUMN] == _HYBRID_FULFILL_NONE).sum())
    assert stats["no_supply_need_skus"] >= 1


def test_dashboard_adds_fulfill_row_only_when_on():
    _, _, _, _, _, on, _ = transform_hybrid(_df(_CONTENTION), real_supply_gate=True)
    _, _, _, _, _, off, _ = transform_hybrid(_df(_CONTENTION))
    on_labels = [c[0] for r in _hybrid_dashboard_metrics(None, on) for c in r]
    off_labels = [c[0] for r in _hybrid_dashboard_metrics(None, off) for c in r]
    assert on_labels[-4:] == ["已補足列數", "部分補足列數", "無貨可調列數", "無貨可調 SKU 數"]
    assert len(on_labels) - len(off_labels) == 4


def test_stats_from_raw_recover_fulfill_counts():
    """補足狀態可由 RAW 反推，開關狀態不可（回退路徑不列，與 nd_priority 準則一致）。"""
    raw, *_ = transform_hybrid(_df(_CONTENTION), real_supply_gate=True)
    fb = _hybrid_stats_from_raw(raw)
    assert fb["no_supply_need_rows"] == int(
        (raw[_HYBRID_FULFILL_COLUMN] == _HYBRID_FULFILL_NONE).sum())
    assert "real_supply_gate" not in fb


# ── 5. 店舖問題 SKU 的缺口說明分流 ───────────────────────────────────────
_ISSUE_ROWS = [
    _row("SKU-S", "Non Tourist", "N1", "OM1", 17, 90, 3),
    _row("SKU-S", "Tourist", "T1", "OM1", 9, 90, 2),    # 部分補足
    _row("SKU-S", "Tourist", "T3", "OM1", 12, 90, 2),   # 無貨可調
]


def test_shop_issue_advice_distinguishes_no_supply():
    raw, *_ = transform_hybrid(_df(_ISSUE_ROWS), real_supply_gate=True)
    issues = _build_hybrid_shop_issues(raw)
    advice = {str(r["Site Code"]): str(r["建議方法"]) for _, r in issues.iterrows()}
    assert "無任何店舖可調出貨品" in advice["T3"] and "D001" in advice["T3"]
    assert "無任何店舖可調出貨品" not in advice["T1"] and "尚有缺口" in advice["T1"]


def test_shop_issue_advice_unchanged_when_off():
    raw, *_ = transform_hybrid(_df(_ISSUE_ROWS))
    advice = " ".join(str(v) for v in _build_hybrid_shop_issues(raw)["建議方法"])
    assert "無任何店舖可調出貨品" not in advice


# ── 6. 說明頁 ──────────────────────────────────────────────────────────
def test_guide_sheet_lists_switch():
    raw, sc, diag, diag_om, ds, stats, _ = transform_hybrid(
        _df(_CONTENTION), real_supply_gate=True)
    ws = load_workbook(generate_excel_hybrid(
        raw, sc, diag, diag_om, ds, stats))["說明"]
    flat = "".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    flat = flat.replace("\n", "").replace("\r", "")
    assert "有貨可調改用實際可轉出量" in flat
    assert "有貨可調改用實際可轉出量（本份已開啟）" in flat
    assert "調貨補足狀態" in flat
