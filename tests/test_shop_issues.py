"""Validation tests for Shop Issue tables (店舖問題 SKU, v1.24.0).

三模式各加一表（只收需行動列，按店舖分組，問題＋建議方法）：
- Stock Checking：需要調入／可轉出供應（缺口在同一列加註）。
- Safety Stock Review：Safety 需增加／減少、ND 退場（New SKU 標記不收錄）。
- Hybrid：只收 Safety 設定不足／過高（設定正常列排除）。
"""
import warnings
warnings.filterwarnings("ignore")
import pandas as pd
from io import BytesIO
from openpyxl import load_workbook
from transformer import (
    transform_stock_checking,
    generate_excel_stock_checking,
    _build_stock_checking_shop_issues,
    _count_transfer_sources,
    _STOCK_CHECKING_SHOP_ISSUE_COLUMNS,
    transform_combined,
    generate_excel_combined,
    _build_combined_shop_issues,
    _COMBINED_SHOP_ISSUE_COLUMNS,
    transform_hybrid,
    generate_excel_hybrid,
    _build_hybrid_shop_issues,
    _HYBRID_SHOP_ISSUE_COLUMNS,
)

SC_BASE = {
    "Buyer": "B1", "Brand": "A", "SKU": 1, "SKU Desc": "Test Product",
    "Item_Shop": "Prod-S01-ShopHK", "Site Code": "S01", "OM": "OM1",
    "Shop Nature": "Tourist", "Vendor MC": "V1", "Item Status": "A",
    "Adj. SST": "SST1",
    "Launch Date": "2024-01-01", "Item Class": "C1", "RP Type": "RF",
    "Shop Stock Status": "Normal", "Adj. Safety Stock": 3.0,
    "Sold Qty (P90D)": 90.0, "Excess Stock Qty": 2.0,
    "Shop Stock Qty (CM)": 5.0,
    "D001 Stock Qty (CM)": 200.0, "D001 Net Stock Qty (CM)": 99.0,
    "Turnday (P90D Qty)": 13.0,
}


def _sc_row(**kw):
    r = dict(SC_BASE)
    r.update(kw)
    return r


def test_sc_shop_issues_only_actionable_and_sorted_by_shop():
    df = pd.DataFrame([
        _sc_row(**{"SKU": 100, "Site Code": "T01", "Shop Nature": "Tourist",
                   "Sold Qty (P90D)": 360.0, "Shop Stock Qty (CM)": 47.0,
                   "Shop Monthly Sold Qty (P90D)": 120.0, "D001 Net Stock Qty (CM)": 10.0}),
        _sc_row(**{"SKU": 100, "Site Code": "N01", "Shop Nature": "Non Tourist",
                   "Sold Qty (P90D)": 90.0, "Shop Stock Qty (CM)": 30.0,
                   "Shop Monthly Sold Qty (P90D)": 30.0}),
        _sc_row(**{"SKU": 101, "Site Code": "M02", "Shop Nature": "Mix",
                   "Shop Monthly Sold Qty (P90D)": 40.0, "Shop Stock Qty (CM)": 10.0}),
    ])
    raw, _, stats, _ = transform_stock_checking(df)
    issues = _build_stock_checking_shop_issues(
        raw, turnday_threshold=stats["turnday_threshold"], d001_threshold=stats["d001_threshold"])
    assert list(issues.columns[:11]) == list(_STOCK_CHECKING_SHOP_ISSUE_COLUMNS[:11])
    assert set(issues["問題類型"]) == {"需要調入", "可轉出供應"}
    assert list(issues["Site Code"]) == sorted(issues["Site Code"].tolist())
    need = issues[issues["Site Code"] == "T01"].iloc[0]
    assert "建議調入" in str(need["建議方法"]) and "間來源店舖" in str(need["建議方法"])
    donor = issues[issues["Site Code"] == "N01"].iloc[0]
    assert "可轉出" in str(donor["建議方法"])
    assert "M02" not in set(issues["Site Code"])
    assert int(raw.loc[raw["需要調貨"] == "是", "調入需求量 (補到14日)"].sum()) == int(
        issues.loc[issues["問題類型"] == "需要調入", "調入需求量 (補到14日)"].sum())


def test_sc_shop_issues_gap_flagged_in_same_row():
    df = pd.DataFrame([
        _sc_row(**{"SKU": 200, "Site Code": "T01", "Shop Nature": "Tourist",
                   "Sold Qty (P90D)": 720.0, "Shop Stock Qty (CM)": 92.0,
                   "D001 Net Stock Qty (CM)": 10.0}),
        _sc_row(**{"SKU": 200, "Site Code": "N01", "Shop Nature": "Non Tourist",
                   "Sold Qty (P90D)": 90.0, "Shop Stock Qty (CM)": 19.0}),
    ])
    df = df.drop(columns=["Shop Monthly Sold Qty (P90D)"], errors="ignore")
    raw, _, stats, _ = transform_stock_checking(df)
    issues = _build_stock_checking_shop_issues(
        raw, turnday_threshold=stats["turnday_threshold"], d001_threshold=stats["d001_threshold"])
    gap_row = issues[issues["Site Code"] == "T01"].iloc[0]
    assert gap_row["問題類型"] == "調入缺口（來源不足）"
    assert "缺口" in str(gap_row["建議方法"])


def test_sc_shop_issues_empty_inputs():
    for frame in (_build_stock_checking_shop_issues(None),
                  _build_stock_checking_shop_issues(pd.DataFrame()),
                  _build_stock_checking_shop_issues(pd.DataFrame([{"Brand": "A"}]))):
        assert frame.empty and list(frame.columns) == list(_STOCK_CHECKING_SHOP_ISSUE_COLUMNS)


def test_sc_excel_contains_shop_issue_sheet():
    df = pd.DataFrame([
        _sc_row(**{"SKU": 100, "Site Code": "T01", "Shop Nature": "Tourist",
                   "Sold Qty (P90D)": 360.0, "Shop Stock Qty (CM)": 47.0,
                   "Shop Monthly Sold Qty (P90D)": 120.0, "D001 Net Stock Qty (CM)": 10.0}),
        _sc_row(**{"SKU": 100, "Site Code": "N01", "Shop Nature": "Non Tourist",
                   "Sold Qty (P90D)": 90.0, "Shop Stock Qty (CM)": 30.0,
                   "Shop Monthly Sold Qty (P90D)": 30.0}),
    ])
    raw, summary, stats, _ = transform_stock_checking(df)
    buf = generate_excel_stock_checking(raw, summary, stats)
    wb = load_workbook(buf, read_only=True)
    assert wb.sheetnames == ["Summary", "RAW", "店舖問題 SKU", "調貨分析摘要"]
    headers = [c.value for c in next(wb["店舖問題 SKU"].iter_rows(min_row=3, max_row=3))]
    assert headers[:11] == list(_STOCK_CHECKING_SHOP_ISSUE_COLUMNS[:11])
    wb2 = load_workbook(buf)
    assert wb2["店舖問題 SKU"].freeze_panes == "A4"


CB_BASE = dict(SC_BASE)
CB_RATES_O = {"tourist": 0.30, "mix": 0.40, "default": 0.50}
CB_RATES_P = {"tourist": 0.50, "mix": 0.50, "default": 0.50}


def _cb_row(**kw):
    r = dict(CB_BASE)
    r.update(kw)
    return r


def test_combined_shop_issues_only_actionable():
    df = pd.DataFrame([
        _cb_row(**{"SKU": "1", "Shop Stock Status": "Normal", "RP Type": "RF",
                   "Shop Monthly Sold Qty (P90D)": 30.0, "Sold Qty (CM - MTD)": 8.0}),
        _cb_row(**{"SKU": "2", "Shop Stock Status": "Overstock (P90D 0 Sales)", "RP Type": "RF",
                   "Shop Monthly Sold Qty (P90D)": 0.0, "Sold Qty (CM - MTD)": 0.0,
                   "Launch Date": "2020-01-01"}),
        _cb_row(**{"SKU": "3", "Shop Stock Status": "Normal", "RP Type": "ND",
                   "Shop Monthly Sold Qty (P90D)": 30.0, "Sold Qty (CM - MTD)": 8.0}),
    ])
    _, raw, _, _ = transform_combined(
        df, rates_overstock=CB_RATES_O, rates_oos=CB_RATES_P,
        mtd_days=15, month_total_days=30, safety_mode="allow_both")
    issues = _build_combined_shop_issues(raw)
    assert list(issues.columns[:11]) == list(_COMBINED_SHOP_ISSUE_COLUMNS[:11])
    assert "New SKU (3 Month)" not in set(
        issues["Recommendation of new Safety Qty / ND"].fillna("").astype(str))
    assert set(issues["問題類型"]) <= {"Safety 設定不足", "Safety 設定過高", "ND 退場"}
    assert len(issues) >= 1
    assert list(issues["Site Code"]) == sorted(issues["Site Code"].tolist())
    assert all(str(v).strip() for v in issues["問題說明"])
    assert all(str(v).strip() for v in issues["建議方法"])


def test_combined_shop_issues_empty_inputs():
    for frame in (_build_combined_shop_issues(None),
                  _build_combined_shop_issues(pd.DataFrame()),
                  _build_combined_shop_issues(pd.DataFrame([{"Brand": "A"}]))):
        assert frame.empty and list(frame.columns) == list(_COMBINED_SHOP_ISSUE_COLUMNS)


def test_combined_excel_contains_shop_issue_sheet():
    df = pd.DataFrame([
        _cb_row(**{"SKU": "1", "Shop Stock Status": "Normal", "RP Type": "RF",
                   "Shop Monthly Sold Qty (P90D)": 30.0, "Sold Qty (CM - MTD)": 8.0}),
        _cb_row(**{"SKU": "2", "Shop Stock Status": "Normal", "RP Type": "ND",
                   "Shop Monthly Sold Qty (P90D)": 30.0, "Sold Qty (CM - MTD)": 8.0}),
    ])
    s, r, c, _ = transform_combined(
        df, rates_overstock=CB_RATES_O, rates_oos=CB_RATES_P,
        mtd_days=15, month_total_days=30, safety_mode="allow_both")
    buf = generate_excel_combined(s, r, c)
    wb = load_workbook(buf, read_only=True)
    assert wb.sheetnames == ["RAW", "Shop Stock Status", "By Shop", "店舖問題 SKU", "SUMMARY", "Safety Stock 數據分析摘要"]


HY_BASE = dict(SC_BASE)


def _hy_row(**kw):
    r = dict(HY_BASE)
    r.update(kw)
    return r


_HY_SCENARIO = [
    {"SKU": 300, "Site Code": "T11", "Shop Nature": "Tourist",
     "Sold Qty (P90D)": 360.0, "Shop Stock Qty (CM)": 47.0,
     "Shop Monthly Sold Qty (P90D)": 120.0, "Adj. Safety Stock": 30,
     "D001 Net Stock Qty (CM)": 10.0},
    {"SKU": 300, "Site Code": "N11", "Shop Nature": "Non Tourist",
     "Sold Qty (P90D)": 90.0, "Shop Stock Qty (CM)": 30.0,
     "Shop Monthly Sold Qty (P90D)": 30.5, "Adj. Safety Stock": 12},
    {"SKU": 300, "Site Code": "M11", "Shop Nature": "Mix",
     "Sold Qty (P90D)": 90.0, "Shop Stock Qty (CM)": 25.0,
     "Shop Monthly Sold Qty (P90D)": 30.5, "Adj. Safety Stock": 10},
]


def test_hybrid_shop_issues_only_fix_rows():
    raw, _, _, _, _, _, _ = transform_hybrid(
        pd.DataFrame([_hy_row(**r) for r in _HY_SCENARIO]),
        safety_days={"Tourist": 10, "Mix": 10, "Non Tourist": 10})
    issues = _build_hybrid_shop_issues(raw)
    assert list(issues.columns[:11]) == list(_HYBRID_SHOP_ISSUE_COLUMNS[:11])
    assert set(issues["問題類型"]) == {"Safety 設定不足", "Safety 設定過高"}
    assert "設定正常" not in " ".join(issues["問題類型"].astype(str).tolist())
    assert "實際調入量 (同OM)" in issues.columns
    assert list(issues["Site Code"]) == sorted(issues["Site Code"].tolist())
    under = issues[issues["問題類型"] == "Safety 設定不足"].iloc[0]
    assert "調整為" in str(under["建議方法"])
    assert "間來源店舖" in str(under["建議方法"])


def test_advice_does_not_repeat_source_sites():
    """建議方法只計涉及店舖數量，不重複列出店號（店號留給調貨來源欄）。"""
    df = pd.DataFrame([
        _sc_row(**{"SKU": 100, "Site Code": "T01", "Shop Nature": "Tourist",
                   "Sold Qty (P90D)": 360.0, "Shop Stock Qty (CM)": 47.0,
                   "Shop Monthly Sold Qty (P90D)": 120.0, "D001 Net Stock Qty (CM)": 10.0}),
        _sc_row(**{"SKU": 100, "Site Code": "N01", "Shop Nature": "Non Tourist",
                   "Sold Qty (P90D)": 90.0, "Shop Stock Qty (CM)": 30.0,
                   "Shop Monthly Sold Qty (P90D)": 30.0}),
    ])
    raw, _, stats, _ = transform_stock_checking(df)
    issues = _build_stock_checking_shop_issues(
        raw, turnday_threshold=stats["turnday_threshold"], d001_threshold=stats["d001_threshold"])
    need_src = str(raw.loc[raw["需要調貨"] == "是", "調貨來源 (Site Code)"].iloc[0])
    need_advice = str(issues.loc[issues["問題類型"] == "需要調入", "建議方法"].iloc[0])
    n_sources = _count_transfer_sources(need_src)
    assert f"涉及 {n_sources} 間來源店舖" in need_advice
    for site in [p.split(":")[0].strip() for p in need_src.split(";") if ":" in p]:
        assert site not in need_advice
    hy_raw, _, _, _, _, _, _ = transform_hybrid(
        pd.DataFrame([_hy_row(**r) for r in _HY_SCENARIO]),
        safety_days={"Tourist": 10, "Mix": 10, "Non Tourist": 10})
    hy_issues = _build_hybrid_shop_issues(hy_raw)
    hy_under = hy_issues[hy_issues["問題類型"] == "Safety 設定不足"].iloc[0]
    assert "間來源店舖" in str(hy_under["建議方法"])


def test_hybrid_shop_issues_empty_inputs():
    for frame in (_build_hybrid_shop_issues(None),
                  _build_hybrid_shop_issues(pd.DataFrame()),
                  _build_hybrid_shop_issues(pd.DataFrame([{"Brand": "A"}]))):
        assert frame.empty and list(frame.columns) == list(_HYBRID_SHOP_ISSUE_COLUMNS)


def test_shop_issues_exclude_noise_columns():
    """被點名不顯示的欄位不得出現在三張新表（即使 RAW 帶有這些欄）。"""
    noise = {
        "Item_Shop": "Prod-X99-ShopX", "Valuation Class": "VC1", "Adj. SST": "SST9",
        "MCH2": "M1", "MC": "MC1", "Shop Type": "ST1",
        "Safety VS Monthly Sold Qty": 1.5, "D001 Stock Qty (LM)": 7,
        "D001 Blocked Qty (CM)": 3, "Shop Stock Qty (LM)": 6,
        "Shop Stock Value (LM)": 60.0, "Shop Stock Value (CM)": 50.0,
        "Excess Stock Qty": 2, "Excess Stock Value": 20.0,
        "Effective Shop Monthly Sold Qty (P90D)": 30.0,
        "D001 Effective Stock Qty (CM)": 9,
    }
    df = pd.DataFrame([
        _sc_row(**{"SKU": 100, "Site Code": "T01", "Shop Nature": "Tourist",
                   "Sold Qty (P90D)": 360.0, "Shop Stock Qty (CM)": 47.0,
                   "Shop Monthly Sold Qty (P90D)": 120.0, "D001 Net Stock Qty (CM)": 10.0, **noise}),
        _sc_row(**{"SKU": 100, "Site Code": "N01", "Shop Nature": "Non Tourist",
                   "Sold Qty (P90D)": 90.0, "Shop Stock Qty (CM)": 30.0,
                   "Shop Monthly Sold Qty (P90D)": 30.0, **noise}),
    ])
    raw, _, stats, _ = transform_stock_checking(df)
    sc_issues = _build_stock_checking_shop_issues(
        raw, turnday_threshold=stats["turnday_threshold"], d001_threshold=stats["d001_threshold"])
    assert list(sc_issues.columns) == [c for c in _STOCK_CHECKING_SHOP_ISSUE_COLUMNS if c in sc_issues.columns]
    for col in noise:
        assert col not in sc_issues.columns, col
    _, cb_raw, _, _ = transform_combined(
        df.rename(columns={"Shop Stock Status": "Shop Stock Status"}),
        rates_overstock={"tourist": 0.30, "mix": 0.40, "default": 0.50},
        rates_oos={"tourist": 0.50, "mix": 0.50, "default": 0.50},
        mtd_days=15, month_total_days=30, safety_mode="allow_both")
    cb_issues = _build_combined_shop_issues(cb_raw)
    for col in noise:
        if col != "Adj. SST":
            assert col not in cb_issues.columns, col
    hy_raw, _, _, _, _, _, _ = transform_hybrid(
        pd.DataFrame([_hy_row(**r, **noise) for r in _HY_SCENARIO]),
        safety_days={"Tourist": 10, "Mix": 10, "Non Tourist": 10})
    hy_issues = _build_hybrid_shop_issues(hy_raw)
    for col in noise:
        assert col not in hy_issues.columns, col


def test_hybrid_excel_contains_shop_issue_sheet():
    raw, sc, d, do, cs, stats, _ = transform_hybrid(
        pd.DataFrame([_hy_row(**r) for r in _HY_SCENARIO]),
        safety_days={"Tourist": 10, "Mix": 10, "Non Tourist": 10})
    buf = generate_excel_hybrid(raw, sc, d, do, cs, stats)
    wb = load_workbook(buf, read_only=True)
    assert wb.sheetnames == [
        "RAW", "Summary", "調貨 Review Safety Stock (不限OM)", "調貨 Review Safety Stock (同OM)",
        "店舖問題 SKU", "Review 摘要",
        "簡易調貨明細 (不限OM)", "簡易調貨明細 (同OM)",
    ]
