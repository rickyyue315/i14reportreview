"""Validation tests for 退場清倉 Review (segment 6)."""
import warnings
warnings.filterwarnings("ignore")
import pandas as pd
from io import BytesIO
from openpyxl import load_workbook
from transformer import (
    transform_phaseout,
    generate_excel_phaseout,
)

BASE = {
    "Buyer": "B1", "Brand": "A", "SKU": 1,
    "Item_Shop": "Prod-S01-ShopHK", "Site Code": "S01", "OM": "OM1",
    "Shop Nature": "Mix", "RP Type": "RF", "Supply Source": "Warehouse",
    "Shop Stock Status": "Normal", "Item Status": "Active", "Item Class": "C1",
    "Shop Stock Qty (CM)": 5.0,
    "Sold Qty (M-2)": 10.0, "Sold Qty (M-1)": 10.0,
    "Sold Qty (CM - MTD)": 5.0, "Sold Qty (P90D)": 30.0,
    "Excess Stock Value": 0.0,
}


def _row(**kw):
    r = dict(BASE)
    r.update(kw)
    return r


def _df(rows):
    return pd.DataFrame([_row(**r) for r in rows])


def test_phaseout_status_warehouse_and_vendor():
    df = _df([
        {"SKU": 1, "Item Status": "Discontinued"},
        {"SKU": 2, "Item Status": "已下市", "Supply Source": "Vendor"},
        {"SKU": 3, "Item Status": "Active"},  # 正常 → 排除
        {"SKU": 4, "Item Status": "停產", "Shop Stock Qty (CM)": 0.0},  # 無庫存 → 排除
    ])
    raw, by_sku, _, _ = transform_phaseout(df)
    got = dict(zip(raw["SKU"].astype(str), raw["清倉建議"]))
    assert got == {"1": "建議退倉", "2": "另有安排"}
    assert int(raw[raw["SKU"] == "1"]["Shop Stock Qty (CM)"].iloc[0]) == 5
    total = by_sku[by_sku["Buyer"] == "總計"].iloc[0]
    assert int(total["建議退倉合計"]) == 5


def test_nd_zero_sales_review():
    df = _df([
        {"SKU": 1, "RP Type": "ND", "Sold Qty (M-2)": 0.0, "Sold Qty (M-1)": 0.0,
         "Sold Qty (CM - MTD)": 0.0, "Sold Qty (P90D)": 0.0},
        {"SKU": 2, "RP Type": "ND", "Sold Qty (M-2)": 0.0, "Sold Qty (M-1)": 0.0,
         "Sold Qty (CM - MTD)": 0.0, "Sold Qty (P90D)": 0.0, "Shop Stock Qty (CM)": 0.0},  # 無庫存 → 排除
    ])
    raw, _, _, _ = transform_phaseout(df)
    assert list(raw["SKU"].astype(str)) == ["1"]
    assert raw["清倉建議"].iloc[0] == "檢視退場"


def test_missing_item_status_fallback():
    df = pd.DataFrame([{
        "Buyer": "B1", "Brand": "A", "SKU": 1, "Shop Stock Status": "Normal",
        "Site Code": "S01", "RP Type": "ND", "Shop Stock Qty (CM)": 4.0,
        "Sold Qty (M-2)": 0.0, "Sold Qty (M-1)": 0.0,
        "Sold Qty (CM - MTD)": 0.0, "Sold Qty (P90D)": 0.0,
    }])
    raw, _, _, warnings = transform_phaseout(df)
    assert raw["清倉建議"].iloc[0] == "檢視退場"
    assert any("Item Status" in w for w in warnings)


def test_missing_site_code_rejected():
    df = pd.DataFrame([{"Buyer": "B1", "Brand": "A", "SKU": 1, "Shop Stock Status": "Normal", "Item Status": "下市"}])
    try:
        transform_phaseout(df)
    except ValueError as e:
        assert "Site Code" in str(e)
    else:
        raise AssertionError("應拒絕缺少 Site Code 的輸入")


def test_by_shop_totals():
    df = _df([
        {"SKU": 1, "Item Status": "下市", "Shop Stock Qty (CM)": 3.0},
        {"SKU": 2, "Item Status": "下市", "Shop Stock Qty (CM)": 7.0},
    ])
    _, _, by_shop, _ = transform_phaseout(df)
    total = by_shop[by_shop["Site Code"] == "總計"].iloc[0]
    assert int(total["涉及 SKU 數"]) == 2
    assert int(total["建議退倉總數"]) == 10


def test_excel_sheets_and_skip_raw():
    df = _df([{"SKU": 1, "Item Status": "下市"}])
    raw, by_sku, by_shop, _ = transform_phaseout(df)
    buf = generate_excel_phaseout(raw, by_sku, by_shop, include_raw=True)
    wb = load_workbook(BytesIO(buf.getvalue()), data_only=True)
    assert wb.sheetnames == ["RAW", "By SKU", "By Shop"]
    buf2 = generate_excel_phaseout(raw, by_sku, by_shop, include_raw=False)
    wb2 = load_workbook(BytesIO(buf2.getvalue()), data_only=True)
    assert wb2["RAW"]["A2"].value is not None


def test_empty_result_warning():
    df = _df([{"SKU": 1}])
    raw, _, _, warnings = transform_phaseout(df)
    assert raw.empty
    assert any("沒有符合" in w for w in warnings)
