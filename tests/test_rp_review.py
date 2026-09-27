"""Validation tests for RP Type 覆核 Review (segment 3)."""
import warnings
warnings.filterwarnings("ignore")
import pandas as pd
from io import BytesIO
from openpyxl import load_workbook
from transformer import (
    transform_rp_review,
    generate_excel_rp_review,
)

BASE = {
    "Buyer": "B1", "Brand": "A", "SKU": 1,
    "Item_Shop": "Prod-S01-ShopHK", "Site Code": "S01", "OM": "OM1",
    "Shop Nature": "Mix", "RP Type": "RF",
    "Shop Stock Status": "Normal",
    "Shop Stock Qty (CM)": 5.0,
    "Sold Qty (M-2)": 10.0, "Sold Qty (M-1)": 10.0,
    "Sold Qty (CM - MTD)": 5.0, "Sold Qty (P90D)": 30.0,
    "Is OOS (CM)": False, "OOS Periods Count": 0,
}


def _row(**kw):
    r = dict(BASE)
    r.update(kw)
    return r


def _df(rows):
    return pd.DataFrame([_row(**r) for r in rows])


def test_nd_with_sales_flagged_to_rf():
    df = _df([
        {"SKU": 1, "RP Type": "ND", "Sold Qty (P90D)": 8.0},
        {"SKU": 2, "RP Type": " nd ", "Sold Qty (M-1)": 3.0, "Sold Qty (P90D)": 0.0,
         "Sold Qty (M-2)": 0.0, "Sold Qty (CM - MTD)": 0.0},
        {"SKU": 3, "RP Type": "ND", "Sold Qty (P90D)": 0.0, "Sold Qty (M-2)": 0.0,
         "Sold Qty (M-1)": 0.0, "Sold Qty (CM - MTD)": 0.0},  # ND 零銷售 → 正常，不收錄
    ])
    raw, _, _, _ = transform_rp_review(df)
    assert set(raw["SKU"].astype(str)) == {"1", "2"}
    assert set(raw["覆核建議"]) == {"疑似應轉 RF"}


def test_rf_zero_sales_with_stock_flagged_to_nd():
    df = _df([
        {"SKU": 1, "RP Type": "RF", "Sold Qty (M-2)": 0.0, "Sold Qty (M-1)": 0.0,
         "Sold Qty (CM - MTD)": 0.0, "Sold Qty (P90D)": 0.0, "Shop Stock Qty (CM)": 4.0},
        {"SKU": 2, "RP Type": "RF", "Sold Qty (M-2)": 0.0, "Sold Qty (M-1)": 0.0,
         "Sold Qty (CM - MTD)": 0.0, "Sold Qty (P90D)": 0.0, "Shop Stock Qty (CM)": 0.0},  # 無庫存 → 不收錄
        {"SKU": 3, "RP Type": "RF"},  # 有銷量 → 不收錄
    ])
    raw, _, _, _ = transform_rp_review(df)
    assert list(raw["SKU"].astype(str)) == ["1"]
    assert raw["覆核建議"].iloc[0] == "疑似應轉 ND"


def test_nd_oos_priority():
    df = _df([
        {"SKU": 1, "RP Type": "ND", "Shop Stock Qty (CM)": 0.0, "Sold Qty (P90D)": 9.0},
    ])
    raw, _, _, _ = transform_rp_review(df)
    assert raw["覆核建議"].iloc[0] == "ND 缺貨跟進"


def test_normal_rows_excluded():
    df = _df([
        {"SKU": 1, "RP Type": "RF"},  # RF 有銷量 → 正常
        {"SKU": 2, "RP Type": "ND", "Sold Qty (M-2)": 0.0, "Sold Qty (M-1)": 0.0,
         "Sold Qty (CM - MTD)": 0.0, "Sold Qty (P90D)": 0.0},  # ND 零銷售 → 正常
    ])
    raw, _, _, warnings = transform_rp_review(df)
    assert raw.empty
    assert any("沒有符合" in w for w in warnings)


def test_nonstandard_rp_treated_as_nd_with_warning():
    df = _df([
        {"SKU": 1, "RP Type": "XX", "Sold Qty (P90D)": 6.0},
    ])
    raw, _, _, warnings = transform_rp_review(df)
    assert raw["覆核建議"].iloc[0] == "疑似應轉 RF"
    assert any("非標準值" in w for w in warnings)


def test_missing_rp_type_rejected():
    df = pd.DataFrame([{"Buyer": "B1", "Brand": "A", "SKU": 1, "Shop Stock Status": "Normal", "Site Code": "S01"}])
    try:
        transform_rp_review(df)
    except ValueError as e:
        assert "RP Type" in str(e)
    else:
        raise AssertionError("應拒絕缺少 RP Type 的輸入")


def test_excel_sheets_and_skip_raw():
    df = _df([
        {"SKU": 1, "RP Type": "ND", "Sold Qty (P90D)": 9.0},
    ])
    raw, by_sku, by_shop, _ = transform_rp_review(df)
    buf = generate_excel_rp_review(raw, by_sku, by_shop, include_raw=True)
    wb = load_workbook(BytesIO(buf.getvalue()), data_only=True)
    assert wb.sheetnames == ["RAW", "By SKU", "By Shop"]
    buf2 = generate_excel_rp_review(raw, by_sku, by_shop, include_raw=False)
    wb2 = load_workbook(BytesIO(buf2.getvalue()), data_only=True)
    assert wb2.sheetnames == ["RAW", "By SKU", "By Shop"]
    assert wb2["RAW"]["A2"].value is not None
