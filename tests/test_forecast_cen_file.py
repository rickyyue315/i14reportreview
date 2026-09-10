"""Regression test for Forecast Review on the real CEN surge file (v1.39.0, v2 demand model).

來源：By SKUxShop_CEN_09Sep2026.xlsx（SKU 111143702001，MTD=2026-09-08）。
預期：該 SKU 判「爆升」（M-2=0、M-1=178），SKU 位置合計＝947
（店存 426＋D001 Net 1＋D001 PO 0＋在途 520）、建議訂貨對象＝D001 PO。
v2 下 M-2 還原與開賣折算未觸發（實檔缺 Shop Stock Qty (M-2) 欄、亦未傳 MTD 日期），
爆升分類與位置欄與 v1.38.0 相同；建議訂量因趨勢外推（S2=0 取增速上限）而高於 v1。
檔案不存在時 skip（不阻塞 CI）。
"""
import pathlib
import warnings
warnings.filterwarnings("ignore")
import pandas as pd
import pytest

CEN_FILE = pathlib.Path(__file__).resolve().parents[1] / "By SKUxShop_CEN_09Sep2026.xlsx"

pytestmark = pytest.mark.skipif(not CEN_FILE.exists(), reason="CEN 實檔不存在，跳過回歸測試")


def _load_cen() -> pd.DataFrame:
    from transformer import _clean_dataframe
    df = pd.read_excel(CEN_FILE, sheet_name=0)
    # 實檔欄名即來源欄名；_clean_dataframe 處理 Total 列與 Applied filters 註腳
    return _clean_dataframe(df)


def test_cen_file_loads_and_cleans():
    df = _load_cen()
    assert len(df) > 80  # 約 89 店列
    assert "Total" not in df.get("Item_Shop", pd.Series(dtype=str)).fillna("").astype(str).str.strip().tolist()


def test_cen_surge_sku_class_position_target():
    from transformer import transform_forecast
    df = _load_cen()
    main, cov, buyer, raw, stats, warns = transform_forecast(df, mtd_days=8, month_total_days=30)
    hit = main[main["SKU"] == "111143702001"]
    assert len(hit) == 1
    row = hit.iloc[0]
    assert row["FC 爆升分類"] == "爆升"  # M-2=0、M-1=178：max(178, 月化)/max(0,10) 遠超 1.5
    assert row["FC M-2 合計"] == 0
    assert row["FC M-1 合計"] == 178
    assert row["FC MTD 合計"] == 1548
    assert row["FC 店舖庫存合計"] == 426
    assert row["FC D001 庫存 (Net)"] == 1
    assert row["FC D001 PO"] == 0
    assert row["FC 在途合計"] == 520
    assert row["FC 位置合計"] == 947  # 426＋1＋0＋520
    assert row["FC 建議訂貨對象"] == "D001 PO"
    assert row["FC 建議訂量"] > 0  # 倉端斷供，必有訂量
    assert row["FC 建議訂量"] == row["FC 缺口"]  # 上限不應綁住（月預估×2 遠大於缺口）
    # 缺貨還原：約三成店列為 CM 缺貨還原列
    assert int(raw["FC 缺貨還原列"].sum()) >= 25
    # 店舖覆蓋：有訂量 → 非空，且含已缺貨列
    assert len(cov) > 0
    assert "已缺貨" in cov["FC 風險旗"].tolist()


def test_cen_excel_roundtrip():
    from transformer import transform_forecast, generate_excel_forecast
    from openpyxl import load_workbook
    df = _load_cen()
    main, cov, buyer, raw, stats, _ = transform_forecast(df, mtd_days=8, month_total_days=30)
    buf = generate_excel_forecast(main, cov, buyer, raw, stats)
    wb = load_workbook(buf)
    assert wb.sheetnames == ["說明", "Forecast 主表", "店舖覆蓋", "按 Buyer 彙總", "RAW", "Forecast 摘要"]
    assert wb.active.title == "說明"
