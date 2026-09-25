"""存貨偏低標示（淺藍底）測試 — 庫存佔比明顯低於銷量佔比時只標示儲存格。"""
import pandas as pd
from openpyxl import load_workbook

from transformer import (
    _hybrid_shortage_thresholds,
    _hybrid_shortage_hint_text,
    _HYBRID_SHORTAGE_THRESHOLDS,
    transform_stock_checking,
    generate_excel_stock_checking,
    transform_hybrid,
    generate_excel_hybrid,
)


def _df(rows):
    base = {
        "Buyer": "B1", "Brand": "BR", "SKU": 999, "SKU Desc": "D",
        "Adj. SST": "S", "Supply Source": "W", "Site Code": "X",
        "Shop Nature": "Tourist", "Sold Qty (P90D)": 90.0,
        "Shop Stock Qty (CM)": 10.0, "Shop Monthly Sold Qty (P90D)": 30.0,
        "Turnday (P90D Qty)": 10.0, "D001 Net Stock Qty (CM)": 200.0,
        "Item_Shop": "x",
    }
    return pd.DataFrame([{**base, **r} for r in rows])


def _fill_rgb(cell):
    try:
        return str(cell.fill.start_color.rgb or "")
    except Exception:
        return ""


def test_shortage_thresholds_registry():
    assert _hybrid_shortage_thresholds("nature") == (0.30, 0.10)
    assert _hybrid_shortage_thresholds("rp") == (0.30, 0.10)
    assert _hybrid_shortage_thresholds("both") == (0.30, 0.10)
    assert _hybrid_shortage_thresholds("om") == (0.20, 0.08)
    assert _hybrid_shortage_thresholds("sales") == (0.20, 0.10)
    assert _hybrid_shortage_thresholds("nature_om") == (0.10, 0.05)
    assert _hybrid_shortage_thresholds("sales_nature") == (0.10, 0.05)
    # 較偏重側寬鬆
    assert _HYBRID_SHORTAGE_THRESHOLDS["nature"][0] < 0.40
    assert _HYBRID_SHORTAGE_THRESHOLDS["nature"][1] < 0.20


def test_shortage_hint_text():
    txt = _hybrid_shortage_hint_text("nature")
    assert "30%" in txt and "10" in txt
    assert "銷量佔比" in txt and "庫存佔比" in txt


def _shortage_rows():
    return [
        {"SKU": 999, "Site Code": "T01", "Shop Nature": "Tourist",
         "Shop Stock Qty (CM)": 5.0, "Shop Monthly Sold Qty (P90D)": 100.0,
         "Sold Qty (P90D)": 300.0},
        {"SKU": 999, "Site Code": "M01", "Shop Nature": "Mix",
         "Shop Stock Qty (CM)": 50.0, "Shop Monthly Sold Qty (P90D)": 10.0,
         "Sold Qty (P90D)": 30.0},
        {"SKU": 999, "Site Code": "N01", "Shop Nature": "Non Tourist",
         "Shop Stock Qty (CM)": 50.0, "Shop Monthly Sold Qty (P90D)": 10.0,
         "Sold Qty (P90D)": 30.0},
    ]


def test_stock_checking_shortage_blue_highlight():
    raw, summary, stats, _ = transform_stock_checking(_df(_shortage_rows()))
    # 只標示儲存格，不新增文字欄 — 欄位與數值契約不變
    assert "存貨偏低提示" not in summary.columns
    assert len(summary.columns) == 29
    buf = generate_excel_stock_checking(raw, summary, stats)
    ws = load_workbook(buf)["Summary"]
    # nature 維度：佔比區起點 20；Tourist 庫存 % 欄 = 20
    assert _fill_rgb(ws.cell(row=6, column=20)).endswith("BDD7EE")
    assert ws.cell(row=6, column=20).font.bold
    assert str(ws.cell(row=6, column=20).font.color.rgb).endswith("1F4E78")


def test_shortage_does_not_overwrite_skew_orange():
    raw, summary, stats, _ = transform_stock_checking(_df(_shortage_rows()))
    buf = generate_excel_stock_checking(raw, summary, stats)
    ws = load_workbook(buf)["Summary"]
    # Mix／Non Tourist 同時觸發偏重 — 維持橘底，不被淺藍覆蓋
    assert _fill_rgb(ws.cell(row=6, column=22)).endswith("FFC000")
    assert _fill_rgb(ws.cell(row=6, column=24)).endswith("FFC000")
    # Tourist 偏低欄不應為橘底
    assert not _fill_rgb(ws.cell(row=6, column=20)).endswith("FFC000")


def test_shortage_absent_when_balanced():
    rows = [
        {"SKU": 999, "Site Code": "T01", "Shop Nature": "Tourist",
         "Shop Stock Qty (CM)": 30.0, "Shop Monthly Sold Qty (P90D)": 30.0,
         "Sold Qty (P90D)": 90.0},
        {"SKU": 999, "Site Code": "M01", "Shop Nature": "Mix",
         "Shop Stock Qty (CM)": 30.0, "Shop Monthly Sold Qty (P90D)": 30.0,
         "Sold Qty (P90D)": 90.0},
        {"SKU": 999, "Site Code": "N01", "Shop Nature": "Non Tourist",
         "Shop Stock Qty (CM)": 30.0, "Shop Monthly Sold Qty (P90D)": 30.0,
         "Sold Qty (P90D)": 90.0},
    ]
    raw, summary, stats, _ = transform_stock_checking(_df(rows))
    buf = generate_excel_stock_checking(raw, summary, stats)
    ws = load_workbook(buf)["Summary"]
    for col in (20, 22, 24):
        assert not _fill_rgb(ws.cell(row=6, column=col)).endswith("BDD7EE")


def test_shortage_absent_when_no_sales():
    rows = [
        {"SKU": 999, "Site Code": "T01", "Shop Nature": "Tourist",
         "Shop Stock Qty (CM)": 0.0, "Shop Monthly Sold Qty (P90D)": 0.0,
         "Sold Qty (P90D)": 0.0},
        {"SKU": 999, "Site Code": "M01", "Shop Nature": "Mix",
         "Shop Stock Qty (CM)": 10.0, "Shop Monthly Sold Qty (P90D)": 0.0,
         "Sold Qty (P90D)": 0.0},
    ]
    raw, summary, stats, _ = transform_stock_checking(_df(rows))
    buf = generate_excel_stock_checking(raw, summary, stats)
    ws = load_workbook(buf)["Summary"]
    for row_cells in ws.iter_rows():
        for cell in row_cells:
            assert not _fill_rgb(cell).endswith("BDD7EE")


def test_hybrid_shortage_blue_highlight():
    rows = [
        {"SKU": 999, "Site Code": "T01", "Shop Nature": "Tourist",
         "Shop Stock Qty (CM)": 5.0, "Shop Monthly Sold Qty (P90D)": 100.0,
         "Sold Qty (P90D)": 300.0, "Adj. Safety Stock": 3,
         "D001 Net Stock Qty (CM)": 10.0},
        {"SKU": 999, "Site Code": "M01", "Shop Nature": "Mix",
         "Shop Stock Qty (CM)": 50.0, "Shop Monthly Sold Qty (P90D)": 10.0,
         "Sold Qty (P90D)": 30.0, "Adj. Safety Stock": 3},
        {"SKU": 999, "Site Code": "N01", "Shop Nature": "Non Tourist",
         "Shop Stock Qty (CM)": 50.0, "Shop Monthly Sold Qty (P90D)": 10.0,
         "Sold Qty (P90D)": 30.0, "Adj. Safety Stock": 3},
    ]
    raw, sc, dg, dg_om, ds, stats, _ = transform_hybrid(_df(rows))
    assert "存貨偏低提示" not in sc.columns
    buf = generate_excel_hybrid(raw, sc, dg, dg_om, ds, stats)
    ws = load_workbook(buf)["Summary"]
    assert _fill_rgb(ws.cell(row=6, column=20)).endswith("BDD7EE")
    # 說明表提及淺藍底
    ws_guide = load_workbook(buf)["說明"]
    text = " ".join(str(ws_guide.cell(row=r, column=c).value)
                    for r in range(1, 80) for c in (1, 2, 3))
    assert "淺藍" in text
