"""Validation tests for 銷售損失量化 Review (segment 1)."""
import warnings
warnings.filterwarnings("ignore")
import pandas as pd
from io import BytesIO
from openpyxl import load_workbook
from transformer import (
    transform_sales_loss,
    generate_excel_sales_loss,
)

BASE = {
    "Buyer": "B1", "Brand": "A", "SKU": 1,
    "Item_Shop": "Prod-S01-ShopHK", "Site Code": "S01", "OM": "OM1",
    "Shop Nature": "Mix", "RP Type": "RF",
    "Shop Stock Status": "Normal",
    "Is OOS (M-2)": False, "Is OOS (M-1)": False, "Is OOS (CM)": False,
    "OOS Periods Count": 0, "OOS Pattern": "",
    "Is Stock Impacted": False,
    "Max Sales Loss (M-2)": 0.0, "Max Sales Loss (P90D)": 0.0,
    "Total Sales Loss (M-2)": 0.0, "Total Sales Loss (P90D)": 0.0,
    "Shop Stock Qty (CM)": 5.0,
    "Sold Qty (M-2)": 10.0, "Sold Qty (M-1)": 10.0,
    "Sold Qty (CM - MTD)": 5.0, "Sold Qty (P90D)": 30.0,
}


def _row(**kw):
    r = dict(BASE)
    r.update(kw)
    return r


def _df(rows):
    return pd.DataFrame([_row(**r) for r in rows])


def test_filter_impacted_and_loss():
    df = _df([
        {"SKU": 1, "Is Stock Impacted": True, "Total Sales Loss (P90D)": 12.0, "Total Sales Loss (M-2)": 5.0},
        {"SKU": 2, "Total Sales Loss (P90D)": 7.0},
        {"SKU": 3},  # 無影響 → 排除
    ])
    raw, by_sku, by_shop, warnings = transform_sales_loss(df)
    assert set(raw["SKU"].astype(str)) == {"1", "2"}
    assert int(by_sku[by_sku["Buyer"] == "總計"]["Total Sales Loss (P90D) 合計"].iloc[0]) == 19


def test_bool_text_normalization():
    df = _df([
        {"SKU": 1, "Is Stock Impacted": "TRUE"},
        {"SKU": 2, "Is Stock Impacted": "Y"},
        {"SKU": 3, "Is Stock Impacted": "否"},
    ])
    raw, _, _, _ = transform_sales_loss(df)
    assert set(raw["SKU"].astype(str)) == {"1", "2"}


def test_by_sku_sort_and_shop_counts():
    df = _df([
        {"SKU": 1, "Site Code": "S01", "Is Stock Impacted": True, "Total Sales Loss (P90D)": 5.0},
        {"SKU": 1, "Site Code": "S02", "Item_Shop": "Prod-S02-ShopKL", "Is Stock Impacted": True, "Total Sales Loss (P90D)": 5.0},
        {"SKU": 2, "Is Stock Impacted": True, "Total Sales Loss (P90D)": 50.0},
    ])
    _, by_sku, by_shop, _ = transform_sales_loss(df)
    data = by_sku[by_sku["Buyer"] != "總計"]
    assert list(data["SKU"].astype(str)) == ["2", "1"]
    assert int(data[data["SKU"] == "1"]["影響店舖數"].iloc[0]) == 2
    assert len(by_shop[by_shop["Site Code"] != "總計"]) == 2


def test_missing_signal_columns_rejected():
    df = pd.DataFrame([{"Buyer": "B1", "Brand": "A", "SKU": 1, "Shop Stock Status": "Normal", "Site Code": "S01"}])
    try:
        transform_sales_loss(df)
    except ValueError as e:
        assert "損失訊號" in str(e)
    else:
        raise AssertionError("應拒絕缺少損失訊號欄位的輸入")


def test_missing_site_code_rejected():
    df = pd.DataFrame([{"Buyer": "B1", "Brand": "A", "SKU": 1, "Shop Stock Status": "Normal", "Is Stock Impacted": True}])
    try:
        transform_sales_loss(df)
    except ValueError as e:
        assert "Site Code" in str(e)
    else:
        raise AssertionError("應拒絕缺少 Site Code 的輸入")


def test_excel_sheets_and_skip_raw():
    df = _df([
        {"SKU": 1, "Is Stock Impacted": True, "Total Sales Loss (P90D)": 9.0},
    ])
    raw, by_sku, by_shop, _ = transform_sales_loss(df)
    buf = generate_excel_sales_loss(raw, by_sku, by_shop, include_raw=True)
    wb = load_workbook(BytesIO(buf.getvalue()), data_only=True)
    assert wb.sheetnames == ["RAW", "By SKU", "By Shop"]
    buf2 = generate_excel_sales_loss(raw, by_sku, by_shop, include_raw=False)
    wb2 = load_workbook(BytesIO(buf2.getvalue()), data_only=True)
    assert wb2.sheetnames == ["RAW", "By SKU", "By Shop"]
    assert wb2["RAW"]["A2"].value is not None


def test_empty_result_warning():
    df = _df([{"SKU": 1}])
    raw, by_sku, by_shop, warnings = transform_sales_loss(df)
    assert raw.empty
    assert any("沒有符合" in w for w in warnings)
