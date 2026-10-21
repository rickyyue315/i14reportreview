"""RAW 明細轉化模式（v1.84.0）驗證測試。

純粹轉化：只做清理（剔除合計列、117 - Marketing）與解析 SKU Desc／Shop Name，
不作任何統計或篩選（OOS／Overstock／P90D 0 Sales 皆不判定、不過濾列）；
輸出單一工作表 RAW，保留來源檔全部欄位。
"""
import warnings
warnings.filterwarnings("ignore")
import importlib
from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from modes import MODE_REGISTRY
from transformer import transform_raw_convert, generate_excel_raw_convert

BASE = {
    "Buyer": "126 - Mona Chan", "Brand": "A", "SKU": 111186903002.0,
    "Item_Shop": "(JDU) 111186903002-HAIR TREAMENT OIL, 01, 40PCHB97-新旺角文華",
    "Site Code": "HB97", "OM": "OM1", "Shop Nature": "Tourist", "RP Type": "RF",
    "Shop Stock Status": "Normal", "Adj. Safety Stock": 3.0,
    "Sold Qty (P90D)": 0.0, "Shop Monthly Sold Qty (P90D)": 0.0,
    "Shop Stock Qty (CM)": 0.0, "Valuation Class": "V1", "Shop Stock Value (CM)": 0.0,
}


def _row(**kw):
    r = dict(BASE)
    r.update(kw)
    return r


def _df(rows):
    return pd.DataFrame([_row(**r) for r in rows])


# ──────────────────────────────────────────────────────────────────────────────
# 1. 註冊與介面
# ──────────────────────────────────────────────────────────────────────────────
def test_registry_entry_resolves_to_functions():
    entry = MODE_REGISTRY["RAW 明細轉化"]
    mod = importlib.import_module(entry["module"])
    assert callable(getattr(mod, entry["transform"]))
    assert callable(getattr(mod, entry["generate_excel"]))


def test_input_dataframe_not_mutated():
    df = _df([{"SKU": 1}, {"SKU": 2, "Buyer": "Total"}])
    before = df.copy()
    transform_raw_convert(df)
    pd.testing.assert_frame_equal(df, before)


# ──────────────────────────────────────────────────────────────────────────────
# 2. 不做統計、不篩選列
# ──────────────────────────────────────────────────────────────────────────────
def test_no_row_filtered_by_stock_or_sales():
    df = _df([
        {"SKU": 1, "Shop Stock Status": "OOS", "Shop Stock Qty (CM)": 0.0},
        {"SKU": 2, "Shop Stock Status": "Overstock", "Shop Stock Qty (CM)": 50.0, "Sold Qty (P90D)": 0.0},
        {"SKU": 3, "Shop Stock Status": "P90D 0 Sales", "Shop Stock Qty (CM)": 4.0, "Sold Qty (P90D)": 0.0},
        {"SKU": 4, "Shop Stock Status": "Normal", "Shop Stock Qty (CM)": 5.0, "Sold Qty (P90D)": 9.0},
    ])
    raw, _ = transform_raw_convert(df)
    assert len(raw) == 4
    assert sorted(raw["SKU"]) == ["1", "2", "3", "4"]


def test_keeps_all_source_columns_and_adds_desc_and_shop_name():
    df = _df([{"SKU": 1}])
    raw, _ = transform_raw_convert(df)
    for col in df.columns:
        assert col in raw.columns
    assert "SKU Desc" in raw.columns
    assert "Shop Name" in raw.columns


def test_no_statistics_columns_added():
    df = _df([{"SKU": 1}])
    raw, _ = transform_raw_convert(df)
    stat_markers = ("Past 90 Days 0 Sales SKU", "Overstock Cap", "OOS Floor", "Projected Monthly", "Recommend")
    assert not any(any(m in str(c) for m in stat_markers) for c in raw.columns)
    assert "銷量帶" not in raw.columns
    assert "銷量存貨分類" not in raw.columns


# ──────────────────────────────────────────────────────────────────────────────
# 3. Item_Shop 解析與清理
# ──────────────────────────────────────────────────────────────────────────────
def test_item_shop_parses_desc_and_shop_name():
    df = _df([
        {"SKU": 111186903002.0, "Site Code": "HB97",
         "Item_Shop": "(JDU) 111186903002-HAIR TREAMENT OIL, 01, 40PCHB97-新旺角文華"},
        {"SKU": 111188002003.0, "Site Code": "HB29",
         "Item_Shop": "(JDU) 111188002003-SILKY HAIR OIL, 02, 6PCHB29-重慶站"},
    ])
    raw, _ = transform_raw_convert(df)
    by_site = raw.set_index("Site Code")
    assert by_site.loc["HB97", "Shop Name"] == "新旺角文華"
    assert by_site.loc["HB97", "SKU Desc"] == "HAIR TREAMENT OIL, 01, 40PC"
    assert by_site.loc["HB29", "Shop Name"] == "重慶站"
    assert by_site.loc["HB29", "SKU Desc"] == "SILKY HAIR OIL, 02, 6PC"


def test_footer_and_marketing_rows_removed():
    df = _df([
        {"SKU": 1},
        {"SKU": 2, "Buyer": "Total"},
        {"SKU": 3, "Buyer": "117 - Marketing"},
        {"SKU": 4, "Buyer": "Applied filters: x"},
    ])
    raw, _ = transform_raw_convert(df)
    assert list(raw["SKU"]) == ["1"]


def test_sku_normalized_to_text():
    df = _df([{"SKU": 111186903002.0}])
    raw, _ = transform_raw_convert(df)
    assert raw["SKU"].iloc[0] == "111186903002"


def test_shop_nature_filter_only_when_selected():
    df = _df([
        {"SKU": 1, "Shop Nature": "Tourist"},
        {"SKU": 2, "Shop Nature": "Mix", "Site Code": "HB29"},
        {"SKU": 3, "Shop Nature": "Non Tourist", "Site Code": "HB50"},
    ])
    raw_all, _ = transform_raw_convert(df)
    assert len(raw_all) == 3
    raw_mix, _ = transform_raw_convert(df, shop_natures=["Mix"])
    assert list(raw_mix["SKU"]) == ["2"]


# ──────────────────────────────────────────────────────────────────────────────
# 4. 必要欄位與警告
# ──────────────────────────────────────────────────────────────────────────────
def test_missing_required_column_raises():
    df = _df([{"SKU": 1}]).drop(columns=["Brand"])
    with pytest.raises(ValueError, match="Brand"):
        transform_raw_convert(df)


def test_missing_item_shop_warns_and_leaves_desc_blank():
    df = _df([{"SKU": 1}]).drop(columns=["Item_Shop"])
    raw, w = transform_raw_convert(df)
    assert len(raw) == 1
    assert raw["SKU Desc"].iloc[0] == ""
    assert raw["Shop Name"].iloc[0] == ""
    assert any("Item_Shop" in x for x in w)


def test_empty_after_filter_warns_and_excel_keeps_header():
    df = _df([{"SKU": 1, "Shop Nature": "Tourist"}])
    raw, w = transform_raw_convert(df, shop_natures=["Mix"])
    assert raw.empty
    assert any("沒有可轉化" in x for x in w)
    buf = generate_excel_raw_convert(raw)
    ws = load_workbook(BytesIO(buf.getvalue()))["RAW"]
    assert ws.max_row == 1
    assert "SKU Desc" in [c.value for c in ws[1]]


# ──────────────────────────────────────────────────────────────────────────────
# 5. Excel 輸出結構
# ──────────────────────────────────────────────────────────────────────────────
def test_excel_has_only_raw_sheet_with_all_columns():
    df = _df([
        {"SKU": 1, "Site Code": "HB97", "Item_Shop": "(JDU) 1-HAIR OIL, 40PCHB97-新旺角文華"},
        {"SKU": 2, "Site Code": "HB29", "Item_Shop": "(JDU) 2-SILKY, 6PCHB29-重慶站",
         "Shop Stock Status": "OOS", "Shop Stock Qty (CM)": 0.0},
    ])
    raw, _ = transform_raw_convert(df)
    buf = generate_excel_raw_convert(raw)
    wb = load_workbook(BytesIO(buf.getvalue()))
    assert wb.sheetnames == ["RAW"]
    ws = wb["RAW"]
    headers = [c.value for c in ws[1]]
    assert "Item_Shop" in headers and "Valuation Class" in headers
    assert "SKU Desc" in headers and "Shop Name" in headers
    assert "Safety Stock (現時)" in headers
    assert "Adj. Safety Stock" not in headers
    assert ws.max_row == 1 + len(raw)


def test_excel_cell_values_match_transform_output():
    df = _df([
        {"SKU": 111186903002.0, "Site Code": "HB97",
         "Item_Shop": "(JDU) 111186903002-HAIR TREAMENT OIL, 01, 40PCHB97-新旺角文華"},
    ])
    raw, _ = transform_raw_convert(df)
    buf = generate_excel_raw_convert(raw)
    ws = load_workbook(BytesIO(buf.getvalue()))["RAW"]
    headers = [c.value for c in ws[1]]
    row = [c.value for c in ws[2]]
    assert row[headers.index("SKU")] == "111186903002"
    assert row[headers.index("Shop Name")] == "新旺角文華"
    assert row[headers.index("SKU Desc")] == "HAIR TREAMENT OIL, 01, 40PC"
