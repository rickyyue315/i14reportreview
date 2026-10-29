"""Validation tests for ND SKU with Stock Review(轉貨模式).

規則：
- 候選轉出店舖：RP Type＝ND 且 Shop Stock Qty (CM) > 0
- 轉出：已選轉出店舖中符合上述條件的列，數量為 Shop Stock Qty (CM) 全數（不看銷售記錄）
- 接收：同 SKU、非轉出店舖、RP Type＝RF 且 Sold Qty (P90D) > 0 的列，不論有無庫存
- 分配：同 SKU 轉出總量按接收店 Sold Qty (P90D) 比例分配（最大餘數法取整數件、全數分完）
- 接收店舖集中模式（可選）：每個 SKU 只取 Sold Qty (P90D) 最高的 2 間或 4 間 RF 接收店，
  再按這批店舖的銷售比例分配；預設不限，維持全數有銷售店舖接收
- RP Type：去除前後空格、大小寫不敏感；缺欄即停止；非 RF／ND 的值一律排除（列 warning）
- MCH 篩選（執行前）：可按 Adj. MCH1 排除分類（預設全部包含），對轉出與接收兩側同時生效
- 分配、配對與彙總與 P90D 0 Sales Review(轉貨模式) 共用同一引擎（nd_mode=True）
"""
import warnings

warnings.filterwarnings("ignore")

import pandas as pd
from openpyxl import load_workbook

from modes.transfer_nd import (
    build_nd_transfer_candidates,
    generate_excel_nd_transfer,
    transform_nd_transfer,
)
import transformer as shim
from modes import MODE_REGISTRY

BASE = {
    "Buyer": "B1", "Brand": "A", "SKU": 1,
    "Item_Shop": "Prod-S01-ShopHK", "Site Code": "S01", "OM": "OM1",
    "Shop Nature": "Mix", "RP Type": "ND",
    "Shop Stock Status": "Overstock (P90D 0 Sales)",
    "Adj. Safety Stock": 3.0, "Adj. MCH1": "01 - SKIN CARE",
    "Supply Source": "Warehouse",
    "Sold Qty (M-2)": 0.0, "Sold Qty (M-1)": 0.0, "Sold Qty (CM - MTD)": 0.0,
    "Sold Qty (P90D)": 0.0,
    "Shop Stock Qty (CM)": 5.0, "Shop Stock Qty (LM)": 4.0,
}


def _row(site, sku=1, sales=0.0, stock=5.0, rp="ND", name=None, om="OM1", nature="Mix",
         brand="A", buyer="B1", mch1="01 - SKIN CARE", mch2=None, supply_source=None):
    r = dict(BASE)
    r.update({
        "Site Code": site,
        "SKU": sku,
        "Item_Shop": f"Prod-{site}-{name or 'Shop' + site}",
        "Sold Qty (P90D)": float(sales),
        "Sold Qty (CM - MTD)": float(sales),
        "Shop Stock Qty (CM)": float(stock),
        "RP Type": rp,
        "OM": om,
        "Shop Nature": nature,
        "Brand": brand,
        "Buyer": buyer,
        "Adj. MCH1": mch1,
    })
    if mch2 is not None:
        r["MCH2"] = mch2
    if supply_source is not None:
        r["Supply Source"] = supply_source
    return r


def _df(rows):
    return pd.DataFrame(rows)


def _lines(plan):
    return {
        (r["轉出 Site Code"], str(r["SKU"]), r["接收 Site Code"]): int(r["轉貨數量 (件)"])
        for _, r in plan.iterrows()
    }


# 1. 候選轉出店舖（RP Type＝ND 且有庫存）

def test_candidates_only_nd_with_stock():
    df = _df([
        _row("S01", sku=1, stock=5, rp="ND"),
        _row("S01", sku=2, stock=9, rp="RF"),   # RF 有庫存 → 不列入候選
        _row("S03", sku=3, stock=0, rp="ND"),   # ND 無庫存 → 不列入候選
        _row("S01", sku=4, stock=2, rp="ND"),
    ])
    cand = build_nd_transfer_candidates(df)
    assert list(cand["Site Code"]) == ["S01"]
    row = cand.iloc[0]
    assert row["Shop Name"] == "ShopS01"
    assert int(row["ND SKU 數"]) == 2
    assert int(row["可轉出庫存 (件)"]) == 7
    assert "0 Sales SKU 數" not in cand.columns


def test_candidates_ignore_sales_and_rp_non_standard():
    df = _df([
        _row("S01", sku=1, sales=8, stock=4, rp="ND"),   # ND 有銷售但仍有庫存 → 仍為候選
        _row("S01", sku=2, sales=0, stock=6, rp=" XX "),  # 非 RF／ND → 排除
        _row("S01", sku=3, sales=0, stock=3, rp=""),      # 空白 → 排除
    ])
    cand = build_nd_transfer_candidates(df)
    assert list(cand["Site Code"]) == ["S01"]
    assert int(cand.iloc[0]["ND SKU 數"]) == 1
    assert int(cand.iloc[0]["可轉出庫存 (件)"]) == 4


# 2. 轉出（ND 有庫存全數，不看銷售）

def test_nd_with_sales_still_transferred():
    df = _df([
        _row("S01", sku=1, sales=9, stock=5, rp="ND", supply_source="Vendor"),
        _row("S02", sku=1, sales=2, stock=0, rp="RF"),
    ])
    plan, _, _, _, _, stats, _ = transform_nd_transfer(df, ["S01"])
    assert _lines(plan) == {("S01", "1", "S02"): 5}
    assert stats["out_qty"] == 5 and stats["alloc_qty"] == 5
    # 轉貨單分類欄：Supply Source（MCH2 已移除）
    assert plan.iloc[0]["Supply Source"] == "Vendor"
    assert "MCH2" not in plan.columns


def test_rf_stock_at_selected_site_is_not_transferred():
    df = _df([
        _row("S01", sku=1, sales=0, stock=7, rp="RF"),
        _row("S02", sku=1, sales=2, stock=0, rp="RF"),
    ])
    plan, _, _, _, _, stats, warns = transform_nd_transfer(df, ["S01"])
    assert plan.empty and stats["out_qty"] == 0
    assert any("沒有符合" in w for w in warns)


# 3. 接收店（RF 且有銷售記錄）

def test_receivers_only_rf_with_sales():
    df = _df([
        _row("S01", sku=1, stock=6, rp="ND"),
        _row("S02", sku=1, sales=4, stock=0, rp="RF"),   # RF 有銷售 → 接收
        _row("S03", sku=1, sales=0, stock=0, rp="RF"),   # RF 無銷售 → 不接收
        _row("S04", sku=1, sales=9, stock=0, rp="ND"),   # ND 有銷售 → 不接收
    ])
    plan, out_df, in_df, sku_df, mch_df, stats, warns = transform_nd_transfer(df, ["S01"])
    assert _lines(plan) == {("S01", "1", "S02"): 6}
    assert stats["in_sites"] == 1 and stats["unalloc_qty"] == 0


def test_multi_receiver_proportional_by_sales():
    df = _df([
        _row("S01", sku=1, stock=10, rp="ND"),
        _row("S02", sku=1, sales=3, stock=0, rp="RF"),
        _row("S03", sku=1, sales=1, stock=0, rp="RF"),
        _row("S04", sku=1, sales=5, stock=0, rp="ND"),
    ])
    plan, _, _, _, _, stats, _ = transform_nd_transfer(df, ["S01"])
    assert _lines(plan) == {("S01", "1", "S02"): 8, ("S01", "1", "S03"): 2}
    assert stats["alloc_qty"] == 10 and stats["unalloc_qty"] == 0


def test_no_receiver_stays_unallocated():
    df = _df([
        _row("S01", sku=1, stock=5, rp="ND"),
        _row("S02", sku=1, sales=0, stock=3, rp="RF"),
        _row("S03", sku=1, sales=6, stock=0, rp="ND"),
    ])
    plan, out_df, in_df, sku_df, mch_df, stats, warns = transform_nd_transfer(df, ["S01"])
    assert plan.empty
    assert stats["unalloc_qty"] == 5 and stats["unalloc_skus"] == 1
    body = sku_df[sku_df["Brand"].astype(str) != "總計"].iloc[0]
    assert int(body["未分配量 (件)"]) == 5
    assert body["狀態"] == "無接收店舖（保留原店）"
    assert any("RP Type＝RF" in w and "未能轉出" in w for w in warns)


def test_selected_site_sales_rows_are_neither_source_nor_receiver():
    df = _df([
        _row("S01", sku=1, stock=5, rp="ND"),
        _row("S01", sku=2, sales=2, stock=4, rp="RF"),
        _row("S02", sku=2, sales=1, stock=0, rp="RF"),
    ])
    plan, _, _, sku_df, _, stats, _ = transform_nd_transfer(df, ["S01"])
    assert plan.empty
    assert stats["out_skus"] == 1
    body = sku_df[sku_df["Brand"].astype(str) != "總計"]
    assert set(body["SKU"].astype(str)) == {"1"}


def test_om_ignored_for_transfer():
    df = _df([
        _row("S01", sku=1, stock=5, rp="ND", om="OM1"),
        _row("S02", sku=1, sales=2, stock=0, rp="RF", om="OM2"),
    ])
    plan, _, _, _, _, stats, _ = transform_nd_transfer(df, ["S01"])
    assert _lines(plan) == {("S01", "1", "S02"): 5}
    assert plan.iloc[0]["接收店 OM"] == "OM2"


def test_shop_nature_filter_applies_to_both_sides():
    df = _df([
        _row("S01", sku=1, stock=6, rp="ND", nature="Mix"),
        _row("S02", sku=1, sales=5, stock=0, rp="RF", nature="Tourist"),
        _row("S03", sku=1, sales=5, stock=0, rp="RF", nature="Non Tourist"),
    ])
    plan, _, _, _, _, _, _ = transform_nd_transfer(
        df, ["S01"], shop_natures=["Mix", "Non Tourist"]
    )
    assert _lines(plan) == {("S01", "1", "S03"): 6}


# 4. RP Type 檢核

def test_missing_rp_type_raises():
    df = _df([_row("S01", sku=1, stock=5, rp="ND"), _row("S02", sku=1, sales=1, stock=0, rp="RF")])
    df = df.drop(columns=["RP Type"])
    try:
        transform_nd_transfer(df, ["S01"])
    except ValueError as exc:
        assert "RP Type" in str(exc)
    else:
        raise AssertionError("expected ValueError for missing RP Type")
    try:
        build_nd_transfer_candidates(df)
    except ValueError as exc:
        assert "RP Type" in str(exc)
    else:
        raise AssertionError("expected ValueError for missing RP Type in candidates")


def test_rp_type_trim_and_case_insensitive():
    df = _df([
        _row("S01", sku=1, stock=5, rp=" nd "),
        _row("S02", sku=1, sales=3, stock=0, rp="rf"),
    ])
    plan, _, _, _, _, stats, warns = transform_nd_transfer(df, ["S01"])
    assert _lines(plan) == {("S01", "1", "S02"): 5}
    assert stats["alloc_qty"] == 5
    assert not any("非 RF／ND" in w for w in warns)


def test_non_standard_rp_values_excluded_with_warning():
    df = _df([
        _row("S01", sku=1, stock=5, rp="ND"),
        _row("S01", sku=2, stock=5, rp="其他"),          # 非標準 → 不轉出
        _row("S02", sku=1, sales=4, stock=0, rp="RF"),
        _row("S03", sku=1, sales=9, stock=0, rp=""),     # 空白 → 不接收
    ])
    plan, _, _, _, _, stats, warns = transform_nd_transfer(df, ["S01"])
    assert _lines(plan) == {("S01", "1", "S02"): 5}
    assert stats["out_qty"] == 5 and stats["out_skus"] == 1
    assert any("非 RF／ND" in w and "2" in w for w in warns)


def test_p90d_mode_does_not_require_rp_type():
    from modes.transfer_p90d import build_transfer_candidates, transform_p90d_transfer
    df = _df([_row("S01", sku=1, sales=0, stock=5), _row("S02", sku=1, sales=1, stock=0)])
    df = df.drop(columns=["RP Type"])
    cand = build_transfer_candidates(df)
    assert list(cand["Site Code"]) == ["S01"]
    plan, _, _, _, _, stats, _ = transform_p90d_transfer(df, ["S01"])
    assert stats["alloc_qty"] == 5


# 5. 重複列、未納入轉出與空結果

def test_duplicate_site_sku_stops_execution():
    """重複列一律停止計算（與 P90D 模式及調貨段同一準則）。"""
    df = _df([
        _row("S01", sku=1, stock=2, rp="ND"),
        _row("S01", sku=1, stock=3, rp="ND"),
        _row("S02", sku=1, sales=1, stock=0, rp="RF"),
    ])
    try:
        transform_nd_transfer(df, ["S01"])
    except ValueError as exc:
        assert "重複" in str(exc) and "S01" in str(exc)
    else:
        raise AssertionError("expected ValueError for duplicate Site Code × SKU")


def test_selected_site_without_nd_stock_is_reported():
    df = _df([
        _row("S01", sku=1, stock=5, rp="ND"),
        _row("S09", sku=1, sales=2, stock=3, rp="RF"),
        _row("S02", sku=1, sales=1, stock=0, rp="RF"),
    ])
    plan, _, _, _, _, _, warns = transform_nd_transfer(df, ["S01", "S09"])
    assert _lines(plan) == {("S01", "1", "S02"): 5}
    assert any("未納入轉出" in w and "S09" in w for w in warns)


def test_no_nd_stock_rows_returns_empty_with_warning():
    df = _df([
        _row("S09", sku=1, sales=2, stock=3, rp="RF"),
        _row("S02", sku=1, sales=1, stock=0, rp="RF"),
    ])
    plan, out_df, in_df, sku_df, mch_df, stats, warns = transform_nd_transfer(df, ["S09"])
    assert plan.empty and out_df.empty and in_df.empty and sku_df.empty and mch_df.empty
    assert stats["out_qty"] == 0
    assert any("RP Type＝ND" in w and "沒有符合" in w for w in warns)


def test_empty_selection_raises():
    df = _df([_row("S01", sku=1, stock=5, rp="ND"), _row("S02", sku=1, sales=1, stock=0, rp="RF")])
    try:
        transform_nd_transfer(df, [])
    except ValueError as exc:
        assert "轉出店舖" in str(exc)
    else:
        raise AssertionError("expected ValueError for empty selection")


# 6. MCH 篩選

def test_mch1_filter_applies_to_both_sides():
    df = _df([
        _row("S01", sku=1, stock=5, rp="ND", mch1="01 - SKIN CARE"),
        _row("S01", sku=2, stock=7, rp="ND", mch1="02 - MAKE UP"),
        _row("S02", sku=1, sales=3, stock=0, rp="RF", mch1="02 - MAKE UP"),
        _row("S03", sku=2, sales=3, stock=0, rp="RF", mch1="01 - SKIN CARE"),
    ])
    plan, _, _, _, _, stats, warns = transform_nd_transfer(
        df, ["S01"], mch1_filter=["01 - SKIN CARE"]
    )
    # SKU 1 屬 SKIN CARE 可轉出，但其接收店 S02 屬 MAKE UP 已被排除 → 未分配
    assert plan.empty
    assert stats["out_qty"] == 5 and stats["unalloc_qty"] == 5
    assert any("MCH 篩選已排除" in w for w in warns)


def test_candidates_respect_mch_filter():
    df = _df([
        _row("S01", sku=1, stock=5, rp="ND", mch1="01 - SKIN CARE"),
        _row("S02", sku=2, stock=9, rp="ND", mch1="02 - MAKE UP"),
    ])
    cand = build_nd_transfer_candidates(df, mch1_filter=["01 - SKIN CARE"])
    assert list(cand["Site Code"]) == ["S01"]


# 7. 彙總與 Excel

def test_summary_totals_match_plan():
    df = _df([
        _row("S01", sku=1, stock=4, rp="ND"),
        _row("S05", sku=1, stock=6, rp="ND"),
        _row("S05", sku=2, stock=3, rp="ND"),
        _row("S02", sku=1, sales=3, stock=2, rp="RF"),
        _row("S03", sku=1, sales=1, stock=0, rp="RF"),
        _row("S03", sku=2, sales=0, stock=0, rp="RF"),
    ])
    plan, out_df, in_df, sku_df, mch_df, stats, _ = transform_nd_transfer(df, ["S01", "S05"])
    out_total = out_df[out_df["Site Code"] == "總計"].iloc[0]
    in_total = in_df[in_df["Site Code"] == "總計"].iloc[0]
    sku_total = sku_df[sku_df["Brand"] == "總計"].iloc[0]
    assert int(out_total["轉出總量 (件)"]) == 13
    assert int(out_total["已分配 (件)"]) == 10
    assert int(out_total["未分配 (件)"]) == 3
    assert int(plan["轉貨數量 (件)"].sum()) == 10
    assert int(in_total["接收總量 (件)"]) == 10
    assert int(sku_total["未分配量 (件)"]) == 3
    mch_total = mch_df[mch_df["Adj. MCH1"] == "總計"].iloc[0]
    assert int(mch_total["轉出總量 (件)"]) == 13
    assert int(mch_total["接收總量 (件)"]) == 10


def test_excel_structure_and_guide_wording():
    df = _df([
        _row("S01", sku=1, stock=10, rp="ND"),
        _row("S02", sku=1, sales=3, stock=0, rp="RF"),
        _row("S03", sku=1, sales=1, stock=0, rp="RF"),
    ])
    plan, out_df, in_df, sku_df, mch_df, stats, _ = transform_nd_transfer(df, ["S01"])
    buf = generate_excel_nd_transfer(
        plan, out_df, in_df, sku_df, mch_df, stats, nature_label="全部"
    )
    wb = load_workbook(buf)
    assert wb.sheetnames == [
        "說明", "轉貨單", "轉出店舖彙總", "接收店舖彙總", "SKU 轉貨彙總", "MCH1 分類彙總",
    ]
    guide = wb["說明"]
    texts = [
        str(cell.value) for row in guide.iter_rows() for cell in row if cell.value is not None
    ]
    assert any("ND SKU with Stock Review（轉貨模式）說明" in t for t in texts)
    assert any("RP Type＝ND" in t for t in texts)
    assert not any("0 Sales 定義" in t for t in texts)
    ws = wb["轉貨單"]
    headers = [c.value for c in ws[1]]
    assert headers == list(plan.columns)
    share_col = headers.index("接收店銷售佔比 (同 SKU)") + 1
    assert ws.cell(row=2, column=share_col).number_format == "0.0%"
    qty_col = headers.index("轉貨數量 (件)") + 1
    assert ws.cell(row=2, column=qty_col).value == 8


def test_excel_empty_plan_writes_placeholder():
    df = _df([_row("S01", sku=1, stock=5, rp="ND"), _row("S02", sku=1, sales=0, stock=3, rp="RF")])
    plan, out_df, in_df, sku_df, mch_df, stats, _ = transform_nd_transfer(df, ["S01"])
    buf = generate_excel_nd_transfer(plan, out_df, in_df, sku_df, mch_df, stats)
    wb = load_workbook(buf)
    assert wb["轉貨單"].cell(row=2, column=1).value == "無符合資料"
    assert wb["MCH1 分類彙總"].cell(row=1, column=1).value == "Adj. MCH1"


def test_excel_plan_group_colors_in_nd_mode():
    """ND 模式共用同一引擎：轉貨單欄位分組著色與 P90D 模式一致。"""
    df = _df([
        _row("S01", sku=1, stock=10, rp="ND"),
        _row("S02", sku=1, sales=3, stock=0, rp="RF"),
    ])
    plan, out_df, in_df, sku_df, mch_df, stats, _ = transform_nd_transfer(df, ["S01"])
    buf = generate_excel_nd_transfer(plan, out_df, in_df, sku_df, mch_df, stats)
    wb = load_workbook(buf)
    ws = wb["轉貨單"]
    headers = [c.value for c in ws[1]]
    _out = headers.index("轉出 Site Code") + 1
    _qty = headers.index("轉貨數量 (件)") + 1
    _in = headers.index("接收店 OM") + 1
    _plain = headers.index("Brand") + 1
    assert ws.cell(row=1, column=_out).fill.start_color.rgb.endswith("FFD966")
    assert ws.cell(row=2, column=_out).fill.start_color.rgb.endswith("FFF2CC")
    assert ws.cell(row=1, column=_qty).fill.start_color.rgb.endswith("FFD966")
    assert ws.cell(row=2, column=_qty).font.bold is True
    assert ws.cell(row=1, column=_in).fill.start_color.rgb.endswith("9DC3E6")
    assert ws.cell(row=2, column=_in).fill.start_color.rgb.endswith("DDEBF7")
    assert ws.cell(row=1, column=_plain).fill.start_color.rgb.endswith("4472C4")
    assert ws.cell(row=2, column=_plain).fill.fill_type is None


def test_p90d_guide_wording_unchanged():
    """P90D 模式的說明頁維持自身措辭（不使用 ND 模式的轉出／接收定義）。"""
    from modes.transfer_p90d import generate_excel_p90d_transfer, transform_p90d_transfer
    df = _df([
        _row("S01", sku=1, sales=0, stock=10, rp="ND"),
        _row("S02", sku=1, sales=3, stock=0, rp="RF"),
    ])
    plan, out_df, in_df, sku_df, mch_df, stats, _ = transform_p90d_transfer(df, ["S01"])
    buf = generate_excel_p90d_transfer(plan, out_df, in_df, sku_df, mch_df, stats)
    wb = load_workbook(buf)
    texts = [
        str(cell.value) for row in wb["說明"].iter_rows() for cell in row if cell.value is not None
    ]
    assert any("P90D 0 Sales Review（轉貨模式）說明" in t for t in texts)
    assert any("0 Sales 定義" in t for t in texts)
    assert not any("ND 定義" in t for t in texts)
    assert not any("RP Type＝RF 且 Sold Qty (P90D) > 0" in t for t in texts)


# 8. 接收店舖集中模式（每個 SKU 最多 2／4 間）

def test_receiver_limit_2_selects_top_sales_stores():
    df = _df([
        _row("S01", sku=1, stock=10, rp="ND"),
        _row("S02", sku=1, sales=1, stock=0, rp="RF"),
        _row("S03", sku=1, sales=3, stock=0, rp="RF"),
        _row("S04", sku=1, sales=5, stock=0, rp="RF"),
        _row("S05", sku=1, sales=2, stock=0, rp="RF"),
    ])
    plan, _, _, sku_df, _, stats, warns = transform_nd_transfer(df, ["S01"], receiver_limit=2)
    # 只取銷量最高的 S04（5）與 S03（3），按 5:3 重新分配 10 件
    assert _lines(plan) == {("S01", "1", "S04"): 6, ("S01", "1", "S03"): 4}
    assert stats["alloc_qty"] == 10 and stats["unalloc_qty"] == 0
    assert stats["receiver_limit"] == 2
    shares = dict(zip(plan["接收 Site Code"], plan["接收店銷售佔比 (同 SKU)"]))
    assert abs(shares["S04"] - 0.625) < 1e-9 and abs(shares["S03"] - 0.375) < 1e-9
    body = sku_df[sku_df["Brand"].astype(str) != "總計"].iloc[0]
    assert int(body["可接收店舖數 (有銷售)"]) == 2
    assert int(body["實際接收店舖數"]) == 2
    assert any("接收店集中模式" in w and "超出上限" in w for w in warns)


def test_receiver_limit_4_keeps_top_four_of_five():
    df = _df([
        _row("S01", sku=1, stock=100, rp="ND"),
        _row("S02", sku=1, sales=1, stock=0, rp="RF"),
        _row("S03", sku=1, sales=2, stock=0, rp="RF"),
        _row("S04", sku=1, sales=3, stock=0, rp="RF"),
        _row("S05", sku=1, sales=4, stock=0, rp="RF"),
        _row("S06", sku=1, sales=5, stock=0, rp="RF"),
    ])
    plan, _, _, sku_df, _, stats, warns = transform_nd_transfer(df, ["S01"], receiver_limit=4)
    assert set(plan["接收 Site Code"]) == {"S03", "S04", "S05", "S06"}
    assert "S02" not in set(plan["接收 Site Code"])
    assert stats["alloc_qty"] == 100 and stats["unalloc_qty"] == 0
    body = sku_df[sku_df["Brand"].astype(str) != "總計"].iloc[0]
    assert int(body["可接收店舖數 (有銷售)"]) == 4
    assert int(body["實際接收店舖數"]) == 4
    assert any("超出上限" in w and "1" in w for w in warns)


def test_receiver_limit_single_receiver_gets_all():
    df = _df([
        _row("S01", sku=1, stock=7, rp="ND"),
        _row("S02", sku=1, sales=4, stock=0, rp="RF"),
    ])
    plan, _, _, _, _, stats, warns = transform_nd_transfer(df, ["S01"], receiver_limit=2)
    assert _lines(plan) == {("S01", "1", "S02"): 7}
    assert stats["alloc_qty"] == 7 and stats["unalloc_qty"] == 0
    assert not any("超出上限" in w for w in warns)


def test_receiver_limit_tie_break_by_site_code():
    df = _df([
        _row("S01", sku=1, stock=6, rp="ND"),
        _row("S04", sku=1, sales=3, stock=0, rp="RF"),
        _row("S03", sku=1, sales=3, stock=0, rp="RF"),
        _row("S02", sku=1, sales=3, stock=0, rp="RF"),
    ])
    plan, _, _, _, _, _, _ = transform_nd_transfer(df, ["S01"], receiver_limit=2)
    assert set(plan["接收 Site Code"]) == {"S02", "S03"}


def test_receiver_limit_none_keeps_all_receivers():
    df = _df([
        _row("S01", sku=1, stock=6, rp="ND"),
        _row("S02", sku=1, sales=3, stock=0, rp="RF"),
        _row("S03", sku=1, sales=1, stock=0, rp="RF"),
    ])
    plan_default, _, _, _, _, stats_default, _ = transform_nd_transfer(df, ["S01"])
    plan_none, _, _, _, _, stats_none, _ = transform_nd_transfer(df, ["S01"], receiver_limit=None)
    assert _lines(plan_default) == _lines(plan_none) == {
        ("S01", "1", "S02"): 5, ("S01", "1", "S03"): 1,
    }
    assert stats_default == stats_none
    assert stats_default["receiver_limit"] is None


def test_invalid_receiver_limit_raises():
    df = _df([_row("S01", sku=1, stock=5, rp="ND"), _row("S02", sku=1, sales=1, stock=0, rp="RF")])
    for bad in (0, -1, "abc"):
        try:
            transform_nd_transfer(df, ["S01"], receiver_limit=bad)
        except ValueError as exc:
            assert "集中上限" in str(exc)
        else:
            raise AssertionError(f"expected ValueError for receiver_limit={bad!r}")


def test_receiver_limit_excel_guide_wording():
    df = _df([
        _row("S01", sku=1, stock=10, rp="ND"),
        _row("S02", sku=1, sales=3, stock=0, rp="RF"),
        _row("S03", sku=1, sales=1, stock=0, rp="RF"),
    ])
    plan, out_df, in_df, sku_df, mch_df, stats, _ = transform_nd_transfer(
        df, ["S01"], receiver_limit=2
    )
    buf = generate_excel_nd_transfer(
        plan, out_df, in_df, sku_df, mch_df, stats, receiver_limit=2
    )
    wb = load_workbook(buf)
    texts = [
        str(cell.value) for row in wb["說明"].iter_rows() for cell in row if cell.value is not None
    ]
    assert any("接收店舖集中模式" in t for t in texts)
    assert any("集中 1–2 間" in t for t in texts)
    assert any("未納入的店舖不會收到該 SKU 的轉貨" in t for t in texts)


# 9. 墊片與模式登記

def test_transformer_shim_and_registry():
    assert shim.build_nd_transfer_candidates is build_nd_transfer_candidates
    assert shim.transform_nd_transfer is transform_nd_transfer
    assert shim.generate_excel_nd_transfer is generate_excel_nd_transfer
    entry = MODE_REGISTRY["ND SKU with Stock Review(轉貨模式)"]
    assert entry["module"] == "modes.transfer_nd"
    assert entry["transform"] == "transform_nd_transfer"
    assert entry["generate_excel"] == "generate_excel_nd_transfer"
