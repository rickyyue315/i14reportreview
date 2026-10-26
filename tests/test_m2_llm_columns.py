"""M-2（LLM）選用庫存欄位的系統認知驗證測試（v1.92.0）。

2026-10-26 起來源報表（By SKUxShop）新增 D001 Stock Qty (LLM)／Shop Stock Qty (LLM)／
Shop Stock Value (LLM)（M-2＝上上月庫存；語意對照 LLM＝上上月、LM＝上月、CM＝當月）。
這些欄位屬「選用」：舊月份檔案沒有；任何模式皆不檢核、不要求；存在時完整保留、
缺失時行為完全不變；現有計算邏輯暫未使用（之後才逐模式加入 M-2 計算）。
"""
import warnings
warnings.filterwarnings("ignore")
from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from transformer import (
    M2_STOCK_LLM_COLUMNS,
    M2_STOCK_LLM_EFFECTIVE_DATE,
    available_m2_stock_columns,
    validate_input,
    _clean_dataframe,
    _reorder_raw_columns,
    transform_raw_convert,
    generate_excel_raw_convert,
)

# 舊月份（2026-10-26 之前）來源檔的欄位形態：只有 LM／CM，沒有任何 LLM 欄
BASE_OLD = {
    "Buyer": "126 - Mona Chan", "Brand": "A", "SKU": 111186903002.0,
    "Item_Shop": "(JDU) 111186903002-HAIR TREAMENT OIL, 01, 40PCHB97-新旺角文華",
    "Site Code": "HB97", "Shop Nature": "Tourist", "Shop Stock Status": "Normal",
    "Adj. Safety Stock": 3.0,
    "Shop Stock Qty (LM)": 9.0, "Shop Stock Qty (CM)": 5.0,
    "Shop Stock Value (LM)": 90.0, "Shop Stock Value (CM)": 50.0,
    "D001 Stock Qty (LM)": 10.0, "D001 Stock Qty (CM)": 8.0,
    "Sold Qty (P90D)": 0.0,
}

# 2026-10-26 起的新欄位（數值與 LM／CM 同性質）
M2_VALUES = {
    "D001 Stock Qty (LLM)": 12.0,
    "Shop Stock Qty (LLM)": 11.0,
    "Shop Stock Value (LLM)": 110.0,
}


def _row(**kw):
    r = dict(BASE_OLD)
    r.update(kw)
    return r


def _df(rows):
    return pd.DataFrame([_row(**r) for r in rows])


def _new_format_df():
    row = dict(BASE_OLD)
    row.update(M2_VALUES)
    return pd.DataFrame([row])


# ──────────────────────────────────────────────────────────────────────────────
# 1. 中央定義（單一事實來源）
# ──────────────────────────────────────────────────────────────────────────────
def test_m2_stock_llm_columns_registry():
    assert M2_STOCK_LLM_COLUMNS == [
        "D001 Stock Qty (LLM)",
        "Shop Stock Qty (LLM)",
        "Shop Stock Value (LLM)",
    ]
    assert M2_STOCK_LLM_EFFECTIVE_DATE == "2026-10-26"


def test_m2_columns_not_required_anywhere():
    # 選用欄位：不得進入任何必要欄位清單
    from modes.common import REQUIRED_COLUMNS
    for col in M2_STOCK_LLM_COLUMNS:
        assert col not in REQUIRED_COLUMNS


# ──────────────────────────────────────────────────────────────────────────────
# 2. 偵測函數 available_m2_stock_columns
# ──────────────────────────────────────────────────────────────────────────────
def test_available_detects_all_columns_in_new_format():
    found = available_m2_stock_columns(_new_format_df())
    assert found == M2_STOCK_LLM_COLUMNS


def test_available_empty_for_old_format():
    assert available_m2_stock_columns(_df([{"SKU": 1}])) == []


def test_available_partial_columns():
    df = _df([{"SKU": 1, "Shop Stock Qty (LLM)": 7.0}])
    assert available_m2_stock_columns(df) == ["Shop Stock Qty (LLM)"]


def test_available_tolerates_surrounding_whitespace():
    # 欄名比對沿用 _resolve_column_name（前後空格容許），回傳檔案內實際欄名
    df = _df([{"SKU": 1}]).rename(columns={"Brand": " Shop Stock Qty (LLM) "})
    df["Brand"] = "A"  # 補回被覆蓋的必要欄
    assert available_m2_stock_columns(df) == [" Shop Stock Qty (LLM) "]


def test_available_none_and_empty_safe():
    assert available_m2_stock_columns(None) == []
    assert available_m2_stock_columns(pd.DataFrame()) == []


# ──────────────────────────────────────────────────────────────────────────────
# 3. 必要欄位檢核不受影響（非必須）
# ──────────────────────────────────────────────────────────────────────────────
def test_validate_input_unaffected_by_new_columns():
    # 新格式（有 LLM 欄）與舊格式（無）同樣通過；不因缺 LLM 欄而失敗
    assert validate_input(_new_format_df()) == []
    assert validate_input(_df([{"SKU": 1}])) == []
    # 必要欄位缺失時的行為與新欄位無關
    missing = _df([{"SKU": 1}]).drop(columns=["Shop Stock Status"])
    assert validate_input(missing) == ["Shop Stock Status"]


# ──────────────────────────────────────────────────────────────────────────────
# 4. 共用管線完整保留（清理 → 欄序重排 → RAW 輸出）
# ──────────────────────────────────────────────────────────────────────────────
def test_clean_and_reorder_preserve_m2_columns():
    df = _new_format_df()
    cleaned = _clean_dataframe(df)
    reordered = _reorder_raw_columns(cleaned)
    for col in M2_STOCK_LLM_COLUMNS:
        assert col in cleaned.columns
        assert col in reordered.columns
        # 數值原樣保留
        assert float(reordered[col].iloc[0]) == M2_VALUES[col]
    # 舊格式不會產生任何 LLM 欄
    old_reordered = _reorder_raw_columns(_clean_dataframe(_df([{"SKU": 1}])))
    assert available_m2_stock_columns(old_reordered) == []


def test_raw_convert_keeps_m2_columns_end_to_end():
    raw, warns = transform_raw_convert(_new_format_df())
    for col in M2_STOCK_LLM_COLUMNS:
        assert col in raw.columns
        assert float(raw[col].iloc[0]) == M2_VALUES[col]

    buf = generate_excel_raw_convert(raw)
    ws = load_workbook(BytesIO(buf.getvalue()), read_only=True)["RAW"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    for col in M2_STOCK_LLM_COLUMNS:
        assert col in header


def test_no_value_variant_drops_only_value_llm():
    # 「不含 Value」變體依既有規則刪除欄名含 Value 的欄位：
    # Shop Stock Value (LLM) 屬金額欄會被刪除，兩個 Qty (LLM) 欄保留
    raw, _ = transform_raw_convert(_new_format_df(), include_value=False)
    assert "Shop Stock Value (LLM)" not in raw.columns
    assert "D001 Stock Qty (LLM)" in raw.columns
    assert "Shop Stock Qty (LLM)" in raw.columns


# ──────────────────────────────────────────────────────────────────────────────
# 5. 實際新格式來源檔偵測（2026-10-26 檔案存在時）
# ──────────────────────────────────────────────────────────────────────────────
def test_real_26oct2026_file_detected():
    path = Path(__file__).resolve().parent.parent / "By SKUxShop_26Oct2026.xlsx"
    if not path.exists():
        import pytest
        pytest.skip("實際來源檔不存在，略過")
    df = pd.read_excel(path)
    assert available_m2_stock_columns(df) == M2_STOCK_LLM_COLUMNS
