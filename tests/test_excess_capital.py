"""Validation tests for 積壓資金 Review (segment 2)."""
import warnings
warnings.filterwarnings("ignore")
import pandas as pd
from io import BytesIO
from openpyxl import load_workbook
from transformer import (
    transform_excess_capital,
    generate_excel_excess_capital,
)

BASE = {
    "Buyer": "B1", "Brand": "A", "SKU": 1,
    "Item_Shop": "Prod-S01-ShopHK", "Site Code": "S01", "OM": "OM1",
    "Shop Nature": "Mix", "RP Type": "RF",
    "Shop Stock Status": "Normal",
    "Excess Stock Qty": 0.0, "Excess Stock Value": 0.0,
    "Shop Stock Qty (CM)": 5.0,
    "Sold Qty (M-2)": 10.0, "Sold Qty (M-1)": 10.0,
    "Sold Qty (CM - MTD)": 5.0, "Sold Qty (P90D)": 30.0,
}


def _row(**kw):
    r = dict(BASE)
    r.update(kw)
    return r


def _df(rows):
    return pd.DataFrame([_row(**r) for r in rows])


def test_filter_excess_qty_or_value():
    df = _df([
        {"SKU": 1, "Excess Stock Qty": 4.0, "Excess Stock Value": 400.0},
        {"SKU": 2, "Excess Stock Value": 100.0},
        {"SKU": 3},  # 無積壓 → 排除
    ])
    raw, by_sku, by_shop, warnings = transform_excess_capital(df)
    assert set(raw["SKU"].astype(str)) == {"1", "2"}
    assert int(by_sku[by_sku["Buyer"] == "總計"]["Excess Stock Value 合計"].iloc[0]) == 500


def test_by_sku_sort_and_shop_counts():
    df = _df([
        {"SKU": 1, "Site Code": "S01", "Excess Stock Qty": 2.0, "Excess Stock Value": 50.0},
        {"SKU": 1, "Site Code": "S02", "Item_Shop": "Prod-S02-ShopKL", "Excess Stock Qty": 2.0, "Excess Stock Value": 50.0},
        {"SKU": 2, "Excess Stock Qty": 1.0, "Excess Stock Value": 500.0},
    ])
    _, by_sku, by_shop, _ = transform_excess_capital(df)
    data = by_sku[by_sku["Buyer"] != "總計"]
    assert list(data["SKU"].astype(str)) == ["2", "1"]
    assert int(data[data["SKU"] == "1"]["影響店舖數"].iloc[0]) == 2
    assert len(by_shop[by_shop["Site Code"] != "總計"]) == 2


def test_missing_signal_columns_rejected():
    df = pd.DataFrame([{"Buyer": "B1", "Brand": "A", "SKU": 1, "Shop Stock Status": "Normal", "Site Code": "S01"}])
    try:
        transform_excess_capital(df)
    except ValueError as e:
        assert "積壓訊號" in str(e)
    else:
        raise AssertionError("應拒絕缺少積壓訊號欄位的輸入")


def test_missing_site_code_rejected():
    df = pd.DataFrame([{"Buyer": "B1", "Brand": "A", "SKU": 1, "Shop Stock Status": "Normal", "Excess Stock Qty": 3.0}])
    try:
        transform_excess_capital(df)
    except ValueError as e:
        assert "Site Code" in str(e)
    else:
        raise AssertionError("應拒絕缺少 Site Code 的輸入")


def test_excel_sheets_and_skip_raw():
    df = _df([
        {"SKU": 1, "Excess Stock Qty": 3.0, "Excess Stock Value": 300.0},
    ])
    raw, by_sku, by_shop, _ = transform_excess_capital(df)
    buf = generate_excel_excess_capital(raw, by_sku, by_shop, include_raw=True)
    wb = load_workbook(BytesIO(buf.getvalue()), data_only=True)
    assert wb.sheetnames == ["RAW", "By SKU", "By Shop"]
    buf2 = generate_excel_excess_capital(raw, by_sku, by_shop, include_raw=False)
    wb2 = load_workbook(BytesIO(buf2.getvalue()), data_only=True)
    assert wb2.sheetnames == ["RAW", "By SKU", "By Shop"]
    assert wb2["RAW"]["A2"].value is not None


def test_empty_result_warning():
    df = _df([{"SKU": 1}])
    raw, by_sku, by_shop, warnings = transform_excess_capital(df)
    assert raw.empty
    assert any("沒有符合" in w for w in warnings)
