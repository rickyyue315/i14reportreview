"""Validation tests for P90D 0 Sales Review(個別店舖).

個別店舖版與店舖回覆版篩選準則相同，差異：
- 分類固定以 MAKE UP 分類
- Top 10 工作表為「店舖90日零銷售Top 10 ALL SKU」，包括全部 SKU，不設「店舖回覆」欄
"""
import warnings
warnings.filterwarnings("ignore")
import pandas as pd
from io import BytesIO
from openpyxl import load_workbook
from transformer import (
    transform_overstock_p90d_shop,
    generate_excel_overstock_p90d_shop,
    _build_op90d_by_sku,
    _build_op90d_by_buyer,
)
from modes.overstock_p90d import (
    _OP90D_TOP10_ALL_SHEET,
    _OP90D_TOP10_ALL_COLUMNS,
    _build_op90d_top10_all,
)

BASE = {
    "Buyer": "B1", "Brand": "A", "SKU": 1,
    "Item_Shop": "Prod-S01-ShopHK", "Site Code": "S01", "OM": "OM1",
    "Shop Nature": "Mix", "RP Type": "RF", "Item Status": "正常",
    "Adj. SST": "SST1", "Vendor MC": "V1", "Adj. MCH1": "01 - SKIN CARE",
    "Shop Stock Status": "Overstock (P90D 0 Sales)",
    "Supply Source": "Warehouse",
    "Item Class": "D",
    "Adj. Safety Stock": 3.0,
    "Shop Stock Qty (CM)": 5.0, "Shop Stock Qty (LM)": 4.0,
    "Sold Qty (P90D)": 0.0, "Sold Qty (M-2)": 0.0,
    "Sold Qty (M-1)": 0.0, "Sold Qty (CM - MTD)": 0.0,
    "Excess Stock Qty": 2.0, "Shop On Order PO Qty (CM)": 0.0,
    "Shop Stock Value (CM)": 500.0, "Excess Stock Value": 100.0,
}


def _row(**kw):
    r = dict(BASE)
    r.update(kw)
    return r


def _df(rows):
    return pd.DataFrame([_row(**r) for r in rows])


def test_shop_filter_same_as_reply_mode():
    df = _df([
        {"SKU": 1},
        {"SKU": 2, "Shop Stock Status": "Normal"},
        {"SKU": 3, "Shop Stock Qty (LM)": 0.0},
    ])
    raw, _, _, _, _, _ = transform_overstock_p90d_shop(df)
    assert set(raw["SKU"].astype(str)) == {"1"}


def test_shop_top10_all_includes_makeup_no_reply():
    rows = [{"SKU": i, "Shop Stock Qty (CM)": float(100 + i),
             "Adj. MCH1": "02 - MAKE UP"} for i in range(1, 6)]
    rows += [{"SKU": i, "Shop Stock Qty (CM)": float(i),
              "Adj. MCH1": "01 - SKIN CARE"} for i in range(6, 12)]
    df = _df(rows)
    _, _, top10, _, _, _ = transform_overstock_p90d_shop(df)
    assert len(top10) == 10
    assert set(top10["MAKE UP 分類"]) == {"MAKE UP", "Non-MAKE UP"}
    assert list(top10["Shop Stock Qty (CM)"])[:5] == [105.0, 104.0, 103.0, 102.0, 101.0]
    assert list(top10["店內庫存排名"]) == list(range(1, 11))
    assert "店舖回覆" not in top10.columns
    assert list(top10.columns) == [c for c in _OP90D_TOP10_ALL_COLUMNS if c in top10.columns]


def test_shop_top10_all_single_shop_limit():
    rows = [{"SKU": i, "Shop Stock Qty (CM)": float(i)} for i in range(1, 13)]
    df = _df(rows)
    raw, _, _, _, _, _ = transform_overstock_p90d_shop(df)
    top10 = _build_op90d_top10_all(raw)
    assert len(top10) == 10
    assert "店舖回覆" not in top10.columns


def test_shop_by_shop_fixed_makeup_columns():
    df = _df([
        {"SKU": 1, "Adj. MCH1": "02 - MAKE UP", "Shop Stock Qty (CM)": 10.0},
        {"SKU": 2, "Adj. MCH1": "01 - SKIN CARE", "Shop Stock Qty (CM)": 5.0},
    ])
    _, by_shop, _, _, _, _ = transform_overstock_p90d_shop(df)
    assert list(by_shop.columns) == [
        "Site Code", "Shop Name", "OM", "Shop Nature", "P90D 0 Sales SKU 總數",
        "Shop Stock Qty (CM) 總數",
        "MAKE UP SKU 數", "MAKE UP Shop Stock Qty (CM) 總數",
        "Non-MAKE UP SKU 數", "Non-MAKE UP Shop Stock Qty (CM) 總數",
    ]
    s01 = by_shop[by_shop["Site Code"] == "S01"].iloc[0]
    assert int(s01["MAKE UP SKU 數"]) == 1
    assert int(s01["Non-MAKE UP SKU 數"]) == 1


def test_shop_excel_sheet_name_and_no_reply_header():
    df = _df([
        {"SKU": 1, "Adj. MCH1": "02 - MAKE UP"},
        {"SKU": 2},
    ])
    raw, by_shop, top10, by_nature, by_om, _ = transform_overstock_p90d_shop(df)
    by_sku = _build_op90d_by_sku(raw)
    by_buyer = _build_op90d_by_buyer(raw)
    buf = generate_excel_overstock_p90d_shop(by_shop, top10, by_nature, by_om, None, by_sku=by_sku, by_buyer=by_buyer)
    wb = load_workbook(BytesIO(buf.getvalue()), data_only=True)
    assert wb.sheetnames == ["By Shop", _OP90D_TOP10_ALL_SHEET, "By Nature", "By OM", "By SKU", "By Buyer"]
    hdr = [c.value for c in wb[_OP90D_TOP10_ALL_SHEET][1]]
    assert "店舖回覆" not in hdr
    assert "MAKE UP 分類" in hdr


def test_shop_empty_result():
    df = _df([{"SKU": 1, "Shop Stock Status": "Normal"}])
    raw, by_shop, top10, by_nature, by_om, warnings = transform_overstock_p90d_shop(df)
    assert raw.empty and top10.empty
    assert any("P90D 0 Sales" in w for w in warnings)
    buf = generate_excel_overstock_p90d_shop(by_shop, top10, by_nature, by_om, None)
    wb = load_workbook(BytesIO(buf.getvalue()), data_only=True)
    assert _OP90D_TOP10_ALL_SHEET in wb.sheetnames
