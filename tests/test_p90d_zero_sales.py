"""Validation tests for Past 90 Days 0 Sales mode (v1.8.6).

篩選：Sold Qty (M-2 / M-1 / CM - MTD) 皆 ≤ 0 且 Shop Stock Qty (LM) 與 (CM) 皆 > 0。
輸出兩頁：RAW（限定欄位白名單）、Summary by Shop。
"""
import warnings
warnings.filterwarnings("ignore")
import numpy as np
import pandas as pd
from io import BytesIO
from openpyxl import load_workbook
from transformer import (
    _p90d_raw_columns,
    _build_p90d_zero_sales_by_shop_summary,
    transform_p90d_zero_sales,
    generate_excel_p90d_zero_sales,
)

BASE = {
    "Buyer": "B1", "Brand": "A", "SKU": 1, "SKU Desc": "Test Product",
    "Item_Shop": "Prod-S01-ShopHK", "Site Code": "S01", "OM": "OM1",
    "Shop Nature": "Mix", "Vendor MC": "V1", "Item Status": "A",
    "Launch Date": "2024-01-01", "Item Class": "C1", "RP Type": "RF",
    "Shop Stock Status": "Normal", "Adj. Safety Stock": 3.0,
    "Sold Qty (M-2)": 0.0, "Sold Qty (M-1)": 0.0, "Sold Qty (CM - MTD)": 0.0,
    "Sold Qty (P90D)": 0.0, "Excess Stock Qty": 2.0, "Shop On Order PO Qty (CM)": 0.0,
    "Shop Stock Qty (LM)": 5.0, "Shop Stock Qty (CM)": 5.0,
    "Shop Stock Value (LM)": 500.0, "Shop Stock Value (CM)": 500.0, "Excess Stock Value": 0.0,
    "D001 Stock Qty (LM)": 0.0, "D001 Stock Qty (CM)": 0.0, "D001 Net Stock Qty (CM)": 0.0,
    "D001 Blocked Qty (CM)": 0.0, "D001 QI Qty (CM)": 0.0, "D001 On Order PO Qty (CM)": 0.0,
}


def _row(**kw):
    r = dict(BASE)
    r.update(kw)
    return r


def _df(rows):
    return pd.DataFrame([_row(**r) for r in rows])


# ──────────────────────────────────────────────────────────────────────────────
# 1. 篩選定義
# ──────────────────────────────────────────────────────────────────────────────
def test_filter_zero_sales_with_stock():
    df = _df([
        {"SKU": 1, "Item_Shop": "A-S01"},  # 符合
        {"SKU": 2, "Sold Qty (M-2)": 3.0},  # M-2 有銷售 → 排除
        {"SKU": 3, "Sold Qty (M-1)": 4.0},  # M-1 有銷售 → 排除
        {"SKU": 4, "Sold Qty (CM - MTD)": 2.0},  # CM 有銷售 → 排除
        {"SKU": 5, "Shop Stock Qty (LM)": 0.0},  # LM=0 → 排除
        {"SKU": 6, "Shop Stock Qty (CM)": 0.0},  # CM=0 → 排除
    ])
    raw, summary, _ = transform_p90d_zero_sales(df)
    assert list(raw["SKU"]) == ["1"]
    assert int(summary[summary["Site Code"] == "總計"]["Past 90 Days 0 Sales SKU 數量"].iloc[0]) == 1


def test_filter_handles_nan_and_negative():
    df = _df([
        {"SKU": 1, "Sold Qty (M-2)": np.nan, "Shop Stock Qty (LM)": np.nan},  # sold NaN=0, stock NaN=0 → 排除
        {"SKU": 2, "Shop Stock Qty (CM)": -1.0},  # 負庫存視為 0 → 排除
        {"SKU": 3, "Sold Qty (M-2)": -5.0, "Shop Stock Qty (LM)": 3.0, "Shop Stock Qty (CM)": 3.0},  # 負銷售 ≤ 0 且庫存>0 → 符合
        {"SKU": 4, "Sold Qty (M-1)": -2.0, "Sold Qty (CM - MTD)": -1.0,
         "Shop Stock Qty (LM)": 4.0, "Shop Stock Qty (CM)": 4.0},  # 三期皆 ≤ 0（含負數）→ 符合
        {"SKU": 5, "Sold Qty (CM - MTD)": 0.5, "Shop Stock Qty (LM)": 3.0, "Shop Stock Qty (CM)": 3.0},  # 正銷售 > 0 → 排除
    ])
    raw, _, _ = transform_p90d_zero_sales(df)
    assert list(raw["SKU"]) == ["3", "4"]


def test_missing_required_column_raises():
    df = _df([{"SKU": 1}]).drop(columns=["Shop Stock Qty (LM)"])
    try:
        transform_p90d_zero_sales(df)
        assert False, "應拋出 ValueError"
    except ValueError as e:
        assert "Shop Stock Qty (LM)" in str(e)


# ──────────────────────────────────────────────────────────────────────────────
# 2. 欄位移除與 Desc./Shop Name 解像
# ──────────────────────────────────────────────────────────────────────────────
def test_p90d_raw_columns_whitelist():
    df = _df([{"SKU": 1}])
    cols = _p90d_raw_columns(df)
    assert not any(("D001" in c or "Value" in c) for c in cols)
    # 只輸出白名單欄位，順序一致（Adj. Safety Stock 表頭顯示為 Safety Stock (現時)）
    assert cols == [
        "Site Code", "Shop Name", "OM", "Shop Nature", "Brand", "SKU", "SKU Desc",
        "Vendor MC", "Item Status", "Launch Date", "Item Class", "RP Type",
        "Shop Stock Status", "Adj. Safety Stock",
        "Sold Qty (M-2)", "Sold Qty (M-1)", "Sold Qty (CM - MTD)", "Sold Qty (P90D)",
        "Shop Stock Qty (CM)", "Shop Stock Qty (LM)",
        "Excess Stock Qty", "Shop On Order PO Qty (CM)", "Buyer",
    ]


def test_p90d_raw_columns_skips_missing_optional():
    # 輸入缺少選填欄位（如 Shop Nature / Adj. Safety Stock）時直接略過，不報錯
    df = pd.DataFrame([{"SKU": 1, "Site Code": "S01", "Brand": "A", "Buyer": "B1"}])
    cols = _p90d_raw_columns(df)
    assert cols == ["Site Code", "Shop Name", "Brand", "SKU", "SKU Desc", "Buyer"]


def test_raw_sheet_has_desc_and_no_d001_value():
    df = _df([
        {"SKU": 1, "Item_Shop": "Eau de Toilette 50ml-S01-ShopHK", "Site Code": "S01"},
        {"SKU": 2, "Item_Shop": "Eau de Parfum 100ml-S02-ShopKowloon", "Site Code": "S02", "Buyer": "B2", "Brand": "B"},
    ])
    raw, _, _ = transform_p90d_zero_sales(df)
    buf = generate_excel_p90d_zero_sales(raw, pd.DataFrame(), None)
    wb = load_workbook(BytesIO(buf.getvalue()))
    ws = wb["RAW"]
    headers = [c.value for c in ws[1]]
    assert "SKU Desc" in headers and "Shop Name" in headers
    assert not any((h and ("D001" in h or "Value" in h)) for h in headers)
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    site_idx = headers.index("Site Code")
    name_idx = headers.index("Shop Name")
    desc_idx = headers.index("SKU Desc")
    by_site = {r[site_idx]: r for r in rows}
    assert by_site["S01"][name_idx] == "ShopHK"
    assert by_site["S02"][name_idx] == "ShopKowloon"
    assert by_site["S02"][desc_idx] == "Eau de Parfum 100ml-"  # 既有 _extract_sku_desc 保留前綴連字號


# ──────────────────────────────────────────────────────────────────────────────
# 3. Summary by Shop
# ──────────────────────────────────────────────────────────────────────────────
def test_summary_by_shop_counts_and_totals():
    df = _df([
        {"SKU": 1, "Site Code": "S01", "Item_Shop": "A-S01-ShopHK", "Shop Stock Qty (LM)": 10.0, "Shop Stock Qty (CM)": 20.0},
        {"SKU": 1, "Site Code": "S01", "Item_Shop": "A-S01-ShopHK", "Shop Stock Qty (LM)": 10.0, "Shop Stock Qty (CM)": 20.0},  # 同 SKU 重複列
        {"SKU": 2, "Site Code": "S01", "Item_Shop": "B-S01-ShopHK", "Shop Stock Qty (LM)": 5.0, "Shop Stock Qty (CM)": 5.0},
        {"SKU": 3, "Site Code": "S02", "Item_Shop": "C-S02-ShopKowloon", "Shop Stock Qty (LM)": 7.0, "Shop Stock Qty (CM)": 8.0},
    ])
    raw, summary, _ = transform_p90d_zero_sales(df)
    s01 = summary[(summary["Site Code"] == "S01") & (summary["Shop Name"] == "ShopHK")]
    assert int(s01["Past 90 Days 0 Sales SKU 數量"].iloc[0]) == 2  # 去重 SKU 1,2
    assert float(s01["Shop Stock Qty (LM) 總數"].iloc[0]) == 10 + 10 + 5  # 未去重加總
    assert float(s01["Shop Stock Qty (CM) 總數"].iloc[0]) == 20 + 20 + 5
    total = summary[summary["Site Code"] == "總計"]
    assert int(total["Past 90 Days 0 Sales SKU 數量"].iloc[0]) == 3  # 全域去重
    assert float(total["Shop Stock Qty (CM) 總數"].iloc[0]) == 20 + 20 + 5 + 8


def test_summary_empty_input():
    df = pd.DataFrame(columns=list(BASE.keys()))
    raw, summary, _ = transform_p90d_zero_sales(df)
    assert raw.empty
    assert summary.empty


# ──────────────────────────────────────────────────────────────────────────────
# 4. Excel 輸出結構
# ──────────────────────────────────────────────────────────────────────────────
def test_excel_sheets_only_raw_and_summary():
    df = _df([
        {"SKU": 1, "Site Code": "S01", "Item_Shop": "A-S01"},
        {"SKU": 2, "Site Code": "S02", "Item_Shop": "B-S02", "Buyer": "B2", "Brand": "B"},
    ])
    raw, summary, _ = transform_p90d_zero_sales(df)
    buf = generate_excel_p90d_zero_sales(raw, summary, None)
    wb = load_workbook(BytesIO(buf.getvalue()))
    assert wb.sheetnames == ["RAW", "Summary by Shop"]
    assert "Shop SKU" not in wb.sheetnames
    assert "_SiteList" not in wb.sheetnames


def test_excel_empty_result_still_has_two_sheets():
    df = _df([{"SKU": 1, "Sold Qty (M-2)": 9.0}])  # 全部排除
    raw, summary, warnings = transform_p90d_zero_sales(df)
    assert raw.empty
    assert len(warnings) == 1
    buf = generate_excel_p90d_zero_sales(raw, summary, None)
    wb = load_workbook(BytesIO(buf.getvalue()))
    assert wb.sheetnames == ["RAW", "Summary by Shop"]


# ──────────────────────────────────────────────────────────────────────────────
# 5. Summary 排序與總計
# ──────────────────────────────────────────────────────────────────────────────
def test_summary_sorted_with_total_last():
    df = _df([
        {"SKU": 1, "Site Code": "S02", "Item_Shop": "A-S02"},
        {"SKU": 2, "Site Code": "S01", "Item_Shop": "B-S01"},
    ])
    raw, summary, _ = transform_p90d_zero_sales(df)
    sites = list(summary["Site Code"])
    assert sites == ["S01", "S02", "總計"]
