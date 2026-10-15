"""Validation tests for P90D 0 Sales Review(轉貨模式).

規則：
- 候選轉出店舖：Sold Qty (P90D) ≤ 0 且 Shop Stock Qty (CM) > 0
- 轉出：已選轉出店舖中符合上述條件的列，數量為 Shop Stock Qty (CM) 全數，不分 OM、不設上限
- 接收：同 SKU、非轉出店舖、Sold Qty (P90D) > 0 的列，不論有無庫存
- 分配：同 SKU 轉出總量按接收店 Sold Qty (P90D) 比例分配（最大餘數法取整數件、全數分完）
- MCH 篩選（執行前）：可按 Adj. MCH1 排除分類（預設全部包含），對轉出與接收兩側同時生效
"""
import random
import warnings

warnings.filterwarnings("ignore")

import pandas as pd
from openpyxl import load_workbook

from modes.transfer_p90d import (
    _allocate_proportional,
    build_mch_options,
    build_transfer_candidates,
    generate_excel_p90d_transfer,
    transform_p90d_transfer,
)
import transformer as shim
from modes import MODE_REGISTRY

BASE = {
    "Buyer": "B1", "Brand": "A", "SKU": 1,
    "Item_Shop": "Prod-S01-ShopHK", "Site Code": "S01", "OM": "OM1",
    "Shop Nature": "Mix", "RP Type": "RF",
    "Shop Stock Status": "Overstock (P90D 0 Sales)",
    "Adj. Safety Stock": 3.0, "Adj. MCH1": "01 - SKIN CARE",
    "Sold Qty (M-2)": 0.0, "Sold Qty (M-1)": 0.0, "Sold Qty (CM - MTD)": 0.0,
    "Sold Qty (P90D)": 0.0,
    "Shop Stock Qty (CM)": 5.0, "Shop Stock Qty (LM)": 4.0,
}


def _row(site, sku=1, sales=0.0, stock=5.0, name=None, om="OM1", nature="Mix",
         brand="A", buyer="B1", mch1="01 - SKIN CARE", mch2=None):
    r = dict(BASE)
    r.update({
        "Site Code": site,
        "SKU": sku,
        "Item_Shop": f"Prod-{site}-{name or 'Shop' + site}",
        "Sold Qty (P90D)": float(sales),
        "Sold Qty (CM - MTD)": float(sales),
        "Shop Stock Qty (CM)": float(stock),
        "OM": om,
        "Shop Nature": nature,
        "Brand": brand,
        "Buyer": buyer,
        "Adj. MCH1": mch1,
    })
    if mch2 is not None:
        r["MCH2"] = mch2
    return r


def _df(rows):
    return pd.DataFrame(rows)


def _lines(plan):
    return {
        (r["轉出 Site Code"], str(r["SKU"]), r["接收 Site Code"]): int(r["轉貨數量 (件)"])
        for _, r in plan.iterrows()
    }


def test_candidates_only_zero_sales_with_stock():
    df = _df([
        _row("S01", sku=1, sales=0, stock=5),
        _row("S02", sku=1, sales=3, stock=5),
        _row("S03", sku=2, sales=0, stock=0),
        _row("S01", sku=2, sales=0, stock=2),
    ])
    cand = build_transfer_candidates(df)
    assert list(cand["Site Code"]) == ["S01"]
    row = cand.iloc[0]
    assert row["Shop Name"] == "ShopS01"
    assert int(row["0 Sales SKU 數"]) == 2
    assert int(row["可轉出庫存 (件)"]) == 7


def test_single_receiver_gets_full_stock():
    df = _df([
        _row("S01", sku=1, sales=0, stock=5),
        _row("S02", sku=1, sales=4, stock=1),
        _row("S03", sku=1, sales=0, stock=9),
    ])
    plan, out_df, in_df, sku_df, mch_df, stats, warnings = transform_p90d_transfer(df, ["S01"])
    assert _lines(plan) == {("S01", "1", "S02"): 5}
    line = plan.iloc[0]
    assert int(line["接收店 Shop Stock Qty (CM) 調前"]) == 1
    assert int(line["接收店同 SKU 分配總量 (件)"]) == 5
    assert float(line["接收店銷售佔比 (同 SKU)"]) == 1.0
    assert stats["alloc_qty"] == 5 and stats["unalloc_qty"] == 0
    assert stats["unalloc_skus"] == 0


def test_multi_receiver_proportional_by_sales():
    df = _df([
        _row("S01", sku=1, sales=0, stock=10),
        _row("S02", sku=1, sales=3, stock=0),
        _row("S03", sku=1, sales=1, stock=0),
    ])
    plan, _, in_df, _, _, stats, _ = transform_p90d_transfer(df, ["S01"])
    assert _lines(plan) == {("S01", "1", "S02"): 8, ("S01", "1", "S03"): 2}
    assert stats["alloc_qty"] == 10 and stats["unalloc_qty"] == 0
    shares = dict(zip(plan["接收 Site Code"], plan["接收店銷售佔比 (同 SKU)"]))
    assert abs(shares["S02"] - 0.75) < 1e-9 and abs(shares["S03"] - 0.25) < 1e-9


def test_multi_out_sites_pooled_and_paired():
    df = _df([
        _row("S01", sku=1, sales=0, stock=4),
        _row("S05", sku=1, sales=0, stock=6),
        _row("S02", sku=1, sales=3, stock=0),
        _row("S03", sku=1, sales=1, stock=0),
    ])
    plan, _, _, _, _, stats, _ = transform_p90d_transfer(df, ["S01", "S05"])
    assert _lines(plan) == {
        ("S01", "1", "S02"): 4,
        ("S05", "1", "S02"): 4,
        ("S05", "1", "S03"): 2,
    }
    assert stats["out_qty"] == 10 and stats["alloc_qty"] == 10


def test_no_receiver_stays_unallocated():
    df = _df([
        _row("S01", sku=1, sales=0, stock=5),
        _row("S02", sku=1, sales=0, stock=3),
    ])
    plan, out_df, in_df, sku_df, mch_df, stats, warnings = transform_p90d_transfer(df, ["S01"])
    assert plan.empty
    assert stats["unalloc_qty"] == 5 and stats["unalloc_skus"] == 1
    body = sku_df[sku_df["SKU"].astype(str) == "1"].iloc[0]
    assert int(body["未分配量 (件)"]) == 5
    assert body["狀態"] == "無接收店舖（保留原店）"
    assert any("未能轉出" in w for w in warnings)


def test_selected_site_sales_rows_are_neither_source_nor_receiver():
    df = _df([
        _row("S01", sku=1, sales=0, stock=5),
        _row("S01", sku=2, sales=2, stock=4),
        _row("S02", sku=2, sales=1, stock=0),
    ])
    plan, _, _, sku_df, _, stats, _ = transform_p90d_transfer(df, ["S01"])
    assert plan.empty
    assert stats["out_skus"] == 1
    body = sku_df[sku_df["Brand"] != "總計"]
    assert set(body["SKU"].astype(str)) == {"1"}


def test_om_ignored_for_transfer():
    df = _df([
        _row("S01", sku=1, sales=0, stock=5, om="OM1"),
        _row("S02", sku=1, sales=2, stock=0, om="OM2"),
    ])
    plan, _, _, _, _, stats, _ = transform_p90d_transfer(df, ["S01"])
    assert _lines(plan) == {("S01", "1", "S02"): 5}
    assert plan.iloc[0]["接收店 OM"] == "OM2"


def test_shop_nature_filter_applies_to_both_sides():
    df = _df([
        _row("S01", sku=1, sales=0, stock=6, nature="Mix"),
        _row("S02", sku=1, sales=5, stock=0, nature="Tourist"),
        _row("S03", sku=1, sales=5, stock=0, nature="Non Tourist"),
    ])
    plan, _, _, _, _, _, _ = transform_p90d_transfer(
        df, ["S01"], shop_natures=["Mix", "Non Tourist"]
    )
    assert _lines(plan) == {("S01", "1", "S03"): 6}


def test_duplicate_site_sku_is_aggregated_with_warning():
    df = _df([
        _row("S01", sku=1, sales=0, stock=2),
        _row("S01", sku=1, sales=0, stock=3),
        _row("S02", sku=1, sales=1, stock=0),
    ])
    plan, _, _, _, _, stats, warnings = transform_p90d_transfer(df, ["S01"])
    assert _lines(plan) == {("S01", "1", "S02"): 5}
    assert any("重複" in w for w in warnings)


def test_selected_site_without_zero_stock_is_reported():
    df = _df([
        _row("S01", sku=1, sales=0, stock=5),
        _row("S09", sku=1, sales=2, stock=3),
        _row("S02", sku=1, sales=1, stock=0),
    ])
    plan, _, _, _, _, _, warnings = transform_p90d_transfer(df, ["S01", "S09"])
    assert _lines(plan) == {("S01", "1", "S02"): 5}
    assert any("未納入轉出" in w and "S09" in w for w in warnings)


def test_no_zero_stock_rows_returns_empty_with_warning():
    df = _df([
        _row("S09", sku=1, sales=2, stock=3),
        _row("S02", sku=1, sales=1, stock=0),
    ])
    plan, out_df, in_df, sku_df, mch_df, stats, warnings = transform_p90d_transfer(df, ["S09"])
    assert plan.empty and out_df.empty and in_df.empty and sku_df.empty and mch_df.empty
    assert stats["out_qty"] == 0
    assert any("沒有符合" in w for w in warnings)


def test_empty_selection_raises():
    df = _df([_row("S01", sku=1, sales=0, stock=5), _row("S02", sku=1, sales=1, stock=0)])
    try:
        transform_p90d_transfer(df, [])
    except ValueError as exc:
        assert "轉出店舖" in str(exc)
    else:
        raise AssertionError("expected ValueError for empty selection")


def test_missing_required_column_raises():
    df = _df([_row("S01", sku=1, sales=0, stock=5)]).drop(columns=["Sold Qty (P90D)"])
    try:
        transform_p90d_transfer(df, ["S01"])
    except ValueError as exc:
        assert "Sold Qty (P90D)" in str(exc)
    else:
        raise AssertionError("expected ValueError for missing column")


# 1. MCH 欄位與分類彙總（只用 Adj. MCH1）

def test_mch_columns_in_plan_and_sku():
    df = _df([
        _row("S01", sku=1, sales=0, stock=5, mch1="02 - MAKE UP", mch2="0201 - FACE"),
        _row("S02", sku=1, sales=3, stock=0),
    ])
    plan, _, _, sku_df, mch_df, _, _ = transform_p90d_transfer(df, ["S01"])
    assert "Adj. MCH1" in plan.columns
    assert plan.iloc[0]["Adj. MCH1"] == "02 - MAKE UP"
    body = sku_df[sku_df["Brand"] != "總計"]
    assert body.iloc[0]["Adj. MCH1"] == "02 - MAKE UP"
    assert "MCH2" not in mch_df.columns


def test_mch_summary_single_level_with_total():
    df = _df([
        _row("S01", sku=1, sales=0, stock=4, mch1="01 - SKIN CARE", mch2="0101 - FACE & NECK"),
        _row("S01", sku=2, sales=0, stock=6, mch1="02 - MAKE UP", mch2="0201 - FACE"),
        _row("S02", sku=1, sales=3, stock=0),
        _row("S03", sku=2, sales=1, stock=0),
    ])
    plan, _, _, _, mch_df, stats, _ = transform_p90d_transfer(df, ["S01"])
    assert list(mch_df.columns) == [
        "Adj. MCH1", "轉出 SKU 數", "接收 SKU 數",
        "轉出總量 (件)", "接收總量 (件)", "未分配量 (件)",
    ]
    assert set(mch_df["Adj. MCH1"]) == {"01 - SKIN CARE", "02 - MAKE UP", "總計"}
    skin = mch_df[mch_df["Adj. MCH1"] == "01 - SKIN CARE"].iloc[0]
    assert int(skin["轉出 SKU 數"]) == 1 and int(skin["轉出總量 (件)"]) == 4
    assert int(skin["接收總量 (件)"]) == 4 and int(skin["未分配量 (件)"]) == 0
    total = mch_df.iloc[-1]
    assert total["Adj. MCH1"] == "總計"
    assert int(total["轉出總量 (件)"]) == 10
    assert int(total["接收總量 (件)"]) == 10
    assert int(total["未分配量 (件)"]) == 0


def test_mch_blank_group():
    df = _df([
        _row("S01", sku=1, sales=0, stock=5, mch1=""),
        _row("S02", sku=1, sales=2, stock=0),
    ])
    _, _, _, _, mch_df, _, _ = transform_p90d_transfer(df, ["S01"])
    blank = mch_df[mch_df["Adj. MCH1"] == "(空白)"]
    assert len(blank) == 1
    assert int(blank.iloc[0]["轉出總量 (件)"]) == 5


# 2. MCH 執行前篩選

def test_mch1_filter_excludes_out_and_receive():
    df = _df([
        _row("S01", sku=1, sales=0, stock=5, mch1="01 - SKIN CARE"),
        _row("S01", sku=2, sales=0, stock=7, mch1="02 - MAKE UP"),
        _row("S02", sku=1, sales=3, stock=0, mch1="02 - MAKE UP"),
        _row("S03", sku=2, sales=3, stock=0, mch1="01 - SKIN CARE"),
    ])
    plan, _, _, _, _, stats, warnings = transform_p90d_transfer(
        df, ["S01"], mch1_filter=["01 - SKIN CARE"]
    )
    # SKU 1 屬 SKIN CARE 可轉出，但其接收店 S02 屬 MAKE UP 已被排除 → 未分配
    assert plan.empty
    assert stats["out_qty"] == 5 and stats["unalloc_qty"] == 5
    assert any("MCH 篩選已排除" in w for w in warnings)


def test_mch1_filter_keeps_selected_only():
    df = _df([
        _row("S01", sku=1, sales=0, stock=5, mch1="01 - SKIN CARE"),
        _row("S01", sku=2, sales=0, stock=7, mch1="02 - MAKE UP"),
        _row("S02", sku=1, sales=3, stock=0, mch1="01 - SKIN CARE"),
        _row("S02", sku=2, sales=3, stock=0, mch1="02 - MAKE UP"),
    ])
    plan, _, _, _, _, stats, _ = transform_p90d_transfer(
        df, ["S01"], mch1_filter=["01 - SKIN CARE"]
    )
    assert _lines(plan) == {("S01", "1", "S02"): 5}
    assert stats["out_qty"] == 5


def test_mch_filter_blank_token():
    df = _df([
        _row("S01", sku=1, sales=0, stock=5, mch1=""),
        _row("S01", sku=2, sales=0, stock=7, mch1="01 - SKIN CARE"),
        _row("S02", sku=1, sales=3, stock=0, mch1=""),
        _row("S02", sku=2, sales=3, stock=0, mch1="01 - SKIN CARE"),
    ])
    plan, _, _, _, _, stats, _ = transform_p90d_transfer(
        df, ["S01"], mch1_filter=["__BLANK__"]
    )
    assert _lines(plan) == {("S01", "1", "S02"): 5}
    assert stats["out_qty"] == 5


def test_mch_filter_excludes_all_returns_empty():
    df = _df([
        _row("S01", sku=1, sales=0, stock=5, mch1="01 - SKIN CARE"),
        _row("S02", sku=1, sales=3, stock=0),
    ])
    plan, _, _, _, _, stats, warnings = transform_p90d_transfer(
        df, ["S01"], mch1_filter=["99 - NOT EXIST"]
    )
    assert plan.empty and stats["out_qty"] == 0
    assert any("MCH 篩選後無有效資料列" in w for w in warnings)


def test_mch_filter_none_means_all():
    df = _df([
        _row("S01", sku=1, sales=0, stock=5, mch1="01 - SKIN CARE"),
        _row("S02", sku=1, sales=3, stock=0),
    ])
    plan_all, _, _, _, _, stats_all, warns_all = transform_p90d_transfer(df, ["S01"])
    plan_none, _, _, _, _, stats_none, _ = transform_p90d_transfer(df, ["S01"], mch1_filter=None)
    assert _lines(plan_all) == _lines(plan_none)
    assert stats_all == stats_none
    assert not any("MCH 篩選" in w for w in warns_all)


def test_candidates_respect_mch_filter():
    df = _df([
        _row("S01", sku=1, sales=0, stock=5, mch1="01 - SKIN CARE"),
        _row("S02", sku=2, sales=0, stock=9, mch1="02 - MAKE UP"),
    ])
    cand = build_transfer_candidates(df, mch1_filter=["01 - SKIN CARE"])
    assert list(cand["Site Code"]) == ["S01"]


def test_build_mch_options():
    df = _df([
        _row("S01", sku=1, sales=0, stock=5, mch1="01 - SKIN CARE", mch2="0101 - FACE & NECK"),
        _row("S01", sku=2, sales=0, stock=5, mch1="01 - SKIN CARE", mch2="0102 - EYES"),
        _row("S02", sku=3, sales=0, stock=5, mch1="02 - MAKE UP", mch2="0201 - FACE"),
        _row("S03", sku=4, sales=0, stock=5, mch1=""),
    ])
    assert build_mch_options(df) == ["01 - SKIN CARE", "02 - MAKE UP", "__BLANK__"]


def test_allocate_proportional_sums_to_total():
    assert _allocate_proportional(10, [3, 1]) == [8, 2]
    alloc = _allocate_proportional(7, [1, 1, 1])
    assert sum(alloc) == 7 and max(alloc) - min(alloc) <= 1
    assert _allocate_proportional(0, [1, 2]) == [0, 0]
    assert _allocate_proportional(5, []) == []
    rng = random.Random(20261014)
    for _ in range(300):
        total = rng.randint(1, 500)
        weights = [float(rng.randint(1, 60)) for _ in range(rng.randint(1, 8))]
        alloc = _allocate_proportional(total, weights)
        assert sum(alloc) == total
        wsum = sum(weights)
        for a, w in zip(alloc, weights):
            ideal = total * w / wsum
            assert abs(a - ideal) < 1.0 + 1e-9


def test_summary_totals_match_plan():
    df = _df([
        _row("S01", sku=1, sales=0, stock=4),
        _row("S05", sku=1, sales=0, stock=6),
        _row("S05", sku=2, sales=0, stock=3),
        _row("S02", sku=1, sales=3, stock=2),
        _row("S03", sku=1, sales=1, stock=0),
        _row("S03", sku=2, sales=0, stock=0),
    ])
    plan, out_df, in_df, sku_df, mch_df, stats, _ = transform_p90d_transfer(df, ["S01", "S05"])
    out_total = out_df[out_df["Site Code"] == "總計"].iloc[0]
    in_total = in_df[in_df["Site Code"] == "總計"].iloc[0]
    sku_total = sku_df[sku_df["Brand"] == "總計"].iloc[0]
    assert int(out_total["轉出總量 (件)"]) == 13
    assert int(out_total["已分配 (件)"]) == 10
    assert int(out_total["未分配 (件)"]) == 3
    assert int(plan["轉貨數量 (件)"].sum()) == 10
    assert int(in_total["接收總量 (件)"]) == 10
    assert int(in_total["接收前庫存 (件)"]) == 2
    assert int(in_total["接收後庫存 (件)"]) == 12
    assert int(sku_total["已分配總量 (件)"]) == 10
    assert int(sku_total["未分配量 (件)"]) == 3
    mch_total = mch_df[mch_df["Adj. MCH1"] == "總計"].iloc[0]
    assert int(mch_total["轉出總量 (件)"]) == 13
    assert int(mch_total["接收總量 (件)"]) == 10
    assert int(mch_total["未分配量 (件)"]) == 3


def test_excel_structure_and_percent_format():
    df = _df([
        _row("S01", sku=1, sales=0, stock=10),
        _row("S02", sku=1, sales=3, stock=0),
        _row("S03", sku=1, sales=1, stock=0),
    ])
    plan, out_df, in_df, sku_df, mch_df, stats, _ = transform_p90d_transfer(df, ["S01"])
    buf = generate_excel_p90d_transfer(
        plan, out_df, in_df, sku_df, mch_df, stats, nature_label="全部"
    )
    wb = load_workbook(buf)
    assert wb.sheetnames == [
        "說明", "轉貨單", "轉出店舖彙總", "接收店舖彙總", "SKU 轉貨彙總", "MCH1 分類彙總",
    ]
    ws = wb["轉貨單"]
    headers = [c.value for c in ws[1]]
    assert headers == list(plan.columns)
    share_col = headers.index("接收店銷售佔比 (同 SKU)") + 1
    assert ws.cell(row=2, column=share_col).number_format == "0.0%"
    qty_col = headers.index("轉貨數量 (件)") + 1
    assert ws.cell(row=2, column=qty_col).value == 8


def test_excel_plan_group_colors():
    """轉貨單欄位分組著色：黃＝轉出店與轉貨數量（表頭 FFD966／資料 FFF2CC、數量加粗）、
    淺藍＝接收店相關欄位（表頭 9DC3E6／資料 DDEBF7）；其餘欄位維持預設深藍表頭、無資料底色。"""
    df = _df([
        _row("S01", sku=1, sales=0, stock=10),
        _row("S02", sku=1, sales=3, stock=0),
        _row("S03", sku=1, sales=1, stock=0),
    ])
    plan, out_df, in_df, sku_df, mch_df, stats, _ = transform_p90d_transfer(df, ["S01"])
    buf = generate_excel_p90d_transfer(plan, out_df, in_df, sku_df, mch_df, stats)
    wb = load_workbook(buf)
    ws = wb["轉貨單"]
    headers = [c.value for c in ws[1]]
    out_cols = ["轉出 Site Code", "轉出 Shop Name", "轉貨數量 (件)"]
    in_cols = [
        "接收 Site Code", "接收 Shop Name", "接收店 Shop Nature", "接收店 OM",
        "接收店 Sold Qty (P90D)", "接收店銷售佔比 (同 SKU)",
        "接收店同 SKU 分配總量 (件)", "接收店 Shop Stock Qty (CM) 調前",
    ]
    plain_cols = ["Brand", "SKU", "SKU Desc", "Buyer", "Adj. MCH1", "MCH2"]
    for c in out_cols:
        i = headers.index(c) + 1
        assert ws.cell(row=1, column=i).fill.start_color.rgb.endswith("FFD966"), c
        assert ws.cell(row=1, column=i).font.color.rgb.endswith("000000"), c
        assert ws.cell(row=2, column=i).fill.start_color.rgb.endswith("FFF2CC"), c
    for c in in_cols:
        i = headers.index(c) + 1
        assert ws.cell(row=1, column=i).fill.start_color.rgb.endswith("9DC3E6"), c
        assert ws.cell(row=1, column=i).font.color.rgb.endswith("000000"), c
        assert ws.cell(row=2, column=i).fill.start_color.rgb.endswith("DDEBF7"), c
    for c in plain_cols:
        i = headers.index(c) + 1
        assert ws.cell(row=1, column=i).fill.start_color.rgb.endswith("4472C4"), c
        assert ws.cell(row=1, column=i).font.color.rgb.endswith("FFFFFF"), c
        assert ws.cell(row=2, column=i).fill.fill_type is None, c
    # 轉貨數量資料列加粗；其餘著色欄資料列不加粗
    qty_i = headers.index("轉貨數量 (件)") + 1
    assert ws.cell(row=2, column=qty_i).font.bold is True
    assert ws.cell(row=2, column=headers.index("接收店 OM") + 1).font.bold is not True
    # 其他工作表維持預設格式（不受分組著色影響）
    for name in ["轉出店舖彙總", "接收店舖彙總", "SKU 轉貨彙總", "MCH1 分類彙總"]:
        assert wb[name].cell(row=1, column=1).fill.start_color.rgb.endswith("4472C4"), name
        assert wb[name].cell(row=2, column=1).fill.fill_type is None, name


def test_excel_plan_group_colors_when_plan_empty():
    """空轉貨單仍保留著色表頭與「無符合資料」提示，不應拋錯。"""
    df = _df([_row("S01", sku=1, sales=0, stock=5), _row("S02", sku=1, sales=0, stock=3)])
    plan, out_df, in_df, sku_df, mch_df, stats, _ = transform_p90d_transfer(df, ["S01"])
    assert plan.empty
    buf = generate_excel_p90d_transfer(plan, out_df, in_df, sku_df, mch_df, stats)
    wb = load_workbook(buf)
    ws = wb["轉貨單"]
    headers = [c.value for c in ws[1]]
    assert ws.cell(row=1, column=headers.index("轉出 Site Code") + 1).fill.start_color.rgb.endswith("FFD966")
    assert ws.cell(row=2, column=1).value == "無符合資料"


def test_excel_empty_plan_writes_placeholder():
    df = _df([_row("S01", sku=1, sales=0, stock=5), _row("S02", sku=1, sales=0, stock=3)])
    plan, out_df, in_df, sku_df, mch_df, stats, _ = transform_p90d_transfer(df, ["S01"])
    buf = generate_excel_p90d_transfer(plan, out_df, in_df, sku_df, mch_df, stats)
    wb = load_workbook(buf)
    assert wb["轉貨單"].cell(row=2, column=1).value == "無符合資料"
    assert wb["MCH1 分類彙總"].cell(row=1, column=1).value == "Adj. MCH1"


def test_transformer_shim_and_registry():
    assert shim.build_transfer_candidates is build_transfer_candidates
    assert shim.build_mch_options is build_mch_options
    assert shim.transform_p90d_transfer is transform_p90d_transfer
    assert shim.generate_excel_p90d_transfer is generate_excel_p90d_transfer
    entry = MODE_REGISTRY["P90D 0 Sales Review(轉貨模式)"]
    assert entry["module"] == "modes.transfer_p90d"
    assert entry["transform"] == "transform_p90d_transfer"
    assert entry["generate_excel"] == "generate_excel_p90d_transfer"
