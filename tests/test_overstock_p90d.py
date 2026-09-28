"""Validation tests for P90D 0 Sales Review(Shop回覆).

覆蓋：狀態篩選、MAKE UP 分類、店舖統計、Top 10、Nature／OM 分佈、
不輸出 Value／COGS／GP 欄位、Excel 四頁結構。
"""
import warnings
warnings.filterwarnings("ignore")
import pandas as pd
from io import BytesIO
from openpyxl import load_workbook
from transformer import (
    transform_overstock_p90d,
    generate_excel_overstock_p90d,
)

BASE = {
    "Buyer": "B1", "Brand": "A", "SKU": 1,
    "Item_Shop": "Prod-S01-ShopHK", "Site Code": "S01", "OM": "OM1",
    "Shop Nature": "Mix", "RP Type": "RF", "Item Status": "正常",
    "Adj. SST": "SST1", "Vendor MC": "V1", "Adj. MCH1": "01 - SKIN CARE",
    "Shop Stock Status": "Overstock (P90D 0 Sales)",
    "Supply Source": "Warehouse",
    "Item Class": "D",
    "Adj. Safety Stock": 3.0,
    "Shop Stock Qty (CM)": 5.0, "Shop Stock Qty (LM)": 4.0,
    "Sold Qty (P90D)": 0.0, "Sold Qty (M-2)": 0.0,
    "Sold Qty (M-1)": 0.0, "Sold Qty (CM - MTD)": 0.0,
    "Excess Stock Qty": 2.0, "Shop On Order PO Qty (CM)": 0.0,
    "Shop Stock Value (CM)": 500.0, "Excess Stock Value": 100.0,
}


def _row(**kw):
    r = dict(BASE)
    r.update(kw)
    return r


def _df(rows):
    return pd.DataFrame([_row(**r) for r in rows])


def _no_banned(df):
    import re
    for c in df.columns:
        lowered = str(c).lower()
        assert "value" not in lowered, c
        assert "cogs" not in lowered, c
        assert re.search(r"\bgp\b", lowered) is None, c


# ── 1. 狀態篩選 ──────────────────────────────────────────────────────────
def test_filter_only_overstock_p90d_status():
    df = _df([
        {"SKU": 1},  # 符合
        {"SKU": 2, "Shop Stock Status": " Overstock (P90D 0 Sales) "},  # 前後空格 → 符合
        {"SKU": 3, "Shop Stock Status": "Normal"},  # 排除
        {"SKU": 4, "Shop Stock Status": "Overstock (>2*Monthly Sold Qty & >2pcs)"},  # 排除
        {"SKU": 5, "Shop Stock Status": None},  # 排除
    ])
    raw, _, _, _, _, _ = transform_overstock_p90d(df)
    assert set(raw["SKU"].astype(str)) == {"1", "2"}


def test_filter_requires_lm_stock():
    import numpy as np
    df = _df([
        {"SKU": 1, "Shop Stock Qty (LM)": 4.0},  # 符合
        {"SKU": 2, "Shop Stock Qty (LM)": 0.0},  # LM 為 0 → 排除
        {"SKU": 3, "Shop Stock Qty (LM)": np.nan},  # LM 空白視為 0 → 排除
        {"SKU": 4, "Shop Stock Qty (LM)": -1.0},  # LM 負數 → 排除
    ])
    raw, _, _, _, _, _ = transform_overstock_p90d(df)
    assert list(raw["SKU"].astype(str)) == ["1"]


def test_empty_result_warning():
    df = _df([{"SKU": 1, "Shop Stock Status": "Normal"}])
    raw, by_shop, top10, by_nature, by_om, warnings = transform_overstock_p90d(df)
    assert raw.empty
    assert by_shop.empty and top10.empty and by_nature.empty and by_om.empty
    assert any("P90D 0 Sales" in w for w in warnings)


def test_missing_required_columns_rejected():
    for col in ("Site Code", "Shop Stock Qty (CM)", "Shop Stock Qty (LM)", "Adj. MCH1"):
        df = _df([{"SKU": 1}]).drop(columns=[col])
        try:
            transform_overstock_p90d(df)
        except ValueError as e:
            assert col in str(e)
        else:
            raise AssertionError(f"缺少 {col} 時應拋出 ValueError")


# ── 2. MAKE UP 分類 ──────────────────────────────────────────────────────
def test_makeup_classification():
    df = _df([
        {"SKU": 1, "Adj. MCH1": "02 - MAKE UP"},
        {"SKU": 2, "Adj. MCH1": " 02 - MAKE UP "},
        {"SKU": 3, "Adj. MCH1": "01 - SKIN CARE"},
        {"SKU": 4, "Adj. MCH1": None},
    ])
    raw, _, _, _, _, _ = transform_overstock_p90d(df)
    by_sku = dict(zip(raw["SKU"].astype(str), raw["MAKE UP 分類"]))
    assert by_sku == {"1": "MAKE UP", "2": "MAKE UP", "3": "Non-MAKE UP", "4": "Non-MAKE UP"}


# ── 3. By Shop ───────────────────────────────────────────────────────────
def test_by_shop_counts_and_totals():
    df = _df([
        {"SKU": 1, "Site Code": "S01", "Adj. MCH1": "02 - MAKE UP", "Shop Stock Qty (CM)": 10.0},
        {"SKU": 1, "Site Code": "S01", "Adj. MCH1": "02 - MAKE UP", "Shop Stock Qty (CM)": 10.0},  # 同 SKU 重複列
        {"SKU": 2, "Site Code": "S01", "Adj. MCH1": "01 - SKIN CARE", "Shop Stock Qty (CM)": 5.0},
        {"SKU": 3, "Site Code": "S02", "Item_Shop": "Prod-S02-ShopKL",
         "OM": "OM2", "Shop Nature": "Tourist", "Adj. MCH1": "02 - MAKE UP", "Shop Stock Qty (CM)": 7.0},
    ])
    _, by_shop, _, _, _, _ = transform_overstock_p90d(df)
    s01 = by_shop[by_shop["Site Code"] == "S01"].iloc[0]
    assert int(s01["P90D 0 Sales SKU 總數"]) == 2  # SKU 1、2（已刪除重複）
    assert int(s01["MAKE UP SKU 數"]) == 1
    assert int(s01["Non-MAKE UP SKU 數"]) == 1
    assert int(s01["Shop Stock Qty (CM) 總數"]) == 25  # 未刪除重複加總
    total = by_shop[by_shop["Site Code"] == "總計"].iloc[0]
    assert int(total["P90D 0 Sales SKU 總數"]) == 3
    assert int(total["MAKE UP SKU 數"]) == 2
    assert int(total["Non-MAKE UP SKU 數"]) == 1
    sites = list(by_shop["Site Code"])
    assert sites[-1] == "總計"
    _no_banned(by_shop)


# ── 4. Top 10 ────────────────────────────────────────────────────────────
def test_top10_per_shop_limit_and_rank():
    rows = [{"SKU": i, "Shop Stock Qty (CM)": float(i)} for i in range(1, 13)]
    df = _df(rows)
    _, _, top10, _, _, _ = transform_overstock_p90d(df)
    assert len(top10) == 10
    assert list(top10["店內庫存排名"]) == list(range(1, 11))
    assert list(top10["Shop Stock Qty (CM)"]) == [12.0, 11.0, 10.0, 9.0, 8.0, 7.0, 6.0, 5.0, 4.0, 3.0]
    assert (top10["店舖回覆"] == "").all()
    assert "MAKE UP 分類" in top10.columns and "Adj. MCH1" in top10.columns
    _no_banned(top10)


def test_top10_columns_supply_lm_no_status_excess():
    df = _df([{"SKU": 1}])
    _, _, top10, _, _, _ = transform_overstock_p90d(df)
    cols = list(top10.columns)
    assert "Supply Source" in cols
    assert "Shop Stock Status" not in cols  # 全表皆為 P90D 0 Sales 狀態，無需顯示
    assert "Excess Stock Qty" not in cols
    assert "Shop Stock Qty (LM)" in cols
    assert cols.index("Shop Stock Qty (LM)") < cols.index("Shop Stock Qty (CM)")
    assert list(top10["Supply Source"]) == ["Warehouse"]


def test_top10_multi_shop():
    df = _df([
        {"SKU": 1, "Site Code": "S01", "Shop Stock Qty (CM)": 9.0},
        {"SKU": 2, "Site Code": "S02", "Item_Shop": "P-S02-KL", "Shop Stock Qty (CM)": 3.0},
    ])
    _, _, top10, _, _, _ = transform_overstock_p90d(df)
    assert len(top10) == 2
    assert set(top10["Site Code"]) == {"S01", "S02"}


def test_top10_excludes_makeup():
    rows = [{"SKU": i, "Shop Stock Qty (CM)": float(100 + i),
             "Adj. MCH1": "02 - MAKE UP"} for i in range(1, 6)]
    rows += [{"SKU": i, "Shop Stock Qty (CM)": float(i),
              "Adj. MCH1": "01 - SKIN CARE"} for i in range(6, 12)]
    df = _df(rows)
    _, _, top10, _, _, _ = transform_overstock_p90d(df)
    assert len(top10) == 6  # 只取 Non-MAKE UP（庫存較低的 6 列）
    assert (top10["MAKE UP 分類"] == "Non-MAKE UP").all()
    assert list(top10["Shop Stock Qty (CM)"]) == [11.0, 10.0, 9.0, 8.0, 7.0, 6.0]
    assert list(top10["店內庫存排名"]) == [1, 2, 3, 4, 5, 6]


def test_top10_all_makeup_returns_empty():
    df = _df([{"SKU": 1, "Adj. MCH1": "02 - MAKE UP"}])
    _, _, top10, _, _, _ = transform_overstock_p90d(df)
    assert top10.empty


def test_top10_include_makeup_option():
    rows = [{"SKU": i, "Shop Stock Qty (CM)": float(100 + i),
             "Adj. MCH1": "02 - MAKE UP"} for i in range(1, 6)]
    rows += [{"SKU": i, "Shop Stock Qty (CM)": float(i),
              "Adj. MCH1": "01 - SKIN CARE"} for i in range(6, 12)]
    df = _df(rows)
    _, _, top10, _, _, _ = transform_overstock_p90d(df, include_makeup=True)
    assert len(top10) == 10  # 勾選後一併計算 MAKE UP（庫存最高的 10 列）
    assert set(top10["MAKE UP 分類"]) == {"MAKE UP", "Non-MAKE UP"}
    assert list(top10["Shop Stock Qty (CM)"])[:5] == [105.0, 104.0, 103.0, 102.0, 101.0]
    assert list(top10["店內庫存排名"]) == list(range(1, 11))


# ── 5. Nature／OM 分佈 ───────────────────────────────────────────────────
def test_nature_and_om_distribution():
    df = _df([
        {"SKU": 1, "Site Code": "S01", "Shop Nature": "Tourist", "OM": "OM1",
         "Adj. MCH1": "02 - MAKE UP", "Shop Stock Qty (CM)": 6.0},
        {"SKU": 2, "Site Code": "S02", "Item_Shop": "P-S02-KL", "Shop Nature": "Mix", "OM": "OM1",
         "Shop Stock Qty (CM)": 4.0},
        {"SKU": 3, "Site Code": "S03", "Item_Shop": "P-S03-NT", "Shop Nature": "Non Tourist", "OM": "OM2",
         "Shop Stock Qty (CM)": 2.0},
    ])
    _, _, _, by_nature, by_om, _ = transform_overstock_p90d(df)
    assert set(by_nature[by_nature["Shop Nature"] != "總計"]["Shop Nature"]) == {"Tourist", "Mix", "Non Tourist"}
    n_total = by_nature[by_nature["Shop Nature"] == "總計"].iloc[0]
    assert int(n_total["P90D 0 Sales 行數"]) == 3
    assert int(n_total["MAKE UP SKU 數"]) == 1
    assert int(n_total["Non-MAKE UP SKU 數"]) == 2
    assert int(n_total["涉及店舖數"]) == 3
    assert set(by_om[by_om["OM"] != "總計"]["OM"]) == {"OM1", "OM2"}
    o_total = by_om[by_om["OM"] == "總計"].iloc[0]
    assert int(o_total["P90D 0 Sales 行數"]) == 3
    assert int(o_total["Shop Stock Qty (CM) 總數"]) == 12
    _no_banned(by_nature)
    _no_banned(by_om)


# ── 5b. 分類維度：Item Class／MCH1 ─────────────────────────────────────────
def test_breakdown_item_class_by_shop():
    df = _df([
        {"SKU": 1, "Site Code": "S01", "Item Class": "A1", "Shop Stock Qty (CM)": 10.0},
        {"SKU": 2, "Site Code": "S01", "Item Class": "B", "Shop Stock Qty (CM)": 5.0},
        {"SKU": 3, "Site Code": "S01", "Item Class": "B", "Shop Stock Qty (CM)": 4.0},
        {"SKU": 4, "Site Code": "S02", "Item_Shop": "P-S02-KL",
         "Item Class": "D", "Shop Stock Qty (CM)": 7.0},
    ])
    _, by_shop, _, by_nature, by_om, warnings = transform_overstock_p90d(df, breakdown="item_class")
    assert list(by_shop.columns) == [
        "Site Code", "Shop Name", "OM", "Shop Nature", "P90D 0 Sales SKU 總數",
        "A1 SKU 數", "B SKU 數", "D SKU 數", "Shop Stock Qty (CM) 總數",
    ]
    s01 = by_shop[by_shop["Site Code"] == "S01"].iloc[0]
    assert int(s01["P90D 0 Sales SKU 總數"]) == 3
    assert int(s01["A1 SKU 數"]) == 1
    assert int(s01["B SKU 數"]) == 2
    assert int(s01["D SKU 數"]) == 0
    total = by_shop[by_shop["Site Code"] == "總計"].iloc[0]
    assert int(total["P90D 0 Sales SKU 總數"]) == 4
    assert int(total["D SKU 數"]) == 1
    assert "MAKE UP SKU 數" not in by_shop.columns
    assert list(by_nature.columns)[4:-1] == ["A1 SKU 數", "B SKU 數", "D SKU 數"]
    assert list(by_om.columns)[4:-1] == ["A1 SKU 數", "B SKU 數", "D SKU 數"]
    assert not warnings  # 無空白 Item Class，不應有提示
    _no_banned(by_shop)
    _no_banned(by_nature)
    _no_banned(by_om)


def test_breakdown_item_class_blank_group_and_warning():
    df = _df([
        {"SKU": 1, "Item Class": "C"},
        {"SKU": 2, "Item Class": None},
    ])
    _, by_shop, _, _, _, warnings = transform_overstock_p90d(df, breakdown="item_class")
    assert list(by_shop.columns)[5:-1] == ["C SKU 數", "空白 SKU 數"]
    assert int(by_shop[by_shop["Site Code"] == "總計"]["空白 SKU 數"].iloc[0]) == 1
    assert any("空白" in w for w in warnings)


def test_breakdown_mch1_sorted_columns():
    df = _df([
        {"SKU": 1, "Adj. MCH1": "03 - FRAGRANCE"},
        {"SKU": 2, "Adj. MCH1": "01 - SKIN CARE"},
        {"SKU": 3, "Adj. MCH1": "02 - MAKE UP"},
    ])
    _, by_shop, _, by_nature, by_om, _ = transform_overstock_p90d(df, breakdown="mch1")
    assert list(by_shop.columns)[5:-1] == [
        "01 - SKIN CARE SKU 數", "02 - MAKE UP SKU 數", "03 - FRAGRANCE SKU 數",
    ]
    total = by_shop[by_shop["Site Code"] == "總計"].iloc[0]
    assert int(total["02 - MAKE UP SKU 數"]) == 1
    assert int(total["P90D 0 Sales SKU 總數"]) == 3
    n_total = by_nature[by_nature["Shop Nature"] == "總計"].iloc[0]
    assert int(n_total["01 - SKIN CARE SKU 數"]) == 1
    o_total = by_om[by_om["OM"] == "總計"].iloc[0]
    assert int(o_total["03 - FRAGRANCE SKU 數"]) == 1
    _no_banned(by_shop)


def test_breakdown_invalid_falls_back_to_makeup():
    df = _df([{"SKU": 1, "Adj. MCH1": "02 - MAKE UP"}])
    _, by_shop, _, by_nature, _, _ = transform_overstock_p90d(df, breakdown="unknown_dim")
    assert "MAKE UP SKU 數" in by_shop.columns
    assert "Non-MAKE UP SKU 數" in by_nature.columns


def test_breakdown_item_class_missing_column_raises():
    df = _df([{"SKU": 1}]).drop(columns=["Item Class"])
    try:
        transform_overstock_p90d(df, breakdown="item_class")
    except ValueError as e:
        assert "Item Class" in str(e)
    else:
        raise AssertionError("缺少 Item Class 時應拋出 ValueError")


# ── 5c. 分類維度：全部（三類合併顯示） ──────────────────────────────────────
def test_breakdown_all_by_shop_sections():
    df = _df([
        {"SKU": 1, "Site Code": "S01", "Item Class": "A1", "Adj. MCH1": "02 - MAKE UP", "Shop Stock Qty (CM)": 10.0},
        {"SKU": 2, "Site Code": "S01", "Item Class": "D", "Adj. MCH1": "01 - SKIN CARE", "Shop Stock Qty (CM)": 5.0},
        {"SKU": 3, "Site Code": "S02", "Item_Shop": "P-S02-KL", "Item Class": "B",
         "Adj. MCH1": "03 - FRAGRANCE", "Shop Stock Qty (CM)": 7.0},
    ])
    _, by_shop, _, by_nature, by_om, warnings = transform_overstock_p90d(df, breakdown="all")
    assert list(by_shop.columns) == [
        "Site Code", "Shop Name", "OM", "Shop Nature", "P90D 0 Sales SKU 總數",
        "MAKE UP SKU 數", "Non-MAKE UP SKU 數",
        "Item Class A1 SKU 數", "Item Class B SKU 數", "Item Class D SKU 數",
        "01 - SKIN CARE SKU 數", "02 - MAKE UP SKU 數", "03 - FRAGRANCE SKU 數",
        "Shop Stock Qty (CM) 總數",
    ]
    s01 = by_shop[by_shop["Site Code"] == "S01"].iloc[0]
    assert int(s01["P90D 0 Sales SKU 總數"]) == 2
    assert int(s01["MAKE UP SKU 數"]) == 1
    assert int(s01["Non-MAKE UP SKU 數"]) == 1
    assert int(s01["Item Class A1 SKU 數"]) == 1
    assert int(s01["Item Class B SKU 數"]) == 0
    assert int(s01["Item Class D SKU 數"]) == 1
    assert int(s01["02 - MAKE UP SKU 數"]) == 1
    assert int(s01["03 - FRAGRANCE SKU 數"]) == 0
    total = by_shop[by_shop["Site Code"] == "總計"].iloc[0]
    assert int(total["P90D 0 Sales SKU 總數"]) == 3
    assert int(total["Item Class B SKU 數"]) == 1
    assert not warnings
    assert list(by_nature.columns)[4:-1] == [
        "MAKE UP SKU 數", "Non-MAKE UP SKU 數",
        "Item Class A1 SKU 數", "Item Class B SKU 數", "Item Class D SKU 數",
        "01 - SKIN CARE SKU 數", "02 - MAKE UP SKU 數", "03 - FRAGRANCE SKU 數",
    ]
    assert list(by_om.columns)[4:-1] == list(by_nature.columns)[4:-1]
    _no_banned(by_shop)
    _no_banned(by_nature)
    _no_banned(by_om)


def test_breakdown_all_blank_groups_no_name_collision():
    df = _df([
        {"SKU": 1, "Item Class": None, "Adj. MCH1": None},
        {"SKU": 2, "Item Class": "C", "Adj. MCH1": "01 - SKIN CARE"},
    ])
    _, by_shop, _, _, _, warnings = transform_overstock_p90d(df, breakdown="all")
    cols = list(by_shop.columns)
    # Item Class 與 Adj. MCH1 的「空白」欄名不相撞
    assert "Item Class 空白 SKU 數" in cols
    assert "空白 SKU 數" in cols
    total = by_shop[by_shop["Site Code"] == "總計"].iloc[0]
    assert int(total["Item Class 空白 SKU 數"]) == 1
    assert int(total["空白 SKU 數"]) == 1
    assert any("Item Class" in w for w in warnings)
    assert any("Adj. MCH1" in w for w in warnings)
    _no_banned(by_shop)


def test_breakdown_all_missing_item_class_column_graceful():
    df = _df([{"SKU": 1}]).drop(columns=["Item Class"])
    _, by_shop, _, _, _, warnings = transform_overstock_p90d(df, breakdown="all")
    assert "Item Class 空白 SKU 數" in by_shop.columns  # 不抛錯，整欄歸入「空白」組
    total = by_shop[by_shop["Site Code"] == "總計"].iloc[0]
    assert int(total["Item Class 空白 SKU 數"]) == 1
    assert any("Item Class" in w for w in warnings)


def test_breakdown_all_excel_two_row_header():
    df = _df([
        {"SKU": 1, "Site Code": "S01", "Item Class": "A1", "Adj. MCH1": "02 - MAKE UP", "Shop Stock Qty (CM)": 10.0},
        {"SKU": 2, "Site Code": "S01", "Item Class": "D", "Adj. MCH1": "01 - SKIN CARE", "Shop Stock Qty (CM)": 5.0},
        {"SKU": 3, "Site Code": "S02", "Item_Shop": "P-S02-KL", "Item Class": "B",
         "Adj. MCH1": "03 - FRAGRANCE", "Shop Stock Qty (CM)": 7.0},
    ])
    _, by_shop, top10, by_nature, by_om, _ = transform_overstock_p90d(df, breakdown="all")
    buf = generate_excel_overstock_p90d(by_shop, top10, by_nature, by_om, None, breakdown="all")
    wb = load_workbook(BytesIO(buf.getvalue()), data_only=True)
    ws = wb["By Shop"]
    # 第一行：節標題與前後欄（上下合併）；第二行：分項表頭
    row1 = [c.value for c in ws[1]]
    row2 = [c.value for c in ws[2]]
    assert row1[0] == "Site Code"  # 前欄第一行顯示、上下合併
    assert "MAKE UP 分類" in row1 and "Item Class" in row1 and "Adj. MCH1" in row1
    assert row2[5:8] == ["MAKE UP SKU 數", "Non-MAKE UP SKU 數", "Item Class A1 SKU 數"]
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert {"A1:A2", "B1:B2", "C1:C2", "D1:D2", "E1:E2", "N1:N2"} <= merged  # 前後欄上下合併
    assert "F1:G1" in merged and "H1:J1" in merged and "K1:M1" in merged  # 三節水平合併
    assert ws.cell(row=3, column=1).value == "S01"  # 資料由第三行開始
    # Top 10 店舖回覆不受合併維度影響，維持單行表頭
    top_row1 = [c.value for c in wb["Top 10 店舖回覆"][1]]
    assert top_row1[0] == "Site Code" and "店舖回覆" in top_row1


def test_breakdown_all_excel_merges_no_overlap():
    """分佈表基礎欄（P90D 0 Sales SKU 數）不得誤併入分節；所有合併範圍不得相交（Excel 拒收重疊合併）。"""
    df = _df([
        {"SKU": 1, "Site Code": "S01", "Shop Nature": "Tourist", "OM": "OM1",
         "Item Class": "A1", "Adj. MCH1": "02 - MAKE UP", "Shop Stock Qty (CM)": 10.0},
        {"SKU": 2, "Site Code": "S02", "Item_Shop": "P-S02-KL", "Shop Nature": "Mix", "OM": "OM1",
         "Item Class": "D", "Adj. MCH1": "01 - SKIN CARE", "Shop Stock Qty (CM)": 5.0},
    ])
    _, by_shop, top10, by_nature, by_om, _ = transform_overstock_p90d(df, breakdown="all")
    buf = generate_excel_overstock_p90d(by_shop, top10, by_nature, by_om, None, breakdown="all")
    wb = load_workbook(BytesIO(buf.getvalue()), data_only=True)
    for name in ("By Shop", "By Nature", "By OM"):
        ws = wb[name]
        ranges = [str(r) for r in ws.merged_cells.ranges]
        cells = set()
        for r in ranges:
            for row in ws[r]:
                for c in row:
                    coord = c.coordinate
                    assert coord not in cells, f"{name} 合併範圍重疊：{r}"
                    cells.add(coord)
        # 分佈表的「P90D 0 Sales SKU 數」屬基礎欄：上下合併而非併入分節
        if name != "By Shop":
            assert "C1:C2" in ranges
            assert ws["C1"].value == "P90D 0 Sales SKU 數"
            assert ws["C2"].value is None


# ── 6. Excel 結構 ────────────────────────────────────────────────────────
def test_excel_four_sheets_no_banned_headers():
    df = _df([
        {"SKU": 1, "Adj. MCH1": "02 - MAKE UP"},
        {"SKU": 2},
    ])
    _, by_shop, top10, by_nature, by_om, _ = transform_overstock_p90d(df)
    buf = generate_excel_overstock_p90d(by_shop, top10, by_nature, by_om, None)
    wb = load_workbook(BytesIO(buf.getvalue()), data_only=True)
    assert wb.sheetnames == ["By Shop", "Top 10 店舖回覆", "By Nature", "By OM"]
    import re
    for ws in wb.worksheets:
        for cell in ws[1]:
            if cell.value:
                lowered = str(cell.value).lower()
                assert "value" not in lowered, cell.value
                assert "cogs" not in lowered, cell.value
                assert re.search(r"\bgp\b", lowered) is None, cell.value


def test_excel_empty_result_still_four_sheets():
    df = _df([{"SKU": 1, "Shop Stock Status": "Normal"}])
    _, by_shop, top10, by_nature, by_om, _ = transform_overstock_p90d(df)
    buf = generate_excel_overstock_p90d(by_shop, top10, by_nature, by_om, None)
    wb = load_workbook(BytesIO(buf.getvalue()), data_only=True)
    assert wb.sheetnames == ["By Shop", "Top 10 店舖回覆", "By Nature", "By OM"]


def test_excel_breakdown_dynamic_headers():
    df = _df([
        {"SKU": 1, "Item Class": "A1", "Adj. MCH1": "02 - MAKE UP"},
        {"SKU": 2, "Item Class": "D", "Adj. MCH1": "01 - SKIN CARE"},
    ])
    _, by_shop, top10, by_nature, by_om, _ = transform_overstock_p90d(df, breakdown="item_class")
    buf = generate_excel_overstock_p90d(by_shop, top10, by_nature, by_om, None)
    wb = load_workbook(BytesIO(buf.getvalue()), data_only=True)
    assert [c.value for c in wb["By Shop"][1]] == [
        "Site Code", "Shop Name", "OM", "Shop Nature", "P90D 0 Sales SKU 總數",
        "A1 SKU 數", "D SKU 數", "Shop Stock Qty (CM) 總數",
    ]
    assert [c.value for c in wb["By Nature"][1]][4:-1] == ["A1 SKU 數", "D SKU 數"]
