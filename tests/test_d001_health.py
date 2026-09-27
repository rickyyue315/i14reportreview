"""Validation tests for D001 健康檢查 Review (segment 4)."""
import warnings
warnings.filterwarnings("ignore")
import pandas as pd
from io import BytesIO
from openpyxl import load_workbook
from transformer import (
    transform_d001_health,
    generate_excel_d001_health,
)

BASE = {
    "Buyer": "B1", "Brand": "A", "SKU": 1, "SKU Desc": "Test",
    "Item_Shop": "Prod-S01-ShopHK", "Site Code": "S01", "OM": "OM1",
    "Shop Nature": "Mix", "Supply Source": "Warehouse",
    "Shop Stock Status": "Normal",
    "D001 Net Stock Qty (CM)": 10.0, "D001 On Order PO Qty (CM)": 0.0,
    "Shop On Order PO Qty (CM)": 0.0, "Shop Pending Receive": 0.0,
    "Shop Stock Qty (CM)": 5.0,
    "Sold Qty (P90D)": 30.0, "Shop Monthly Sold Qty (P90D)": 10.0,
}


def _row(**kw):
    r = dict(BASE)
    r.update(kw)
    return r


def _df(rows):
    return pd.DataFrame([_row(**r) for r in rows])


def test_shortage_overstock_low_cover():
    df = _df([
        {"SKU": 1, "D001 Net Stock Qty (CM)": 0.0},  # 缺貨（月銷 10）
        {"SKU": 2, "D001 Net Stock Qty (CM)": 20.0, "Sold Qty (P90D)": 0.0,
         "Shop Monthly Sold Qty (P90D)": 0.0},  # 積壓
        {"SKU": 3, "D001 Net Stock Qty (CM)": 5.0},  # 5 < 月銷 10 → 覆蓋偏低
        {"SKU": 4, "D001 Net Stock Qty (CM)": 50.0},  # 充足 → 不收錄
    ])
    check, by_buyer, raw, _ = transform_d001_health(df)
    got = dict(zip(check["SKU"].astype(str), check["檢查結論"]))
    assert got == {"1": "D001 缺貨需補", "2": "倉庫積壓", "3": "覆蓋偏低"}
    assert set(raw["SKU"].astype(str)) == {"1", "2", "3"}


def test_vendor_skus_excluded():
    df = _df([
        {"SKU": 1, "Supply Source": "Vendor", "D001 Net Stock Qty (CM)": 0.0},
    ])
    check, _, raw, _ = transform_d001_health(df)
    assert check.empty
    assert raw.empty


def test_missing_required_rejected():
    df = pd.DataFrame([{"Buyer": "B1", "Brand": "A", "SKU": 1, "Shop Stock Status": "Normal",
                        "Site Code": "S01", "Shop Stock Qty (CM)": 5.0, "Sold Qty (P90D)": 9.0}])
    try:
        transform_d001_health(df)
    except ValueError as e:
        assert "D001 Net Stock Qty (CM)" in str(e)
    else:
        raise AssertionError("應拒絕缺少 D001 庫存欄的輸入")


def test_pipeline_counts_as_cover():
    df = _df([
        {"SKU": 1, "D001 Net Stock Qty (CM)": 5.0, "D001 On Order PO Qty (CM)": 8.0},  # 5+8 ≥ 10 → 不收錄
    ])
    check, _, _, _ = transform_d001_health(df)
    assert check.empty


def test_by_buyer_totals():
    df = _df([
        {"SKU": 1, "D001 Net Stock Qty (CM)": 0.0},
        {"SKU": 2, "Buyer": "B2", "D001 Net Stock Qty (CM)": 0.0},
    ])
    _, by_buyer, _, _ = transform_d001_health(df)
    total = by_buyer[by_buyer["Buyer"] == "總計"].iloc[0]
    assert int(total["需跟進 SKU 數"]) == 2
    assert int(total["D001 缺貨數"]) == 2


def test_excel_sheets_and_skip_raw():
    df = _df([
        {"SKU": 1, "D001 Net Stock Qty (CM)": 0.0},
    ])
    check, by_buyer, raw, _ = transform_d001_health(df)
    buf = generate_excel_d001_health(check, by_buyer, raw, include_raw=True)
    wb = load_workbook(BytesIO(buf.getvalue()), data_only=True)
    assert wb.sheetnames == ["D001 檢查", "By Buyer", "RAW"]
    buf2 = generate_excel_d001_health(check, by_buyer, raw, include_raw=False)
    wb2 = load_workbook(BytesIO(buf2.getvalue()), data_only=True)
    assert wb2.sheetnames == ["D001 檢查", "By Buyer", "RAW"]
    assert wb2["RAW"]["A2"].value is not None


def test_empty_result_warning():
    df = _df([{"SKU": 1}])
    check, _, _, warnings = transform_d001_health(df)
    assert check.empty
    assert any("沒有符合" in w for w in warnings)
