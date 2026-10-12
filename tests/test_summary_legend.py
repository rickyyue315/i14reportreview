"""Summary 底色圖例（第 1–2 列）測試 — 色塊僅列出本次工作表實際會啟用的顏色。"""
import pandas as pd
from openpyxl import load_workbook

from transformer import (
    transform_stock_checking,
    generate_excel_stock_checking,
    transform_hybrid,
    generate_excel_hybrid,
)


def _df(rows):
    base = {
        "Buyer": "B1", "Brand": "BR", "SKU": 999, "SKU Desc": "D",
        "Adj. SST": "S", "Adj. Safety Stock": 3, "Supply Source": "W", "Site Code": "X",
        "Shop Nature": "Tourist", "Sold Qty (P90D)": 90.0,
        "Shop Stock Qty (CM)": 10.0, "Shop Monthly Sold Qty (P90D)": 30.0,
        "Turnday (P90D Qty)": 10.0, "D001 Net Stock Qty (CM)": 200.0,
        "Item_Shop": "x",
    }
    return pd.DataFrame([{**base, **r} for r in rows])


def _rows():
    return [
        {"SKU": 999, "Site Code": "T01", "Shop Nature": "Tourist",
         "Shop Stock Qty (CM)": 5.0, "Shop Monthly Sold Qty (P90D)": 100.0,
         "Sold Qty (P90D)": 300.0},
        {"SKU": 999, "Site Code": "M01", "Shop Nature": "Mix",
         "Shop Stock Qty (CM)": 50.0, "Shop Monthly Sold Qty (P90D)": 10.0,
         "Sold Qty (P90D)": 30.0},
        {"SKU": 999, "Site Code": "N01", "Shop Nature": "Non Tourist",
         "Shop Stock Qty (CM)": 50.0, "Shop Monthly Sold Qty (P90D)": 10.0,
         "Sold Qty (P90D)": 30.0},
    ]


def test_stock_checking_legend_shows_skew_and_shortage_only():
    raw, summary, stats, _ = transform_stock_checking(_df(_rows()))
    ws = load_workbook(generate_excel_stock_checking(raw, summary, stats))["Summary"]
    assert ws["A1"].value == "存貨偏重"
    assert ws["B1"].value == "存貨偏低"
    assert ws["C1"].value is None
    assert str(ws["A1"].fill.start_color.rgb).endswith("FFC000")
    assert str(ws["B1"].fill.start_color.rgb).endswith("BDD7EE")
    assert ws["A2"].value.startswith("顏色說明：")
    assert "橘底＝存貨偏重" in ws["A2"].value and "淺藍底＝存貨偏低" in ws["A2"].value
    assert "紅底" not in ws["A2"].value
    assert ws.freeze_panes == "A6"
    assert ws["A5"].value == "Brand"


def test_hybrid_legend_zero_chip_only_for_sales_dimensions():
    raw, sc, dg, dg_om, ds, stats, _ = transform_hybrid(_df(_rows()), dimension="nature")
    ws = load_workbook(generate_excel_hybrid(raw, sc, dg, dg_om, ds, stats))["Summary"]
    assert ws["A1"].value == "存貨偏重"
    assert ws["B1"].value == "存貨偏低"
    assert ws["C1"].value is None

    raw, sc, dg, dg_om, ds, stats, _ = transform_hybrid(_df(_rows()), dimension="sales")
    ws = load_workbook(generate_excel_hybrid(raw, sc, dg, dg_om, ds, stats))["Summary"]
    assert ws["C1"].value == "零銷仍有庫存"
    assert str(ws["C1"].fill.start_color.rgb).endswith("FFC7CE")
    assert "紅底＝零銷組仍有庫存" in ws["A2"].value
