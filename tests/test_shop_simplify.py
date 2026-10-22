"""店舖精簡SKU模式驗證測試。

核心契約：
- 過去90日有無賣過：Sold Qty (P90D) > 0 →「有」，其餘（0／負數／空白）→「無賣」；
- 過去60日有無存貨：Shop Stock Qty (LM) 或 (CM) > 0 →「有存貨」，兩者同時為 0 →「無存貨」；
- SCT建議、Shop 回覆新的NDRF、如果RF的Safety Qty 一律留空；
- 不篩選列，保留來源欄序，五個輸出欄固定排最後；輸出 Summary＋Export 兩個工作表，
  並依 Brand／SKU（文字）排序、套用範本欄寬／隱藏欄／數字格式與顏色標示。
"""
import warnings
warnings.filterwarnings("ignore")
import importlib
from io import BytesIO

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from modes import MODE_REGISTRY
from transformer import (
    transform_shop_simplify,
    generate_excel_shop_simplify,
    shop_simplify_stats,
    _SHOP_SIMPLIFY_SOLD_COLUMN,
    _SHOP_SIMPLIFY_STOCK_COLUMN,
    _SHOP_SIMPLIFY_SCT_COLUMN,
    _SHOP_SIMPLIFY_NDRF_COLUMN,
    _SHOP_SIMPLIFY_SAFETY_COLUMN,
    _SHOP_SIMPLIFY_REPLY_COLUMNS,
    _SHOP_SIMPLIFY_SOLD_YES,
    _SHOP_SIMPLIFY_SOLD_NO,
    _SHOP_SIMPLIFY_STOCK_YES,
    _SHOP_SIMPLIFY_STOCK_NO,
)

BASE = {
    "Buyer": "127 - Grace Wong", "Brand": "SUI - SUISSE PROGRAMME", "SKU": 1103775,
    "Item_Shop": "(SUI) 1103775-CAV PRE LIFT ESSENCE WATERHA20-新香港仔",
    "Valuation Class": "Normal Trade Item", "Desc.": "CAV PRE LIFT ESSENCE WATER",
    "Shop Stock Status": "Normal", "Site Code": "HA20", "OM": "OM1",
    "Shop Nature": "Tourist", "RP Type": "RF", "Adj. Safety Stock": 6,
    "Sold Qty (P90D)": 40, "Shop Stock Qty (LM)": 28, "Shop Stock Qty (CM)": 18,
}


def _row(**kw):
    r = dict(BASE)
    r.update(kw)
    return r


def _df(rows):
    return pd.DataFrame([_row(**r) for r in rows])


def _status(df, column):
    return list(df[column])


# ──────────────────────────────────────────────────────────────────────────────
# 1. 註冊與基本行為
# ──────────────────────────────────────────────────────────────────────────────
def test_registry_entry_resolves_to_functions():
    entry = MODE_REGISTRY["店舖精簡SKU模式"]
    mod = importlib.import_module(entry["module"])
    assert callable(getattr(mod, entry["transform"]))
    assert callable(getattr(mod, entry["generate_excel"]))


def test_output_columns_contract_names():
    assert _SHOP_SIMPLIFY_SOLD_COLUMN == "過去90日有無賣過"
    assert _SHOP_SIMPLIFY_STOCK_COLUMN == "過去60日有無存貨"
    assert _SHOP_SIMPLIFY_SCT_COLUMN == "SCT建議"
    assert _SHOP_SIMPLIFY_NDRF_COLUMN == "Shop 回覆新的NDRF"
    assert _SHOP_SIMPLIFY_SAFETY_COLUMN == "如果RF的Safety Qty"
    assert _SHOP_SIMPLIFY_REPLY_COLUMNS == [
        "SCT建議", "Shop 回覆新的NDRF", "如果RF的Safety Qty",
    ]


def test_input_dataframe_not_mutated():
    df = _df([{"SKU": 1}, {"SKU": 2}])
    before = df.copy()
    transform_shop_simplify(df)
    pd.testing.assert_frame_equal(df, before)


# ──────────────────────────────────────────────────────────────────────────────
# 2. 過去90日有無賣過
# ──────────────────────────────────────────────────────────────────────────────
def test_sold_positive_is_yes():
    raw, _ = transform_shop_simplify(_df([{"Sold Qty (P90D)": 1}, {"Sold Qty (P90D)": 293.0}]))
    assert _status(raw, _SHOP_SIMPLIFY_SOLD_COLUMN) == [_SHOP_SIMPLIFY_SOLD_YES] * 2


def test_sold_zero_negative_and_blank_is_no_sale():
    raw, _ = transform_shop_simplify(_df([
        {"Sold Qty (P90D)": 0},
        {"Sold Qty (P90D)": -3},
        {"Sold Qty (P90D)": None},
        {"Sold Qty (P90D)": ""},
    ]))
    assert _status(raw, _SHOP_SIMPLIFY_SOLD_COLUMN) == [_SHOP_SIMPLIFY_SOLD_NO] * 4


# ──────────────────────────────────────────────────────────────────────────────
# 3. 過去60日有無存貨
# ──────────────────────────────────────────────────────────────────────────────
def test_stock_any_positive_is_in_stock():
    raw, _ = transform_shop_simplify(_df([
        {"Shop Stock Qty (LM)": 2, "Shop Stock Qty (CM)": 0},
        {"Shop Stock Qty (LM)": 0, "Shop Stock Qty (CM)": 1},
        {"Shop Stock Qty (LM)": 7, "Shop Stock Qty (CM)": 9},
    ]))
    assert _status(raw, _SHOP_SIMPLIFY_STOCK_COLUMN) == [_SHOP_SIMPLIFY_STOCK_YES] * 3


def test_stock_both_zero_is_no_stock():
    raw, _ = transform_shop_simplify(_df([
        {"Shop Stock Qty (LM)": 0, "Shop Stock Qty (CM)": 0},
        {"Shop Stock Qty (LM)": None, "Shop Stock Qty (CM)": None},
        {"Shop Stock Qty (LM)": -2, "Shop Stock Qty (CM)": 0},
    ]))
    assert _status(raw, _SHOP_SIMPLIFY_STOCK_COLUMN) == [_SHOP_SIMPLIFY_STOCK_NO] * 3


# ──────────────────────────────────────────────────────────────────────────────
# 4. 三個回覆欄一律留空
# ──────────────────────────────────────────────────────────────────────────────
def test_reply_columns_are_blank():
    raw, _ = transform_shop_simplify(_df([{}, {}]))
    for col in _SHOP_SIMPLIFY_REPLY_COLUMNS:
        assert col in raw.columns
        assert list(raw[col]) == ["", ""]


def test_existing_reply_values_are_replaced_by_blank():
    df = _df([{
        "SCT建議": "需要keep",
        "Shop 回覆新的NDRF": "ND",
        "如果RF的Safety Qty": 99,
    }])
    raw, w = transform_shop_simplify(df)
    for col in _SHOP_SIMPLIFY_REPLY_COLUMNS:
        assert list(raw[col]) == [""]
    assert any("一律留空" in x for x in w)


def test_existing_output_columns_are_replaced_and_moved_to_end():
    df = _df([{
        _SHOP_SIMPLIFY_SOLD_COLUMN: "無賣",
        _SHOP_SIMPLIFY_STOCK_COLUMN: "無存貨",
        "SCT建議": "自行決定",
    }])
    raw, _ = transform_shop_simplify(df)
    # 五個輸出欄各只出現一次，且固定排最後
    assert list(raw.columns)[-5:] == [
        _SHOP_SIMPLIFY_SOLD_COLUMN,
        _SHOP_SIMPLIFY_STOCK_COLUMN,
        *_SHOP_SIMPLIFY_REPLY_COLUMNS,
    ]
    assert sum(c == _SHOP_SIMPLIFY_SOLD_COLUMN for c in raw.columns) == 1
    # 舊值被本次計算覆蓋
    assert raw[_SHOP_SIMPLIFY_SOLD_COLUMN].iloc[0] == _SHOP_SIMPLIFY_SOLD_YES
    assert raw[_SHOP_SIMPLIFY_STOCK_COLUMN].iloc[0] == _SHOP_SIMPLIFY_STOCK_YES


# ──────────────────────────────────────────────────────────────────────────────
# 5. 欄序、清理與篩選
# ──────────────────────────────────────────────────────────────────────────────
def test_source_column_order_preserved():
    df = _df([{"SKU": 1}])
    raw, _ = transform_shop_simplify(df)
    cols = list(raw.columns)
    # 解析欄插入慣用位置：SKU Desc 於 Brand 後、Shop Name 於 Site Code 後
    assert cols[:4] == ["Buyer", "Brand", "SKU Desc", "SKU"]
    assert cols[cols.index("Site Code") + 1] == "Shop Name"
    # 其餘來源欄位相對順序不變，五個輸出欄固定排最後
    source_cols = [c for c in cols if c in df.columns and c not in ("SKU Desc", "Shop Name")]
    assert source_cols == list(df.columns)
    assert cols[-5:] == [
        _SHOP_SIMPLIFY_SOLD_COLUMN,
        _SHOP_SIMPLIFY_STOCK_COLUMN,
        *_SHOP_SIMPLIFY_REPLY_COLUMNS,
    ]


def test_value_columns_removed_but_valuation_class_kept():
    df = _df([{
        "Shop Stock Value (LM)": 100.0,
        "Shop Stock Value (CM)": 50.0,
        "Excess Stock Value": 0.0,
        "Valuation Class": "Normal Trade Item",
    }])
    raw, _ = transform_shop_simplify(df)
    for col in ("Shop Stock Value (LM)", "Shop Stock Value (CM)", "Excess Stock Value"):
        assert col not in raw.columns
    # 大小寫敏感：Valuation Class 不含「Value」，必須保留
    assert "Valuation Class" in raw.columns
    assert "SKU Desc" in raw.columns and "Shop Name" in raw.columns
    assert list(raw.columns)[-5] == _SHOP_SIMPLIFY_SOLD_COLUMN

    ws = load_workbook(BytesIO(generate_excel_shop_simplify(raw).getvalue()))["Export"]
    headers = [c.value for c in ws[1]]
    assert not any(h and "Value" in str(h) for h in headers)
    assert "Valuation Class" in headers


def test_item_shop_parses_desc_and_shop_name():
    df = _df([
        {"SKU": 111186903002.0, "Site Code": "HB97",
         "Item_Shop": "(JDU) 111186903002-HAIR TREAMENT OIL, 01, 40PCHB97-新旺角文華"},
        {"SKU": 111188002003.0, "Site Code": "HB29",
         "Item_Shop": "(JDU) 111188002003-SILKY HAIR OIL, 02, 6PCHB29-重慶站"},
    ])
    raw, _ = transform_shop_simplify(df)
    by_site = raw.set_index("Site Code")
    assert by_site.loc["HB97", "Shop Name"] == "新旺角文華"
    assert by_site.loc["HB97", "SKU Desc"] == "HAIR TREAMENT OIL, 01, 40PC"
    assert by_site.loc["HB29", "Shop Name"] == "重慶站"
    assert by_site.loc["HB29", "SKU Desc"] == "SILKY HAIR OIL, 02, 6PC"


def test_missing_item_shop_warns_and_leaves_derived_columns_blank():
    df = _df([{"SKU": 1}]).drop(columns=["Item_Shop"])
    raw, w = transform_shop_simplify(df)
    assert raw["SKU Desc"].iloc[0] == ""
    assert raw["Shop Name"].iloc[0] == ""
    assert any("Item_Shop" in x for x in w)


def test_sku_normalized_to_text():
    raw, _ = transform_shop_simplify(_df([{"SKU": 1103775.0}]))
    assert raw["SKU"].iloc[0] == "1103775"


def test_footer_and_marketing_rows_removed():
    df = _df([
        {"SKU": 1},
        {"SKU": 2, "Buyer": "Total"},
        {"SKU": 3, "Buyer": "117 - Marketing"},
        {"SKU": 4, "Buyer": "Applied filters: x"},
    ])
    raw, _ = transform_shop_simplify(df)
    assert list(raw["SKU"]) == ["1"]


def test_shop_nature_filter_only_when_selected():
    df = _df([
        {"SKU": 1, "Shop Nature": "Tourist"},
        {"SKU": 2, "Shop Nature": "Mix", "Site Code": "HB29"},
        {"SKU": 3, "Shop Nature": "Non Tourist", "Site Code": "HB50"},
    ])
    raw_all, _ = transform_shop_simplify(df)
    assert len(raw_all) == 3
    raw_mix, _ = transform_shop_simplify(df, shop_natures=["Mix"])
    assert list(raw_mix["SKU"]) == ["2"]


# ──────────────────────────────────────────────────────────────────────────────
# 6. 必要欄位與警告
# ──────────────────────────────────────────────────────────────────────────────
def test_missing_sold_column_raises():
    df = _df([{"SKU": 1}]).drop(columns=["Sold Qty (P90D)"])
    with pytest.raises(ValueError, match="Sold Qty \\(P90D\\)"):
        transform_shop_simplify(df)


def test_missing_both_stock_columns_raises():
    df = _df([{"SKU": 1}]).drop(columns=["Shop Stock Qty (LM)", "Shop Stock Qty (CM)"])
    with pytest.raises(ValueError, match="Shop Stock Qty"):
        transform_shop_simplify(df)


def test_missing_lm_warns_and_uses_cm():
    df = _df([{"SKU": 1, "Shop Stock Qty (CM)": 3}]).drop(columns=["Shop Stock Qty (LM)"])
    raw, w = transform_shop_simplify(df)
    assert raw[_SHOP_SIMPLIFY_STOCK_COLUMN].iloc[0] == _SHOP_SIMPLIFY_STOCK_YES
    assert any("Shop Stock Qty (LM)" in x for x in w)


def test_missing_cm_warns_and_uses_lm():
    df = _df([{"SKU": 1, "Shop Stock Qty (LM)": 0}]).drop(columns=["Shop Stock Qty (CM)"])
    raw, w = transform_shop_simplify(df)
    assert raw[_SHOP_SIMPLIFY_STOCK_COLUMN].iloc[0] == _SHOP_SIMPLIFY_STOCK_NO
    assert any("Shop Stock Qty (CM)" in x for x in w)


def test_empty_after_filter_warns():
    df = _df([{"SKU": 1, "Shop Nature": "Tourist"}])
    raw, w = transform_shop_simplify(df, shop_natures=["Mix"])
    assert raw.empty
    assert any("沒有可輸出" in x for x in w)


# ──────────────────────────────────────────────────────────────────────────────
# 7. 統計與 Excel 輸出
# ──────────────────────────────────────────────────────────────────────────────
def test_stats_counts_sold_stock_both():
    df = _df([
        {"SKU": 1, "Sold Qty (P90D)": 5, "Shop Stock Qty (LM)": 1, "Shop Stock Qty (CM)": 1},
        {"SKU": 2, "Sold Qty (P90D)": 0, "Shop Stock Qty (LM)": 1, "Shop Stock Qty (CM)": 0},
        {"SKU": 3, "Sold Qty (P90D)": 2, "Shop Stock Qty (LM)": 0, "Shop Stock Qty (CM)": 0},
        {"SKU": 4, "Sold Qty (P90D)": 0, "Shop Stock Qty (LM)": 0, "Shop Stock Qty (CM)": 0},
    ])
    raw, _ = transform_shop_simplify(df)
    stats = shop_simplify_stats(raw)
    assert stats["rows"] == 4
    assert stats["sold_no"] == 2
    assert stats["stock_no"] == 2
    assert stats["both_no"] == 1
    assert stats["sites"] == 1
    assert stats["skus"] == 4


def test_excel_has_summary_and_export_sheets():
    df = _df([
        {"SKU": 1, "Adj. MCH1": "01 - SKIN CARE", "Item Status": "Active"},
        {"SKU": 2, "Adj. MCH1": "01 - SKIN CARE", "Item Status": "Discontinued"},
        {"SKU": 3, "Adj. MCH1": "02 - MAKE UP", "Item Status": "Active"},
    ])
    raw, _ = transform_shop_simplify(df)
    buf = generate_excel_shop_simplify(raw)
    wb = load_workbook(BytesIO(buf.getvalue()))
    assert wb.sheetnames == ["Summary", "Export"]
    assert wb.active.title == "Export"

    ws = wb["Export"]
    headers = [c.value for c in ws[1]]
    assert headers == list(raw.columns)
    assert "SKU Desc" in headers and "Shop Name" in headers
    assert headers[-5:] == [
        "過去90日有無賣過", "過去60日有無存貨",
        "SCT建議", "Shop 回覆新的NDRF", "如果RF的Safety Qty",
    ]
    assert ws.max_row == 1 + len(raw)
    # 來源表頭不改名（Adj. Safety Stock 不轉成 Safety Stock (現時)）
    assert "Adj. Safety Stock" in headers

    summary = wb["Summary"]
    assert summary["A3"].value == "計數 - SKU"
    assert summary["B3"].value == "欄標籤"
    assert summary["A4"].value == "列標籤"
    assert [summary.cell(row=4, column=c).value for c in range(2, 5)] == [
        "Active", "Discontinued", "總計",
    ]
    assert summary["A5"].value == "01 - SKIN CARE"
    assert [summary.cell(row=5, column=c).value for c in range(2, 5)] == [1, 1, 2]
    assert summary["A6"].value == "02 - MAKE UP"
    assert [summary.cell(row=6, column=c).value for c in range(2, 5)] == [1, None, 1]
    assert summary["A7"].value == "總計"
    assert [summary.cell(row=7, column=c).value for c in range(2, 5)] == [2, 1, 3]


def test_summary_blank_groups_sorted_last():
    df = _df([
        {"SKU": 1, "Adj. MCH1": None, "Item Status": "Active"},
        {"SKU": 2, "Adj. MCH1": "01 - SKIN CARE", "Item Status": ""},
        {"SKU": 3, "Adj. MCH1": "01 - SKIN CARE", "Item Status": "Active"},
    ])
    raw, _ = transform_shop_simplify(df)
    summary = load_workbook(BytesIO(generate_excel_shop_simplify(raw).getvalue()))["Summary"]
    assert summary["A5"].value == "01 - SKIN CARE"
    assert [summary.cell(row=5, column=c).value for c in range(2, 5)] == [1, 1, 2]
    assert summary["A6"].value == "(空白)"
    assert [summary.cell(row=6, column=c).value for c in range(2, 5)] == [1, None, 1]
    assert summary["A7"].value == "總計"
    assert [summary.cell(row=7, column=c).value for c in range(2, 5)] == [2, 1, 3]


def test_output_sorted_by_brand_then_sku_text():
    df = _df([
        {"Brand": "B - B", "SKU": 200},
        {"Brand": "A - A", "SKU": 1000},
        {"Brand": "A - A", "SKU": 999},
    ])
    raw, _ = transform_shop_simplify(df)
    # SKU 以文字排序："1000" < "999"
    assert list(zip(raw["Brand"], raw["SKU"])) == [
        ("A - A", "1000"), ("A - A", "999"), ("B - B", "200"),
    ]


def test_launch_date_converted_to_int():
    raw, _ = transform_shop_simplify(_df([
        {"Launch Date": 20120520.0},
        {"Launch Date": "2020-08-28"},
        {"Launch Date": None},
    ]))
    values = list(raw["Launch Date"])
    assert values[0] == 20120520 and isinstance(values[0], int)
    assert values[1] == 20200828 and isinstance(values[1], int)
    assert pd.isna(values[2])


def test_infinite_turnday_becomes_blank_in_excel():
    raw, _ = transform_shop_simplify(_df([
        {"Turnday (P90D Qty)": np.inf},
        {"Turnday (P90D Qty)": 12.5},
    ]))
    assert pd.isna(raw["Turnday (P90D Qty)"].iloc[0])
    ws = load_workbook(BytesIO(generate_excel_shop_simplify(raw).getvalue()))["Export"]
    headers = [c.value for c in ws[1]]
    col = headers.index("Turnday (P90D Qty)") + 1
    assert ws.cell(row=2, column=col).value is None


def test_numeric_values_keep_full_precision():
    value = 7.455555555555556
    raw, _ = transform_shop_simplify(_df([{"Shop Monthly Sold Qty (P90D)": value}]))
    ws = load_workbook(BytesIO(generate_excel_shop_simplify(raw).getvalue()))["Export"]
    headers = [c.value for c in ws[1]]
    written = ws.cell(row=2, column=headers.index("Shop Monthly Sold Qty (P90D)") + 1).value
    assert abs(float(written) - value) < 1e-12


def test_excel_cell_values_and_blank_reply_cells():
    df = _df([
        {"SKU": 1103775, "Sold Qty (P90D)": 40, "Shop Stock Qty (LM)": 0, "Shop Stock Qty (CM)": 0},
    ])
    raw, _ = transform_shop_simplify(df)
    buf = generate_excel_shop_simplify(raw)
    ws = load_workbook(BytesIO(buf.getvalue()))["Export"]
    headers = [c.value for c in ws[1]]
    row = [c.value for c in ws[2]]
    assert row[headers.index(_SHOP_SIMPLIFY_SOLD_COLUMN)] == "有"
    assert row[headers.index(_SHOP_SIMPLIFY_STOCK_COLUMN)] == "無存貨"
    for col in _SHOP_SIMPLIFY_REPLY_COLUMNS:
        assert row[headers.index(col)] is None


def test_excel_format_matches_template():
    df = pd.DataFrame([{
        "Item_Shop": "(SUI) 1103775-CAV PRE LIFT ESSENCE WATERHA20-新香港仔",
        "SKU": 1103775, "Valuation Class": "Normal Trade Item",
        "Brand": "SUI - SUISSE PROGRAMME", "Buyer": "127 - Grace Wong",
        "Shop Stock Status": "Normal", "Site Code": "HA20", "Shop Nature": "Tourist",
        "RP Type": "RF", "Adj. Safety Stock": 6,
        "Adj. MCH1": "01 - SKIN CARE", "Item Status": "Active",
        "Shop Monthly Sold Qty (P90D)": 7.455555555555556,
        "Safety VS Monthly Sold Qty": 0.94,
        "Sold Qty (P90D)": 40, "Shop Stock Qty (LM)": 28, "Shop Stock Qty (CM)": 18,
    }, {
        "Item_Shop": "(SUI) 1106553-THE RICH SERUM 50MLHA20-新香港仔",
        "SKU": 1106553, "Valuation Class": "Normal Trade Item",
        "Brand": "SUI - SUISSE PROGRAMME", "Buyer": "127 - Grace Wong",
        "Shop Stock Status": "Normal", "Site Code": "HA20", "Shop Nature": "Tourist",
        "RP Type": "RF", "Adj. Safety Stock": 6,
        "Adj. MCH1": "01 - SKIN CARE", "Item Status": "Active",
        "Shop Monthly Sold Qty (P90D)": 1.0, "Safety VS Monthly Sold Qty": 0.0,
        "Sold Qty (P90D)": 0, "Shop Stock Qty (LM)": 0, "Shop Stock Qty (CM)": 0,
    }])
    raw, _ = transform_shop_simplify(df)
    ws = load_workbook(BytesIO(generate_excel_shop_simplify(raw).getvalue()))["Export"]
    headers = [c.value for c in ws[1]]
    col = {name: i + 1 for i, name in enumerate(headers)}

    # 表頭：列高、換行、凍結 F2、縮放 85%
    assert ws.row_dimensions[1].height == 70.5
    assert ws.freeze_panes == "F2"
    assert ws.sheet_view.zoomScale == 85
    header = ws.cell(row=1, column=col["Item_Shop"])
    assert header.alignment.wrap_text is True
    assert header.alignment.horizontal == "center"
    assert header.alignment.vertical == "top"
    assert header.fill.fgColor.rgb[-6:] == "D9E2F3"
    assert ws.cell(row=1, column=col["RP Type"]).fill.fgColor.rgb[-6:] == "CCFF99"
    assert ws.cell(row=1, column=col["Sold Qty (P90D)"]).fill.fgColor.rgb[-6:] == "E4DFEC"
    assert ws.cell(row=1, column=col["Shop 回覆新的NDRF"]).fill.fgColor.rgb[-6:] == "CCFF99"

    # 資料格重點色
    assert ws.cell(row=2, column=col["RP Type"]).fill.fgColor.rgb[-6:] == "CCFF99"
    assert ws.cell(row=2, column=col["Shop Stock Qty (CM)"]).fill.fgColor.rgb[-6:] == "CCFF99"
    assert ws.cell(row=2, column=col["Sold Qty (P90D)"]).fill.fgColor.rgb[-6:] == "E4DFEC"
    # 第二列 Sold Qty = 0 → 無賣 → 淡紅
    sold_cell = ws.cell(row=3, column=col[_SHOP_SIMPLIFY_SOLD_COLUMN])
    assert sold_cell.value == "無賣"
    assert sold_cell.fill.fgColor.rgb[-6:] == "F2DCDB"

    # 欄寬與隱藏欄
    assert ws.column_dimensions["A"].hidden is True
    assert ws.column_dimensions["C"].hidden is True
    assert round(ws.column_dimensions["A"].width, 2) == 79.71
    assert round(ws.column_dimensions["E"].width, 2) == 46.43

    # 數字格式
    assert ws.cell(row=2, column=col["Shop Monthly Sold Qty (P90D)"]).number_format.startswith("_(* #,##0")
    assert ws.cell(row=2, column=col["Safety VS Monthly Sold Qty"]).number_format == "0%"


def test_excel_freeze_panes_after_sku_desc():
    df = _df([{"SKU": 1}])
    raw, _ = transform_shop_simplify(df)
    buf = generate_excel_shop_simplify(raw)
    ws = load_workbook(BytesIO(buf.getvalue()))["Export"]
    from openpyxl.utils import get_column_letter
    headers = [c.value for c in ws[1]]
    expected = f"{get_column_letter(headers.index('SKU Desc') + 2)}2"
    assert ws.freeze_panes == expected
    assert ws.auto_filter.ref == ws.dimensions


def test_empty_dataframe_still_writes_header():
    raw = pd.DataFrame(columns=list(BASE) + [
        _SHOP_SIMPLIFY_SOLD_COLUMN, _SHOP_SIMPLIFY_STOCK_COLUMN, *_SHOP_SIMPLIFY_REPLY_COLUMNS,
    ])
    buf = generate_excel_shop_simplify(raw)
    wb = load_workbook(BytesIO(buf.getvalue()))
    ws = wb["Export"]
    assert ws.max_row == 1
    assert [c.value for c in ws[1]][-5] == _SHOP_SIMPLIFY_SOLD_COLUMN
