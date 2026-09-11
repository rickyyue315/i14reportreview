"""Validation tests for Safety Stock x NetStock Review ND 可調出規劃（v1.43.0）。

- 新欄「ND 建議調出量（參考）」：RP Type＝ND 店舖×SKU 列，只扣 14 日周轉所需
  （＝max(0, Shop Stock Qty (CM) − Sold Qty (P90D) ÷ 90 × 14)，零銷有存＝全數可調出），
  不設 3 件保底、不套池級門檻；RF 列留空
- 只作規劃參考：不影響調貨分配、Safety Review 與建議調整量，亦不加入兩張 Review 主表
- Excel 新增第 7 張工作表「ND 可調出規劃」（共 10 工作表），RAW 新增同欄
- 缺「RP Type」欄（全列按 RF 處理）→ 欄全空、工作表「無符合資料」
"""
import warnings
warnings.filterwarnings("ignore")
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from transformer import (
    transform_hybrid,
    generate_excel_hybrid,
    _build_hybrid_nd_clearance,
    _hybrid_raw_display_columns,
    _HYBRID_DIAG_COLUMNS,
    _HYBRID_ND_CLEARANCE_COLUMN,
    _HYBRID_ND_CLEARANCE_COLUMNS,
)

BASE = {
    "Buyer": "B1", "Brand": "A", "SKU": 1, "SKU Desc": "Test Product",
    "Item_Shop": "Prod-S01-ShopHK", "Site Code": "S01", "OM": "OM1",
    "Shop Nature": "Tourist", "Vendor MC": "V1", "Item Status": "A",
    "Adj. SST": "SST1",
    "Launch Date": "2024-01-01", "Item Class": "C1", "RP Type": "RF",
    "Shop Stock Status": "Normal", "Adj. Safety Stock": 3.0,
    "Sold Qty (P90D)": 90.0, "Sold Qty (M-2)": 30.0, "Sold Qty (M-1)": 30.0,
    "Excess Stock Qty": 2.0,
    "Shop Stock Qty (CM)": 5.0,
    "Shop Monthly Sold Qty (P90D)": 30.5,
    "Sold Qty (CM - MTD)": 8.0,
    "D001 Stock Qty (CM)": 200.0, "D001 Net Stock Qty (CM)": 99.0,
    "Turnday (P90D Qty)": 13.0,
}


def _row(**kw):
    r = dict(BASE)
    r.update(kw)
    return r


def _df(rows):
    return pd.DataFrame([_row(**r) for r in rows])


def _run(rows, **kwargs):
    return transform_hybrid(_df(rows), **kwargs)


def _value(raw, site, col):
    return raw.loc[raw["Site Code"] == site, col].iloc[0]


# ──────────────────────────────────────────────────────────────────────────────
# 1. 新欄公式與範圍（RP Type＝ND；14 日線；零銷有存全數；RF 留空）
# ──────────────────────────────────────────────────────────────────────────────

def test_nd_clearance_formula_and_scope():
    raw, *_ = _run([
        {"Site Code": "L01", "RP Type": "ND", "Shop Nature": "Non Tourist",
         "Shop Stock Qty (CM)": 20.0, "Sold Qty (P90D)": 90.0, "Shop Monthly Sold Qty (P90D)": 30.5},
        {"Site Code": "L02", "RP Type": "ND", "Shop Nature": "Non Tourist",
         "Shop Stock Qty (CM)": 10.0, "Sold Qty (P90D)": 6.0, "Shop Monthly Sold Qty (P90D)": 2.0},
        {"Site Code": "L03", "RP Type": "ND", "Shop Nature": "Non Tourist",
         "Shop Stock Qty (CM)": 4.0, "Sold Qty (P90D)": 0.0, "Shop Monthly Sold Qty (P90D)": 0.0},
        {"Site Code": "R01", "RP Type": "RF", "Shop Nature": "Tourist",
         "Shop Stock Qty (CM)": 50.0, "Sold Qty (P90D)": 0.0, "Shop Monthly Sold Qty (P90D)": 0.0},
    ])
    # L01：20 − 90/90×14 = 6
    assert int(_value(raw, "L01", _HYBRID_ND_CLEARANCE_COLUMN)) == 6
    # L02：10 − 0.9333 = 9.0667 → floor 9（只扣 14 日線，不設 3 件保底）
    assert int(_value(raw, "L02", _HYBRID_ND_CLEARANCE_COLUMN)) == 9
    # L03：零銷有存 → 全數可調出
    assert int(_value(raw, "L03", _HYBRID_ND_CLEARANCE_COLUMN)) == 4
    # RF 列留空（非 ND 不輸出）
    assert pd.isna(_value(raw, "R01", _HYBRID_ND_CLEARANCE_COLUMN))
    # 庫存 ≤ 14 日所需 → 0（不入規劃表）
    raw2, *_ = _run([
        {"Site Code": "L04", "RP Type": "ND", "Shop Nature": "Non Tourist",
         "Shop Stock Qty (CM)": 5.0, "Sold Qty (P90D)": 90.0, "Shop Monthly Sold Qty (P90D)": 30.5},
    ])
    assert int(_value(raw2, "L04", _HYBRID_ND_CLEARANCE_COLUMN)) == 0
    assert _build_hybrid_nd_clearance(raw2).empty


# ──────────────────────────────────────────────────────────────────────────────
# 2. 不套池級門檻、與分類維度無關
# ──────────────────────────────────────────────────────────────────────────────

_POOL_ROWS = [
    {"Site Code": "N01", "SKU": 2, "RP Type": "ND", "Shop Nature": "Non Tourist",
     "Shop Stock Qty (CM)": 100.0, "Sold Qty (P90D)": 30.0, "Shop Monthly Sold Qty (P90D)": 30.5},
    {"Site Code": "N02", "SKU": 2, "RP Type": "RF", "Shop Nature": "Non Tourist",
     "Shop Stock Qty (CM)": 0.0, "Sold Qty (P90D)": 3000.0, "Shop Monthly Sold Qty (P90D)": 1016.7},
]


def test_nd_clearance_ignores_pool_gating():
    raw, *_ = _run(_POOL_ROWS)
    # Nature 池加總 TurnDay = 100 ÷ 1047.2 × 30.5 ≈ 2.91 < 14 → 現行可轉出量池級歸零
    assert int(_value(raw, "N01", "可轉出貨量 (Non Tourist)")) == 0
    # ND 規劃不套池級門檻：100 − 30/90×14 = 95.33 → 95
    assert int(_value(raw, "N01", _HYBRID_ND_CLEARANCE_COLUMN)) == 95


def test_nd_clearance_dimension_independent():
    raw_free, *_ = _run(_POOL_ROWS)
    raw_rp, *_ = _run(_POOL_ROWS, dimension="rp")
    for site in ("N01", "N02"):
        a = _value(raw_free, site, _HYBRID_ND_CLEARANCE_COLUMN)
        b = _value(raw_rp, site, _HYBRID_ND_CLEARANCE_COLUMN)
        assert (pd.isna(a) and pd.isna(b)) or float(a) == float(b), site
    assert int(_value(raw_rp, "N01", _HYBRID_ND_CLEARANCE_COLUMN)) == 95


def test_nd_clearance_negative_sales_treated_as_zero():
    """負銷售按 0（與 transform_hybrid 同一口徑）：14 日周轉所需不得為負、不得與建議調出量矛盾。"""
    raw, *_ = _run([
        {"Site Code": "G01", "RP Type": "ND", "Shop Nature": "Non Tourist",
         "Shop Stock Qty (CM)": 8.0, "Sold Qty (P90D)": -76.0, "Shop Monthly Sold Qty (P90D)": 0.0},
    ])
    # 負銷售按 0 → 零銷有存 → 全數可調出
    assert int(_value(raw, "G01", _HYBRID_ND_CLEARANCE_COLUMN)) == 8
    out = _build_hybrid_nd_clearance(raw)
    assert len(out) == 1
    assert float(out["Sold Qty (P90D)"].iloc[0]) == -76.0  # 原值保留供稽核
    assert float(out["14 日周轉所需"].iloc[0]) == 0.0      # 不顯示負需求


# ──────────────────────────────────────────────────────────────────────────────
# 3. Builder：篩選、欄序契約與排序
# ──────────────────────────────────────────────────────────────────────────────

def test_builder_filters_sorts_and_columns():
    raw, *_ = _run([
        {"Site Code": "B01", "RP Type": "ND", "Shop Nature": "Mix",
         "Shop Stock Qty (CM)": 10.0, "Sold Qty (P90D)": 0.0, "Shop Monthly Sold Qty (P90D)": 0.0},
        {"Site Code": "A01", "RP Type": "ND", "Shop Nature": "Non Tourist",
         "Shop Stock Qty (CM)": 20.0, "Sold Qty (P90D)": 90.0, "Shop Monthly Sold Qty (P90D)": 30.5},
        {"Site Code": "C01", "RP Type": "RF", "Shop Nature": "Non Tourist",
         "Shop Stock Qty (CM)": 99.0, "Sold Qty (P90D)": 0.0, "Shop Monthly Sold Qty (P90D)": 0.0},
    ])
    out = _build_hybrid_nd_clearance(raw)
    assert list(out.columns) == _HYBRID_ND_CLEARANCE_COLUMNS
    assert out["Site Code"].tolist() == ["A01", "B01"]
    assert out["RP Type"].tolist() == ["ND", "ND"]
    assert out[_HYBRID_ND_CLEARANCE_COLUMN].tolist() == [6, 10]
    assert float(out.loc[out["Site Code"] == "A01", "14 日周轉所需"].iloc[0]) == 14.0
    assert float(out.loc[out["Site Code"] == "B01", "14 日周轉所需"].iloc[0]) == 0.0
    # RF 店 C01 不入表
    assert "C01" not in out["Site Code"].tolist()
    # 14 日周轉所需保留 1 位小數：30 ÷ 90 × 14 = 4.7
    raw2, *_ = _run(_POOL_ROWS)
    out2 = _build_hybrid_nd_clearance(raw2)
    assert float(out2["14 日周轉所需"].iloc[0]) == 4.7


# ──────────────────────────────────────────────────────────────────────────────
# 4. 缺「RP Type」欄（來源全列按 RF）→ 欄全空、builder 空表
# ──────────────────────────────────────────────────────────────────────────────

def test_missing_rp_type_yields_empty_clearance():
    df = _df([
        {"Site Code": "X01", "Shop Nature": "Mix",
         "Shop Stock Qty (CM)": 50.0, "Sold Qty (P90D)": 0.0, "Shop Monthly Sold Qty (P90D)": 0.0},
    ]).drop(columns=["RP Type"])
    raw, sc, diag, diag_om, cross, stats, warns = transform_hybrid(df)
    assert _HYBRID_ND_CLEARANCE_COLUMN in raw.columns
    assert raw[_HYBRID_ND_CLEARANCE_COLUMN].isna().all()
    assert _build_hybrid_nd_clearance(raw).empty
    assert any("RP Type" in w for w in warns)


# ──────────────────────────────────────────────────────────────────────────────
# 5. Excel 契約：10 工作表、新表位置／表頭、RAW 欄序與 Review 主表不變
# ──────────────────────────────────────────────────────────────────────────────

def test_excel_nd_clearance_sheet_and_raw_column():
    raw, sc, diag, diag_om, cross, stats, _ = _run([
        {"Site Code": "A01", "RP Type": "ND", "Shop Nature": "Non Tourist",
         "Shop Stock Qty (CM)": 20.0, "Sold Qty (P90D)": 90.0, "Shop Monthly Sold Qty (P90D)": 30.5},
        {"Site Code": "T01", "RP Type": "RF", "Shop Nature": "Tourist",
         "Shop Stock Qty (CM)": 1.0, "Sold Qty (P90D)": 90.0, "Shop Monthly Sold Qty (P90D)": 30.5},
    ])
    buf = generate_excel_hybrid(raw, sc, diag, diag_om, cross, stats)
    wb = load_workbook(buf, read_only=True)
    assert wb.sheetnames == [
        "說明", "RAW", "Summary", "調貨 Review Safety Stock (不限OM)", "調貨 Review Safety Stock (同OM)",
        "店舖問題 SKU", "ND 可調出規劃", "Review 摘要",
        "簡易調貨明細 (不限OM)", "簡易調貨明細 (同OM)",
    ]
    raw_headers = [c.value for c in next(wb["RAW"].iter_rows(min_row=1, max_row=1))]
    assert _HYBRID_ND_CLEARANCE_COLUMN in raw_headers
    ws = wb["ND 可調出規劃"]
    headers = [c.value for c in next(ws.iter_rows(min_row=3, max_row=3))]
    assert headers == _HYBRID_ND_CLEARANCE_COLUMNS
    assert "建議調出參考" in str(ws.cell(row=1, column=1).value)
    first_row = [c.value for c in next(ws.iter_rows(min_row=4, max_row=4))]
    idx = {h: i for i, h in enumerate(headers)}
    assert first_row[idx[_HYBRID_ND_CLEARANCE_COLUMN]] == 6
    # 14 日周轉所需以 0.0 格式顯示（不四捨五入成整數）
    assert ws.cell(row=4, column=idx["14 日周轉所需"] + 1).number_format == "0.0"


def test_excel_nd_clearance_empty_placeholder():
    raw, sc, diag, diag_om, cross, stats, _ = _run([
        {"Site Code": "T01", "RP Type": "RF", "Shop Nature": "Tourist",
         "Shop Stock Qty (CM)": 1.0, "Sold Qty (P90D)": 90.0, "Shop Monthly Sold Qty (P90D)": 30.5},
    ])
    buf = generate_excel_hybrid(raw, sc, diag, diag_om, cross, stats)
    wb = load_workbook(buf)
    assert wb["ND 可調出規劃"].cell(row=3, column=1).value == "無符合資料"


def test_raw_display_order_and_review_tables_untouched():
    raw, sc, diag, diag_om, cross, stats, _ = _run(_POOL_ROWS)
    cols = _hybrid_raw_display_columns(raw)
    assert cols.index("可轉出貨量 (Non Tourist)") < cols.index(_HYBRID_ND_CLEARANCE_COLUMN)
    assert cols.index(_HYBRID_ND_CLEARANCE_COLUMN) < cols.index("實際調入量")
    # 主表契約不變：新欄不入 `_HYBRID_DIAG_COLUMNS`，兩張主表亦不含
    assert _HYBRID_ND_CLEARANCE_COLUMN not in _HYBRID_DIAG_COLUMNS
    assert _HYBRID_ND_CLEARANCE_COLUMN not in diag.columns
    assert _HYBRID_ND_CLEARANCE_COLUMN not in diag_om.columns
