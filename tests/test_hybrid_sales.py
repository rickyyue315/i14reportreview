"""Validation tests for Safety Stock x NetStock Review 銷量分類與銷量再平衡（v1.33.0）。

- 銷量帶：同 SKU 內按 Sold Qty (P90D) 平均排名分帶（>75%＝高銷、≤25%＝低銷、其餘＝中銷、P90D≤0＝零銷；
  有銷列 < 4 不產生高銷／低銷）
- 銷量存貨分類：高銷缺貨（周轉 < 缺貨線）、低銷積壓（周轉 ≥ 積壓線且庫存 ≥ 保底）、零銷積壓（庫存 ≥ 保底）、正常
- 銷量再平衡（預設關閉）：高銷缺貨 RF 列豁免 Nature 池門檻入需求方；積壓列經既有轉出判定提供來源；
  跟隨單向／雙向框架；ND 不作需求、池級歸零與同 Nature／同OM／自我調貨守衛保留
- 分類維度第六選項 sales：計算同 nature，僅 Summary／交叉表按銷量帶分組
- 分類維度第七選項 sales_nature（v1.34.0）：計算同 nature，Summary／交叉表按 Nature×銷量帶 12 組分組
- Review 摘要區塊⑤：銷量 × 存貨分類交叉表（母體為全部列，無調貨事件時仍輸出）
"""
import warnings
warnings.filterwarnings("ignore")
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from transformer import (
    transform_hybrid,
    generate_excel_hybrid,
    _hybrid_dashboard_metrics,
    _HYBRID_DIAG_COLUMNS,
    _HYBRID_DIMENSION_LABELS,
    _HYBRID_SALES_BANDS,
    _HYBRID_SALES_BAND_COLUMN,
    _HYBRID_SALES_CLASS_COLUMN,
    _HYBRID_SALES_CLASS_ORDER,
    _HYBRID_SALES_NATURE_GROUPS,
    _HYBRID_SALES_OOS_DAYS_DEFAULT,
    _HYBRID_SALES_OVERSTOCK_DAYS_DEFAULT,
    _HYBRID_SALES_OVERSTOCK_MIN_UNITS_DEFAULT,
    _HYBRID_SALES_SUMMARY_COLUMNS,
    _HYBRID_SKEW_THRESHOLDS,
    _hybrid_sales_bands,
    _hybrid_sales_classification,
    _hybrid_sales_params,
    _hybrid_dimension,
    _hybrid_dimension_groups,
    _hybrid_group_key_series,
    _hybrid_skew_thresholds,
    _hybrid_skew_hint_text,
    _build_hybrid_sales_summary,
)

BASE = {
    "Buyer": "B1", "Brand": "A", "SKU": 1, "SKU Desc": "Test Product",
    "Item_Shop": "Prod-S01-ShopHK", "Site Code": "S01", "OM": "OM1",
    "Shop Nature": "Tourist", "Vendor MC": "V1", "Item Status": "A",
    "Adj. SST": "SST1",
    "Launch Date": "2024-01-01", "Item Class": "C1", "RP Type": "RF",
    "Shop Stock Status": "Normal", "Adj. Safety Stock": 3.0,
    "Sold Qty (P90D)": 90.0, "Sold Qty (M-2)": 30.0, "Sold Qty (M-1)": 30.0,
    "Excess Stock Qty": 2.0,
    "Shop Stock Qty (CM)": 5.0,
    "Shop Monthly Sold Qty (P90D)": 30.5,
    "Sold Qty (CM - MTD)": 8.0,
    "D001 Stock Qty (CM)": 200.0, "D001 Net Stock Qty (CM)": 99.0,
    "Turnday (P90D Qty)": 13.0,
}


def _row(**kw):
    r = dict(BASE)
    r.update(kw)
    return r


def _df(rows):
    return pd.DataFrame([_row(**r) for r in rows])


# ──────────────────────────────────────────────────────────────────────────────
# 1. 銷量帶分帶規則（_hybrid_sales_bands 純函數）
# ──────────────────────────────────────────────────────────────────────────────

def _bands(sold, sku):
    s = pd.Series(sold, dtype=float)
    k = pd.Series(sku)
    return _hybrid_sales_bands(s, k).tolist()


def test_bands_rank_quarters():
    sold = [1, 2, 3, 4, 5, 6, 7, 8]
    sku = ["S"] * 8
    out = _bands(sold, sku)
    assert out[:2] == ["低銷", "低銷"]          # rank/n <= 0.25
    assert out[6:] == ["高銷", "高銷"]          # rank/n > 0.75
    assert set(out[2:6]) == {"中銷"}


def test_bands_zero_sales_band():
    out = _bands([0, 5, 5, 5, -3], ["S"] * 5)
    assert out[0] == "零銷" and out[4] == "零銷"
    assert set(out[1:4]) == {"中銷"}            # 有銷列數 3 < 4 → 不產生高／低


def test_bands_small_sample_guard():
    out = _bands([5, 10, 20], ["S"] * 3)
    assert out == ["中銷", "中銷", "中銷"]
    out4 = _bands([5, 10, 20, 40], ["S"] * 4)
    assert out4 == ["低銷", "中銷", "中銷", "高銷"]


def test_bands_ties_share_band():
    # [1,1,2,3]：平均排名 1.5,1.5,3,4 → frac .375,.375,.75,1 → 僅 3 為高銷、無低銷
    out = _bands([1, 1, 2, 3], ["S"] * 4)
    assert out == ["中銷", "中銷", "中銷", "高銷"]
    # 全同值 → 全中銷（無相對高低）
    assert _bands([2, 2, 2, 2], ["S"] * 4) == ["中銷"] * 4


def test_bands_per_sku_independent():
    out = _bands([1, 2, 3, 4, 1, 2, 3, 4], ["S1", "S1", "S1", "S1", "S2", "S2", "S2", "S2"])
    assert out[0] == out[4] and out[-1] == out[3]


# ──────────────────────────────────────────────────────────────────────────────
# 2. 銷量存貨分類邊界（_hybrid_sales_classification 純函數）
# ──────────────────────────────────────────────────────────────────────────────

_PARAMS = {"oos_days": 7, "overstock_days": 30, "overstock_min_units": 5}


def _cls(bands, td, stock):
    return _hybrid_sales_classification(
        pd.Series(bands), np.array(td, dtype=float), np.array(stock, dtype=float), _PARAMS,
    ).tolist()


def test_classification_boundaries():
    assert _cls(["高銷"], [6.99], [10]) == ["高銷缺貨"]
    assert _cls(["高銷"], [7.0], [10]) == ["正常"]        # 缺貨用「低於」
    assert _cls(["低銷"], [30.0], [5]) == ["低銷積壓"]     # 積壓用「不少於」、庫存含等號
    assert _cls(["低銷"], [30.0], [4]) == ["正常"]        # 未達保底件數
    assert _cls(["低銷"], [29.9], [99]) == ["正常"]       # 未達積壓線
    assert _cls(["零銷"], [np.inf], [5]) == ["零銷積壓"]
    assert _cls(["零銷"], [np.inf], [4]) == ["正常"]      # 零銷有存未達保底
    assert _cls(["零銷"], [np.inf], [0]) == ["正常"]
    assert _cls(["中銷"], [1.0], [100]) == ["正常"]       # 中銷不判缺貨／積壓


# ──────────────────────────────────────────────────────────────────────────────
# 3. 銷量參數正規化
# ──────────────────────────────────────────────────────────────────────────────

def test_sales_params_defaults():
    reb, params, warns = _hybrid_sales_params(False, None, None, None)
    assert reb is False and warns == []
    assert params == {
        "oos_days": _HYBRID_SALES_OOS_DAYS_DEFAULT,
        "overstock_days": _HYBRID_SALES_OVERSTOCK_DAYS_DEFAULT,
        "overstock_min_units": _HYBRID_SALES_OVERSTOCK_MIN_UNITS_DEFAULT,
    }


def test_sales_params_clamped_and_rejected():
    _, params, warns = _hybrid_sales_params(True, 20, 5, 999)
    assert params == {"oos_days": 14, "overstock_days": 15, "overstock_min_units": 100}
    assert len(warns) == 3
    _, params2, _ = _hybrid_sales_params(True, "8", 40, "6")
    assert params2 == {"oos_days": 8, "overstock_days": 40, "overstock_min_units": 6}
    for bad in (True, 7.5):
        try:
            _hybrid_sales_params(True, bad, None, None)
            raise AssertionError(f"應拒絕 {bad!r}")
        except ValueError:
            pass


# ──────────────────────────────────────────────────────────────────────────────
# 4. transform 整合——兩欄恆輸出、開關關閉時調貨數值零變動
# ──────────────────────────────────────────────────────────────────────────────

# SKU 500：Tourist 池健康（加總 TurnDay ≈16.3 ≥ 14）但 TH 高銷缺貨；
# N501（中銷）與 M502（低銷積壓）為 Mix／NT 有貨列
_SCENARIO_SALES = [
    {"SKU": 500, "Site Code": "TH", "Shop Nature": "Tourist",
     "Sold Qty (P90D)": 360.0, "Shop Stock Qty (CM)": 20.0,
     "Shop Monthly Sold Qty (P90D)": 120.0, "Adj. Safety Stock": 30,
     "D001 Net Stock Qty (CM)": 10.0},
    {"SKU": 500, "Site Code": "TO", "Shop Nature": "Tourist",
     "Sold Qty (P90D)": 90.0, "Shop Stock Qty (CM)": 60.0,
     "Shop Monthly Sold Qty (P90D)": 30.5, "Adj. Safety Stock": 10},
    {"SKU": 500, "Site Code": "N501", "Shop Nature": "Non Tourist",
     "Sold Qty (P90D)": 45.0, "Shop Stock Qty (CM)": 60.0,
     "Shop Monthly Sold Qty (P90D)": 15.25, "Adj. Safety Stock": 10},
    {"SKU": 500, "Site Code": "M502", "Shop Nature": "Mix", "RP Type": "ND",
     "Sold Qty (P90D)": 30.0, "Shop Stock Qty (CM)": 40.0,
     "Shop Monthly Sold Qty (P90D)": 10.17, "Adj. Safety Stock": 10},
]

_TRANSFER_NUMERIC_COLS = [
    "需要調貨", "調入需求量 (補到14日)", "實際調入量", "調入缺口", "轉出量",
    "可轉出貨量 (Mix)", "可轉出貨量 (Non Tourist)", "實際調入量 (同OM)",
    "轉出量 (同OM)", "調貨後庫存", "調貨後 TurnDay",
]


def test_sales_columns_always_present():
    raw, _, diagnosis, _, _, stats, _ = transform_hybrid(_df(_SCENARIO_SALES))
    assert _HYBRID_SALES_BAND_COLUMN in raw.columns
    assert _HYBRID_SALES_CLASS_COLUMN in raw.columns
    assert list(diagnosis.columns) == _HYBRID_DIAG_COLUMNS
    assert stats["sales_rebalance"] is False
    assert stats["sales_need_rows"] == 0
    # 開關關閉：池健康 → 無調貨事件
    assert (raw["需要調貨"] != "是").all()


def test_bands_and_classes_in_scenario():
    raw, _, _, _, _, _, _ = transform_hybrid(_df(_SCENARIO_SALES))
    bands = raw.set_index("Site Code")[_HYBRID_SALES_BAND_COLUMN].to_dict()
    assert bands["TH"] == "高銷" and bands["TO"] == "中銷"
    assert bands["N501"] == "中銷" and bands["M502"] == "低銷"
    cls = raw.set_index("Site Code")[_HYBRID_SALES_CLASS_COLUMN].to_dict()
    assert cls["TH"] == "高銷缺貨"      # td = 20/120*30.5 ≈ 5.1 < 7
    assert cls["M502"] == "低銷積壓"    # td ≈ 120 ≥ 30 且庫存 40 ≥ 5
    assert cls["TO"] == "正常" and cls["N501"] == "正常"


def test_rebalance_off_numeric_contract_unchanged():
    df = _df(_SCENARIO_SALES)
    a, *_rest_a, _ = transform_hybrid(df)
    b, *_rest_b, _ = transform_hybrid(df, sales_rebalance=False, sales_oos_days=1,
                                      sales_overstock_days=15, sales_overstock_min_units=1)
    pd.testing.assert_frame_equal(a[_TRANSFER_NUMERIC_COLS], b[_TRANSFER_NUMERIC_COLS])
    c, *_rest_c, _ = transform_hybrid(df, sales_rebalance=False)
    pd.testing.assert_frame_equal(a[_TRANSFER_NUMERIC_COLS], c[_TRANSFER_NUMERIC_COLS])


# ──────────────────────────────────────────────────────────────────────────────
# 5. 銷量再平衡通道
# ──────────────────────────────────────────────────────────────────────────────

def test_rebalance_adds_high_sales_oos_demand():
    raw, _, diagnosis, _, _, stats, warns = transform_hybrid(
        _df(_SCENARIO_SALES), sales_rebalance=True)
    th = raw[raw["Site Code"] == "TH"].iloc[0]
    assert th["需要調貨"] == "是" and th["事件類型"] == "需要調入"
    assert int(th["調入需求量 (補到14日)"]) == 36  # ceil(56) − 20
    # 積壓列 ND 仍可作供應方（ND 只禁需求）
    assert raw[raw["Site Code"] == "M502"].iloc[0]["事件類型"] in ("實際轉出", "可轉出未使用")
    assert stats["sales_need_rows"] == 1
    assert stats["sales_rebalance"] is True
    assert any("銷量再平衡" in w for w in warns)
    # 實際調入自 NT／Mix 池（NT 優先）
    assert int(th["實際調入量"]) > 0


def test_rebalance_unidirectional_keeps_direction_framework():
    # SKU 510：高銷缺貨列是 Mix（MH）——單向時不作需求
    rows = [
        {"SKU": 510, "Site Code": "MH", "Shop Nature": "Mix",
         "Sold Qty (P90D)": 360.0, "Shop Stock Qty (CM)": 20.0,
         "Shop Monthly Sold Qty (P90D)": 120.0, "D001 Net Stock Qty (CM)": 10.0},
        {"SKU": 510, "Site Code": "MO", "Shop Nature": "Mix",
         "Sold Qty (P90D)": 90.0, "Shop Stock Qty (CM)": 60.0,
         "Shop Monthly Sold Qty (P90D)": 30.5},
        {"SKU": 510, "Site Code": "N511", "Shop Nature": "Non Tourist",
         "Sold Qty (P90D)": 45.0, "Shop Stock Qty (CM)": 60.0,
         "Shop Monthly Sold Qty (P90D)": 15.25},
        {"SKU": 510, "Site Code": "T512", "Shop Nature": "Tourist",
         "Sold Qty (P90D)": 30.0, "Shop Stock Qty (CM)": 100.0,
         "Shop Monthly Sold Qty (P90D)": 10.17},
    ]
    raw_off, *_ , _ = transform_hybrid(_df(rows))
    assert (raw_off["需要調貨"] != "是").all()
    raw_on, _, _, _, _, stats_on, _ = transform_hybrid(_df(rows), sales_rebalance=True)
    assert (raw_on["需要調貨"] != "是").all()   # 單向：Mix 高銷缺貨不放寬
    assert stats_on["sales_need_rows"] == 0
    raw_bi, _, _, _, _, stats_bi, _ = transform_hybrid(
        _df(rows), sales_rebalance=True, bidirectional=True)
    mh = raw_bi[raw_bi["Site Code"] == "MH"].iloc[0]
    assert mh["需要調貨"] == "是" and int(mh["實際調入量"]) == 36
    assert stats_bi["sales_need_rows"] == 1


def test_rebalance_nd_never_demand():
    rows = [
        {"SKU": 520, "Site Code": "TD", "Shop Nature": "Tourist", "RP Type": "ND",
         "Sold Qty (P90D)": 360.0, "Shop Stock Qty (CM)": 20.0,
         "Shop Monthly Sold Qty (P90D)": 120.0, "D001 Net Stock Qty (CM)": 10.0},
        {"SKU": 520, "Site Code": "TO", "Shop Nature": "Tourist",
         "Sold Qty (P90D)": 90.0, "Shop Stock Qty (CM)": 60.0,
         "Shop Monthly Sold Qty (P90D)": 30.5},
        {"SKU": 520, "Site Code": "N521", "Shop Nature": "Non Tourist",
         "Sold Qty (P90D)": 45.0, "Shop Stock Qty (CM)": 60.0,
         "Shop Monthly Sold Qty (P90D)": 15.25},
        {"SKU": 520, "Site Code": "M522", "Shop Nature": "Mix",
         "Sold Qty (P90D)": 30.0, "Shop Stock Qty (CM)": 40.0,
         "Shop Monthly Sold Qty (P90D)": 10.17},
    ]
    raw, _, _, _, _, _, _ = transform_hybrid(_df(rows), sales_rebalance=True)
    td = raw[raw["Site Code"] == "TD"].iloc[0]
    assert td[_HYBRID_SALES_CLASS_COLUMN] == "高銷缺貨"
    assert td["需要調貨"] != "是" and td["事件類型"] == ""


def test_rebalance_pool_zeroing_guard_kept():
    # SKU 530：Mix 池缺貨（加總 TurnDay < 14）但池內 M532 屬低銷積壓列——仍不可轉出
    rows = [
        {"SKU": 530, "Site Code": "T531", "Shop Nature": "Tourist",
         "Sold Qty (P90D)": 360.0, "Shop Stock Qty (CM)": 20.0,
         "Shop Monthly Sold Qty (P90D)": 120.0, "D001 Net Stock Qty (CM)": 10.0},
        {"SKU": 530, "Site Code": "T532", "Shop Nature": "Tourist",
         "Sold Qty (P90D)": 90.0, "Shop Stock Qty (CM)": 60.0,
         "Shop Monthly Sold Qty (P90D)": 30.5},
        {"SKU": 530, "Site Code": "M531", "Shop Nature": "Mix",
         "Sold Qty (P90D)": 900.0, "Shop Stock Qty (CM)": 10.0,
         "Shop Monthly Sold Qty (P90D)": 300.0},
        {"SKU": 530, "Site Code": "M532", "Shop Nature": "Mix",
         "Sold Qty (P90D)": 30.0, "Shop Stock Qty (CM)": 60.0,
         "Shop Monthly Sold Qty (P90D)": 10.17},
        {"SKU": 530, "Site Code": "N531", "Shop Nature": "Non Tourist",
         "Sold Qty (P90D)": 45.0, "Shop Stock Qty (CM)": 60.0,
         "Shop Monthly Sold Qty (P90D)": 15.25},
    ]
    raw, _, _, _, _, stats, _ = transform_hybrid(_df(rows), sales_rebalance=True)
    t531 = raw[raw["Site Code"] == "T531"].iloc[0]
    assert t531["需要調貨"] == "是"
    m532 = raw[raw["Site Code"] == "M532"].iloc[0]
    assert m532[_HYBRID_SALES_CLASS_COLUMN] == "低銷積壓"
    assert int(m532["可轉出貨量 (Mix)"]) == 0 and int(m532["轉出量"]) == 0
    assert m532["事件類型"] == ""
    # 單向：Mix 高銷缺貨 M531 不作需求
    assert raw[raw["Site Code"] == "M531"].iloc[0]["需要調貨"] != "是"
    # NT 池健康 → 由 N531 轉出
    assert int(raw[raw["Site Code"] == "N531"].iloc[0]["轉出量"]) > 0
    assert stats["sales_need_rows"] == 1


def test_rebalance_respects_d001_gate():
    rows = [dict(r, **{"D001 Net Stock Qty (CM)": 500.0, "D001 Stock Qty (CM)": 500.0})
            for r in _SCENARIO_SALES]
    raw, _, _, _, _, stats, _ = transform_hybrid(_df(rows), sales_rebalance=True)
    assert (raw["需要調貨"] != "是").all()   # D001 充足 → 倉庫補貨，不調貨
    assert stats["sales_need_rows"] == 0


# ──────────────────────────────────────────────────────────────────────────────
# 6. 分類維度 sales
# ──────────────────────────────────────────────────────────────────────────────

def test_sales_dimension_registry():
    assert _hybrid_dimension("sales") == "sales"
    assert _hybrid_dimension_groups("sales") == _HYBRID_SALES_BANDS
    assert _hybrid_skew_thresholds("sales") == (0.30, 0.15)
    assert _HYBRID_SKEW_THRESHOLDS["sales"] == (0.30, 0.15)
    assert "≥ 30%" in _hybrid_skew_hint_text("sales")


def test_sales_dimension_groups_summary_and_cross():
    df = _df(_SCENARIO_SALES)
    base, *_ , _ = transform_hybrid(df)
    raw, sc_summary, _, _, diag_summary, stats, _ = transform_hybrid(df, dimension="sales")
    assert stats["dimension"] == "sales"
    for g in _HYBRID_SALES_BANDS:
        assert f"{g} 加總 - Shop Stock Qty (CM)" in sc_summary.columns
        assert f"由{g} 店舖預計調出數量" in sc_summary.columns
        assert f"{g} 列數" in diag_summary.columns
    # 計算與預設相同（調貨數值零變動）
    pd.testing.assert_frame_equal(base[_TRANSFER_NUMERIC_COLS], raw[_TRANSFER_NUMERIC_COLS])


def test_group_key_sales_fallback_without_band_column():
    df = pd.DataFrame({"Shop Nature": ["Tourist", "Mix"], "RP Type": ["RF", "RF"]})
    key = _hybrid_group_key_series(df, "sales")
    assert key.tolist() == ["Tourist", "Mix"]   # 缺銷量帶欄 → 顯式退回 Shop Nature 鍵


# ──────────────────────────────────────────────────────────────────────────────
# 6b. 分類維度 sales_nature（銷量＋Shop Nature 合併分組，v1.34.0）
# ──────────────────────────────────────────────────────────────────────────────

def test_sales_nature_dimension_registry():
    assert _hybrid_dimension("sales_nature") == "sales_nature"
    groups = _hybrid_dimension_groups("sales_nature")
    assert groups == _HYBRID_SALES_NATURE_GROUPS
    assert groups == tuple(
        f"{n}-{b}" for n in ("Tourist", "Mix", "Non Tourist") for b in _HYBRID_SALES_BANDS
    )
    assert len(groups) == 12
    assert _hybrid_skew_thresholds("sales_nature") == (0.12, 0.08)
    assert _HYBRID_SKEW_THRESHOLDS["sales_nature"] == (0.12, 0.08)
    assert "≥ 12%" in _hybrid_skew_hint_text("sales_nature")
    assert _HYBRID_DIMENSION_LABELS["sales_nature"] == "銷量＋Shop Nature"


def test_sales_nature_dimension_groups_summary_and_cross():
    df = _df(_SCENARIO_SALES)
    base, *_, _ = transform_hybrid(df)
    raw, sc_summary, _, _, diag_summary, stats, _ = transform_hybrid(df, dimension="sales_nature")
    assert stats["dimension"] == "sales_nature"
    # Summary／交叉表按 Nature-銷量帶 12 組分欄
    for g in _HYBRID_SALES_NATURE_GROUPS:
        assert f"{g} 列數" in diag_summary.columns
    for g in ("Tourist-高銷", "Mix-低銷", "Non Tourist-零銷"):
        assert f"{g} 加總 - Shop Stock Qty (CM)" in sc_summary.columns
        assert f"由{g} 店舖預計調出數量" in sc_summary.columns
    # 分組加總對帳：TH＝Tourist-高銷（庫存 20）、M502＝Mix-低銷（庫存 40）
    assert int(sc_summary["Tourist-高銷 加總 - Shop Stock Qty (CM)"].iloc[0]) == 20
    assert int(sc_summary["Mix-低銷 加總 - Shop Stock Qty (CM)"].iloc[0]) == 40
    assert int(sc_summary["Tourist-中銷 加總 - Shop Stock Qty (CM)"].iloc[0]) == 60
    # 計算與預設相同（調貨數值零變動）
    pd.testing.assert_frame_equal(base[_TRANSFER_NUMERIC_COLS], raw[_TRANSFER_NUMERIC_COLS])


def test_group_key_sales_nature():
    raw, *_ = transform_hybrid(_df(_SCENARIO_SALES))
    key = _hybrid_group_key_series(raw, "sales_nature")
    by_site = dict(zip(raw["Site Code"], key))
    assert by_site["TH"] == "Tourist-高銷"
    assert by_site["TO"] == "Tourist-中銷"
    assert by_site["N501"] == "Non Tourist-中銷"
    assert by_site["M502"] == "Mix-低銷"


def test_group_key_sales_nature_fallback_without_band_column():
    df = pd.DataFrame({"Shop Nature": ["Tourist", "Mix"], "RP Type": ["RF", "RF"]})
    key = _hybrid_group_key_series(df, "sales_nature")
    assert key.tolist() == ["Tourist", "Mix"]   # 缺銷量帶欄 → 顯式退回 Shop Nature 鍵


def test_sales_nature_excel_outputs():
    raw, sc, dg, dg_om, ds, stats, _ = transform_hybrid(
        _df(_SCENARIO_SALES), dimension="sales_nature", sales_rebalance=True)
    assert not ds.empty  # 銷量再平衡開啟 → TH 高銷缺貨入事件，交叉表有內容
    wb = load_workbook(generate_excel_hybrid(raw, sc, dg, dg_om, ds, stats))
    ws_sum = wb["Summary"]
    row4 = [ws_sum.cell(row=4, column=c).value for c in range(1, ws_sum.max_column + 1)]
    assert "Tourist-高銷" in row4 and "Non Tourist-零銷" in row4
    ws_rev = wb["Review 摘要"]
    col_a = [ws_rev.cell(row=r, column=1).value for r in range(1, 80)]
    assert any(isinstance(v, str) and v.startswith("Safety Review × 銷量＋Shop Nature 交叉表")
               for v in col_a)


# ──────────────────────────────────────────────────────────────────────────────
# 7. 區塊⑤與指標卡第 6 排
# ──────────────────────────────────────────────────────────────────────────────

def test_sales_summary_block_population_all_rows():
    raw, *_ , _ = transform_hybrid(_df(_SCENARIO_SALES))
    block = _build_hybrid_sales_summary(raw)
    assert list(block.columns) == _HYBRID_SALES_SUMMARY_COLUMNS
    assert block["銷量存貨分類"].tolist() == [*_HYBRID_SALES_CLASS_ORDER, "總計"]
    total = block[block["銷量存貨分類"] == "總計"].iloc[0]
    assert int(total["總計 列數"]) == len(raw)
    assert int(total["總計 店舖數"]) == raw["Site Code"].nunique()
    under = block[block["銷量存貨分類"] == "高銷缺貨"].iloc[0]
    assert int(under["Tourist 列數"]) == 1 and int(under["總計 列數"]) == 1
    over = block[block["銷量存貨分類"] == "低銷積壓"].iloc[0]
    assert int(over["庫存合計（件）"]) == 40


def test_excel_block5_written_even_without_events():
    raw, sc, dg, dg_om, ds, stats, _ = transform_hybrid(_df(_SCENARIO_SALES))
    assert ds.empty  # 開關關閉、池健康 → 無事件，①–④ 無內容
    ws = load_workbook(generate_excel_hybrid(raw, sc, dg, dg_om, ds, stats))["Review 摘要"]
    col_a = [ws.cell(row=r, column=1).value for r in range(1, 60)]
    assert any(isinstance(v, str) and v.startswith("無符合資料（①–④") for v in col_a)
    assert any(isinstance(v, str) and v.startswith("銷量 × 存貨分類交叉表") for v in col_a)


def test_excel_block5_after_event_blocks_position():
    raw, sc, dg, dg_om, ds, stats, _ = transform_hybrid(
        _df(_SCENARIO_SALES), sales_rebalance=True)
    ws = load_workbook(generate_excel_hybrid(raw, sc, dg, dg_om, ds, stats))["Review 摘要"]
    col_a = [ws.cell(row=r, column=1).value for r in range(1, 80)]
    idx_5 = next(i for i, v in enumerate(col_a)
                 if isinstance(v, str) and v.startswith("銷量 × 存貨分類交叉表"))
    idx_cross = next(i for i, v in enumerate(col_a)
                     if isinstance(v, str) and v.startswith("Safety Review ×"))
    assert idx_cross < idx_5
    # 銷量分類概況指標卡排（單向無第 5 排時為第 5 排：卡名列 19；有則第 6 排：卡名列 23）
    idx_card = next(i for i, v in enumerate(col_a) if v == "高銷缺貨列數")
    assert idx_card + 1 in (19, 23)  # 1-based 卡名列號
    assert int(ws.cell(row=idx_card + 2, column=1).value) == 1
    # 標題含門檻
    assert "缺貨線 7 日" in col_a[idx_5] and "積壓線 30 日" in col_a[idx_5]
    # 雙向時銷量排恆為第 6 排（卡名在第 23 列）
    raw_b, sc_b, dg_b, dgo_b, ds_b, stats_b, _ = transform_hybrid(
        _df(_SCENARIO_SALES), sales_rebalance=True, bidirectional=True)
    ws_b = load_workbook(generate_excel_hybrid(raw_b, sc_b, dg_b, dgo_b, ds_b, stats_b))["Review 摘要"]
    assert ws_b.cell(row=23, column=1).value == "高銷缺貨列數"
    assert int(ws_b.cell(row=24, column=1).value) == 1


def test_dashboard_metrics_sales_row_only_when_rebalance():
    raw, _, _, _, _, stats, _ = transform_hybrid(_df(_SCENARIO_SALES), sales_rebalance=True)
    rows = _hybrid_dashboard_metrics(raw, stats)
    assert len(rows) == 5  # 單向無第 5 排 → 銷量排為末排
    assert [m[0] for m in rows[-1]] == [
        "高銷缺貨列數", "高銷缺貨 SKU 數", "積壓列數", "積壓庫存件數"]
    assert int(rows[-1][0][1]) == 1
    raw_b, _, _, _, _, stats_b, _ = transform_hybrid(
        _df(_SCENARIO_SALES), sales_rebalance=True, bidirectional=True)
    rows_b = _hybrid_dashboard_metrics(raw_b, stats_b)
    assert len(rows_b) == 6
    assert [m[0] for m in rows_b[5]] == [
        "高銷缺貨列數", "高銷缺貨 SKU 數", "積壓列數", "積壓庫存件數"]
    rows_off = _hybrid_dashboard_metrics(raw, {**stats, "sales_rebalance": False})
    assert len(rows_off) == 4
    # RAW 回退路徑（stats=None）無法反推開關 → 不顯示銷量排
    assert len(_hybrid_dashboard_metrics(raw, None)) == 4


def test_stats_from_raw_sales_fallback():
    raw, *_ , _ = transform_hybrid(_df(_SCENARIO_SALES), sales_rebalance=True)
    from transformer import _hybrid_stats_from_raw
    st = _hybrid_stats_from_raw(raw)
    assert st["sales_oos_rows"] == 1
    assert st["sales_overstock_rows"] == 1
    assert st["sales_overstock_units"] == 40
    assert st["sales_oos_skus"] == 1


# ──────────────────────────────────────────────────────────────────────────────
# 8. 參數快照向後相容
# ──────────────────────────────────────────────────────────────────────────────

def test_params_snapshot_backward_compatible():
    from views.analyzer_helpers import build_hybrid_params, hybrid_params_changed
    old = build_hybrid_params(14, 100, {"Tourist": 7, "Mix": 9, "Non Tourist": 10},
                              False, "nature", None, False)
    new_default = build_hybrid_params(14, 100, {"Tourist": 7, "Mix": 9, "Non Tourist": 10},
                                      False, "nature", None, False,
                                      False, 7, 30, 2)
    assert not hybrid_params_changed(old, new_default)
    new_on = build_hybrid_params(14, 100, {"Tourist": 7, "Mix": 9, "Non Tourist": 10},
                                 False, "nature", None, False,
                                 True, 7, 30, 2)
    assert hybrid_params_changed(old, new_on)
    new_dim = build_hybrid_params(14, 100, {"Tourist": 7, "Mix": 9, "Non Tourist": 10},
                                  False, "sales", None, False,
                                  False, 7, 30, 2)
    assert hybrid_params_changed(old, new_dim)
