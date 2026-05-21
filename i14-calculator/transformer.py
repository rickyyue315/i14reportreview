import pandas as pd
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

REQUIRED_COLUMNS = ["Buyer", "Brand", "SKU", "Shop Stock Status"]


def validate_input(df: pd.DataFrame) -> list[str]:
    missing = [col for col in REQUIRED_COLUMNS if col not in df.columns]
    return missing


def transform(df_input: pd.DataFrame):
    missing = validate_input(df_input)
    if missing:
        raise ValueError(f"\u7f3a\u5c11\u5fc5\u8981\u6b04\u4f4d: {', '.join(missing)}")

    df = df_input.dropna(subset=["Buyer"]).copy()

    status_col = df["Shop Stock Status"].fillna("")

    categories = [
        "Overstock (P90D 0 Sales)",
        "Overstock (>2*Monthly Sold Qty & >2pcs)",
    ]

    group_cols = ["Buyer", "Brand", "SKU"]

    counts = {}
    for cat in categories:
        mask = status_col.str.contains(cat, regex=False, na=False)
        counts[cat] = (
            df[mask]
            .groupby(group_cols, observed=True)
            .size()
            .rename(cat)
        )

    summary = pd.concat(counts.values(), axis=1).fillna(0).astype(int).reset_index()

    summary["SKU"] = summary["SKU"].apply(lambda x: str(int(x)) if pd.notna(x) else x)
    summary["\u7e3d\u8a08"] = summary[categories].sum(axis=1)

    buyer_order = summary.groupby("Buyer")["\u7e3d\u8a08"].sum().sort_values(ascending=False).index.tolist()
    summary["_buyer_key"] = summary["Buyer"].map({b: i for i, b in enumerate(buyer_order)})
    summary = summary.sort_values(["_buyer_key", "Brand"], ascending=[True, True]).drop(columns=["_buyer_key"]).reset_index(drop=True)

    total_row = pd.DataFrame(
        {
            "Buyer": ["\u7e3d\u8a08"],
            "Brand": [None],
            "SKU": [None],
            **{cat: [int(summary[cat].sum())] for cat in categories},
            "\u7e3d\u8a08": [int(summary["\u7e3d\u8a08"].sum())],
        }
    )
    summary = pd.concat([summary, total_row], ignore_index=True)

    raw = df.copy()
    raw["SKU"] = raw["SKU"].apply(lambda x: str(int(x)) if pd.notna(x) else x)

    return summary, raw, categories


def generate_excel(summary_df: pd.DataFrame, raw_df: pd.DataFrame, categories: list[str]) -> BytesIO:
    wb = Workbook()

    ws_raw = wb.active
    ws_raw.title = "RAW"

    raw_cols = list(raw_df.columns)
    ws_raw.append(raw_cols)

    raw_values = raw_df.values.tolist()
    for row_vals in raw_values:
        cleaned = [None if isinstance(v, float) and pd.isna(v) else v for v in row_vals]
        ws_raw.append(cleaned)

    header_fill_raw = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    ncols = len(raw_cols)
    for col_idx in range(1, ncols + 1):
        cell = ws_raw.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        cell.fill = header_fill_raw
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    ws_raw.auto_filter.ref = ws_raw.dimensions

    ws_summary = wb.create_sheet(title="Shop Stock Status")

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    bold_font = Font(bold=True, size=11)

    ws_summary["A3"] = "\u884c\u6578 - Site Code"
    ws_summary["D3"] = "Shop Stock Status"

    headers = ["Buyer", "Brand", "SKU"] + categories + ["\u7e3d\u8a08"]
    for col_idx, h in enumerate(headers, 1):
        cell = ws_summary.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    data_start = 5

    for row_idx, row_data in summary_df.iterrows():
        excel_row = data_start + row_idx
        is_total = row_data["Buyer"] == "\u7e3d\u8a08"

        for col_idx, h in enumerate(headers, 1):
            val = row_data.get(h)
            if pd.isna(val):
                val = None
            cell = ws_summary.cell(row=excel_row, column=col_idx, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")

            if is_total:
                cell.font = bold_font

            if isinstance(val, (int, float)) and col_idx >= 4 and h != "SKU":
                cell.number_format = "#,##0"

    _auto_width(ws_summary)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _auto_width(ws, min_width=8, max_width=50):
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value is not None:
                val_len = len(str(cell.value))
                if val_len > max_len:
                    max_len = val_len
        adjusted = min(max(max_len + 2, min_width), max_width)
        ws.column_dimensions[col_letter].width = adjusted
