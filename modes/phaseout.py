# -*- coding: utf-8 -*-
"""退場清倉 Review：找出已停產或疑似應退場品項的殘留庫存，建議清倉安排。

只引用 modes.common，不引用其他模式模組，避免循環引用。
"""
import gc
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


from modes.common import (
    _auto_width,
    _clean_dataframe,
    _compute_raw_extra_columns,
    _extract_shop_name,
    _reorder_raw_columns,
    _resolve_column_name,
    _safe_default_series,
    _safe_numeric,
    _safe_sku,
    _write_raw_placeholder,
    _write_raw_sheet,
    validate_input,
)

__all__ = [
    "transform_phaseout",
    "generate_excel_phaseout",
    "_build_po_by_sku",
    "_build_po_by_shop",
]

_PO_REVIEW_COLUMN = "清倉建議"
_PO_RETURN = "建議退倉"
_PO_VENDOR = "另有安排"
_PO_REVIEW = "檢視退場"

# Item Status 退場關鍵字（小寫子字串比對；中文與英文常見寫法）
_PO_STATUS_KEYWORDS = (
    "下市", "停產", "退場", "停止", "結束", "下架",
    "discont", "phase out", "phaseout", "obsolete", "delist", "dead",
)

_PO_RAW_OUTPUT_COLUMNS = [
    "Site Code",
    "Shop Name",
    "OM",
    "Shop Nature",
    "Buyer",
    "Brand",
    "SKU",
    "SKU Desc",
    "Item Status",
    "Item Class",
    _PO_REVIEW_COLUMN,
    "RP Type",
    "Supply Source",
    "Shop Stock Status",
    "Shop Stock Qty (CM)",
    "Sold Qty (M-2)",
    "Sold Qty (M-1)",
    "Sold Qty (CM - MTD)",
    "Sold Qty (P90D)",
    "Excess Stock Value",
]

_PO_BY_SKU_COLUMNS = [
    "Buyer",
    "Brand",
    "SKU",
    "SKU Desc",
    "涉及店舖數",
    "庫存合計",
    "建議退倉合計",
]

_PO_BY_SHOP_COLUMNS = [
    "Site Code",
    "Shop Name",
    "OM",
    "涉及 SKU 數",
    "庫存總數",
    "建議退倉總數",
]


def _po_is_phaseout_status(value) -> bool:
    """Item Status 是否標示退場：關鍵字子字串比對（大小寫不敏感）。"""
    if pd.isna(value):
        return False
    txt = str(value).strip().lower()
    if not txt:
        return False
    return any(k in txt for k in _PO_STATUS_KEYWORDS)


def _po_raw_columns(raw: pd.DataFrame) -> list[str]:
    r = _reorder_raw_columns(_compute_raw_extra_columns(raw))
    available = set(r.columns)
    cols = [c for c in _PO_RAW_OUTPUT_COLUMNS if _resolve_column_name(c, available) is not None]
    if _PO_REVIEW_COLUMN not in cols:
        cols.insert(10, _PO_REVIEW_COLUMN)
    return cols


def _build_po_by_sku(raw: pd.DataFrame) -> pd.DataFrame:
    cols = ["Buyer", "Brand", "SKU", "SKU Desc"]
    if raw is None or raw.empty:
        return pd.DataFrame(columns=_PO_BY_SKU_COLUMNS)
    df = raw.copy()
    df["SKU"] = df["SKU"].apply(_safe_sku)
    df["_stock"] = _safe_numeric(df.get("Shop Stock Qty (CM)", _safe_default_series(0, df.index)))
    df["_return"] = _safe_numeric(df.get("_po_return_qty", _safe_default_series(0, df.index)))
    grouped = (
        df.groupby(cols, dropna=False, observed=True)
        .agg(
            **{
                "涉及店舖數": ("Site Code", "nunique"),
                "庫存合計": ("_stock", "sum"),
                "建議退倉合計": ("_return", "sum"),
            },
        )
        .reset_index()
    )
    for c in ["涉及店舖數", "庫存合計", "建議退倉合計"]:
        grouped[c] = pd.to_numeric(grouped[c], errors="coerce").fillna(0).round().astype(int)
    grouped = grouped.sort_values(["建議退倉合計", "庫存合計", "Brand", "SKU"], ascending=[False, False, True, True], kind="stable").reset_index(drop=True)
    if not grouped.empty:
        total_row = pd.DataFrame([{
            "Buyer": "總計",
            "Brand": None,
            "SKU": None,
            "SKU Desc": None,
            "涉及店舖數": int(df["Site Code"].dropna().nunique()),
            "庫存合計": int(grouped["庫存合計"].sum()),
            "建議退倉合計": int(grouped["建議退倉合計"].sum()),
        }])
        grouped = pd.concat([grouped, total_row], ignore_index=True)
    return grouped[_PO_BY_SKU_COLUMNS]


def _build_po_by_shop(raw: pd.DataFrame) -> pd.DataFrame:
    if raw is None or raw.empty:
        return pd.DataFrame(columns=_PO_BY_SHOP_COLUMNS)
    df = raw.copy()
    site_code = df.get("Site Code", _safe_default_series("", df.index)).fillna("").astype(str).str.strip()
    item_shop = df.get("Item_Shop", _safe_default_series("", df.index))
    df["Site Code"] = site_code
    df["Shop Name"] = [_extract_shop_name(item_shop.iloc[i], site_code.iloc[i]) for i in range(len(df))]
    df["OM"] = df.get("OM", _safe_default_series("", df.index)).fillna("").astype(str).str.strip()
    df["SKU"] = df["SKU"].apply(_safe_sku)
    df["_stock"] = _safe_numeric(df.get("Shop Stock Qty (CM)", _safe_default_series(0, df.index)))
    df["_return"] = _safe_numeric(df.get("_po_return_qty", _safe_default_series(0, df.index)))
    grouped = (
        df.groupby(["Site Code", "Shop Name", "OM"], dropna=False, observed=True)
        .agg(
            **{
                "涉及 SKU 數": ("SKU", "nunique"),
                "庫存總數": ("_stock", "sum"),
                "建議退倉總數": ("_return", "sum"),
            },
        )
        .reset_index()
    )
    for c in ["涉及 SKU 數", "庫存總數", "建議退倉總數"]:
        grouped[c] = pd.to_numeric(grouped[c], errors="coerce").fillna(0).round().astype(int)
    grouped = grouped.sort_values(["建議退倉總數", "庫存總數", "Site Code"], ascending=[False, False, True], kind="stable").reset_index(drop=True)
    if not grouped.empty:
        total_row = pd.DataFrame([{
            "Site Code": "總計",
            "Shop Name": "",
            "OM": "",
            "涉及 SKU 數": int(df["SKU"].dropna().nunique()),
            "庫存總數": int(grouped["庫存總數"].sum()),
            "建議退倉總數": int(grouped["建議退倉總數"].sum()),
        }])
        grouped = pd.concat([grouped, total_row], ignore_index=True)
    return grouped[_PO_BY_SHOP_COLUMNS]


def transform_phaseout(
    df_input: pd.DataFrame,
    progress_callback=None,
    shop_natures: list[str] | None = None,
):
    """退場清倉轉換：標示三類清倉建議，輸出 RAW、By SKU、By Shop 與提示訊息。"""
    if progress_callback:
        progress_callback(0.0, "驗證輸入資料...")
    missing = validate_input(df_input)
    if missing:
        raise ValueError(f"缺少必要欄位: {', '.join(missing)}")
    avail = set(df_input.columns)
    if _resolve_column_name("Site Code", avail) is None:
        raise ValueError("缺少必需欄位: Site Code")

    if progress_callback:
        progress_callback(0.1, "清理資料中...")
    df = _clean_dataframe(df_input, shop_natures=shop_natures, progress_callback=progress_callback)
    if df.empty:
        raise ValueError("過濾後無有效資料列，請確認來源檔案中的 Buyer 欄位是否包含有效數據")

    if progress_callback:
        progress_callback(0.2, "解像 SKU Desc / Shop Name...")
    df = _compute_raw_extra_columns(df)
    df["SKU"] = df["SKU"].apply(_safe_sku)

    has_status = "Item Status" in df.columns
    is_phaseout = df["Item Status"].apply(_po_is_phaseout_status) if has_status else pd.Series(False, index=df.index)
    rp = df.get("RP Type", _safe_default_series("", df.index)).fillna("").astype(str).str.strip().str.upper()
    is_nd = (rp == "ND").fillna(False)
    sold_m2 = _safe_numeric(df.get("Sold Qty (M-2)", _safe_default_series(0, df.index)))
    sold_m1 = _safe_numeric(df.get("Sold Qty (M-1)", _safe_default_series(0, df.index)))
    sold_mtd = _safe_numeric(df.get("Sold Qty (CM - MTD)", _safe_default_series(0, df.index)))
    sold_p90d = _safe_numeric(df.get("Sold Qty (P90D)", _safe_default_series(0, df.index)))
    stock_cm = _safe_numeric(df.get("Shop Stock Qty (CM)", _safe_default_series(0, df.index)))
    all_zero = ((sold_m2 <= 0) & (sold_m1 <= 0) & (sold_mtd <= 0) & (sold_p90d <= 0)).fillna(False)
    has_stock = (stock_cm > 0).fillna(False)
    supply = df.get("Supply Source", _safe_default_series("", df.index)).fillna("").astype(str).str.strip().str.lower()
    is_vendor = (supply == "vendor").fillna(False)

    rec = pd.Series("", index=df.index, dtype=object)
    # 優先次序：狀態退場 ＞ ND 零銷檢視
    rec[is_phaseout & has_stock & is_vendor] = _PO_VENDOR
    rec[is_phaseout & has_stock & ~is_vendor] = _PO_RETURN
    rec[~is_phaseout & is_nd & all_zero & has_stock] = _PO_REVIEW
    df[_PO_REVIEW_COLUMN] = rec.values
    df["_po_return_qty"] = [int(s) if r == _PO_RETURN else 0 for r, s in zip(rec.values, stock_cm.values)]

    mask = (rec != "").fillna(False)
    raw = df[mask.to_numpy()].copy()
    raw.drop(columns=["_po_return_qty"], inplace=True, errors="ignore")

    if progress_callback:
        progress_callback(0.35, "建立彙總...")
    # 彙總需退倉量：以內部欄重建（RAW 已移除）
    df["_po_return_qty"] = [int(s) if r == _PO_RETURN else 0 for r, s in zip(rec.values, stock_cm.values)]
    agg_src = df[mask.to_numpy()].copy()
    by_sku = _build_po_by_sku(agg_src)
    by_shop = _build_po_by_shop(agg_src)

    warnings = []
    if not has_status:
        warnings.append("來源檔案缺少「Item Status」，只按 ND 長期零銷售有庫存檢視退場候選。")
    if "Supply Source" not in df.columns:
        warnings.append("來源檔案缺少「Supply Source」，全部列按 Warehouse 貨處理（可退倉）。")
    if raw.empty:
        warnings.append("沒有符合「退場清倉（狀態退場有庫存／ND 長期零銷售有庫存）」的資料。")

    if progress_callback:
        progress_callback(0.4, "完成資料轉換...")
    return raw, by_sku, by_shop, warnings


def _write_po_table(ws, headers, data: pd.DataFrame, progress_callback=None, progress_start=0.0, progress_end=0.1, sheet_name=""):
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    bold_font = Font(bold=True, size=11)
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    if data is not None and not data.empty:
        total = len(data)
        for row_idx, (_, row_data) in enumerate(data.iterrows()):
            excel_row = row_idx + 2
            is_total = row_data.iloc[0] == "總計"
            for col_idx, h in enumerate(headers, 1):
                val = row_data.get(h)
                if pd.isna(val):
                    val = None
                cell = ws.cell(row=excel_row, column=col_idx, value=val)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if is_total:
                    cell.font = bold_font
                if isinstance(val, (int, float)) and col_idx > 4:
                    cell.number_format = "#,##0"
            if progress_callback and (row_idx % 500 == 0 or row_idx == total - 1):
                p = progress_start + (progress_end - progress_start) * ((row_idx + 1) / max(total, 1))
                progress_callback(p, f"寫入 {sheet_name} ({row_idx+1}/{total})...")
    else:
        ws.cell(row=2, column=1, value="無符合資料").font = Font(size=12)
    _auto_width(ws)


def generate_excel_phaseout(
    raw: pd.DataFrame,
    by_sku: pd.DataFrame,
    by_shop: pd.DataFrame,
    progress_callback=None,
    include_raw: bool = True,
) -> BytesIO:
    wb = Workbook()
    if progress_callback:
        progress_callback(0.0, "建立 Excel 結構...")

    out_cols = _po_raw_columns(raw if raw is not None else pd.DataFrame())

    ws_raw = wb.active
    ws_raw.title = "RAW"
    if include_raw:
        _write_raw_sheet(ws_raw, raw, progress_callback, progress_start=0.0, progress_end=0.55, columns=out_cols)
    else:
        _write_raw_placeholder(ws_raw)

    if progress_callback:
        progress_callback(0.6, "建立 By SKU...")
    ws_sku = wb.create_sheet(title="By SKU")
    _write_po_table(ws_sku, _PO_BY_SKU_COLUMNS, by_sku, progress_callback, 0.6, 0.75, "By SKU")

    if progress_callback:
        progress_callback(0.78, "建立 By Shop...")
    ws_shop = wb.create_sheet(title="By Shop")
    _write_po_table(ws_shop, _PO_BY_SHOP_COLUMNS, by_shop, progress_callback, 0.78, 0.9, "By Shop")

    if progress_callback:
        progress_callback(0.95, "正在儲存 Excel 檔案...")

    gc.collect()
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    if progress_callback:
        progress_callback(1.0, "完成!")
    return buf
