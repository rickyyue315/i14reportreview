# -*- coding: utf-8 -*-
"""新品追蹤 Review：追蹤上市 90 天內新品的鋪貨與起量情況。

只引用 modes.common，不引用其他模式模組，避免循環引用。
"""
import datetime
import gc
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


from modes.common import (
    _auto_width,
    _clean_dataframe,
    _compute_raw_extra_columns,
    _extract_shop_name,
    _normalize_ref_date,
    _parse_launch_dates,
    _reorder_raw_columns,
    _resolve_column_name,
    _safe_default_series,
    _safe_numeric,
    _safe_sku,
    _write_raw_placeholder,
    _write_raw_sheet,
    validate_input,
)

__all__ = [
    "transform_new_sku",
    "generate_excel_new_sku",
    "_build_ns_by_sku",
    "_build_ns_by_shop",
]

_NS_WINDOW_DAYS = 90
_NS_AGE_COLUMN = "上市日數"
_NS_STATUS_COLUMN = "起量狀態"
_NS_STARTED = "已起量"
_NS_NOT_STARTED = "未起量"

_NS_RAW_OUTPUT_COLUMNS = [
    "Site Code",
    "Shop Name",
    "OM",
    "Shop Nature",
    "Buyer",
    "Brand",
    "SKU",
    "SKU Desc",
    "Launch Date",
    _NS_AGE_COLUMN,
    _NS_STATUS_COLUMN,
    "Item Status",
    "Item Class",
    "RP Type",
    "Shop Stock Status",
    "Shop Stock Qty (CM)",
    "Sold Qty (M-2)",
    "Sold Qty (M-1)",
    "Sold Qty (CM - MTD)",
    "Sold Qty (P90D)",
    "Shop Monthly Sold Qty (P90D)",
]

_NS_BY_SKU_COLUMNS = [
    "Buyer",
    "Brand",
    "SKU",
    "SKU Desc",
    "Launch Date",
    "上市日數",
    "鋪貨店舖數",
    "有銷店舖數",
    "總庫存",
    "P90D 總銷量",
]

_NS_BY_SHOP_COLUMNS = [
    "Site Code",
    "Shop Name",
    "OM",
    "新品 SKU 數",
    "已起量 SKU 數",
    "未起量 SKU 數",
]


def _ns_raw_columns(raw: pd.DataFrame) -> list[str]:
    r = _reorder_raw_columns(_compute_raw_extra_columns(raw))
    available = set(r.columns)
    return [c for c in _NS_RAW_OUTPUT_COLUMNS if _resolve_column_name(c, available) is not None]


def _build_ns_by_sku(raw: pd.DataFrame) -> pd.DataFrame:
    cols = ["Buyer", "Brand", "SKU", "SKU Desc"]
    if raw is None or raw.empty:
        return pd.DataFrame(columns=_NS_BY_SKU_COLUMNS)
    df = raw.copy()
    df["SKU"] = df["SKU"].apply(_safe_sku)
    df["_started"] = (df[_NS_STATUS_COLUMN] == _NS_STARTED).astype(int)
    df["_stock"] = _safe_numeric(df.get("Shop Stock Qty (CM)", _safe_default_series(0, df.index)))
    df["_p90d"] = _safe_numeric(df.get("Sold Qty (P90D)", _safe_default_series(0, df.index)))
    grouped = (
        df.groupby(cols, dropna=False, observed=True)
        .agg(
            **{
                "Launch Date": ("Launch Date", "first"),
                "上市日數": ("上市日數", "min"),
                "鋪貨店舖數": ("Site Code", "nunique"),
                "有銷店舖數": ("_started", "sum"),
                "總庫存": ("_stock", "sum"),
                "P90D 總銷量": ("_p90d", "sum"),
            },
        )
        .reset_index()
    )
    grouped["上市日數"] = pd.to_numeric(grouped["上市日數"], errors="coerce").fillna(0).astype(int)
    for c in ["鋪貨店舖數", "有銷店舖數", "總庫存", "P90D 總銷量"]:
        grouped[c] = pd.to_numeric(grouped[c], errors="coerce").fillna(0).round().astype(int)
    grouped = grouped.sort_values(["Launch Date", "Brand", "SKU"], ascending=[False, True, True], kind="stable").reset_index(drop=True)
    if not grouped.empty:
        total_row = pd.DataFrame([{
            "Buyer": "總計",
            "Brand": None,
            "SKU": None,
            "SKU Desc": None,
            "Launch Date": "",
            "上市日數": None,
            "鋪貨店舖數": int(df["Site Code"].dropna().nunique()),
            "有銷店舖數": int(df[df["_started"] == 1]["Site Code"].dropna().nunique()),
            "總庫存": int(grouped["總庫存"].sum()),
            "P90D 總銷量": int(grouped["P90D 總銷量"].sum()),
        }])
        grouped = pd.concat([grouped, total_row], ignore_index=True)
    return grouped[_NS_BY_SKU_COLUMNS]


def _build_ns_by_shop(raw: pd.DataFrame) -> pd.DataFrame:
    if raw is None or raw.empty:
        return pd.DataFrame(columns=_NS_BY_SHOP_COLUMNS)
    df = raw.copy()
    site_code = df.get("Site Code", _safe_default_series("", df.index)).fillna("").astype(str).str.strip()
    item_shop = df.get("Item_Shop", _safe_default_series("", df.index))
    df["Site Code"] = site_code
    df["Shop Name"] = [_extract_shop_name(item_shop.iloc[i], site_code.iloc[i]) for i in range(len(df))]
    df["OM"] = df.get("OM", _safe_default_series("", df.index)).fillna("").astype(str).str.strip()
    df["SKU"] = df["SKU"].apply(_safe_sku)
    df["_started"] = (df[_NS_STATUS_COLUMN] == _NS_STARTED).astype(int)
    df["_not_started"] = (df[_NS_STATUS_COLUMN] == _NS_NOT_STARTED).astype(int)
    grouped = (
        df.groupby(["Site Code", "Shop Name", "OM"], dropna=False, observed=True)
        .agg(
            新品SKU數=("SKU", "nunique"),
            已起量SKU數=("_started", "sum"),
            未起量SKU數=("_not_started", "sum"),
        )
        .reset_index()
    )
    grouped = grouped.rename(columns={
        "新品SKU數": "新品 SKU 數",
        "已起量SKU數": "已起量 SKU 數",
        "未起量SKU數": "未起量 SKU 數",
    })
    for c in ["新品 SKU 數", "已起量 SKU 數", "未起量 SKU 數"]:
        grouped[c] = pd.to_numeric(grouped[c], errors="coerce").fillna(0).astype(int)
    grouped = grouped.sort_values(["新品 SKU 數", "Site Code"], ascending=[False, True], kind="stable").reset_index(drop=True)
    if not grouped.empty:
        total_row = pd.DataFrame([{
            "Site Code": "總計",
            "Shop Name": "",
            "OM": "",
            "新品 SKU 數": int(df["SKU"].dropna().nunique()),
            "已起量 SKU 數": int(grouped["已起量 SKU 數"].sum()),
            "未起量 SKU 數": int(grouped["未起量 SKU 數"].sum()),
        }])
        grouped = pd.concat([grouped, total_row], ignore_index=True)
    return grouped[_NS_BY_SHOP_COLUMNS]


def transform_new_sku(
    df_input: pd.DataFrame,
    progress_callback=None,
    shop_natures: list[str] | None = None,
    ref_date=None,
):
    """新品追蹤轉換：篩選上市 90 天內品項，輸出 RAW、By SKU、By Shop 與提示訊息。"""
    if progress_callback:
        progress_callback(0.0, "驗證輸入資料...")
    missing = validate_input(df_input)
    if missing:
        raise ValueError(f"缺少必要欄位: {', '.join(missing)}")
    avail = set(df_input.columns)
    for col in ["Site Code", "Launch Date"]:
        if _resolve_column_name(col, avail) is None:
            raise ValueError(f"缺少必需欄位: {col}")
    ref = _normalize_ref_date(ref_date)

    if progress_callback:
        progress_callback(0.1, "清理資料中...")
    df = _clean_dataframe(df_input, shop_natures=shop_natures, progress_callback=progress_callback)
    if df.empty:
        raise ValueError("過濾後無有效資料列，請確認來源檔案中的 Buyer 欄位是否包含有效數據")

    if progress_callback:
        progress_callback(0.2, "解像 SKU Desc / Shop Name...")
    df = _compute_raw_extra_columns(df)
    df["SKU"] = df["SKU"].apply(_safe_sku)

    launch_dt = _parse_launch_dates(df["Launch Date"].fillna("").astype(str).str.strip())
    ref_ts = pd.Timestamp(ref)
    age = (ref_ts - launch_dt).dt.days
    is_new = launch_dt.notna() & (age >= 0) & (age <= _NS_WINDOW_DAYS)

    raw = df[is_new.fillna(False).to_numpy()].copy()
    if not raw.empty:
        raw_age = age[is_new].astype(int)
        raw[_NS_AGE_COLUMN] = raw_age.values
        sold_m2 = _safe_numeric(raw.get("Sold Qty (M-2)", _safe_default_series(0, raw.index)))
        sold_m1 = _safe_numeric(raw.get("Sold Qty (M-1)", _safe_default_series(0, raw.index)))
        sold_mtd = _safe_numeric(raw.get("Sold Qty (CM - MTD)", _safe_default_series(0, raw.index)))
        sold_p90d = _safe_numeric(raw.get("Sold Qty (P90D)", _safe_default_series(0, raw.index)))
        started = ((sold_m2 > 0) | (sold_m1 > 0) | (sold_mtd > 0) | (sold_p90d > 0)).fillna(False)
        raw[_NS_STATUS_COLUMN] = [_NS_STARTED if v else _NS_NOT_STARTED for v in started.to_numpy()]

    if progress_callback:
        progress_callback(0.35, "建立彙總...")
    by_sku = _build_ns_by_sku(raw)
    by_shop = _build_ns_by_shop(raw)

    warnings = []
    bad_mask = df["Launch Date"].fillna("").astype(str).str.strip() != ""
    bad_count = int((bad_mask & launch_dt.isna()).sum())
    if bad_count:
        warnings.append(f"有 {bad_count} 列「Launch Date」無法解析，已視為非新品。")
    future_count = int((launch_dt.notna() & (age < 0)).sum())
    if future_count:
        warnings.append(f"有 {future_count} 列「Launch Date」晚於參考日期，已視為非新品。")
    if raw.empty:
        warnings.append(f"沒有符合「上市 {_NS_WINDOW_DAYS} 天內（參考日期 {ref.isoformat()}）」的新品資料。")

    if progress_callback:
        progress_callback(0.4, "完成資料轉換...")
    return raw, by_sku, by_shop, warnings


def _write_ns_table(ws, headers, data: pd.DataFrame, progress_callback=None, progress_start=0.0, progress_end=0.1, sheet_name=""):
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    bold_font = Font(bold=True, size=11)
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    if data is not None and not data.empty:
        total = len(data)
        for row_idx, (_, row_data) in enumerate(data.iterrows()):
            excel_row = row_idx + 2
            is_total = row_data.iloc[0] == "總計"
            for col_idx, h in enumerate(headers, 1):
                val = row_data.get(h)
                if pd.isna(val):
                    val = None
                cell = ws.cell(row=excel_row, column=col_idx, value=val)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if is_total:
                    cell.font = bold_font
                if isinstance(val, (int, float)) and col_idx > 6:
                    cell.number_format = "#,##0"
            if progress_callback and (row_idx % 500 == 0 or row_idx == total - 1):
                p = progress_start + (progress_end - progress_start) * ((row_idx + 1) / max(total, 1))
                progress_callback(p, f"寫入 {sheet_name} ({row_idx+1}/{total})...")
    else:
        ws.cell(row=2, column=1, value="無符合資料").font = Font(size=12)
    _auto_width(ws)


def generate_excel_new_sku(
    raw: pd.DataFrame,
    by_sku: pd.DataFrame,
    by_shop: pd.DataFrame,
    progress_callback=None,
    include_raw: bool = True,
) -> BytesIO:
    wb = Workbook()
    if progress_callback:
        progress_callback(0.0, "建立 Excel 結構...")

    out_cols = _ns_raw_columns(raw if raw is not None else pd.DataFrame())

    ws_raw = wb.active
    ws_raw.title = "RAW"
    if include_raw:
        _write_raw_sheet(ws_raw, raw, progress_callback, progress_start=0.0, progress_end=0.55, columns=out_cols)
    else:
        _write_raw_placeholder(ws_raw)

    if progress_callback:
        progress_callback(0.6, "建立 By SKU...")
    ws_sku = wb.create_sheet(title="By SKU")
    _write_ns_table(ws_sku, _NS_BY_SKU_COLUMNS, by_sku, progress_callback, 0.6, 0.75, "By SKU")

    if progress_callback:
        progress_callback(0.78, "建立 By Shop...")
    ws_shop = wb.create_sheet(title="By Shop")
    _write_ns_table(ws_shop, _NS_BY_SHOP_COLUMNS, by_shop, progress_callback, 0.78, 0.9, "By Shop")

    if progress_callback:
        progress_callback(0.95, "正在儲存 Excel 檔案...")

    gc.collect()
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    if progress_callback:
        progress_callback(1.0, "完成!")
    return buf
