# -*- coding: utf-8 -*-
"""銷售趨勢 Review：比較 M-2 與 M-1 銷量變動，找出上升與下跌品項。

趨勢按兩個完整月份（M-2 → M-1）判定；CM-MTD 為未完整月份，只作參考顯示。
只引用 modes.common，不引用其他模式模組，避免循環引用。
"""
import gc
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


from modes.common import (
    _auto_width,
    _clean_dataframe,
    _compute_raw_extra_columns,
    _extract_shop_name,
    _reorder_raw_columns,
    _resolve_column_name,
    _safe_default_series,
    _safe_numeric,
    _safe_sku,
    _write_raw_placeholder,
    _write_raw_sheet,
    validate_input,
)

__all__ = [
    "transform_sales_trend",
    "generate_excel_sales_trend",
    "_build_tr_by_sku",
    "_build_tr_by_shop",
]

_TR_UP = "上升"
_TR_DOWN = "下跌"
_TR_FLAT = "持平"
_TR_COLUMN = "趨勢分類"
_TR_RATE_COLUMN = "變動率 %"
_TR_THRESHOLD = 0.20

_TR_RAW_OUTPUT_COLUMNS = [
    "Site Code",
    "Shop Name",
    "OM",
    "Shop Nature",
    "Buyer",
    "Brand",
    "SKU",
    "SKU Desc",
    _TR_COLUMN,
    "Sold Qty (M-2)",
    "Sold Qty (M-1)",
    "Sold Qty (CM - MTD)",
    "Sold Qty (P90D)",
    "Shop Monthly Sold Qty (P90D)",
    "Shop Stock Qty (CM)",
]

_TR_BY_SKU_COLUMNS = [
    "Buyer",
    "Brand",
    "SKU",
    "SKU Desc",
    "趨勢分類",
    "M-2 合計",
    "M-1 合計",
    "變動率 %",
    "CM-MTD 合計",
    "P90D 月均",
    "店舖總庫存",
    "覆蓋日數",
]

_TR_BY_SHOP_COLUMNS = [
    "Site Code",
    "Shop Name",
    "OM",
    "涉及 SKU 數",
    "上升 SKU 數",
    "下跌 SKU 數",
    "持平 SKU 數",
]


def _tr_classify(m2: float, m1: float) -> tuple[str, float | None]:
    """M-2 → M-1 趨勢分類與變動率（小數）。"""
    if m2 <= 0:
        if m1 > 0:
            return _TR_UP, None
        return _TR_FLAT, None
    rate = (m1 - m2) / m2
    if rate >= _TR_THRESHOLD:
        return _TR_UP, rate
    if rate <= -_TR_THRESHOLD:
        return _TR_DOWN, rate
    return _TR_FLAT, rate


def _tr_raw_columns(raw: pd.DataFrame) -> list[str]:
    r = _reorder_raw_columns(_compute_raw_extra_columns(raw))
    available = set(r.columns)
    cols = [c for c in _TR_RAW_OUTPUT_COLUMNS if _resolve_column_name(c, available) is not None]
    if _TR_COLUMN not in cols:
        cols.insert(8, _TR_COLUMN)
    return cols


def _build_tr_by_sku(raw: pd.DataFrame) -> pd.DataFrame:
    cols = ["Buyer", "Brand", "SKU", "SKU Desc"]
    if raw is None or raw.empty:
        return pd.DataFrame(columns=_TR_BY_SKU_COLUMNS)
    df = raw.copy()
    df["SKU"] = df["SKU"].apply(_safe_sku)
    rows = []
    for (buyer, brand, sku), group in df.groupby(["Buyer", "Brand", "SKU"], dropna=False, observed=True):
        m2 = float(_safe_numeric(group.get("Sold Qty (M-2)", _safe_default_series(0, group.index))).sum())
        m1 = float(_safe_numeric(group.get("Sold Qty (M-1)", _safe_default_series(0, group.index))).sum())
        mtd = float(_safe_numeric(group.get("Sold Qty (CM - MTD)", _safe_default_series(0, group.index))).sum())
        p90d = float(_safe_numeric(group.get("Sold Qty (P90D)", _safe_default_series(0, group.index))).sum())
        if "Shop Monthly Sold Qty (P90D)" in group.columns:
            monthly = float(_safe_numeric(group["Shop Monthly Sold Qty (P90D)"]).sum())
        else:
            monthly = p90d / 90.0 * 30.5
        stock = float(_safe_numeric(group.get("Shop Stock Qty (CM)", _safe_default_series(0, group.index))).sum())
        cat, rate = _tr_classify(m2, m1)
        desc = group["SKU Desc"].dropna().astype(str)
        rows.append({
            "Buyer": buyer,
            "Brand": brand,
            "SKU": _safe_sku(sku),
            "SKU Desc": desc.iloc[0] if len(desc) else "",
            "趨勢分類": cat,
            "M-2 合計": int(round(m2)),
            "M-1 合計": int(round(m1)),
            "變動率 %": None if rate is None else int(round(rate * 100)),
            "CM-MTD 合計": int(round(mtd)),
            "P90D 月均": int(round(monthly if "Shop Monthly Sold Qty (P90D)" in group.columns else p90d / 90.0 * 30.5)),
            "店舖總庫存": int(round(stock)),
            "覆蓋日數": None if monthly <= 0 else int(round(stock / monthly * 30.5)),
        })
    by_sku = pd.DataFrame(rows, columns=_TR_BY_SKU_COLUMNS)
    order = {_TR_UP: 0, _TR_DOWN: 1, _TR_FLAT: 2}
    by_sku["_rank"] = by_sku["趨勢分類"].map(order).fillna(9).astype(int)
    by_sku["_abs"] = pd.to_numeric(by_sku["變動率 %"], errors="coerce").abs().fillna(0)
    by_sku = by_sku.sort_values(["_rank", "_abs", "Brand", "SKU"], ascending=[True, False, True, True], kind="stable")
    by_sku = by_sku.drop(columns=["_rank", "_abs"]).reset_index(drop=True)
    if not by_sku.empty:
        data = by_sku
        total_row = pd.DataFrame([{
            "Buyer": "總計",
            "Brand": None,
            "SKU": None,
            "SKU Desc": None,
            "趨勢分類": "",
            "M-2 合計": int(data["M-2 合計"].sum()),
            "M-1 合計": int(data["M-1 合計"].sum()),
            "變動率 %": None,
            "CM-MTD 合計": int(data["CM-MTD 合計"].sum()),
            "P90D 月均": int(data["P90D 月均"].sum()),
            "店舖總庫存": int(data["店舖總庫存"].sum()),
            "覆蓋日數": None,
        }])
        by_sku = pd.concat([by_sku, total_row], ignore_index=True)
    return by_sku[_TR_BY_SKU_COLUMNS]


def _build_tr_by_shop(raw: pd.DataFrame) -> pd.DataFrame:
    if raw is None or raw.empty:
        return pd.DataFrame(columns=_TR_BY_SHOP_COLUMNS)
    df = raw.copy()
    site_code = df.get("Site Code", _safe_default_series("", df.index)).fillna("").astype(str).str.strip()
    item_shop = df.get("Item_Shop", _safe_default_series("", df.index))
    df["Site Code"] = site_code
    df["Shop Name"] = [_extract_shop_name(item_shop.iloc[i], site_code.iloc[i]) for i in range(len(df))]
    df["OM"] = df.get("OM", _safe_default_series("", df.index)).fillna("").astype(str).str.strip()
    df["SKU"] = df["SKU"].apply(_safe_sku)
    rec = df[_TR_COLUMN].fillna("").astype(str)
    # 同一店舖同一 SKU 只計一次（RAW 為店舖×SKU 粒度，同一組合不應重複，防呆去重）
    site_sku = df.assign(_rec=rec.values).drop_duplicates(subset=["Site Code", "SKU"])
    grouped = (
        site_sku.groupby(["Site Code", "Shop Name", "OM"], dropna=False, observed=True)
        .agg(
            **{
                "涉及 SKU 數": ("SKU", "nunique"),
                "上升 SKU 數": ("_rec", lambda s: int((s == _TR_UP).sum())),
                "下跌 SKU 數": ("_rec", lambda s: int((s == _TR_DOWN).sum())),
                "持平 SKU 數": ("_rec", lambda s: int((s == _TR_FLAT).sum())),
            },
        )
        .reset_index()
    )
    for c in ["涉及 SKU 數", "上升 SKU 數", "下跌 SKU 數", "持平 SKU 數"]:
        grouped[c] = pd.to_numeric(grouped[c], errors="coerce").fillna(0).astype(int)
    grouped = grouped.sort_values(["涉及 SKU 數", "Site Code"], ascending=[False, True], kind="stable").reset_index(drop=True)
    if not grouped.empty:
        sku_once = df.drop_duplicates(subset=["SKU"])
        sku_rec = sku_once[_TR_COLUMN].fillna("").astype(str)
        total_row = pd.DataFrame([{
            "Site Code": "總計",
            "Shop Name": "",
            "OM": "",
            "涉及 SKU 數": int(df["SKU"].dropna().nunique()),
            "上升 SKU 數": int((sku_rec == _TR_UP).sum()),
            "下跌 SKU 數": int((sku_rec == _TR_DOWN).sum()),
            "持平 SKU 數": int((sku_rec == _TR_FLAT).sum()),
        }])
        grouped = pd.concat([grouped, total_row], ignore_index=True)
    return grouped[_TR_BY_SHOP_COLUMNS]


def transform_sales_trend(
    df_input: pd.DataFrame,
    progress_callback=None,
    shop_natures: list[str] | None = None,
):
    """銷售趨勢轉換：SKU 級 M-2 → M-1 分類，輸出 RAW、By SKU、By Shop 與提示訊息。"""
    if progress_callback:
        progress_callback(0.0, "驗證輸入資料...")
    missing = validate_input(df_input)
    if missing:
        raise ValueError(f"缺少必要欄位: {', '.join(missing)}")
    avail = set(df_input.columns)
    for col in ["Site Code", "Sold Qty (M-2)", "Sold Qty (M-1)", "Sold Qty (CM - MTD)"]:
        if _resolve_column_name(col, avail) is None:
            raise ValueError(f"缺少必需欄位: {col}")

    if progress_callback:
        progress_callback(0.1, "清理資料中...")
    df = _clean_dataframe(df_input, shop_natures=shop_natures, progress_callback=progress_callback)
    if df.empty:
        raise ValueError("過濾後無有效資料列，請確認來源檔案中的 Buyer 欄位是否包含有效數據")

    if progress_callback:
        progress_callback(0.2, "解像 SKU Desc / Shop Name...")
    df = _compute_raw_extra_columns(df)
    df["SKU"] = df["SKU"].apply(_safe_sku)

    m2 = _safe_numeric(df.get("Sold Qty (M-2)", _safe_default_series(0, df.index)))
    m1 = _safe_numeric(df.get("Sold Qty (M-1)", _safe_default_series(0, df.index)))
    mtd = _safe_numeric(df.get("Sold Qty (CM - MTD)", _safe_default_series(0, df.index)))
    sku_key = df["SKU"].fillna("").astype(str)
    sku_m2 = m2.groupby(sku_key, sort=False).transform("sum")
    sku_m1 = m1.groupby(sku_key, sort=False).transform("sum")
    sku_mtd = mtd.groupby(sku_key, sort=False).transform("sum")
    keep = ((sku_m2 > 0) | (sku_m1 > 0) | (sku_mtd > 0)).fillna(False)

    raw = df[keep.to_numpy()].copy()
    if not raw.empty:
        cats = [_tr_classify(a, b)[0] for a, b in zip(
            sku_m2[keep].values, sku_m1[keep].values,
        )]
        raw[_TR_COLUMN] = cats

    if progress_callback:
        progress_callback(0.35, "建立彙總...")
    by_sku = _build_tr_by_sku(raw)
    by_shop = _build_tr_by_shop(raw)

    warnings = []
    for col in ["Shop Stock Qty (CM)", "Sold Qty (P90D)", "Shop Monthly Sold Qty (P90D)"]:
        if col not in df.columns:
            warnings.append(f"來源檔案缺少「{col}」，相關參考欄以空值顯示。")
    if raw.empty:
        warnings.append("沒有符合「近三期任一期有銷售」的 SKU，所有品項皆長期零銷售。")

    if progress_callback:
        progress_callback(0.4, "完成資料轉換...")
    return raw, by_sku, by_shop, warnings


def _write_tr_table(ws, headers, data: pd.DataFrame, progress_callback=None, progress_start=0.0, progress_end=0.1, sheet_name=""):
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    bold_font = Font(bold=True, size=11)
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    if data is not None and not data.empty:
        total = len(data)
        for row_idx, (_, row_data) in enumerate(data.iterrows()):
            excel_row = row_idx + 2
            is_total = row_data.iloc[0] == "總計"
            for col_idx, h in enumerate(headers, 1):
                val = row_data.get(h)
                if pd.isna(val):
                    val = None
                cell = ws.cell(row=excel_row, column=col_idx, value=val)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if is_total:
                    cell.font = bold_font
                if isinstance(val, (int, float)) and col_idx > 5:
                    cell.number_format = "#,##0"
            if progress_callback and (row_idx % 500 == 0 or row_idx == total - 1):
                p = progress_start + (progress_end - progress_start) * ((row_idx + 1) / max(total, 1))
                progress_callback(p, f"寫入 {sheet_name} ({row_idx+1}/{total})...")
    else:
        ws.cell(row=2, column=1, value="無符合資料").font = Font(size=12)
    _auto_width(ws)


def generate_excel_sales_trend(
    raw: pd.DataFrame,
    by_sku: pd.DataFrame,
    by_shop: pd.DataFrame,
    progress_callback=None,
    include_raw: bool = True,
) -> BytesIO:
    wb = Workbook()
    if progress_callback:
        progress_callback(0.0, "建立 Excel 結構...")

    out_cols = _tr_raw_columns(raw if raw is not None else pd.DataFrame())

    ws_raw = wb.active
    ws_raw.title = "RAW"
    if include_raw:
        _write_raw_sheet(ws_raw, raw, progress_callback, progress_start=0.0, progress_end=0.55, columns=out_cols)
    else:
        _write_raw_placeholder(ws_raw)

    if progress_callback:
        progress_callback(0.6, "建立 By SKU...")
    ws_sku = wb.create_sheet(title="By SKU")
    _write_tr_table(ws_sku, _TR_BY_SKU_COLUMNS, by_sku, progress_callback, 0.6, 0.75, "By SKU")

    if progress_callback:
        progress_callback(0.78, "建立 By Shop...")
    ws_shop = wb.create_sheet(title="By Shop")
    _write_tr_table(ws_shop, _TR_BY_SHOP_COLUMNS, by_shop, progress_callback, 0.78, 0.9, "By Shop")

    if progress_callback:
        progress_callback(0.95, "正在儲存 Excel 檔案...")

    gc.collect()
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    if progress_callback:
        progress_callback(1.0, "完成!")
    return buf
