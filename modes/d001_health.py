# -*- coding: utf-8 -*-
"""D001 健康檢查 Review：檢查倉庫備貨與店舖需求是否脫節。

SKU 級檢查：D001 缺貨需補／倉庫積壓／覆蓋偏低三類結論。
只引用 modes.common，不引用其他模式模組，避免循環引用。
"""
import gc
from io import BytesIO

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


from modes.common import (
    _auto_width,
    _clean_dataframe,
    _compute_raw_extra_columns,
    _extract_shop_name,
    _reorder_raw_columns,
    _resolve_column_name,
    _safe_default_series,
    _safe_numeric,
    _safe_sku,
    _write_raw_placeholder,
    _write_raw_sheet,
    validate_input,
)

__all__ = [
    "transform_d001_health",
    "generate_excel_d001_health",
    "_build_d001_check_table",
    "_build_d001_by_buyer",
]

_D001_SHORTAGE = "D001 缺貨需補"
_D001_OVERSTOCK = "倉庫積壓"
_D001_LOW_COVER = "覆蓋偏低"
_D001_CHECK_COLUMN = "檢查結論"

_D001_RAW_OUTPUT_COLUMNS = [
    "Site Code",
    "Shop Name",
    "OM",
    "Shop Nature",
    "Buyer",
    "Brand",
    "SKU",
    "SKU Desc",
    "Supply Source",
    "D001 Net Stock Qty (CM)",
    "D001 On Order PO Qty (CM)",
    "Shop On Order PO Qty (CM)",
    "Shop Pending Receive",
    "Shop Stock Qty (CM)",
    "Sold Qty (P90D)",
    "Shop Monthly Sold Qty (P90D)",
]

_D001_CHECK_COLUMNS = [
    "Buyer",
    "Brand",
    "SKU",
    "SKU Desc",
    "Supply Source",
    _D001_CHECK_COLUMN,
    "D001 Net Stock Qty (CM)",
    "D001 On Order PO Qty (CM)",
    "在途在訂合計",
    "店舖總庫存",
    "店舖總月銷",
    "D001 覆蓋日數",
]

_D001_BY_BUYER_COLUMNS = [
    "Buyer",
    "需跟進 SKU 數",
    "D001 缺貨數",
    "倉庫積壓數",
    "覆蓋偏低數",
]


def _d001_raw_columns(raw: pd.DataFrame) -> list[str]:
    r = _reorder_raw_columns(_compute_raw_extra_columns(raw))
    available = set(r.columns)
    return [c for c in _D001_RAW_OUTPUT_COLUMNS if _resolve_column_name(c, available) is not None]


def _d001_sku_value(group: pd.DataFrame, column: str) -> float:
    """D001 欄逐 SKU 取值：同 SKU 各列重複填報，取最大值（空白視為 0）。"""
    if column not in group.columns:
        return 0.0
    vals = pd.to_numeric(group[column], errors="coerce").fillna(0.0)
    return float(vals.max()) if len(vals) else 0.0


def _build_d001_check_table(df: pd.DataFrame) -> pd.DataFrame:
    """SKU 級檢查主表：只收錄三類異常 SKU，按結論嚴重度與 D001 覆蓋排序。"""
    if df is None or df.empty:
        return pd.DataFrame(columns=_D001_CHECK_COLUMNS)
    rows = []
    for (buyer, brand, sku), group in df.groupby(["Buyer", "Brand", "SKU"], dropna=False, observed=True):
        sku_text = _safe_sku(sku)
        desc = group["SKU Desc"].dropna().astype(str)
        desc_text = desc.iloc[0] if len(desc) else ""
        supply = group.get("Supply Source", _safe_default_series("", group.index)).fillna("").astype(str).str.strip()
        supply_mode = supply[supply != ""].mode()
        supply_text = str(supply_mode.iloc[0]) if len(supply_mode) else ""
        if supply_text.lower() == "vendor":
            continue
        d001 = _d001_sku_value(group, "D001 Net Stock Qty (CM)")
        d001_po = _d001_sku_value(group, "D001 On Order PO Qty (CM)")
        shop_po = float(_safe_numeric(group.get("Shop On Order PO Qty (CM)", _safe_default_series(0, group.index))).sum())
        pending = 0.0
        for cand in ("Shop Pending Receive (±7 days)", "Shop Pending Receive"):
            if cand in group.columns:
                pending = float(_safe_numeric(group[cand]).sum())
                break
        shop_stock = float(_safe_numeric(group.get("Shop Stock Qty (CM)", _safe_default_series(0, group.index))).sum())
        if "Shop Monthly Sold Qty (P90D)" in group.columns:
            monthly = float(_safe_numeric(group["Shop Monthly Sold Qty (P90D)"]).sum())
        else:
            monthly = float(_safe_numeric(group.get("Sold Qty (P90D)", _safe_default_series(0, group.index))).sum()) / 90.0 * 30.5
        pipeline = d001_po + shop_po + pending
        if monthly > 0:
            cover_days = (d001 / monthly) * 30.5
        else:
            cover_days = np.inf
        conclusion = ""
        if d001 <= 0 and monthly > 0:
            conclusion = _D001_SHORTAGE
        elif d001 > 0 and monthly <= 0:
            conclusion = _D001_OVERSTOCK
        elif monthly > 0 and (d001 + pipeline) < monthly:
            conclusion = _D001_LOW_COVER
        if not conclusion:
            continue
        rows.append({
            "Buyer": buyer,
            "Brand": brand,
            "SKU": sku_text,
            "SKU Desc": desc_text,
            "Supply Source": supply_text,
            _D001_CHECK_COLUMN: conclusion,
            "D001 Net Stock Qty (CM)": int(round(d001)),
            "D001 On Order PO Qty (CM)": int(round(d001_po)),
            "在途在訂合計": int(round(pipeline)),
            "店舖總庫存": int(round(shop_stock)),
            "店舖總月銷": int(round(monthly)),
            "D001 覆蓋日數": None if not np.isfinite(cover_days) else int(round(cover_days)),
        })
    check = pd.DataFrame(rows, columns=_D001_CHECK_COLUMNS)
    if check.empty:
        return check
    order = {_D001_SHORTAGE: 0, _D001_OVERSTOCK: 1, _D001_LOW_COVER: 2}
    check["_rank"] = check[_D001_CHECK_COLUMN].map(order).fillna(9).astype(int)
    check = check.sort_values(["_rank", "D001 覆蓋日數", "Brand", "SKU"], ascending=[True, True, True, True], kind="stable")
    return check.drop(columns=["_rank"]).reset_index(drop=True)


def _build_d001_by_buyer(check: pd.DataFrame) -> pd.DataFrame:
    if check is None or check.empty:
        return pd.DataFrame(columns=_D001_BY_BUYER_COLUMNS)
    rec = check[_D001_CHECK_COLUMN].fillna("").astype(str)
    check = check.copy()
    check["_shortage"] = (rec == _D001_SHORTAGE).astype(int)
    check["_overstock"] = (rec == _D001_OVERSTOCK).astype(int)
    check["_low"] = (rec == _D001_LOW_COVER).astype(int)
    grouped = (
        check.groupby(["Buyer"], dropna=False, observed=True)
        .agg(
            需跟進SKU數=("SKU", "nunique"),
            D001缺貨數=("_shortage", "sum"),
            倉庫積壓數=("_overstock", "sum"),
            覆蓋偏低數=("_low", "sum"),
        )
        .reset_index()
    )
    grouped = grouped.rename(columns={
        "需跟進SKU數": "需跟進 SKU 數",
        "D001缺貨數": "D001 缺貨數",
        "倉庫積壓數": "倉庫積壓數",
        "覆蓋偏低數": "覆蓋偏低數",
    })
    for c in ["需跟進 SKU 數", "D001 缺貨數", "倉庫積壓數", "覆蓋偏低數"]:
        grouped[c] = pd.to_numeric(grouped[c], errors="coerce").fillna(0).astype(int)
    grouped = grouped.sort_values(["需跟進 SKU 數", "Buyer"], ascending=[False, True], kind="stable").reset_index(drop=True)
    total_row = pd.DataFrame([{
        "Buyer": "總計",
        "需跟進 SKU 數": int(check["SKU"].dropna().nunique()),
        "D001 缺貨數": int(grouped["D001 缺貨數"].sum()),
        "倉庫積壓數": int(grouped["倉庫積壓數"].sum()),
        "覆蓋偏低數": int(grouped["覆蓋偏低數"].sum()),
    }])
    return pd.concat([grouped, total_row], ignore_index=True)[_D001_BY_BUYER_COLUMNS]


def transform_d001_health(
    df_input: pd.DataFrame,
    progress_callback=None,
    shop_natures: list[str] | None = None,
):
    """D001 健康檢查轉換：SKU 級三類結論，輸出檢查表、By Buyer、RAW 與提示訊息。"""
    if progress_callback:
        progress_callback(0.0, "驗證輸入資料...")
    missing = validate_input(df_input)
    if missing:
        raise ValueError(f"缺少必要欄位: {', '.join(missing)}")
    avail = set(df_input.columns)
    for col in ["Site Code", "D001 Net Stock Qty (CM)", "Shop Stock Qty (CM)", "Sold Qty (P90D)"]:
        if _resolve_column_name(col, avail) is None:
            raise ValueError(f"缺少必需欄位: {col}")

    if progress_callback:
        progress_callback(0.1, "清理資料中...")
    df = _clean_dataframe(df_input, shop_natures=shop_natures, progress_callback=progress_callback)
    if df.empty:
        raise ValueError("過濾後無有效資料列，請確認來源檔案中的 Buyer 欄位是否包含有效數據")

    if progress_callback:
        progress_callback(0.2, "解像 SKU Desc / Shop Name...")
    df = _compute_raw_extra_columns(df)
    df["SKU"] = df["SKU"].apply(_safe_sku)

    if progress_callback:
        progress_callback(0.3, "檢查倉庫備貨...")
    check = _build_d001_check_table(df)
    by_buyer = _build_d001_by_buyer(check)

    if not check.empty:
        hit = set(check["SKU"].dropna().astype(str))
        raw = df[df["SKU"].astype(str).isin(hit)].copy()
    else:
        raw = df.iloc[0:0].copy()

    warnings = []
    for col in ["D001 On Order PO Qty (CM)", "Shop On Order PO Qty (CM)", "Shop Pending Receive"]:
        if col not in df.columns and not (col == "Shop Pending Receive" and "Shop Pending Receive (±7 days)" in df.columns):
            warnings.append(f"來源檔案缺少「{col}」，在途在訂合計以 0 計。")
    if "Supply Source" not in df.columns:
        warnings.append("來源檔案缺少「Supply Source」，全部 SKU 按 Warehouse 貨檢查。")
    if check.empty:
        warnings.append("沒有符合「D001 缺貨需補／倉庫積壓／覆蓋偏低」的 SKU，倉庫備貨與店舖需求一致。")

    if progress_callback:
        progress_callback(0.4, "完成資料轉換...")
    return check, by_buyer, raw, warnings


def _write_d001_table(ws, headers, data: pd.DataFrame, progress_callback=None, progress_start=0.0, progress_end=0.1, sheet_name=""):
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    bold_font = Font(bold=True, size=11)
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    if data is not None and not data.empty:
        total = len(data)
        for row_idx, (_, row_data) in enumerate(data.iterrows()):
            excel_row = row_idx + 2
            is_total = row_data.iloc[0] == "總計"
            for col_idx, h in enumerate(headers, 1):
                val = row_data.get(h)
                if pd.isna(val):
                    val = None
                cell = ws.cell(row=excel_row, column=col_idx, value=val)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if is_total:
                    cell.font = bold_font
                if isinstance(val, (int, float)) and col_idx > 5:
                    cell.number_format = "#,##0"
            if progress_callback and (row_idx % 500 == 0 or row_idx == total - 1):
                p = progress_start + (progress_end - progress_start) * ((row_idx + 1) / max(total, 1))
                progress_callback(p, f"寫入 {sheet_name} ({row_idx+1}/{total})...")
    else:
        ws.cell(row=2, column=1, value="無符合資料").font = Font(size=12)
    _auto_width(ws)


def generate_excel_d001_health(
    check: pd.DataFrame,
    by_buyer: pd.DataFrame,
    raw: pd.DataFrame,
    progress_callback=None,
    include_raw: bool = True,
) -> BytesIO:
    wb = Workbook()
    if progress_callback:
        progress_callback(0.0, "建立 Excel 結構...")

    ws_check = wb.active
    ws_check.title = "D001 檢查"
    _write_d001_table(ws_check, _D001_CHECK_COLUMNS, check, progress_callback, 0.0, 0.3, "D001 檢查")

    if progress_callback:
        progress_callback(0.35, "建立 By Buyer...")
    ws_buyer = wb.create_sheet(title="By Buyer")
    _write_d001_table(ws_buyer, _D001_BY_BUYER_COLUMNS, by_buyer, progress_callback, 0.35, 0.5, "By Buyer")

    out_cols = _d001_raw_columns(raw if raw is not None else pd.DataFrame())
    ws_raw = wb.create_sheet(title="RAW")
    if include_raw:
        _write_raw_sheet(ws_raw, raw, progress_callback, progress_start=0.5, progress_end=0.9, columns=out_cols)
    else:
        _write_raw_placeholder(ws_raw)

    if progress_callback:
        progress_callback(0.95, "正在儲存 Excel 檔案...")

    gc.collect()
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    if progress_callback:
        progress_callback(1.0, "完成!")
    return buf
