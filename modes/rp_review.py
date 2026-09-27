# -*- coding: utf-8 -*-
"""RP Type 覆核 Review：檢查 RF／ND 設定與實際銷量是否相符，找出疑似誤設列。

只引用 modes.common，不引用其他模式模組，避免循環引用。
"""
import gc
from io import BytesIO

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


from modes.common import (
    _auto_width,
    _clean_dataframe,
    _compute_raw_extra_columns,
    _extract_shop_name,
    _reorder_raw_columns,
    _resolve_column_name,
    _safe_default_series,
    _safe_numeric,
    _safe_sku,
    _write_raw_placeholder,
    _write_raw_sheet,
    validate_input,
)

__all__ = [
    "transform_rp_review",
    "generate_excel_rp_review",
    "_build_rp_by_sku",
    "_build_rp_by_shop",
]

_RP_REVIEW_COLUMN = "覆核建議"
_RP_TO_RF = "疑似應轉 RF"
_RP_TO_ND = "疑似應轉 ND"
_RP_ND_OOS = "ND 缺貨跟進"

_RP_RAW_OUTPUT_COLUMNS = [
    "Site Code",
    "Shop Name",
    "OM",
    "Shop Nature",
    "Buyer",
    "Brand",
    "SKU",
    "SKU Desc",
    "RP Type",
    _RP_REVIEW_COLUMN,
    "Shop Stock Status",
    "Shop Stock Qty (CM)",
    "Sold Qty (M-2)",
    "Sold Qty (M-1)",
    "Sold Qty (CM - MTD)",
    "Sold Qty (P90D)",
    "Shop Monthly Sold Qty (P90D)",
    "Is OOS (CM)",
    "OOS Periods Count",
]

_RP_BY_SKU_COLUMNS = [
    "Buyer",
    "Brand",
    "SKU",
    "SKU Desc",
    "涉及店舖數",
    "應轉 RF 店數",
    "應轉 ND 店數",
    "ND 缺貨店數",
]

_RP_BY_SHOP_COLUMNS = [
    "Site Code",
    "Shop Name",
    "OM",
    "涉及 SKU 數",
    "應轉 RF 數",
    "應轉 ND 數",
    "ND 缺貨數",
]


def _rp_as_true(series: pd.Series) -> pd.Series:
    """OOS 旗標正規化：接受布林、數值與常見文字（TRUE/Y/YES/1/是）。"""
    if series.dtype == bool:
        return series.fillna(False)
    txt = series.fillna("").astype(str).str.strip().str.lower()
    return txt.isin({"true", "t", "y", "yes", "1", "是"})


def _rp_raw_columns(raw: pd.DataFrame) -> list[str]:
    r = _reorder_raw_columns(_compute_raw_extra_columns(raw))
    available = set(r.columns)
    cols = [c for c in _RP_RAW_OUTPUT_COLUMNS if _resolve_column_name(c, available) is not None]
    if _RP_REVIEW_COLUMN not in cols:
        cols.insert(9, _RP_REVIEW_COLUMN)
    return cols


def _build_rp_by_sku(raw: pd.DataFrame) -> pd.DataFrame:
    cols = ["Buyer", "Brand", "SKU", "SKU Desc"]
    if raw is None or raw.empty:
        return pd.DataFrame(columns=_RP_BY_SKU_COLUMNS)
    df = raw.copy()
    df["SKU"] = df["SKU"].apply(_safe_sku)
    rec = df[_RP_REVIEW_COLUMN].fillna("").astype(str)
    df["_to_rf"] = (rec == _RP_TO_RF).astype(int)
    df["_to_nd"] = (rec == _RP_TO_ND).astype(int)
    df["_nd_oos"] = (rec == _RP_ND_OOS).astype(int)
    grouped = (
        df.groupby(cols, dropna=False, observed=True)
        .agg(
            涉及店舖數=("Site Code", "nunique"),
            應轉RF店數=("_to_rf", "sum"),
            應轉ND店數=("_to_nd", "sum"),
            ND缺貨店數=("_nd_oos", "sum"),
        )
        .reset_index()
    )
    grouped = grouped.rename(columns={
        "應轉RF店數": "應轉 RF 店數",
        "應轉ND店數": "應轉 ND 店數",
        "ND缺貨店數": "ND 缺貨店數",
    })
    for c in ["涉及店舖數", "應轉 RF 店數", "應轉 ND 店數", "ND 缺貨店數"]:
        grouped[c] = pd.to_numeric(grouped[c], errors="coerce").fillna(0).astype(int)
    grouped = grouped.sort_values(
        ["涉及店舖數", "應轉 RF 店數", "Brand", "SKU"],
        ascending=[False, False, True, True],
        kind="stable",
    ).reset_index(drop=True)
    if not grouped.empty:
        total_row = pd.DataFrame([{
            "Buyer": "總計",
            "Brand": None,
            "SKU": None,
            "SKU Desc": None,
            "涉及店舖數": int(df["Site Code"].dropna().nunique()),
            "應轉 RF 店數": int(grouped["應轉 RF 店數"].sum()),
            "應轉 ND 店數": int(grouped["應轉 ND 店數"].sum()),
            "ND 缺貨店數": int(grouped["ND 缺貨店數"].sum()),
        }])
        grouped = pd.concat([grouped, total_row], ignore_index=True)
    return grouped[_RP_BY_SKU_COLUMNS]


def _build_rp_by_shop(raw: pd.DataFrame) -> pd.DataFrame:
    if raw is None or raw.empty:
        return pd.DataFrame(columns=_RP_BY_SHOP_COLUMNS)
    df = raw.copy()
    site_code = df.get("Site Code", _safe_default_series("", df.index)).fillna("").astype(str).str.strip()
    item_shop = df.get("Item_Shop", _safe_default_series("", df.index))
    df["Site Code"] = site_code
    df["Shop Name"] = [_extract_shop_name(item_shop.iloc[i], site_code.iloc[i]) for i in range(len(df))]
    df["OM"] = df.get("OM", _safe_default_series("", df.index)).fillna("").astype(str).str.strip()
    df["SKU"] = df["SKU"].apply(_safe_sku)
    rec = df[_RP_REVIEW_COLUMN].fillna("").astype(str)
    df["_to_rf"] = (rec == _RP_TO_RF).astype(int)
    df["_to_nd"] = (rec == _RP_TO_ND).astype(int)
    df["_nd_oos"] = (rec == _RP_ND_OOS).astype(int)
    grouped = (
        df.groupby(["Site Code", "Shop Name", "OM"], dropna=False, observed=True)
        .agg(
            涉及SKU數=("SKU", "nunique"),
            應轉RF數=("_to_rf", "sum"),
            應轉ND數=("_to_nd", "sum"),
            ND缺貨數=("_nd_oos", "sum"),
        )
        .reset_index()
    )
    grouped = grouped.rename(columns={
        "涉及SKU數": "涉及 SKU 數",
        "應轉RF數": "應轉 RF 數",
        "應轉ND數": "應轉 ND 數",
        "ND缺貨數": "ND 缺貨數",
    })
    for c in ["涉及 SKU 數", "應轉 RF 數", "應轉 ND 數", "ND 缺貨數"]:
        grouped[c] = pd.to_numeric(grouped[c], errors="coerce").fillna(0).astype(int)
    grouped = grouped.sort_values(
        ["涉及 SKU 數", "應轉 RF 數", "Site Code"],
        ascending=[False, False, True],
        kind="stable",
    ).reset_index(drop=True)
    if not grouped.empty:
        total_row = pd.DataFrame([{
            "Site Code": "總計",
            "Shop Name": "",
            "OM": "",
            "涉及 SKU 數": int(df["SKU"].dropna().nunique()),
            "應轉 RF 數": int(grouped["應轉 RF 數"].sum()),
            "應轉 ND 數": int(grouped["應轉 ND 數"].sum()),
            "ND 缺貨數": int(grouped["ND 缺貨數"].sum()),
        }])
        grouped = pd.concat([grouped, total_row], ignore_index=True)
    return grouped[_RP_BY_SHOP_COLUMNS]


def transform_rp_review(
    df_input: pd.DataFrame,
    progress_callback=None,
    shop_natures: list[str] | None = None,
):
    """RP Type 覆核轉換：標示三類疑似誤設，輸出 RAW、By SKU、By Shop 與提示訊息。"""
    if progress_callback:
        progress_callback(0.0, "驗證輸入資料...")
    missing = validate_input(df_input)
    if missing:
        raise ValueError(f"缺少必要欄位: {', '.join(missing)}")
    avail = set(df_input.columns)
    if _resolve_column_name("Site Code", avail) is None:
        raise ValueError("缺少必需欄位: Site Code")
    if _resolve_column_name("RP Type", avail) is None:
        raise ValueError("缺少必需欄位: RP Type（本模式覆核 RF／ND 設定，缺少時無法執行）")

    if progress_callback:
        progress_callback(0.1, "清理資料中...")
    df = _clean_dataframe(df_input, shop_natures=shop_natures, progress_callback=progress_callback)
    if df.empty:
        raise ValueError("過濾後無有效資料列，請確認來源檔案中的 Buyer 欄位是否包含有效數據")

    if progress_callback:
        progress_callback(0.2, "解像 SKU Desc / Shop Name...")
    df = _compute_raw_extra_columns(df)
    df["SKU"] = df["SKU"].apply(_safe_sku)

    rp_raw = df["RP Type"].fillna("").astype(str).str.strip().str.upper()
    rp = np.where(rp_raw == "RF", "RF", "ND")
    is_rf = rp_raw == "RF"
    is_nd = ~is_rf

    sold_m2 = _safe_numeric(df.get("Sold Qty (M-2)", _safe_default_series(0, df.index)))
    sold_m1 = _safe_numeric(df.get("Sold Qty (M-1)", _safe_default_series(0, df.index)))
    sold_mtd = _safe_numeric(df.get("Sold Qty (CM - MTD)", _safe_default_series(0, df.index)))
    sold_p90d = _safe_numeric(df.get("Sold Qty (P90D)", _safe_default_series(0, df.index)))
    stock_cm = _safe_numeric(df.get("Shop Stock Qty (CM)", _safe_default_series(0, df.index)))
    if "Is OOS (CM)" in df.columns:
        oos_cm = _rp_as_true(df["Is OOS (CM)"])
    else:
        oos_cm = pd.Series(False, index=df.index)

    has_recent_sales = ((sold_p90d > 0) | (sold_m1 > 0) | (sold_mtd > 0)).fillna(False)
    all_zero_sales = ((sold_m2 <= 0) & (sold_m1 <= 0) & (sold_mtd <= 0) & (sold_p90d <= 0)).fillna(False)
    is_oos = ((stock_cm <= 0) | oos_cm).fillna(False)

    rec = np.full(len(df), "", dtype=object)
    # 優先次序：ND 缺貨跟進 優先於 ND 有銷量（缺貨且 ND 屬高風險，先跟進供貨安排）
    rec[is_nd & is_oos] = _RP_ND_OOS
    rec[is_nd & ~is_oos & has_recent_sales.values] = _RP_TO_RF
    rec[is_rf & all_zero_sales.values & (stock_cm > 0).values] = _RP_TO_ND
    df[_RP_REVIEW_COLUMN] = rec

    mask = pd.Series(rec != "", index=df.index)
    raw = df[mask].copy()

    if progress_callback:
        progress_callback(0.35, "建立彙總...")
    by_sku = _build_rp_by_sku(raw)
    by_shop = _build_rp_by_shop(raw)

    warnings = []
    mapped = int((~rp_raw.isin(["RF", "ND"])).sum())
    if mapped:
        warnings.append(f"「RP Type」有 {mapped} 列為非標準值（非 RF／ND），已按 ND 處理（保守視為不可補貨）。")
    if raw.empty:
        warnings.append("沒有符合「疑似誤設（應轉 RF／應轉 ND／ND 缺貨）」的資料，設定與銷量一致。")

    if progress_callback:
        progress_callback(0.4, "完成資料轉換...")
    return raw, by_sku, by_shop, warnings


def _write_rp_table(ws, headers, data: pd.DataFrame, progress_callback=None, progress_start=0.0, progress_end=0.1, sheet_name=""):
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    bold_font = Font(bold=True, size=11)
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    if data is not None and not data.empty:
        total = len(data)
        for row_idx, (_, row_data) in enumerate(data.iterrows()):
            excel_row = row_idx + 2
            is_total = row_data.iloc[0] == "總計"
            for col_idx, h in enumerate(headers, 1):
                val = row_data.get(h)
                if pd.isna(val):
                    val = None
                cell = ws.cell(row=excel_row, column=col_idx, value=val)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if is_total:
                    cell.font = bold_font
                if isinstance(val, (int, float)) and col_idx > 4:
                    cell.number_format = "#,##0"
            if progress_callback and (row_idx % 500 == 0 or row_idx == total - 1):
                p = progress_start + (progress_end - progress_start) * ((row_idx + 1) / max(total, 1))
                progress_callback(p, f"寫入 {sheet_name} ({row_idx+1}/{total})...")
    else:
        ws.cell(row=2, column=1, value="無符合資料").font = Font(size=12)
    _auto_width(ws)


def generate_excel_rp_review(
    raw: pd.DataFrame,
    by_sku: pd.DataFrame,
    by_shop: pd.DataFrame,
    progress_callback=None,
    include_raw: bool = True,
) -> BytesIO:
    wb = Workbook()
    if progress_callback:
        progress_callback(0.0, "建立 Excel 結構...")

    out_cols = _rp_raw_columns(raw if raw is not None else pd.DataFrame())

    ws_raw = wb.active
    ws_raw.title = "RAW"
    if include_raw:
        _write_raw_sheet(ws_raw, raw, progress_callback, progress_start=0.0, progress_end=0.55, columns=out_cols)
    else:
        _write_raw_placeholder(ws_raw)

    if progress_callback:
        progress_callback(0.6, "建立 By SKU...")
    ws_sku = wb.create_sheet(title="By SKU")
    _write_rp_table(ws_sku, _RP_BY_SKU_COLUMNS, by_sku, progress_callback, 0.6, 0.75, "By SKU")

    if progress_callback:
        progress_callback(0.78, "建立 By Shop...")
    ws_shop = wb.create_sheet(title="By Shop")
    _write_rp_table(ws_shop, _RP_BY_SHOP_COLUMNS, by_shop, progress_callback, 0.78, 0.9, "By Shop")

    if progress_callback:
        progress_callback(0.95, "正在儲存 Excel 檔案...")

    gc.collect()
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    if progress_callback:
        progress_callback(1.0, "完成!")
    return buf
