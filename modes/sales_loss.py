# -*- coding: utf-8 -*-
"""銷售損失量化 Review：彙總缺貨造成的銷售損失，排列補貨優先次序。

只引用 modes.common，不引用其他模式模組，避免循環引用。
"""
import gc
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


from modes.common import (
    _auto_width,
    _clean_dataframe,
    _compute_raw_extra_columns,
    _extract_shop_name,
    _reorder_raw_columns,
    _resolve_column_name,
    _safe_default_series,
    _safe_numeric,
    _safe_sku,
    _write_raw_placeholder,
    _write_raw_sheet,
    validate_input,
)
__all__ = [
    "transform_sales_loss",
    "generate_excel_sales_loss",
    "_build_sales_loss_by_sku",
    "_build_sales_loss_by_shop",
]

_LOSS_SIGNAL_COLUMNS = [
    "Is Stock Impacted",
    "OOS Periods Count",
    "Total Sales Loss (M-2)",
    "Total Sales Loss (P90D)",
    "Max Sales Loss (M-2)",
    "Max Sales Loss (P90D)",
]

_SALES_LOSS_RAW_OUTPUT_COLUMNS = [
    "Site Code",
    "Shop Name",
    "OM",
    "Shop Nature",
    "Buyer",
    "Brand",
    "SKU",
    "SKU Desc",
    "RP Type",
    "Shop Stock Status",
    "Is OOS (M-2)",
    "Is OOS (M-1)",
    "Is OOS (CM)",
    "OOS Periods Count",
    "OOS Pattern",
    "Is Stock Impacted",
    "Max Sales Loss (M-2)",
    "Max Sales Loss (P90D)",
    "Total Sales Loss (M-2)",
    "Total Sales Loss (P90D)",
    "Shop Stock Qty (CM)",
    "Sold Qty (M-2)",
    "Sold Qty (M-1)",
    "Sold Qty (CM - MTD)",
    "Sold Qty (P90D)",
]

_BY_SKU_COLUMNS = [
    "Buyer",
    "Brand",
    "SKU",
    "SKU Desc",
    "影響店舖數",
    "OOS Periods 合計",
    "Total Sales Loss (M-2) 合計",
    "Total Sales Loss (P90D) 合計",
]

_BY_SHOP_COLUMNS = [
    "Site Code",
    "Shop Name",
    "OM",
    "影響 SKU 數",
    "Total Sales Loss (M-2) 總數",
    "Total Sales Loss (P90D) 總數",
]


def _as_bool(series: pd.Series) -> pd.Series:
    """Is Stock Impacted 正規化：接受布林、數值與常見文字（TRUE/Y/YES/1/是）。"""
    if series.dtype == bool:
        return series.fillna(False)
    txt = series.fillna("").astype(str).str.strip().str.lower()
    return txt.isin({"true", "t", "y", "yes", "1", "是"})


def _sales_loss_raw_columns(raw: pd.DataFrame) -> list[str]:
    r = _reorder_raw_columns(_compute_raw_extra_columns(raw))
    available = set(r.columns)
    return [c for c in _SALES_LOSS_RAW_OUTPUT_COLUMNS if _resolve_column_name(c, available) is not None]


def _build_sales_loss_by_sku(raw: pd.DataFrame) -> pd.DataFrame:
    cols = ["Buyer", "Brand", "SKU", "SKU Desc"]
    if raw is None or raw.empty:
        return pd.DataFrame(columns=_BY_SKU_COLUMNS)
    df = raw.copy()
    df["SKU"] = df["SKU"].apply(_safe_sku)
    df["_oos_periods"] = _safe_numeric(df.get("OOS Periods Count", _safe_default_series(0, df.index)))
    df["_loss_m2"] = _safe_numeric(df.get("Total Sales Loss (M-2)", _safe_default_series(0, df.index)))
    df["_loss_p90d"] = _safe_numeric(df.get("Total Sales Loss (P90D)", _safe_default_series(0, df.index)))
    grouped = (
        df.groupby(cols, dropna=False, observed=True)
        .agg(
            影響店舖數=("Site Code", "nunique"),
            OOS_Periods=("_oos_periods", "sum"),
            Loss_M2=("_loss_m2", "sum"),
            Loss_P90D=("_loss_p90d", "sum"),
        )
        .reset_index()
    )
    grouped = grouped.rename(columns={
        "OOS_Periods": "OOS Periods 合計",
        "Loss_M2": "Total Sales Loss (M-2) 合計",
        "Loss_P90D": "Total Sales Loss (P90D) 合計",
    })
    for c in ["影響店舖數", "OOS Periods 合計", "Total Sales Loss (M-2) 合計", "Total Sales Loss (P90D) 合計"]:
        grouped[c] = pd.to_numeric(grouped[c], errors="coerce").fillna(0).round().astype(int)
    grouped = grouped.sort_values(
        ["Total Sales Loss (P90D) 合計", "Total Sales Loss (M-2) 合計", "Brand", "SKU"],
        ascending=[False, False, True, True],
        kind="stable",
    ).reset_index(drop=True)
    if not grouped.empty:
        total_row = pd.DataFrame([{
            "Buyer": "總計",
            "Brand": None,
            "SKU": None,
            "SKU Desc": None,
            "影響店舖數": int(df["Site Code"].dropna().nunique()),
            "OOS Periods 合計": int(grouped["OOS Periods 合計"].sum()),
            "Total Sales Loss (M-2) 合計": int(grouped["Total Sales Loss (M-2) 合計"].sum()),
            "Total Sales Loss (P90D) 合計": int(grouped["Total Sales Loss (P90D) 合計"].sum()),
        }])
        grouped = pd.concat([grouped, total_row], ignore_index=True)
    return grouped[_BY_SKU_COLUMNS]


def _build_sales_loss_by_shop(raw: pd.DataFrame) -> pd.DataFrame:
    if raw is None or raw.empty:
        return pd.DataFrame(columns=_BY_SHOP_COLUMNS)
    df = raw.copy()
    site_code = df.get("Site Code", _safe_default_series("", df.index)).fillna("").astype(str).str.strip()
    item_shop = df.get("Item_Shop", _safe_default_series("", df.index))
    df["Site Code"] = site_code
    df["Shop Name"] = [_extract_shop_name(item_shop.iloc[i], site_code.iloc[i]) for i in range(len(df))]
    df["OM"] = df.get("OM", _safe_default_series("", df.index)).fillna("").astype(str).str.strip()
    df["SKU"] = df["SKU"].apply(_safe_sku)
    df["_loss_m2"] = _safe_numeric(df.get("Total Sales Loss (M-2)", _safe_default_series(0, df.index)))
    df["_loss_p90d"] = _safe_numeric(df.get("Total Sales Loss (P90D)", _safe_default_series(0, df.index)))
    grouped = (
        df.groupby(["Site Code", "Shop Name", "OM"], dropna=False, observed=True)
        .agg(
            影響SKU數=("SKU", "nunique"),
            Loss_M2=("_loss_m2", "sum"),
            Loss_P90D=("_loss_p90d", "sum"),
        )
        .reset_index()
    )
    grouped = grouped.rename(columns={
        "影響SKU數": "影響 SKU 數",
        "Loss_M2": "Total Sales Loss (M-2) 總數",
        "Loss_P90D": "Total Sales Loss (P90D) 總數",
    })
    grouped["影響 SKU 數"] = pd.to_numeric(grouped["影響 SKU 數"], errors="coerce").fillna(0).astype(int)
    for c in ["Total Sales Loss (M-2) 總數", "Total Sales Loss (P90D) 總數"]:
        grouped[c] = pd.to_numeric(grouped[c], errors="coerce").fillna(0).round().astype(int)
    grouped = grouped.sort_values(
        ["Total Sales Loss (P90D) 總數", "Total Sales Loss (M-2) 總數", "Site Code"],
        ascending=[False, False, True],
        kind="stable",
    ).reset_index(drop=True)
    if not grouped.empty:
        total_row = pd.DataFrame([{
            "Site Code": "總計",
            "Shop Name": "",
            "OM": "",
            "影響 SKU 數": int(df["SKU"].dropna().nunique()),
            "Total Sales Loss (M-2) 總數": int(grouped["Total Sales Loss (M-2) 總數"].sum()),
            "Total Sales Loss (P90D) 總數": int(grouped["Total Sales Loss (P90D) 總數"].sum()),
        }])
        grouped = pd.concat([grouped, total_row], ignore_index=True)
    return grouped[_BY_SHOP_COLUMNS]


def transform_sales_loss(
    df_input: pd.DataFrame,
    progress_callback=None,
    shop_natures: list[str] | None = None,
):
    """銷售損失量化轉換：篩選受缺貨影響列，輸出 RAW、By SKU、By Shop 與提示訊息。"""
    if progress_callback:
        progress_callback(0.0, "驗證輸入資料...")
    missing = validate_input(df_input)
    if missing:
        raise ValueError(f"缺少必要欄位: {', '.join(missing)}")
    avail = set(df_input.columns)
    if _resolve_column_name("Site Code", avail) is None:
        raise ValueError("缺少必需欄位: Site Code")
    if all(_resolve_column_name(c, avail) is None for c in _LOSS_SIGNAL_COLUMNS):
        raise ValueError(f"缺少損失訊號欄位（需至少一個）: {'、'.join(_LOSS_SIGNAL_COLUMNS)}")

    if progress_callback:
        progress_callback(0.1, "清理資料中...")
    df = _clean_dataframe(df_input, shop_natures=shop_natures, progress_callback=progress_callback)
    if df.empty:
        raise ValueError("過濾後無有效資料列，請確認來源檔案中的 Buyer 欄位是否包含有效數據")

    if progress_callback:
        progress_callback(0.2, "解像 SKU Desc / Shop Name...")
    df = _compute_raw_extra_columns(df)
    df["SKU"] = df["SKU"].apply(_safe_sku)

    if "Is Stock Impacted" in df.columns:
        impacted_flag = _as_bool(df["Is Stock Impacted"])
    else:
        impacted_flag = pd.Series(False, index=df.index)
    loss_m2 = _safe_numeric(df.get("Total Sales Loss (M-2)", _safe_default_series(0, df.index)))
    loss_p90d = _safe_numeric(df.get("Total Sales Loss (P90D)", _safe_default_series(0, df.index)))
    mask = (impacted_flag | (loss_m2 > 0) | (loss_p90d > 0)).fillna(False)

    raw = df[mask].copy()

    if progress_callback:
        progress_callback(0.35, "建立彙總...")
    by_sku = _build_sales_loss_by_sku(raw)
    by_shop = _build_sales_loss_by_shop(raw)

    warnings = []
    if "Is Stock Impacted" not in df.columns:
        warnings.append("來源檔案缺少「Is Stock Impacted」，改以損失金額大於 0 判定受影響列。")
    for col in ["Total Sales Loss (M-2)", "Total Sales Loss (P90D)"]:
        if col not in df.columns:
            warnings.append(f"來源檔案缺少「{col}」，相關合計以 0 計。")
    if raw.empty:
        warnings.append("沒有符合「受缺貨影響（Is Stock Impacted 或銷售損失大於 0）」的資料。")

    if progress_callback:
        progress_callback(0.4, "完成資料轉換...")
    return raw, by_sku, by_shop, warnings


def _write_simple_table(ws, headers, data: pd.DataFrame, progress_callback=None, progress_start=0.0, progress_end=0.1, sheet_name=""):
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    bold_font = Font(bold=True, size=11)
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    if data is not None and not data.empty:
        total = len(data)
        for row_idx, (_, row_data) in enumerate(data.iterrows()):
            excel_row = row_idx + 2
            is_total = row_data.iloc[0] == "總計"
            for col_idx, h in enumerate(headers, 1):
                val = row_data.get(h)
                if pd.isna(val):
                    val = None
                cell = ws.cell(row=excel_row, column=col_idx, value=val)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if is_total:
                    cell.font = bold_font
                if isinstance(val, (int, float)) and col_idx > 4:
                    cell.number_format = "#,##0"
            if progress_callback and (row_idx % 500 == 0 or row_idx == total - 1):
                p = progress_start + (progress_end - progress_start) * ((row_idx + 1) / max(total, 1))
                progress_callback(p, f"寫入 {sheet_name} ({row_idx+1}/{total})...")
    else:
        ws.cell(row=2, column=1, value="無符合資料").font = Font(size=12)
    _auto_width(ws)


def generate_excel_sales_loss(
    raw: pd.DataFrame,
    by_sku: pd.DataFrame,
    by_shop: pd.DataFrame,
    progress_callback=None,
    include_raw: bool = True,
) -> BytesIO:
    wb = Workbook()
    if progress_callback:
        progress_callback(0.0, "建立 Excel 結構...")

    out_cols = _sales_loss_raw_columns(raw if raw is not None else pd.DataFrame())

    ws_raw = wb.active
    ws_raw.title = "RAW"
    if include_raw:
        _write_raw_sheet(ws_raw, raw, progress_callback, progress_start=0.0, progress_end=0.55, columns=out_cols)
    else:
        _write_raw_placeholder(ws_raw)

    if progress_callback:
        progress_callback(0.6, "建立 By SKU...")
    ws_sku = wb.create_sheet(title="By SKU")
    _write_simple_table(ws_sku, _BY_SKU_COLUMNS, by_sku, progress_callback, 0.6, 0.75, "By SKU")

    if progress_callback:
        progress_callback(0.78, "建立 By Shop...")
    ws_shop = wb.create_sheet(title="By Shop")
    _write_simple_table(ws_shop, _BY_SHOP_COLUMNS, by_shop, progress_callback, 0.78, 0.9, "By Shop")

    if progress_callback:
        progress_callback(0.95, "正在儲存 Excel 檔案...")

    gc.collect()
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    if progress_callback:
        progress_callback(1.0, "完成!")
    return buf
