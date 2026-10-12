# -*- coding: utf-8 -*-
"""Stock Checking模式：由 transformer.py 拆分而來。
共用 helper 由 transformer（最終為 modes.common／issues／framework）引用；
transformer 以墊片 re-export 本模組，舊有 import 路徑不變。
"""
import datetime
import gc
import re
import unicodedata
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


from modes.common import (
    _auto_width,
    _clean_dataframe,
    _compute_raw_extra_columns,
    _iter_header_rows,
    _resolve_column_name,
    _safe_default_series,
    _safe_numeric,
    _safe_sku,
    _cached_safe_sku_series,
    _write_guide_sheet,
    _write_metric_cards,
    _write_raw_placeholder,
    _write_raw_sheet,
    validate_input,
)


from modes.framework import (
    _HYBRID_DIAG_NATURES,
    _HYBRID_DIMENSION_LABELS,
    _HYBRID_OM_BLANK_LABEL,
    _HYBRID_SALES_BANDS,
    _HYBRID_SALES_BAND_COLUMN,
    _HYBRID_SALES_OOS_DAYS_DEFAULT,
    _HYBRID_SALES_OVERSTOCK_DAYS_DEFAULT,
    _HYBRID_SALES_OVERSTOCK_MIN_UNITS_DEFAULT,
    _attach_sales_classification_columns,
    _hybrid_dimension,
    _hybrid_dimension_groups,
    _hybrid_om_groups_from_values,
    _hybrid_sales_params,
    _hybrid_shortage_hint_text,
    _hybrid_shortage_thresholds,
    _hybrid_skew_hint_text,
    _hybrid_skew_thresholds,
    _hybrid_zero_sales_groups,
    _transfer_input_consistency_warnings,
)


from modes.issues import (
    _build_stock_checking_shop_issues,
    _write_shop_issue_sheet,
)
__all__ = [
    "_STOCK_CHECKING_RAW_DISPLAY_GROUPS",
    "_STOCK_CHECKING_REQUIRED",
    "_stock_checking_buyer_stats",
    "_stock_checking_dashboard_metrics",
    "_stock_checking_raw_display_columns",
    "_write_stock_checking_buyer_table",
    "_write_stock_checking_guide_sheet",
    "generate_excel_stock_checking",
    "transform_stock_checking",
    "_STOCK_CHECKING_SUMMARY_COLUMNS",
    "_build_stock_checking_summary",
    "_hybrid_om_groups_from_summary_columns",
    "_stock_checking_new_columns",
    "_stock_checking_stats_empty",
    "_stock_checking_summary_columns",
    "_write_stock_checking_summary_sheet",
]


def _write_stock_checking_guide_sheet(ws, stats):
    """Stock Checking「說明」工作表（工作簿首位）：列出本次執行參數、
    講解調貨邏輯，並逐頁導讀各工作表應留意的部分。"""
    st_ = stats or {}
    dim = _hybrid_dimension(st_.get("dimension"))
    td = int(st_.get("turnday_threshold", 14))
    d001 = int(st_.get("d001_threshold", 100))

    params = [
        ("分類維度", f"{_HYBRID_DIMENSION_LABELS[dim]}（只影響 Summary 分組，不影響調貨計算）"),
        ("調貨門檻", f"{td} 日（SKU 級 Nature 池加總 TurnDay）"),
        ("D001 庫存門檻", f"{d001} 件（D001 Net Stock Qty (CM) 低於此值才標示需要調貨）"),
        ("調貨方向", "固定單向——只向遊客區（Tourist）調入，Mix／Non Tourist 為來源"),
        ("銷量分類門檻",
         f'缺貨線 {st_.get("sales_oos_days", _HYBRID_SALES_OOS_DAYS_DEFAULT)} 日／'
         f'積壓線 {st_.get("sales_overstock_days", _HYBRID_SALES_OVERSTOCK_DAYS_DEFAULT)} 日／'
         f'保底 {st_.get("sales_overstock_min_units", _HYBRID_SALES_OVERSTOCK_MIN_UNITS_DEFAULT)} 件'
         "（銷量帶與銷量存貨分類恆按此計算）"),
    ]

    logic = [
        f"觸發條件（SKU 級，同時符合才標示）：該 SKU 的「Tourist 加總 TurnDay」"
        f"（各 Tourist 店 Shop Stock Qty (CM) 加總 ÷ Shop Monthly Sold Qty (P90D) 加總 × 30.5）< {td} 日，"
        f"且 D001 Net Stock Qty (CM) < {d001} 件，且 Mix／Non Tourist 至少一個 Nature 的加總 TurnDay ≥ {td} 日（有貨可調）。",
        "調入需求量 = ceil(P90D ÷ 90 × 14) − Shop Stock Qty (CM)（補到 14 日周轉）；"
        "只為已標示 SKU 且個別 Tourist 店需求 > 0 的列計算。",
        "可轉出貨量（來源店）= max(0, Shop Stock Qty (CM) − P90D ÷ 90 × 14)；「無銷售但有存貨」全數可釋放；"
        "RF 與 ND 店皆可作來源；該 Nature 的 SKU 級加總 TurnDay < 門檻時視為無貨可調（可轉出量歸零）。",
        "分配：按 SKU 彙總需求，Non Tourist 店優先轉出、不足由 Mix 店補足；需求量大→小依序分配；"
        "來源店與需求店同為一店（Site Code 相同）一律排除，防止自我調貨。",
        "分類維度只改 Summary 分組與顯示欄，調貨觸發、分配與數量一律按 Shop Nature 準則計算；"
        "銷量帶與銷量存貨分類（準則見上方 ①）恆附於 RAW 與店舖問題 SKU。",
        "存貨偏重提示：某分組庫存佔比 ≥ 門檻且（庫存佔比 − 銷量佔比）≥ 差距門檻 → 以醒目橘底標示；"
        "門檻按分類維度切換（nature／rp／both 40%／20 個百分點、om 25%／12、sales 30%／15、"
        "nature_om／sales_nature 12%／8）。",
        f"存貨偏低標示：{_hybrid_shortage_hint_text(dim)} → 以淺藍底標示該分組的庫存佔比欄；"
        "門檻較偏重側寬鬆，只標示儲存格，不新增文字欄；"
        "門檻按分類維度切換（nature／rp／both 30%／10 個百分點、om 20%／8、sales 20%／10、"
        "nature_om／sales_nature 10%／5）。",
    ]

    sheet_rows = [
        ("說明（本頁）", "調貨邏輯、本次參數與導讀", "執行任何行動前，先核對上方 ① 的參數是否符合預期"),
        ("Summary", "SKU 級庫存／銷量分組（按分類維度）",
         "留意「存貨偏重提示」（橘底）——哪些分組存貨偏重而銷量不高；另留意淺藍底——庫存佔比明顯低於銷量佔比的分組（潛在缺貨風險）；TurnDay = 加總庫存 ÷ 加總月銷 × 30.5"),
        ("RAW", "全部店舖×SKU 逐列明細與判定欄",
         "留意「需要調貨」「調入需求量 (補到14日)」「實際調入量」「調入缺口」「轉出量」"
         "「調貨來源 (Site Code)」「銷量帶」「銷量存貨分類」"),
        ("店舖問題 SKU", "每間店舖需行動列（問題＋建議方法）",
         "店務同事可先從這一頁著手；含需要調入／可轉出供應與缺口跟進"
         "（Warehouse＋D001 有貨由 D001 補貨、D001 無貨寫 D001缺貨、Vendor 貨向供應商追貨）"),
        ("調貨分析摘要", "指標卡＋Buyer 需調貨 SKU 統計",
         "管理層概覽；需要調貨列數／SKU 數、調入需求量、缺口"),
    ]

    tips = [
        "本模式調貨方向固定單向（只向 Tourist 店舖調入）；「允許雙向調貨」只適用於 Safety Stock x NetStock Review。",
        "「需要調貨」為 SKU 級判定：Tourist 加總 TurnDay 達標時，個別 Tourist 店缺貨不會標示。",
        "來源店與需求店同一店舖（Site Code 相同）時一律排除，避免自我調貨。",
        "Summary 橘底＝觸發存貨偏重提示的分組；淺藍底＝存貨偏低分組（庫存佔比明顯低於銷量佔比，潛在缺貨風險）；分類維度只改變分組顯示，不影響調貨計算。",
        "缺口代表無足夠可轉出貨量；跟進方法見「店舖問題 SKU」的建議方法。",
    ]

    _write_guide_sheet(
        ws, title="Stock Checking 說明與導讀",
        subtitle="本頁講解此報告的調貨邏輯與本次執行參數，並逐頁說明應留意的部分；各項判定以 Summary 與 RAW 逐列欄位為準。",
        params=params, logic=logic, sheet_rows=sheet_rows, tips=tips,
    )



_STOCK_CHECKING_REQUIRED = [
    "Shop Nature",
    "Sold Qty (P90D)",
    "Shop Stock Qty (CM)",
    "Site Code",
]



def _stock_checking_buyer_stats(raw: pd.DataFrame, need_series: pd.Series | None = None) -> list[tuple[str, int]]:
    """Buyer 級統計：各 Buyer 涉及多少個「需要調貨」SKU（已刪除重複）。

    - Buyer 取同 SKU 首個非空值（與 Summary 同一準則）；全空歸入「(未填 Buyer)」
    - 回傳 [(Buyer, 需調貨 SKU 數)]，按數量降序、同數按 Buyer 名排序
    """
    if raw is None or raw.empty or "SKU" not in raw.columns:
        return []
    if need_series is None:
        need_series = raw.get("需要調貨", _safe_default_series("", raw.index)).astype(str) == "是"
    need_skus = set(raw.loc[need_series, "SKU"].dropna().unique())
    if not need_skus:
        return []
    buyer = (
        raw["Buyer"].fillna("").astype(str).str.strip()
        if "Buyer" in raw.columns
        else _safe_default_series("", raw.index)
    )
    tmp = pd.DataFrame({"_sku": raw["SKU"], "_buyer": buyer})
    tmp = tmp[tmp["_sku"].isin(need_skus)]

    # P0：空字串遮成 NaN＋內建 first 取代自訂 agg；全空組回填「(未填 Buyer)」，語意等價
    try:
        _masked = tmp["_buyer"].mask(tmp["_buyer"] == "")
        per_sku_buyer = _masked.groupby(tmp["_sku"], sort=False).first().fillna("(未填 Buyer)")
    except Exception:
        def _first_nonempty(s: pd.Series) -> str:
            s = s[s != ""]
            return s.iloc[0] if len(s) else "(未填 Buyer)"

        per_sku_buyer = tmp.groupby("_sku", sort=False)["_buyer"].agg(_first_nonempty)
    counts = per_sku_buyer.value_counts()
    return sorted(((str(b), int(c)) for b, c in counts.items()), key=lambda kv: (-kv[1], kv[0]))



def transform_stock_checking(
    df_input: pd.DataFrame,
    turnday_threshold: int = 14,
    d001_threshold: int = 100,
    progress_callback=None,
    dimension: str = "nature",
    sales_oos_days=None,
    sales_overstock_days=None,
    sales_overstock_min_units=None,
):
    """Stock Checking（調貨檢查）核心轉換。

    D001 倉庫少貨（低於門檻）且遊客區（Tourist）店舖缺貨時的檢查（SKU 級判定）：

    - 觸發條件（同時符合才標示）：該 SKU 的「Tourist 加總 TurnDay」
      （= 各 Tourist 店 Shop Stock Qty (CM) 加總 ÷ Shop Monthly Sold Qty (P90D) 加總 × 30.5，與 Summary 同一準則）
      < Turnday 門檻，且 D001 < 門檻，且 Mix / Non Tourist 至少一個 Nature 的加總 TurnDay ≥ 門檻（有貨可調）。
      加總 TurnDay ≥ 門檻的 SKU 一律無須調貨；Mix 與 Non Tourist 加總 TurnDay 皆低於門檻時視為無貨可調，不標示。
    - 需求端：已標示 SKU 的 Tourist 店列，補到 14 日周轉的需求 > 0
      → 標示「需要調貨」，調入需求量 = ceil(P90D / 90 × 14) − Shop Stock Qty (CM)；RAW 只為已標示 SKU 計算調貨數量。
    - 供應端：Mix / Non Tourist 店列的可轉出貨量（14 日超額標準，與 Stock Impact 模式一致）：
      可轉出 = max(0, Shop Stock Qty (CM) − P90D / 90 × 14)；「無銷售但有存貨」全數可釋放。
      該 Nature 的 SKU 級加總 TurnDay < 門檻時，該 Nature 視為無貨可調（可轉出量歸零）。RP Type 為 RF 與 ND 皆可作為來源。
    - 分配：按 SKU 彙總需求，Non Tourist 店優先，不足再由 Mix 店補足；需求列按需求量大→小依序分配；
      來源店與需求店同為一店（Site Code 相同）時一律排除，防止重複列資料異常造成自我調貨。
    - 分類維度（dimension）與銷量帶／銷量存貨分類（常開顯示欄，門檻見 sales_* 參數）只改
      Summary 分組與顯示欄，調貨計算一律按 Shop Nature 準則，不隨維度改變。
    """
    if progress_callback:
        progress_callback(0.0, "驗證輸入資料...")
    validate_input(df_input)
    avail = set(df_input.columns)
    missing = [name for name in _STOCK_CHECKING_REQUIRED if _resolve_column_name(name, avail) is None]
    if missing:
        raise ValueError(f"缺少必需欄位: {', '.join(missing)}")

    if progress_callback:
        progress_callback(0.1, "清理資料中...")
    df = _clean_dataframe(df_input, shop_natures=None, progress_callback=progress_callback)

    if progress_callback:
        progress_callback(0.2, "解像 SKU Desc / Shop Name...")
    df = _compute_raw_extra_columns(df)

    idx = df.index
    warnings: list[str] = []

    # ── 欄位正規化 ──────────────────────────────────────────────────────────
    df["Shop Nature"] = df.get("Shop Nature", _safe_default_series("", idx)).fillna("").astype(str).str.strip()
    nature = df["Shop Nature"]
    if "SKU" in df.columns:
        df["SKU"] = df["SKU"].apply(_safe_sku)
    warnings.extend(_transfer_input_consistency_warnings(df))

    # D001：優先 D001 Net Stock Qty (CM)，退用 D001 Stock Qty (CM)，兩者皆缺視為 0
    d001_col = None
    for cand in ("D001 Net Stock Qty (CM)", "D001 Stock Qty (CM)"):
        if _resolve_column_name(cand, avail):
            d001_col = cand
            break
    if d001_col is None:
        warnings.append("來源檔案缺少「D001 Net Stock Qty (CM)」與「D001 Stock Qty (CM)」，D001 一律視為 0（D001 低於門檻的判定將全數成立）。")
    d001_cm = _safe_numeric(df.get(d001_col, _safe_default_series(0, idx)) if d001_col else _safe_default_series(0, idx))

    sold_p90d = _safe_numeric(df.get("Sold Qty (P90D)", _safe_default_series(0, idx)))
    cm_stock = _safe_numeric(df.get("Shop Stock Qty (CM)", _safe_default_series(0, idx)))

    # ── 每列計算欄位 ────────────────────────────────────────────────────────
    is_mix = (nature == "Mix").to_numpy()
    is_nt = (nature == "Non Tourist").to_numpy()
    is_tourist = (nature == "Tourist").to_numpy()
    cm_stock_val = cm_stock.to_numpy(dtype=float)
    sold_val = sold_p90d.to_numpy(dtype=float)
    no_sale_has_stock = (sold_val <= 0) & (cm_stock_val > 0)
    ideal_14d = sold_val / 90.0 * 14

    # ── SKU 級各 Nature 加總 TurnDay（與 Summary 同一準則：加總庫存 ÷ 加總月銷 × 30.5）──
    if _resolve_column_name("Shop Monthly Sold Qty (P90D)", avail) is not None:
        msold_monthly = _safe_numeric(df["Shop Monthly Sold Qty (P90D)"])
    else:
        # 缺月銷欄位時以 Sold Qty (P90D) 換算（P90D 為 90 日總量 → 月銷 = /90 × 30.5）
        msold_monthly = sold_p90d / 90.0 * 30.5
        warnings.append("來源檔案缺少「Shop Monthly Sold Qty (P90D)」，加總 TurnDay 改以 Sold Qty (P90D) 換算。")
    msold_val = msold_monthly.to_numpy(dtype=float)
    sku_key = df["SKU"] if "SKU" in df.columns else _safe_default_series(None, idx)

    # ── 分類維度與銷量帶／銷量存貨分類（常開顯示欄；只改 Summary 分組，不影響調貨計算）──
    dim = _hybrid_dimension(dimension)
    _sales_params, _sales_param_warnings = _hybrid_sales_params(
        False, sales_oos_days, sales_overstock_days, sales_overstock_min_units
    )[1:]
    warnings.extend(_sales_param_warnings)
    df, _sales_col_warnings = _attach_sales_classification_columns(
        df, _sales_params, monthly_sold=msold_monthly,
    )
    warnings.extend(_sales_col_warnings)
    om_groups = None
    if dim in ("om", "nature_om"):
        om_groups = _hybrid_om_groups_from_values(
            df["OM"] if "OM" in df.columns else _safe_default_series("", idx)
        )

    def _sku_nature_turnday(nature_mask: np.ndarray) -> np.ndarray:
        s_stock = pd.Series(np.where(nature_mask, cm_stock_val, 0.0), index=df.index)
        s_msold = pd.Series(np.where(nature_mask, msold_val, 0.0), index=df.index)
        agg_stock = s_stock.groupby(sku_key).transform("sum")
        agg_msold = s_msold.groupby(sku_key).transform("sum")
        with np.errstate(divide="ignore", invalid="ignore"):
            return np.where(agg_msold.to_numpy() > 0,
                            agg_stock.to_numpy() / agg_msold.to_numpy() * 30.5,
                            np.inf)

    sku_tourist_turnday = _sku_nature_turnday(is_tourist)
    sku_mix_turnday = _sku_nature_turnday(is_mix)
    sku_nt_turnday = _sku_nature_turnday(is_nt)
    for label, td in (("Tourist", sku_tourist_turnday), ("Mix", sku_mix_turnday), ("Non Tourist", sku_nt_turnday)):
        df[f"{label} 加總 TurnDay"] = np.round(np.where(np.isfinite(td), td, np.nan), 2)

    # 供應門檻：Nature 加總 TurnDay < 門檻 → 該 Nature 理論上無貨可調，可轉出量歸零
    mix_pool_ok = sku_mix_turnday >= turnday_threshold   # inf（無銷售）→ 有存貨可釋放
    nt_pool_ok = sku_nt_turnday >= turnday_threshold
    mix_excess = np.where(mix_pool_ok, np.where(is_mix, np.where(no_sale_has_stock, cm_stock_val, np.maximum(0, cm_stock_val - ideal_14d)), 0), 0)
    nt_excess = np.where(nt_pool_ok, np.where(is_nt, np.where(no_sale_has_stock, cm_stock_val, np.maximum(0, cm_stock_val - ideal_14d)), 0), 0)
    df["可轉出貨量 (Mix)"] = mix_excess.astype(int)
    df["可轉出貨量 (Non Tourist)"] = nt_excess.astype(int)

    demand = np.maximum(0, np.ceil(ideal_14d) - cm_stock_val)
    d001_np = d001_cm.to_numpy(dtype=float)
    ttd_np = pd.Series(sku_tourist_turnday, index=df.index).to_numpy(dtype=float)
    need_mask = (
        is_tourist
        & (ttd_np < turnday_threshold)   # Tourist 加總 TurnDay ≥ 門檻（含月銷 0 → inf、NaN）→ 無須調貨
        & (d001_np < d001_threshold)
        & (demand > 0)
        & (mix_pool_ok | nt_pool_ok)     # Mix 與 Non Tourist 加總 TurnDay 皆低於門檻 → 無貨可調，不標示
    )
    need_series = pd.Series(need_mask, index=df.index)

    df["需要調貨"] = np.where(need_mask, "是", "")
    df["調入需求量 (補到14日)"] = np.where(need_mask, demand, 0).astype(int)
    df["實際調入量"] = 0
    df["調入缺口"] = 0
    df["轉出量"] = 0
    df["調貨來源 (Site Code)"] = ""

    # ── 跨店調貨分配（按 SKU：Non Tourist 優先 → Mix 補足）───────────────────
    if need_mask.any():
        site_arr = df.get("Site Code", _safe_default_series("", idx)).fillna("").astype(str).str.strip().to_numpy(dtype=object)
        sku_arr = df["SKU"].to_numpy(dtype=object)
        nt_avail = df["可轉出貨量 (Non Tourist)"].to_numpy(dtype=float).copy()
        mix_avail = df["可轉出貨量 (Mix)"].to_numpy(dtype=float).copy()
        demand_np = demand.copy()
        donate_np = np.zeros(len(df), dtype=int)
        actual_np = np.zeros(len(df), dtype=int)
        gap_np = np.zeros(len(df), dtype=int)
        desc_np = np.full(len(df), "", dtype=object)

        # P0：預建 SKU→列索引，避免每 SKU 全表 sku_arr==sku 掃描；集合內容與順序與原實作一致
        sku_to_idx: dict = {}
        for _p in range(len(df)):
            _k = sku_arr[_p]
            try:
                if pd.isna(_k):
                    continue
            except Exception:
                pass
            sku_to_idx.setdefault(_k, []).append(int(_p))
        need_sku_set = set(sku_arr[need_mask])
        for sku in need_sku_set:
            try:
                if pd.isna(sku):
                    continue
            except Exception:
                pass
            _idx_list = sku_to_idx.get(sku, [])
            if not _idx_list:
                continue
            need_rows = np.array([p for p in _idx_list if bool(need_mask[p])], dtype=int)
            if need_rows.size == 0:
                continue
            nt_src = [int(p) for p in _idx_list if float(nt_avail[p]) > 0]
            mix_src = [int(p) for p in _idx_list if float(mix_avail[p]) > 0]
            order = np.argsort(-demand_np[need_rows], kind="stable")
            remain_nt = {int(p): float(nt_avail[p]) for p in nt_src}
            remain_mix = {int(p): float(mix_avail[p]) for p in mix_src}
            for n in order:
                row_pos = int(need_rows[n])
                rem = float(demand_np[row_pos])
                parts: list[str] = []
                donated_actual = 0
                for pool, _label in ((remain_nt, "Non Tourist"), (remain_mix, "Mix")):
                    for spos in list(pool.keys()):
                        if rem <= 0:
                            break
                        if site_arr[spos] == site_arr[row_pos]:
                            continue  # 自我調貨守衛：來源店＝需求店時排除（重複列資料異常防護）
                        avail = pool[spos]
                        take = min(rem, avail)
                        if take > 0:
                            pool[spos] -= take
                            rem -= take
                            donate_np[spos] += int(take)
                            donated_actual += int(take)
                            parts.append(f"{site_arr[spos]}:{int(take)}")
                    if rem <= 0:
                        break
                actual_np[row_pos] = donated_actual
                gap_np[row_pos] = int(rem)
                desc_np[row_pos] = "; ".join(parts)

        df["實際調入量"] = actual_np
        df["調入缺口"] = gap_np
        df["轉出量"] = donate_np
        df["調貨來源 (Site Code)"] = desc_np

    # ── 統計 ────────────────────────────────────────────────────────────────
    stats = {
        "turnday_threshold": int(turnday_threshold),
        "d001_threshold": int(d001_threshold),
        "need_rows": int(need_series.sum()),
        "need_skus": int(df.loc[need_series, "SKU"].nunique()) if need_series.any() else 0,
        "total_demand": int(df.loc[need_series, "調入需求量 (補到14日)"].sum()) if need_series.any() else 0,
        "total_actual": int(df["實際調入量"].sum()),
        "total_gap": int(df["調入缺口"].sum()),
        "nt_transfer": int(df.loc[df["Shop Nature"] == "Non Tourist", "轉出量"].sum()),
        "mix_transfer": int(df.loc[df["Shop Nature"] == "Mix", "轉出量"].sum()),
        "buyer_need_skus": _stock_checking_buyer_stats(df, need_series),
        "dimension": dim,
        "om_groups": om_groups,
        "sales_oos_days": int(_sales_params["oos_days"]),
        "sales_overstock_days": int(_sales_params["overstock_days"]),
        "sales_overstock_min_units": int(_sales_params["overstock_min_units"]),
    }

    if not need_mask.any():
        warnings.append("沒有符合調貨條件的 Tourist SKU（Turnday < 門檻 且 D001 < 門檻 且補到 14 日周轉的需求 > 0）。")

    # ── SKU 級 Summary ──────────────────────────────────────────────────────
    if progress_callback:
        progress_callback(0.35, "建立 SKU Summary...")
    summary = _build_stock_checking_summary(df, dimension=dim, om_groups=om_groups)

    if progress_callback:
        progress_callback(0.4, "完成資料轉換...")
    return df, summary, stats, warnings



def _stock_checking_dashboard_metrics(raw: pd.DataFrame, stats: dict | None = None):
    """Stock Checking Dashboard 指標（與 Streamlit Transfer Analysis Dashboard 同一準則）。

    回傳 (row1_metrics, row2_metrics)，供 _write_metric_cards 使用。
    stats 缺失時（舊呼叫相容）直接由 raw 匯總。
    """
    total_rows = len(raw) if raw is not None else 0
    if stats is None:
        stats = {}
    need_col = raw.get("需要調貨", pd.Series(dtype=object)) if raw is not None else pd.Series(dtype=object)
    need_mask = (need_col.astype(str) == "是")
    need_rows = int(stats.get("need_rows", int(need_mask.sum())))
    if "need_skus" in stats:
        need_skus = int(stats["need_skus"])
    elif raw is not None and "SKU" in raw.columns and need_mask.any():
        need_skus = int(raw.loc[need_mask, "SKU"].nunique())
    else:
        need_skus = 0
    demand_col = "調入需求量 (補到14日)"
    actual_col = "實際調入量"
    gap_col = "調入缺口"
    out_col = "轉出量"
    total_demand = int(stats.get("total_demand", int(_safe_numeric(raw.get(demand_col, _safe_default_series(0, raw.index))).sum()) if raw is not None else 0))
    total_actual = int(stats.get("total_actual", int(_safe_numeric(raw.get(actual_col, _safe_default_series(0, raw.index))).sum()) if raw is not None else 0))
    total_gap = int(stats.get("total_gap", int(_safe_numeric(raw.get(gap_col, _safe_default_series(0, raw.index))).sum()) if raw is not None else 0))
    if "nt_transfer" in stats and "mix_transfer" in stats:
        nt_transfer = int(stats["nt_transfer"])
        mix_transfer = int(stats["mix_transfer"])
    elif raw is not None and "Shop Nature" in raw.columns and out_col in raw.columns:
        _out = _safe_numeric(raw.get(out_col, _safe_default_series(0, raw.index)))
        nt_transfer = int(_out[raw["Shop Nature"] == "Non Tourist"].sum())
        mix_transfer = int(_out[raw["Shop Nature"] == "Mix"].sum())
    else:
        nt_transfer = int(stats.get("nt_transfer", 0))
        mix_transfer = int(stats.get("mix_transfer", 0))

    row1_metrics = [
        ("需要調貨列數", need_rows, "Tourist 加總 TurnDay & D001 皆低於門檻的 SKU", "EF4444"),
        ("需調貨 SKU 數", need_skus, "符合條件的 SKU 總數（已刪除重複）", "F59E0B"),
        ("調入需求量", total_demand, "補到 14 日周轉的總需求（件）", "3B82F6"),
        ("實際調入量", total_actual, "NT 優先、Mix 補足後的總調入（件）", "10B981"),
    ]
    row2_metrics = [
        ("Non Tourist 轉出量", nt_transfer, "優先由 Non Tourist 店舖轉出（件）", "8B5CF6"),
        ("Mix 轉出量", mix_transfer, "NT 不足時由 Mix 店舖補足（件）", "06B6D4"),
        ("無法足額缺口", total_gap, "來源可轉出量不足的缺口（件）", "F97316"),
        ("總明細行數", total_rows, "RAW 完整明細總行數", "64748B"),
    ]
    return row1_metrics, row2_metrics



def _write_stock_checking_buyer_table(ws, buyer_stats, thin_border, start_row: int = 11):
    """Dashboard 指標卡下方的「Buyer 需調貨 SKU 統計」表（Buyer / 需調貨 SKU 數 / 佔比 + 總計）。"""
    if not buyer_stats:
        return
    center = Alignment(horizontal="center", vertical="center")
    header_fill = PatternFill(start_color="0EA5E9", end_color="0EA5E9", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    total_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    total_skus = sum(c for _, c in buyer_stats)

    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=3)
    tc = ws.cell(row=start_row, column=1, value="Buyer 需調貨 SKU 統計（按需調貨 SKU 數排序）")
    tc.font = Font(bold=True, size=12, color="1E40AF")
    tc.alignment = Alignment(horizontal="left", vertical="center")

    hr = start_row + 1
    for col_idx, h in enumerate(("Buyer", "需調貨 SKU 數", "佔比"), 1):
        cell = ws.cell(row=hr, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin_border
    r = hr + 1
    for buyer, cnt in buyer_stats:
        c1 = ws.cell(row=r, column=1, value=buyer)
        c1.border = thin_border
        c2 = ws.cell(row=r, column=2, value=cnt)
        c2.number_format = "#,##0"
        c2.alignment = center
        c2.border = thin_border
        c3 = ws.cell(row=r, column=3, value=cnt / total_skus if total_skus else None)
        c3.number_format = "0.0%"
        c3.alignment = center
        c3.border = thin_border
        r += 1
    for col_idx, v in ((1, "總計"), (2, total_skus), (3, 1.0 if total_skus else None)):
        cell = ws.cell(row=r, column=col_idx, value=v)
        cell.font = Font(bold=True, size=11)
        cell.fill = total_fill
        cell.border = thin_border
        cell.alignment = center
        if col_idx == 2:
            cell.number_format = "#,##0"
        elif col_idx == 3:
            cell.number_format = "0.0%"
    _auto_width(ws)



# Stock Checking RAW 顯示欄序分組（識別→品項屬性→庫存與銷售→SKU 級判定→調貨執行→其他）
_STOCK_CHECKING_RAW_DISPLAY_GROUPS = [
    ["Buyer", "Brand", "SKU", "SKU Desc", "Site Code", "Shop Name", "OM", "Shop Nature"],
    ["Adj. SST", "Adj. MCH1", "Vendor MC", "MCH2", "Item Status", "Item Class", "RP Type",
     "Shop Stock Status", "Supply Source", "Launch Date"],
    ["Shop Stock Qty (CM)", "Adj. Safety Stock", "Sold Qty (P90D)", "Shop Monthly Sold Qty (P90D)",
     "銷量帶", "銷量存貨分類",
     "Turnday (P90D Qty)", "Sold Qty (M-2)", "Sold Qty (M-1)", "Sold Qty (CM - MTD)",
     "Net Sales (M-2)", "Net Sales (M-1)", "Net Sales (CM - MTD)", "Excess Stock Qty",
     "Shop On Order PO Qty (CM)", "Shop Pending Receive",
     "D001 Net Stock Qty (CM)", "D001 Stock Qty (CM)", "D001 Effective Stock Qty (CM)"],
    ["Tourist 加總 TurnDay", "Mix 加總 TurnDay", "Non Tourist 加總 TurnDay",
     "需要調貨", "調入需求量 (補到14日)"],
    ["可轉出貨量 (Non Tourist)", "可轉出貨量 (Mix)", "實際調入量", "調入缺口", "轉出量",
     "調貨來源 (Site Code)"],
]



def _stock_checking_raw_display_columns(raw: pd.DataFrame) -> list[str]:
    """Stock Checking RAW 顯示欄序：不顯示「Item_Shop」與所有含「Value」的金額欄，
    其餘欄位（含 _compute_raw_extra_columns 衍生的 Shop Name / SKU Desc）按
    _STOCK_CHECKING_RAW_DISPLAY_GROUPS 分組排列，未歸組的原始欄位按原序置於最後。"""
    cols = [str(c) for c in raw.columns] + ["Shop Name", "SKU Desc"]
    seen: set[str] = set()
    kept: list[str] = []
    for c in cols:
        if c in seen or c == "Item_Shop" or "value" in c.lower():
            continue
        seen.add(c)
        kept.append(c)
    kept_set = set(kept)
    ordered: list[str] = []
    for group in _STOCK_CHECKING_RAW_DISPLAY_GROUPS:
        for c in group:
            if c in kept_set:
                ordered.append(c)
                kept_set.discard(c)
    ordered += [c for c in kept if c in kept_set]
    return ordered



def generate_excel_stock_checking(
    raw: pd.DataFrame,
    summary: pd.DataFrame | None = None,
    stats: dict | None = None,
    progress_callback=None,
    include_raw: bool = True,
) -> BytesIO:
    """Stock Checking 輸出 5 工作表：說明（首頁導讀：本次執行參數＋核心邏輯＋工作表導讀＋實用提示，
    開啟 Excel 時預設顯示本頁）+ Summary（SKU 級彙總，按 stats["dimension"] 分類維度分組，
    偏重提示以醒目橘底加強顯示）+ RAW（重排欄序提升可讀性，不顯示 Item_Shop 與 Value 金額欄）
    + 店舖問題 SKU（每間店舖需行動列；空表寫「無符合資料」）+ 調貨分析摘要（Dashboard）。"""
    wb = Workbook()
    if progress_callback:
        progress_callback(0.0, "建立 Excel 結構...")

    _dim = _hybrid_dimension((stats or {}).get("dimension"))
    _om_groups = (stats or {}).get("om_groups")
    if summary is None:
        summary = _build_stock_checking_summary(raw, dimension=_dim, om_groups=_om_groups)

    # ── Sheet 0: 說明（邏輯、本次執行參數與導讀；置於工作簿首位）──
    ws_guide = wb.active
    ws_guide.title = "說明"
    _write_stock_checking_guide_sheet(ws_guide, stats)

    ws_summary = wb.create_sheet(title="Summary")
    _write_stock_checking_summary_sheet(ws_summary, summary, progress_callback, highlight_skew=True,
                                        dimension=_dim, om_groups=_om_groups)

    if progress_callback:
        progress_callback(0.35, "建立 RAW...")
    ws_raw = wb.create_sheet(title="RAW")
    (_write_raw_sheet(ws_raw, raw, progress_callback, progress_start=0.35, progress_end=0.80,
                     columns=_stock_checking_raw_display_columns(raw)) if include_raw else _write_raw_placeholder(ws_raw))

    if progress_callback:
        progress_callback(0.82, "建立店舖問題 SKU...")
    ws_shop_issue = wb.create_sheet(title="店舖問題 SKU")
    _write_shop_issue_sheet(
        ws_shop_issue,
        _build_stock_checking_shop_issues(
            raw,
            turnday_threshold=int((stats or {}).get("turnday_threshold", 14)),
            d001_threshold=int((stats or {}).get("d001_threshold", 100)),
        ),
        title="店舖問題 SKU — 每間店舖需行動列（問題＋建議方法）",
        progress_callback=progress_callback, progress_start=0.82, progress_end=0.86,
    )

    if progress_callback:
        progress_callback(0.88, "建立調貨分析摘要工作表...")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    desc_font = Font(size=9, color="6B7280", italic=True)
    ws_dashboard = wb.create_sheet(title="調貨分析摘要")
    row1_metrics, row2_metrics = _stock_checking_dashboard_metrics(raw, stats)
    _write_metric_cards(
        ws_dashboard, "調貨分析摘要 (Transfer Analysis Dashboard)",
        row1_metrics, row2_metrics, thin_border, desc_font, progress_callback,
    )
    buyer_stats = (stats or {}).get("buyer_need_skus")
    if buyer_stats is None:
        buyer_stats = _stock_checking_buyer_stats(raw)
    _write_stock_checking_buyer_table(ws_dashboard, buyer_stats, thin_border)

    if progress_callback:
        progress_callback(0.95, "正在儲存 Excel 檔案...")

    gc.collect()
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    if progress_callback:
        progress_callback(1.0, "完成!")

    return buf




def _stock_checking_new_columns(raw: pd.DataFrame) -> list[str]:
    """Stock Checking 新增的計算欄位清單（存在才輸出，供測試與 Excel 使用）。"""
    cols = [
        "Tourist 加總 TurnDay",
        "Mix 加總 TurnDay",
        "Non Tourist 加總 TurnDay",
        "需要調貨",
        "調入需求量 (補到14日)",
        "實際調入量",
        "調入缺口",
        "可轉出貨量 (Mix)",
        "可轉出貨量 (Non Tourist)",
        "轉出量",
        "調貨來源 (Site Code)",
    ]
    avail = set(raw.columns)
    return [c for c in cols if _resolve_column_name(c, avail) is not None]



def _stock_checking_stats_empty() -> dict:
    return {
        "need_rows": 0,
        "need_skus": 0,
        "total_demand": 0,
        "total_actual": 0,
        "total_gap": 0,
        "nt_transfer": 0,
        "mix_transfer": 0,
        "buyer_need_skus": [],
        "dimension": "nature",
        "om_groups": None,
    }



# Stock Checking／混合 nature 維度的 Summary A–AC 契約欄序基準
# （由 _stock_checking_summary_columns("nature") 產生，測試對帳用）
_STOCK_CHECKING_SUMMARY_COLUMNS = [
    "Brand", "SKU", "SKU Desc", "Buyer", "Adj. SST", "Supply Source", "D001 Net Stock Qty (CM)",
    "Tourist 加總 - Shop Stock Qty (CM)", "Tourist 加總 - Shop Monthly Sold Qty (P90D)", "Tourist TurnDay",
    "Tourist 涉及店舖數",
    "Mix 加總 - Shop Stock Qty (CM)", "Mix 加總 - Shop Monthly Sold Qty (P90D)", "Mix TurnDay",
    "Mix 涉及店舖數",
    "Non Tourist 加總 - Shop Stock Qty (CM)", "Non Tourist 加總 - Shop Monthly Sold Qty (P90D)", "Non Tourist TurnDay",
    "Non Tourist 涉及店舖數",
    "Tourist 加總 - Shop Stock Qty (CM) %", "Tourist 加總 - Shop Monthly Sold Qty (P90D) %",
    "Mix 加總 - Shop Stock Qty (CM) %", "Mix 加總 - Shop Monthly Sold Qty (P90D) %",
    "Non Tourist 加總 - Shop Stock Qty (CM) %", "Non Tourist 加總 - Shop Monthly Sold Qty (P90D) %",
    "強制調貨提示", "由Non Tourist 店舖預計調出數量", "由Mix 店舖預計調出數量",
    "存貨偏重提示",
]



def _stock_checking_summary_columns(dimension: str = "nature", bidirectional: bool = False, om_groups=None) -> list[str]:
    """Summary 欄序（按分類維度切換分組欄位）：
    nature＝三 Nature（現行契約，含雙向時末尾「由Tourist 店舖預計調出數量」）；
    rp＝RF／ND 兩組；both＝六組（Tourist-RF … Non Tourist-ND）；
    om＝動態 OM 名單；nature_om＝三 Nature × OM 笛卡兒積；
    sales＝四銷量帶；sales_nature＝三 Nature × 四銷量帶 12 組。
    每組加總區為加總庫存／加總月銷／TurnDay／涉及店舖數四欄（店舖數＝同 SKU 同組
    Site Code 不重複計數，空白 Site Code 不計）。"""
    dim = _hybrid_dimension(dimension)
    groups = _hybrid_dimension_groups(dim, om_groups=om_groups)
    cols = ["Brand", "SKU", "SKU Desc", "Buyer", "Adj. SST", "Supply Source", "D001 Net Stock Qty (CM)"]
    for gname in groups:
        cols += [
            f"{gname} 加總 - Shop Stock Qty (CM)",
            f"{gname} 加總 - Shop Monthly Sold Qty (P90D)",
            f"{gname} TurnDay",
            f"{gname} 涉及店舖數",
        ]
    for gname in groups:
        cols += [
            f"{gname} 加總 - Shop Stock Qty (CM) %",
            f"{gname} 加總 - Shop Monthly Sold Qty (P90D) %",
        ]
    cols += ["強制調貨提示"]
    if dim == "nature":
        cols += ["由Non Tourist 店舖預計調出數量", "由Mix 店舖預計調出數量"]
    else:
        cols += [f"由{gname} 店舖預計調出數量" for gname in groups]
    cols += ["存貨偏重提示"]
    if dim == "nature" and bidirectional:
        cols += ["由Tourist 店舖預計調出數量"]
    return cols



def _build_stock_checking_summary(
    raw: pd.DataFrame, bidirectional: bool = False, dimension: str = "nature", om_groups=None
) -> pd.DataFrame:
    """SKU 級 Summary（參考 By SKUxShop_Stock Checking_Output 格式）：

    - 每個 SKU 一列：Brand / SKU / SKU Desc / Buyer / Adj. SST / Supply Source / D001 Net Stock Qty (CM)
      （Buyer、Adj. SST 取同 SKU 首個非空值作代表）
    - 按分類維度分組（dimension：nature 三組／rp 兩組 RF／ND／both 六組／
      om 動態 OM 名單／nature_om 三 Nature × OM 笛卡兒積／sales 四銷量帶／
      sales_nature 三 Nature × 四帶 12 組；om_groups 缺省時由 raw 即時推導）加總
      Shop Stock Qty (CM) 與 Shop Monthly Sold Qty (P90D)，
      TurnDay = 加總庫存 ÷ 加總月銷 × 30.5（月銷為 0 時留空）；除佔比外一律取整顯示
    - 各組涉及店舖數：同 SKU 同組的 Site Code 不重複計數（空白 Site Code 不計），
      與加總同母體，置於該組 TurnDay 之後
    - 各組佔比 (%)：該組加總 ÷ 全部分組加總合計（0-1 小數，Excel 以 0.00% 顯示；
      合計為 0 時留空；nature 維度僅計 Tourist / Mix / Non Tourist 精確比對）
    - 強制調貨提示：該 SKU 有任何店舖列被標示「需要調貨」時顯示「需要調貨」
    - 存貨偏重提示：按維度門檻（nature／rp／both：庫存佔比 ≥ 40% 且差距 ≥ 20 個百分點；
      om：≥ 25% 且差距 ≥ 12 個百分點；nature_om／sales_nature：≥ 12% 且差距 ≥ 8 個百分點；
      sales：≥ 30% 且差距 ≥ 15 個百分點）→
      「存貨偏重 {分組}（庫存 x% / 銷量 y%）」，多組以；分隔（差距大者在前）
    - 由 X 店舖預計調出數量 = 該 SKU 對應分組來源店的「轉出量」加總；
      nature 維度：Non Tourist / Mix 固定、bidirectional=True 時另加「由Tourist 店舖預計調出數量」
      （置於欄序最後，A–AC 契約欄位不動）；rp／both／om／nature_om／sales／sales_nature 維度：每個分組各一欄
    """
    if raw is None or raw.empty:
        return pd.DataFrame()
    dim = _hybrid_dimension(dimension)
    if dim in ("om", "nature_om") and om_groups is None:
        om_groups = _hybrid_om_groups_from_values(
            raw["OM"] if "OM" in raw.columns else _safe_default_series("", raw.index)
        )
    groups = _hybrid_dimension_groups(dim, om_groups=om_groups)
    skew_share, skew_gap = _hybrid_skew_thresholds(dim)
    df = raw.copy()
    df["_sku"] = df["SKU"]
    df = df[df["_sku"].notna()]
    if df.empty:
        return pd.DataFrame()
    df["_nature"] = df.get("Shop Nature", _safe_default_series("", df.index)).fillna("").astype(str).str.strip()
    if dim == "nature":
        df["_grp"] = df["_nature"]
    elif dim in ("sales", "sales_nature"):
        if _HYBRID_SALES_BAND_COLUMN in df.columns:
            _band = df[_HYBRID_SALES_BAND_COLUMN].fillna("").astype(str).str.strip()
            _band = _band.where(_band.isin(_HYBRID_SALES_BANDS), _HYBRID_SALES_BANDS[1])
            df["_grp"] = df["_nature"] + "-" + _band if dim == "sales_nature" else _band
        else:
            df["_grp"] = df["_nature"]
    elif dim in ("om", "nature_om"):
        _om = (
            df["OM"].fillna("").astype(str).str.strip()
            if "OM" in df.columns
            else _safe_default_series("", df.index)
        )
        _om = pd.Series(np.where(_om == "", _HYBRID_OM_BLANK_LABEL, _om), index=df.index)
        df["_grp"] = _om if dim == "om" else df["_nature"] + "-" + _om
    else:
        rp = (
            df["RP Type"].fillna("").astype(str).str.strip().str.upper()
            if "RP Type" in df.columns
            else _safe_default_series("RF", df.index)
        )
        df["_grp"] = pd.Series(np.where(rp == "RF", "RF", "ND"), index=df.index)
        if dim == "both":
            df["_grp"] = df["_nature"] + "-" + df["_grp"]
    df["_brand"] = df.get("Brand", _safe_default_series("", df.index)).fillna("").astype(str).str.strip()
    df["_stock"] = _safe_numeric(df.get("Shop Stock Qty (CM)", _safe_default_series(0, df.index)))
    df["_msold"] = _safe_numeric(df.get("Shop Monthly Sold Qty (P90D)", _safe_default_series(0, df.index)))
    df["_transfer"] = _safe_numeric(df.get("轉出量", _safe_default_series(0, df.index)))
    df["_need"] = df.get("需要調貨", _safe_default_series("", df.index)).astype(str) == "是"
    df["_site"] = df.get("Site Code", _safe_default_series("", df.index)).fillna("").astype(str).str.strip()

    def _first_text(series):
        s = series.dropna().astype(str).str.strip()
        s = s[s != ""]
        return s.iloc[0] if len(s) else ""

    rows = []
    for sku, g in df.groupby("_sku", sort=False):
        row = {
            "SKU": sku,
            "Brand": _first_text(g["_brand"]),
        }
        if "SKU Desc" in g.columns:
            row["SKU Desc"] = _first_text(g["SKU Desc"])
        # SKU Desc 之後：Buyer、Adj. SST（同 SKU 取首個非空值作代表）
        if "Buyer" in g.columns:
            row["Buyer"] = _first_text(g["Buyer"].fillna("").astype(str))
        if "Adj. SST" in g.columns:
            row["Adj. SST"] = _first_text(g["Adj. SST"].fillna("").astype(str))
        if "Supply Source" in g.columns:
            row["Supply Source"] = _first_text(g["Supply Source"])
        d001_source = (
            "D001 Effective Stock Qty (CM)"
            if "D001 Effective Stock Qty (CM)" in g.columns
            else "D001 Net Stock Qty (CM)"
        )
        if d001_source in g.columns:
            d001 = g[d001_source].dropna()
            row["D001 Net Stock Qty (CM)"] = d001.iloc[0] if len(d001) else None
        _sums: dict[str, tuple[float, float]] = {}
        for grp_name in groups:
            gn = g[g["_grp"] == grp_name]
            stock_sum = float(gn["_stock"].sum())
            msold_sum = float(gn["_msold"].sum())
            _sums[grp_name] = (stock_sum, msold_sum)
            # 整數顯示（% 除外）：庫存本為整數；月銷四捨五入到整數；TurnDay 四捨五入到整數日
            row[f"{grp_name} 加總 - Shop Stock Qty (CM)"] = int(round(stock_sum))
            row[f"{grp_name} 加總 - Shop Monthly Sold Qty (P90D)"] = int(round(msold_sum))
            row[f"{grp_name} TurnDay"] = int(round(stock_sum / msold_sum * 30.5)) if msold_sum > 0 else None
            row[f"{grp_name} 涉及店舖數"] = int(gn.loc[gn["_site"] != "", "_site"].nunique())
            # nature 維度維持現行契約（Tourist 欄僅雙向時存在）；rp／both 每組各一欄
            if dim != "nature" or grp_name != "Tourist" or bidirectional:
                row[f"由{grp_name} 店舖預計調出數量"] = int(gn["_transfer"].sum())
        # 各組佔比（分母 = 全部分組加總合計；合計為 0 留空）
        _total_stock = sum(v[0] for v in _sums.values())
        _total_msold = sum(v[1] for v in _sums.values())
        for grp_name in groups:
            stock_sum, msold_sum = _sums[grp_name]
            row[f"{grp_name} 加總 - Shop Stock Qty (CM) %"] = (
                round(stock_sum / _total_stock, 4) if _total_stock > 0 else None
            )
            row[f"{grp_name} 加總 - Shop Monthly Sold Qty (P90D) %"] = (
                round(msold_sum / _total_msold, 4) if _total_msold > 0 else None
            )
        row["強制調貨提示"] = "需要調貨" if bool(g["_need"].any()) else ""
        # 存貨偏重提示：按維度門檻（見 _HYBRID_SKEW_THRESHOLDS）
        _hints: list[tuple[float, str]] = []
        if _total_stock > 0:
            for grp_name in groups:
                stock_sum, msold_sum = _sums[grp_name]
                s_share = stock_sum / _total_stock
                v_share = (msold_sum / _total_msold) if _total_msold > 0 else 0.0
                if s_share >= skew_share and (s_share - v_share) >= skew_gap:
                    _hints.append((s_share - v_share,
                                   f"存貨偏重 {grp_name}（庫存 {s_share:.1%} / 銷量 {v_share:.1%}）"))
        row["存貨偏重提示"] = "；".join(h for _, h in sorted(_hints, key=lambda t: -t[0]))
        rows.append(row)

    summary = pd.DataFrame(rows)
    summary = summary.sort_values(["Brand", "SKU"], kind="stable").reset_index(drop=True)
    ordered = _stock_checking_summary_columns(dim, bidirectional=bidirectional, om_groups=om_groups)
    ordered = [c for c in ordered if c in summary.columns]
    trailing = [c for c in summary.columns if c not in ordered]
    return summary[ordered + trailing]



def _hybrid_om_groups_from_summary_columns(summary: pd.DataFrame, dimension: str) -> tuple[str, ...] | None:
    """由已產生的 Summary 欄位反推 OM 名單（Excel 寫入回退路徑用）。

    om 維度欄形如「{OM} 加總 - Shop Stock Qty (CM)」；nature_om 維度欄形如
    「{Nature}-{OM} 加總 - Shop Stock Qty (CM)」。無法反推時回傳 None。
    """
    if summary is None or summary.empty:
        return None
    dim = _hybrid_dimension(dimension)
    if dim not in ("om", "nature_om"):
        return None
    suffix = " 加總 - Shop Stock Qty (CM)"
    prefixes = [str(c)[:-len(suffix)] for c in summary.columns
                if isinstance(c, str) and c.endswith(suffix)]
    if not prefixes:
        return None
    if dim == "om":
        return tuple(sorted(set(prefixes)))
    oms: set[str] = set()
    for p in prefixes:
        for n in _HYBRID_DIAG_NATURES:
            if p.startswith(n + "-"):
                oms.add(p[len(n) + 1:])
                break
    return tuple(sorted(oms)) if oms else None



def _write_summary_color_legend(ws, dimension, highlight_skew: bool = False, highlight_zero_sales: bool = False, highlight_shortage: bool = True):
    """Summary 圖例：第 1 列為色塊名稱，第 2 列為各色判定規則；位於凍結表頭上方，隨捲動固定可見。
    色塊與門檻與本工作表實際標示一致（僅列出本次輸出會啟用的顏色）。須於 _auto_width 之後呼叫，避免圖例文字撐寬欄位。"""
    dim = _hybrid_dimension(dimension)
    chips = []
    notes = []
    if highlight_skew:
        chips.append(("存貨偏重", "FFC000", "000000"))
        notes.append(f"橘底＝存貨偏重（{_hybrid_skew_hint_text(dim)}），留意積壓風險")
    if highlight_shortage:
        chips.append(("存貨偏低", "BDD7EE", "1F4E78"))
        notes.append(f"淺藍底＝存貨偏低（{_hybrid_shortage_hint_text(dim)}），留意缺貨風險")
    if highlight_zero_sales:
        chips.append(("零銷仍有庫存", "FFC7CE", "9C0006"))
        notes.append("紅底＝零銷組仍有庫存（零銷售而仍有存貨，理論上全屬 Overstock）")
    if not chips:
        return
    for i, (text, fill, color) in enumerate(chips, 1):
        cell = ws.cell(row=1, column=i, value=text)
        cell.fill = PatternFill(start_color=fill, end_color=fill, fill_type="solid")
        cell.font = Font(bold=True, size=10, color=color)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        _col = get_column_letter(i)
        ws.column_dimensions[_col].width = max(ws.column_dimensions[_col].width or 0, 13)
    ws.cell(row=2, column=1, value="顏色說明：" + "；".join(notes)).font = Font(size=9, color="374151")
    ws.row_dimensions[1].height = 20


def _write_stock_checking_summary_sheet(ws, summary, progress_callback=None, progress_start: float = 0.05, progress_end: float = 0.30, bidirectional: bool = False, highlight_skew: bool = False, dimension: str = "nature", om_groups=None, highlight_zero_sales: bool = False, highlight_shortage: bool = True):
    """Summary 工作表：SKU 級彙總（雙列表頭模仿 By SKUxShop_Stock Checking_Output 樞紐佈局）。

    欄位配置按分類維度分組數 G（nature 3／rp 2／both 6／om N／nature_om 3×N／
    sales 4／sales_nature 12）動態排列：
    - A-G：Brand / SKU / SKU Desc / Buyer / Adj. SST / Supply Source / D001
    - H 起：各分組加總庫存 / 加總月銷 / TurnDay / 涉及店舖數 ×G（整數顯示，TurnDay 活公式四捨五入到整數日）
    - 其後：各分組佔比 (%) ×G（活公式，0.00% 顯示；合計 0 留空）
    - 其後：強制調貨提示 / 各分組預計調出數量（nature 維度＝NT、Mix 兩欄固定順序，
      雙向時「由Tourist」欄附於存貨偏重提示之後，A–AC 契約欄序不動；
      rp／both／om／nature_om／sales／sales_nature 維度＝每組各一欄）
    - 最後：存貨偏重提示

    nature 維度（預設）之欄位號與 v1.22 版一致再各組加一欄涉及店舖數
    （加總區每組 4 欄；佔比 T–Y、提示 Z、調出 AA–AB、偏重 AC、雙向 AD）。
    highlight_skew=True 時，存貨偏重提示非空的列以醒目橘底（FFC000）加強顯示
    提示欄與偏重分組的庫存佔比欄。
    highlight_zero_sales=True 時（僅 sales／sales_nature 維度生效），零銷組的
    庫存佔比欄（加總 - Shop Stock Qty (CM) %）只要該組庫存 > 0（即佔比 > 0%）
    即以醒目紅底（FFC7CE）加粗深紅字標示——零銷售而仍有存貨，理論上全屬 Overstock，
    不設佔比門檻；同格與存貨偏重橘底重疊時以紅底優先。Stock Checking 不傳本參數
    （預設 False，輸出零變動）。
    highlight_shortage=True 時（預設開啟，兩模式全維度生效），存貨偏低分組
    （銷量佔比達標且銷量佔比 − 庫存佔比達差距門檻，門檻較偏重側寬鬆，
    見 _HYBRID_SHORTAGE_THRESHOLDS）的庫存佔比欄以淺藍底（BDD7EE）加粗深藍字標示——
    庫存佔比明顯低於銷量佔比，提示潛在缺貨風險；只標示儲存格，不新增文字欄；
    與橘底／紅底互斥（同一儲存格不會同時觸發），欄位、數值與排序一律不變（純標示）。
    om_groups 缺省時由 summary 欄位反推（om／nature_om 維度）；無法反推時退回單一群組。
    """
    dim = _hybrid_dimension(dimension)
    if dim in ("om", "nature_om") and not om_groups and summary is not None and not summary.empty:
        om_groups = _hybrid_om_groups_from_summary_columns(summary, dim)
    groups = _hybrid_dimension_groups(dim, om_groups=om_groups)
    G = len(groups)
    base_col = 8
    pct_start = base_col + 4 * G
    tip_col = pct_start + 2 * G
    if dim == "nature":
        out_keys = ["由Non Tourist 店舖預計調出數量", "由Mix 店舖預計調出數量"]
        tail_keys = ["由Tourist 店舖預計調出數量"] if bidirectional else []
        block_label, pct_label = "Shop Nature", "各 Nature 佔比 (%)"
    else:
        out_keys = [f"由{g} 店舖預計調出數量" for g in groups]
        tail_keys = []
        block_label = _HYBRID_DIMENSION_LABELS[dim]
        pct_label = f"各 {block_label} 佔比 (%)"
    hint_col = tip_col + 1 + len(out_keys)
    _zero_idx = (
        [i for i, g in enumerate(groups) if g in set(_hybrid_zero_sales_groups(dim))]
        if highlight_zero_sales else []
    )

    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    header_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
    nature_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    pct_fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
    band_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    skew_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
    skew_font = Font(bold=True)
    zero_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    zero_font = Font(bold=True, color="9C0006")
    shortage_fill = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
    shortage_font = Font(bold=True, color="1F4E78")
    _short_share, _short_gap = _hybrid_shortage_thresholds(dim)
    header_font = Font(bold=True, size=11)
    nature_font = Font(bold=True, size=11, color="FFFFFF")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def _style_merge(row, c1, c2, value, font, fill):
        ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
        for cc in range(c1, c2 + 1):
            cell = ws.cell(row=row, column=cc)
            cell.border = thin_border
            cell.alignment = center
            cell.fill = fill
            cell.font = font
        ws.cell(row=row, column=c1, value=value)

    # Row 3: 分組群組標籤 + 佔比群組
    _style_merge(3, base_col, pct_start - 1, block_label, header_font, PatternFill(fill_type=None))
    _style_merge(3, pct_start, tip_col - 1, pct_label, header_font, PatternFill(fill_type=None))

    # Row 4: 各分組標籤 — 加總區（橫向合併 4 欄）+ 佔比區（橫向合併 2 欄）
    for i, gname in enumerate(groups):
        _style_merge(4, base_col + i * 4, base_col + i * 4 + 3, gname, nature_font, nature_fill)
    for i, gname in enumerate(groups):
        _style_merge(4, pct_start + i * 2, pct_start + i * 2 + 1, gname, nature_font, pct_fill)

    # Row 5: 欄位標題
    _headers = [
        "Brand", "SKU", "SKU Desc", "Buyer", "Adj. SST", "Supply Source", "D001 Net Stock Qty (CM)",
    ]
    for _ in groups:
        _headers += ["加總 - Shop Stock Qty (CM)", "加總 - Shop Monthly Sold Qty (P90D)", "TurnDay", "涉及店舖數"]
    for _ in groups:
        _headers += ["加總 - Shop Stock Qty (CM) %", "加總 - Shop Monthly Sold Qty (P90D) %"]
    _headers += ["強制調貨提示"] + out_keys + ["存貨偏重提示"] + tail_keys
    for col_idx, h in enumerate(_headers, 1):
        cell = ws.cell(row=5, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin_border
    ws.row_dimensions[3].height = 20
    ws.row_dimensions[4].height = 28
    ws.row_dimensions[5].height = 45

    if summary is None or summary.empty:
        ws.cell(row=6, column=1, value="無符合資料").font = Font(size=12)
        _auto_width(ws)
        _write_summary_color_legend(ws, dimension, highlight_skew, bool(_zero_idx), highlight_shortage)
        return

    _text_keys = {"Brand", "SKU Desc", "Buyer", "Adj. SST"}
    _center_text_keys = {"SKU", "Supply Source"}

    total = len(summary)
    n_cols = len(_headers)
    _stock_letters = [get_column_letter(base_col + i * 4) for i in range(G)]
    _sold_letters = [get_column_letter(base_col + i * 4 + 1) for i in range(G)]
    # P0：TurnDay 公式欄字母只與分組有關，與列無關，預先計算避免逐列重算；值與公式語意不變
    _s_letters = [get_column_letter(base_col + i * 4) for i in range(G)]
    _m_letters = [get_column_letter(base_col + i * 4 + 1) for i in range(G)]
    for r_idx, (_, row) in enumerate(summary.iterrows()):
        excel_row = r_idx + 6
        is_band = (r_idx % 2 == 1)
        # A-G：Brand / SKU / SKU Desc / Buyer / Adj. SST / Supply Source / D001
        for col_idx, key in enumerate(
            ["Brand", "SKU", "SKU Desc", "Buyer", "Adj. SST", "Supply Source", "D001 Net Stock Qty (CM)"], 1
        ):
            val = row.get(key)
            cell = ws.cell(row=excel_row, column=col_idx, value=None if pd.isna(val) else val)
            cell.border = thin_border
            if key in _text_keys:
                cell.alignment = left_wrap
            else:
                cell.alignment = center
            if col_idx == 7:
                cell.number_format = "#,##0"
            if is_band:
                cell.fill = band_fill
        # 各分組加總、TurnDay 與涉及店舖數（整數；TurnDay 活公式四捨五入）
        for i, gname in enumerate(groups):
            bcol = base_col + i * 4
            stock = row.get(f"{gname} 加總 - Shop Stock Qty (CM)", 0)
            msold = row.get(f"{gname} 加總 - Shop Monthly Sold Qty (P90D)", 0)
            nshop = row.get(f"{gname} 涉及店舖數", 0)
            c_stock = ws.cell(row=excel_row, column=bcol,
                              value=int(stock) if pd.notna(stock) else None)
            c_sold = ws.cell(row=excel_row, column=bcol + 1,
                             value=int(msold) if pd.notna(msold) else None)
            c_nshop = ws.cell(row=excel_row, column=bcol + 3,
                              value=int(nshop) if pd.notna(nshop) else None)
            for cc in (c_stock, c_sold, c_nshop):
                cc.number_format = "#,##0"
                cc.border = thin_border
                cc.alignment = center
                if is_band:
                    cc.fill = band_fill
            s_letter = _s_letters[i]
            m_letter = _m_letters[i]
            if msold is not None and pd.notna(msold) and float(msold) > 0:
                c_td = ws.cell(row=excel_row, column=bcol + 2,
                               value=f"=ROUND({s_letter}{excel_row}/{m_letter}{excel_row}*30.5,0)")
            else:
                c_td = ws.cell(row=excel_row, column=bcol + 2, value=None)
            c_td.number_format = "#,##0"
            c_td.border = thin_border
            c_td.alignment = center
            if is_band:
                c_td.fill = band_fill
        # 各分組佔比 (%) 活公式（分母 = 全部分組合計；合計 0 留空）
        _stock_total = "(" + "+".join(f"{L}{excel_row}" for L in _stock_letters) + ")"
        _sold_total = "(" + "+".join(f"{L}{excel_row}" for L in _sold_letters) + ")"
        for i, gname in enumerate(groups):
            for j, (num_letter, denom) in enumerate(((_stock_letters[i], _stock_total), (_sold_letters[i], _sold_total))):
                denom_bare = denom.strip("()")
                c_pct = ws.cell(row=excel_row, column=pct_start + i * 2 + j,
                                value=f'=IF({denom_bare}=0,"",{num_letter}{excel_row}/{denom})')
                c_pct.number_format = "0.00%"
                c_pct.border = thin_border
                c_pct.alignment = center
                if is_band:
                    c_pct.fill = band_fill
        # 強制調貨提示
        tip = row.get("強制調貨提示", "")
        cell = ws.cell(row=excel_row, column=tip_col, value=tip if tip else None)
        cell.border = thin_border
        cell.alignment = center
        if is_band:
            cell.fill = band_fill
        # 各分組預計調出數量
        for j, key in enumerate(out_keys):
            val = row.get(key, 0)
            cell = ws.cell(row=excel_row, column=tip_col + 1 + j,
                           value=int(val) if pd.notna(val) else None)
            cell.number_format = "#,##0"
            cell.border = thin_border
            cell.alignment = center
            if is_band:
                cell.fill = band_fill
        # 存貨偏重提示
        hint = row.get("存貨偏重提示", "")
        hint = "" if pd.isna(hint) else str(hint)
        cell = ws.cell(row=excel_row, column=hint_col, value=hint if hint else None)
        cell.border = thin_border
        cell.alignment = left_wrap
        if is_band:
            cell.fill = band_fill
        # 偏重數值以醒目橘底加強顯示（提示欄 + 偏重分組的庫存佔比欄）
        if hint and highlight_skew:
            cell.fill = skew_fill
            cell.font = skew_font
            for i, gname in enumerate(groups):
                if f"存貨偏重 {gname}" in hint:
                    _pc = ws.cell(row=excel_row, column=pct_start + i * 2)
                    _pc.fill = skew_fill
                    _pc.font = skew_font
        # 零銷組仍有庫存（庫存佔比 > 0%）以醒目紅底標示——零銷售而仍有存貨，
        # 理論上全屬 Overstock；同格與存貨偏重橘底重疊時以紅底優先
        _zero_cells = set()
        for i in _zero_idx:
            _zstock = row.get(f"{groups[i]} 加總 - Shop Stock Qty (CM)", 0)
            if pd.notna(_zstock) and float(_zstock) > 0:
                _zc = ws.cell(row=excel_row, column=pct_start + i * 2)
                _zc.fill = zero_fill
                _zc.font = zero_font
                _zero_cells.add(i)
        # 存貨偏低標示（庫存佔比明顯低於銷量佔比）以淺藍底標示庫存佔比欄——
        # 銷量佔比達標且銷量佔比 − 庫存佔比達差距門檻（門檻較偏重側寬鬆）；
        # 只標示儲存格，不新增文字欄；與橘底／紅底互斥（已標示橘底／紅底的儲存格不重疊）。
        if highlight_shortage:
            _stock_sums = []
            _sold_sums = []
            for gname in groups:
                _s = row.get(f"{gname} 加總 - Shop Stock Qty (CM)", 0)
                _m = row.get(f"{gname} 加總 - Shop Monthly Sold Qty (P90D)", 0)
                _stock_sums.append(0.0 if pd.isna(_s) else float(_s))
                _sold_sums.append(0.0 if pd.isna(_m) else float(_m))
            _total_stock = sum(_stock_sums)
            _total_sold = sum(_sold_sums)
            if _total_sold > 0:
                for i, gname in enumerate(groups):
                    if i in _zero_cells:
                        continue
                    _pc_col = pct_start + i * 2
                    _pc_cell = ws.cell(row=excel_row, column=_pc_col)
                    # 已有橘底（偏重）不重疊
                    try:
                        _is_skew = (
                            _pc_cell.fill.start_color.rgb is not None
                            and str(_pc_cell.fill.start_color.rgb).endswith("FFC000")
                        )
                    except Exception:
                        _is_skew = False
                    if _is_skew:
                        continue
                    s_share = (_stock_sums[i] / _total_stock) if _total_stock > 0 else 0.0
                    v_share = _sold_sums[i] / _total_sold
                    if v_share >= _short_share and (v_share - s_share) >= _short_gap:
                        _pc_cell.fill = shortage_fill
                        _pc_cell.font = shortage_font
        # nature 維度雙向：由Tourist 店舖預計調出數量（附於最後，A–AC 契約欄序不動）
        for j, key in enumerate(tail_keys):
            val = row.get(key, 0)
            cell = ws.cell(row=excel_row, column=hint_col + 1 + j,
                           value=int(val) if pd.notna(val) else None)
            cell.number_format = "#,##0"
            cell.border = thin_border
            cell.alignment = center
            if is_band:
                cell.fill = band_fill
        if progress_callback and (r_idx % 200 == 0 or r_idx == total - 1):
            p = progress_start + (progress_end - progress_start) * ((r_idx + 1) / max(total, 1))
            progress_callback(p, f"寫入 Summary ({r_idx+1}/{total})...")

    last_row = 5 + total
    ws.auto_filter.ref = f"A5:{get_column_letter(n_cols)}{last_row}"
    ws.freeze_panes = "A6"
    _auto_width(ws)
    _write_summary_color_legend(ws, dimension, highlight_skew, bool(_zero_idx), highlight_shortage)
