# -*- coding: utf-8 -*-
"""Overstock Review模式：由 transformer.py 拆分而來。
共用 helper 由 transformer（最終為 modes.common／issues／framework）引用；
transformer 以墊片 re-export 本模組，舊有 import 路徑不變。
"""
import datetime
import gc
import re
import unicodedata
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


from modes.common import (
    _ADJ_SAFETY_SOURCE_COL,
    _auto_width,
    _cached_shop_names,
    _clean_dataframe,
    _compute_raw_extra_columns,
    _display_safety_header,
    _extract_shop_name,
    _iter_header_rows,
    _nature_rate_arr,
    _resolve_column_name,
    _safe_default_series,
    _safe_numeric,
    _safe_sku,
    _cached_safe_sku_series,
    _write_metric_cards,
    _write_raw_placeholder,
    _write_raw_sheet,
    validate_input,
)
__all__ = [
    "_add_recommendation_columns",
    "_build_by_shop_summary",
    "generate_excel",
    "transform",
]


def _build_by_shop_summary(raw: pd.DataFrame) -> pd.DataFrame:
    df = raw.copy()

    site_code = df.get("Site Code", _safe_default_series("", df.index)).fillna("").astype(str).str.strip()
    item_shop = df.get("Item_Shop", _safe_default_series("", df.index))
    # P0：快取解析（同一店舖×品項只執行一次正則），輸出與逐列解析完全一致
    shop_name = _cached_shop_names(item_shop, site_code)

    df["Site Code"] = site_code
    df["Shop Name"] = shop_name
    df["OM"] = df.get("OM", _safe_default_series("", df.index)).fillna("").astype(str).str.strip()
    df["SKU"] = df["SKU"].apply(_safe_sku)

    status_col = df["Shop Stock Status"].fillna("")
    rec_col = df["Recommendation of new Safety Qty / ND"].fillna("").astype(str)
    rec_num = pd.to_numeric(rec_col, errors="coerce")
    transfer_col = "After new Safety Qty / ND can transfer out of QTY"
    transfer_num = pd.to_numeric(df.get(transfer_col, _safe_default_series(0, df.index)), errors="coerce").fillna(0)

    group_cols = ["Site Code", "Shop Name", "OM"]
    df["_is_0_sales"] = status_col == "Overstock (P90D 0 Sales)"
    df["_is_2x"] = status_col == "Overstock (>2*Monthly Sold Qty & >2pcs)"
    df["_is_nd_change"] = df["_is_0_sales"] & (rec_col == "ND")
    df["_is_safety_change"] = df["_is_2x"] & rec_num.notna()
    df["_sku_0_sales"] = df["SKU"].where(df["_is_0_sales"])
    df["_sku_2x"] = df["SKU"].where(df["_is_2x"])
    df["_transfer_num"] = transfer_num

    by_shop = (
        df.groupby(group_cols, dropna=False, observed=True)
        .agg(
            更改NDRF行數=("_is_nd_change", "sum"),
            更改Safety行數=("_is_safety_change", "sum"),
            預計總轉出貨量=("_transfer_num", "sum"),
            **{
                "P90D 0 Sales SKU數量": ("_sku_0_sales", "nunique"),
                ">2*Monthly Sold Qty & >2pcs數量": ("_sku_2x", "nunique"),
            },
        )
        .reset_index()
    )

    by_shop["更改NDRF行數"] = by_shop["更改NDRF行數"].fillna(0).astype(int)
    by_shop["更改Safety行數"] = by_shop["更改Safety行數"].fillna(0).astype(int)
    by_shop["預計總轉出貨量"] = by_shop["預計總轉出貨量"].fillna(0).astype(int)
    by_shop["P90D 0 Sales SKU數量"] = by_shop["P90D 0 Sales SKU數量"].fillna(0).astype(int)
    by_shop[">2*Monthly Sold Qty & >2pcs數量"] = by_shop[">2*Monthly Sold Qty & >2pcs數量"].fillna(0).astype(int)

    if not by_shop.empty:
        total_row = pd.DataFrame(
            {
                "Site Code": ["總計"],
                "Shop Name": [""],
                "OM": [""],
                "更改NDRF行數": [int(by_shop["更改NDRF行數"].sum())],
                "更改Safety行數": [int(by_shop["更改Safety行數"].sum())],
                "預計總轉出貨量": [int(by_shop["預計總轉出貨量"].sum())],
                "P90D 0 Sales SKU數量": [int(df["_sku_0_sales"].dropna().nunique())],
                ">2*Monthly Sold Qty & >2pcs數量": [int(df["_sku_2x"].dropna().nunique())],
            }
        )
        by_shop = pd.concat([by_shop.sort_values(["Site Code", "Shop Name", "OM"], kind="stable"), total_row], ignore_index=True)

    return by_shop



def _add_recommendation_columns(df: pd.DataFrame, rates: dict | None = None) -> pd.DataFrame:
    status_col = df["Shop Stock Status"].fillna("")

    adj_ss = _safe_numeric(df.get("Adj. Safety Stock", _safe_default_series(0, df.index)))
    monthly_sold = _safe_numeric(df.get("Shop Monthly Sold Qty (P90D)", _safe_default_series(0, df.index)))

    # Normalize RP Type (strip + upper) so values like " rf " or "nd" are recognized,
    # consistent with the OOS / Carry review paths.
    rp_type = (
        df.get("RP Type", _safe_default_series("", df.index))
        .fillna("")
        .astype(str)
        .str.strip()
        .str.upper()
    )

    cond_overstock_2x = (status_col == "Overstock (>2*Monthly Sold Qty & >2pcs)")

    cond_overstock_0 = (
        (status_col == "Overstock (P90D 0 Sales)") & rp_type.isin(["ND", "RF"])
    )

    rec_values = np.full(len(df), "", dtype=object)

    if cond_overstock_2x.any():
        shop_nature = df.get("Shop Nature", _safe_default_series("", df.index)).fillna("").astype(str).str.lower()
        _rates = rates if rates is not None else {}
        _default_rate = _rates.get("default", 0.50)
        _tourist_rate = _rates.get("tourist", 0.30)
        _mix_rate = _rates.get("mix", 0.40)
        rate = _nature_rate_arr(shop_nature, _default_rate, _tourist_rate, _mix_rate)
        # Shop Monthly Sold Qty (P90D) 已是月均（P90D 總量 ÷ 3），直接乘 rate
        rec_2x = np.floor(monthly_sold[cond_overstock_2x] * rate[cond_overstock_2x]).clip(lower=3).astype(int)
        # If new safety qty > Adj. Safety Stock, cap at Adj. Safety Stock
        adj_ss_2x = adj_ss[cond_overstock_2x].values
        over_adj = rec_2x > adj_ss_2x
        rec_2x[over_adj] = adj_ss_2x[over_adj].astype(int)
        rec_values[cond_overstock_2x.values] = rec_2x.astype(str)

    if cond_overstock_0.any():
        is_rf = cond_overstock_0 & (rp_type == "RF")
        is_nd = cond_overstock_0 & (rp_type == "ND")
        rec_values[is_rf.values] = "ND"
        rec_values[is_nd.values] = "0"
    df["Recommendation of new Safety Qty / ND"] = rec_values

    shop_stock = _safe_numeric(df.get("Shop Stock Qty (CM)", _safe_default_series(0, df.index)))

    rec = df["Recommendation of new Safety Qty / ND"]

    # Numeric recommendation (from >2x): transfer = shop_stock - rec_num
    rec_num = pd.to_numeric(rec, errors="coerce")
    mask_rec_num = rec_num.notna() & (shop_stock > rec_num)

    # ND recommendation: transfer = all shop stock
    mask_nd = (rec == "ND") & (shop_stock > 0)

    transfer_values = np.full(len(df), "", dtype=object)
    if mask_rec_num.any():
        transfer_values[mask_rec_num.values] = (shop_stock[mask_rec_num] - rec_num[mask_rec_num]).astype(int).values
    if mask_nd.any():
        transfer_values[mask_nd.values] = shop_stock[mask_nd].astype(int).values

    mask_rec_zero = rec_num.notna() & (shop_stock <= rec_num)
    if mask_rec_zero.any():
        transfer_values[mask_rec_zero.values] = 0

    mask_nd_zero = (rec == "ND") & (shop_stock == 0)
    if mask_nd_zero.any():
        transfer_values[mask_nd_zero.values] = 0

    df["After new Safety Qty / ND can transfer out of QTY"] = transfer_values

    df["Buyer reply"] = ""

    return df



def transform(df_input: pd.DataFrame, progress_callback=None, shop_natures=None, rates=None):
    missing = validate_input(df_input)
    if missing:
        raise ValueError(f"\u7f3a\u5c11\u5fc5\u8981\u6b04\u4f4d: {', '.join(missing)}")

    warnings = []

    # Safety 來源欄相容檢查：顯示名一律視為來源名（Safety Stock (現時) = 來源檔案的 Adj. Safety Stock）
    _adj_avail = set(df_input.columns) | {_resolve_column_name(_ADJ_SAFETY_SOURCE_COL, set(df_input.columns))}
    _optional_cols = ["Adj. Safety Stock", "Shop Stock Qty (CM)", "Shop Monthly Sold Qty (P90D)"]
    for _col in _optional_cols:
        _label = _display_safety_header(_col)
        if _resolve_column_name(_col, _adj_avail) is None:
            warnings.append(f"缺少可選欄位「{_label}」，建議值與轉移數量將全為 0")

    if progress_callback:
        progress_callback(0.05, "\u9a57\u8b49\u901a\u904e")

    df = _clean_dataframe(df_input, shop_natures=shop_natures, progress_callback=progress_callback)

    if df.empty:
        raise ValueError("過濾後無有效資料列，請確認來源檔案中的 Buyer 欄位是否包含有效數據（非空白或總計列）")

    df["SKU"] = df["SKU"].apply(_safe_sku)

    if "SKU Desc" not in df.columns:
        df = _compute_raw_extra_columns(df)

    status_col = df["Shop Stock Status"].fillna("")

    categories = [
        "P90D 0 Sales",
        "Overstock (>2*Monthly Sold Qty & >2pcs)",
    ]

    _STATUS_VALUE = {
        "P90D 0 Sales": "Overstock (P90D 0 Sales)",
        "Overstock (>2*Monthly Sold Qty & >2pcs)": "Overstock (>2*Monthly Sold Qty & >2pcs)",
    }

    group_cols = ["Buyer", "Brand", "SKU", "SKU Desc"]

    if progress_callback:
        progress_callback(0.2, "\u5206\u985e\u7d71\u8a08\u4e2d...")

    # 預計算 int 旗標（單次 groupby 取代多次 filter+groupby + transfer merge）
    _agg_dict = {}
    for cat in categories:
        mask = (status_col == _STATUS_VALUE.get(cat, cat))
        key = f"_cat_{abs(hash(cat)) % 10_000_000}"
        df[key] = mask.astype(int)
        _agg_dict[cat] = (key, "sum")

    if progress_callback:
        progress_callback(0.5, "計算庫存建議中...")

    raw = _add_recommendation_columns(df, rates=rates)

    transfer_col = "After new Safety Qty / ND can transfer out of QTY"
    raw["_transfer_num"] = pd.to_numeric(raw.get(transfer_col, _safe_default_series(0, raw.index)), errors="coerce").fillna(0)
    _agg_dict[transfer_col] = ("_transfer_num", "sum")

    if progress_callback:
        progress_callback(0.55, "排序彙總中...")

    summary = raw.groupby(group_cols, observed=True).agg(**_agg_dict).fillna(0).astype(int).reset_index()

    summary["SKU"] = summary["SKU"].apply(_safe_sku)
    summary["總計"] = summary[categories].sum(axis=1)

    buyer_order = summary.groupby("Buyer")["總計"].sum().sort_values(ascending=False).index.tolist()
    summary["_buyer_key"] = summary["Buyer"].map({b: i for i, b in enumerate(buyer_order)})
    summary = summary.sort_values(["_buyer_key", "Brand", "SKU", "SKU Desc"], ascending=[True, True, True, True]).drop(columns=["_buyer_key"]).reset_index(drop=True)

    transfer_total = int(raw["_transfer_num"].sum())
    del raw["_transfer_num"]

    total_overstock = int(summary["總計"].sum())
    total_row = pd.DataFrame(
        {
            "Buyer": ["總計"],
            "Brand": [None],
            "SKU": [None],
            "SKU Desc": [None],
            **{cat: [int(summary[cat].sum())] for cat in categories},
            "\u7e3d\u8a08": [total_overstock],
            transfer_col: [transfer_total],
        }
    )
    summary = pd.concat([summary, total_row], ignore_index=True)

    # 清理暫存欄位
    raw.drop(columns=[c for c in raw.columns if c.startswith("_cat_") or c == "_transfer_num"], inplace=True, errors="ignore")

    if total_overstock == 0:
        warnings.append(
            "\u672a\u5075\u6e2c\u5230\u4efb\u4f55\u7b26\u5408\u689d\u4ef6\u7684\u5eab\u5b58\u7570\u5e38\u8cc7\u6599\u3002\u8acb\u78ba\u8a8d\u4f86\u6e90\u6a94\u6848\u4e2d\u7684\u300cShop Stock Status\u300d\u6b04\u4f4d\u503c\u662f\u5426\u5305\u542b\u4ee5\u4e0b\u5206\u985e\uff1a"
            + ", ".join(categories)
        )

    if progress_callback:
        progress_callback(0.7, "\u8f49\u63db\u5b8c\u6210")

    return summary, raw, categories, warnings



def generate_excel(summary_df: pd.DataFrame, raw_df: pd.DataFrame, categories: list[str], progress_callback=None, include_raw: bool = True) -> BytesIO:
    wb = Workbook()

    if progress_callback:
        progress_callback(0.0, "建立 Excel 結構...")

    ws_raw = wb.active
    ws_raw.title = "RAW"

    (_write_raw_sheet(ws_raw, raw_df, progress_callback, progress_start=0.0, progress_end=0.3) if include_raw else _write_raw_placeholder(ws_raw))

    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    if progress_callback:
        progress_callback(0.35, "建立摘要工作表...")

    ws_summary = wb.create_sheet(title="Shop Stock Status")

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    bold_font = Font(bold=True, size=11)

    ws_summary["A3"] = "\u884c\u6578 - Site Code"
    ws_summary["D3"] = "Shop Stock Status"

    transfer_col = "After new Safety Qty / ND can transfer out of QTY"
    headers = ["Buyer", "Brand", "SKU", "SKU Desc"] + categories + ["總計"] + [transfer_col]
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col_idx, h in enumerate(headers, 1):
        cell = ws_summary.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    data_start = 5
    total_summary = len(summary_df)
    # P0：itertuples＋共用對齊取代 iterrows＋逐格新建 Alignment；值與格式語意不變
    center_align = Alignment(horizontal="center", vertical="center")

    for row_idx, _vals in enumerate(_iter_header_rows(summary_df, headers)):
        excel_row = data_start + row_idx
        is_total = _vals[0] == "總計"

        for col_idx, h in enumerate(headers, 1):
            val = _vals[col_idx - 1]
            if pd.isna(val):
                val = None
            cell = ws_summary.cell(row=excel_row, column=col_idx, value=val)
            cell.border = thin_border
            cell.alignment = center_align

            if is_total:
                cell.font = bold_font

            if isinstance(val, (int, float)) and col_idx >= 4 and h != "SKU":
                cell.number_format = "#,##0"

        if progress_callback and (row_idx % 500 == 0 or row_idx == total_summary - 1):
            progress_callback(0.6 + 0.3 * ((row_idx + 1) / total_summary), f"\u5beb\u5165\u6458\u8981 ({row_idx+1}/{total_summary})...")

    _auto_width(ws_summary)

    if progress_callback:
        progress_callback(0.72, "\u5efa\u7acb By Shop \u5206\u6790\u5de5\u4f5c\u8868...")

    by_shop_df = _build_by_shop_summary(raw_df)

    ws_by_shop = wb.create_sheet(title="By Shop")
    by_shop_headers = [
        "Site Code",
        "Shop Name",
        "OM",
        "更改NDRF行數",
        "更改Safety行數",
        "預計總轉出貨量",
        "P90D 0 Sales SKU數量",
        ">2*Monthly Sold Qty & >2pcs數量",
    ]

    for col_idx, h in enumerate(by_shop_headers, 1):
        cell = ws_by_shop.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    for row_idx, _vals in enumerate(_iter_header_rows(by_shop_df, by_shop_headers)):
        excel_row = row_idx + 2
        is_total = _vals[0] == "總計"
        for col_idx, h in enumerate(by_shop_headers, 1):
            val = _vals[col_idx - 1]
            if pd.isna(val):
                val = None
            cell = ws_by_shop.cell(row=excel_row, column=col_idx, value=val)
            cell.border = thin_border
            cell.alignment = center_align
            if is_total:
                cell.font = bold_font
            if isinstance(val, (int, float)) and col_idx >= 4:
                cell.number_format = "#,##0"

        if progress_callback and (row_idx % 500 == 0 or row_idx == len(by_shop_df) - 1):
            progress_callback(0.82 + 0.08 * ((row_idx + 1) / max(len(by_shop_df), 1)), f"\u5beb\u5165 By Shop ({row_idx+1}/{len(by_shop_df)})...")

    _auto_width(ws_by_shop)

    if progress_callback:
        progress_callback(0.92, "\u5efa\u7acb\u5206\u6790\u6458\u8981\u5de5\u4f5c\u8868...")

    ws_dashboard = wb.create_sheet(title="異常庫存數據分析摘要")

    _total_row = summary_df[summary_df["Buyer"] == "總計"]
    if _total_row.empty:
        raise ValueError("摘要資料中未找到總計列，無法生成分析看板")
    total_overstock = int(_total_row["總計"].values[0])
    total_buyers = len(summary_df[summary_df["Buyer"] != "總計"]["Buyer"].dropna().unique())
    total_brands = len(summary_df[summary_df["Buyer"] != "總計"]["Brand"].dropna().unique())
    total_skus = raw_df["SKU"].dropna().nunique()
    cat1_val = int(_total_row[categories[0]].values[0])
    cat2_val = int(_total_row[categories[1]].values[0])

    transfer_col = _resolve_column_name("After new Safety Qty / ND can transfer out of QTY", set(raw_df.columns))
    transfer_qty_total = 0
    if transfer_col:
        transfer_qty_total = int(pd.to_numeric(raw_df[transfer_col], errors="coerce").fillna(0).sum())

    row1_metrics = [
        ("總 Overstock 行數", total_overstock, "所有庫存異常之總計行數", "EF4444"),
        ("影響採購人員數 (Buyers)", total_buyers, "位採購人員名下有異常品項", "3B82F6"),
        ("影響品牌總數 (Brands)", total_brands, "個品牌包含異常庫存商品", "8B5CF6"),
        ("影響SKU總數", total_skus, "個包含異常庫存商品(SKUs)", "10B981"),
    ]

    row2_metrics = [
        ("P90D 0 Sales", cat1_val, "「過去 90 天內零銷售」之異常行數", "F59E0B"),
        ("Overstock (>2*Monthly Sold Qty & >2pcs)", cat2_val, "「大於月均銷售量兩倍且庫存大於 2 件」之異常行數", "EC4899"),
        ("調整後的預計可退倉總數", transfer_qty_total, "After new Safety Qty / ND can transfer out of TOTAL QTY", "06B6D4"),
    ]

    desc_font = Font(size=9, color="6B7280", italic=True)
    _write_metric_cards(ws_dashboard, "異常庫存數據分析摘要 (Analysis Dashboard)", row1_metrics, row2_metrics, thin_border, desc_font, progress_callback)

    if progress_callback:
        progress_callback(0.95, "正在儲存 Excel 檔案...")

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    if progress_callback:
        progress_callback(1.0, "完成!")

    return buf

