# -*- coding: utf-8 -*-
"""P90D 0 Sales Review(轉貨模式)：指定轉出店舖（0 Sales 且有庫存）轉貨至有銷售記錄的店舖。

流程：
1. ``build_transfer_candidates``：列出候選轉出店舖（Sold Qty (P90D) ≤ 0 且 Shop Stock Qty (CM) > 0；
   ``Shop Stock Qty (LM)``／``(LLM)`` 欄存在時另須 > 0——90 日有貨準則；``Launch Date`` 欄存在時
   排除 90 天內新品），供介面勾選指定轉出店舖。
2. ``transform_p90d_transfer``：依勾選的轉出店舖產生轉貨單與三張彙總表。

規則（不分 OM 群組，轉貨數量不設上限）：
- 轉出：已選轉出店舖中符合上述 0 Sales（含 90 日有貨準則與新品除外）且有庫存的列，
  轉出數量為 Shop Stock Qty (CM) 全數。
- 接收：同 SKU、非轉出店舖、Sold Qty (P90D) > 0 的列（有銷售記錄），不論有無庫存；
  ``RP Type`` 欄存在時排除 RP Type＝ND 的店舖（ND 不再補貨，不作調入需求方）。
- 分配：同 SKU 的轉出總量按接收店 Sold Qty (P90D) 比例分配，以最大餘數法取整數件、全數分完；
  只有一間接收店時全數給該店；無接收店時該 SKU 不轉出，列入未分配。
- 接收量上限（可選，「補到 14 日」封頂）：每間接收店最多接收其 14 日需求
  （Sold Qty (P90D) ÷ 90 × 14 無條件進位 − 調前庫存，最低 0）；超出上限的餘量在尚有
  需求的接收店之間按銷售比例重新分配，仍未派完者列未分配（保留原店）。
- MCH 篩選（選填）：執行前可按 Adj. MCH1 排除部分分類（傳 None 或包含全部值＝不排除），
  對轉出與接收兩側同時生效；報表另附 MCH1 分類彙總（按 Adj. MCH1）。
- 資料粒度：「Site Code × SKU」重複列一律停止計算（重複列會重複計算庫存與銷售），
  與調貨段（Safety Stock x NetStock Review）同一準則。

``nd_mode=True`` 時改為 ND SKU with Stock Review(轉貨模式) 的判定（轉出＝RP Type＝ND 且有庫存、
接收＝RP Type＝RF 且 Sold Qty (P90D) > 0），分配、配對與彙總沿用同一引擎；
對外介面由 modes/transfer_nd.py 提供（90 日有貨準則只適用 P90D 模式的轉出判定，
ND 模式不看銷售記錄，不受此準則約束）。

只引用 modes.common，不引用其他模式模組，避免循環引用。
"""
import math
from io import BytesIO

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from modes.common import (
    _auto_width,
    _clean_dataframe,
    _compute_raw_extra_columns,
    _normalize_ref_date,
    _parse_launch_dates,
    _resolve_column_name,
    _safe_default_series,
    _safe_numeric,
    _safe_sku,
    _write_guide_sheet,
    validate_input,
)

__all__ = [
    "build_transfer_candidates",
    "build_mch_options",
    "build_mch_breakdown",
    "receiver_limit_label",
    "transform_p90d_transfer",
    "generate_excel_p90d_transfer",
]

_TRANSFER_REQUIRED = ["Site Code", "SKU", "Sold Qty (P90D)", "Shop Stock Qty (CM)"]
# ND 模式（RP Type＝ND 轉出、RP Type＝RF 接收）另要求 RP Type：缺欄即停止（無法辨識轉出列與接收店）
_TRANSFER_ND_REQUIRED = ["RP Type"]

# 選用欄位：存在時套用 90 日有貨準則（LM／LLM）與新品除外（Launch Date）；
# P90D 模式接收店排除 ND 亦以 RP Type 欄存在為前提（缺欄時維持原有接收定義）。
_TRANSFER_LM_COL = "Shop Stock Qty (LM)"
_TRANSFER_LLM_COL = "Shop Stock Qty (LLM)"
_TRANSFER_LAUNCH_COL = "Launch Date"
_TRANSFER_NEW_PRODUCT_DAYS = 90
# 14 日需求線（接收量封頂用）：與調貨段同一公式 ceil(Sold Qty (P90D) ÷ 90 × 14)
_TRANSFER_CAP_DAYS = 14

_TRANSFER_MCH1_COL = "Adj. MCH1"
_TRANSFER_MCH_BLANK_LABEL = "(空白)"
_TRANSFER_MCH_SHEET = "MCH1 分類彙總"
# 空白（含 NaN、空字串）MCH1 值在介面與報表中的表示法
_TRANSFER_MCH_BLANK_TOKEN = "__BLANK__"

# RP Type 比對：去除前後空格、大小寫不敏感；非 RF／ND（空白或其他值）一律不納入轉出與接收
_TRANSFER_RP_COL = "RP Type"
_TRANSFER_RP_RF = "RF"
_TRANSFER_RP_ND = "ND"

_TRANSFER_MCH_SUMMARY_COLUMNS = [
    "Adj. MCH1", "轉出 SKU 數", "接收 SKU 數",
    "轉出總量 (件)", "接收總量 (件)", "未分配量 (件)",
]

_TRANSFER_CANDIDATE_COLUMNS = [
    "Site Code", "Shop Name", "Shop Nature", "OM", "0 Sales SKU 數", "可轉出庫存 (件)",
]

_TRANSFER_ND_CANDIDATE_COLUMNS = [
    "Site Code", "Shop Name", "Shop Nature", "OM", "ND SKU 數", "有銷售 ND SKU 數", "可轉出庫存 (件)",
]

_TRANSFER_PLAN_COLUMNS = [
    "轉出 Site Code", "轉出 Shop Name", "轉出店 Sold Qty (P90D)", "接收 Site Code", "接收 Shop Name",
    "Brand", "SKU", "SKU Desc", "Buyer", "Adj. MCH1", "Supply Source", "轉貨數量 (件)",
    "接收店 Shop Nature", "接收店 OM", "接收店 Sold Qty (P90D)", "接收店銷售佔比 (同 SKU)",
    "接收店同 SKU 分配總量 (件)", "接收店 Shop Stock Qty (CM) 調前", "接收店 14 日需求 (件)",
]

_TRANSFER_OUT_COLUMNS = [
    "Site Code", "Shop Name", "Shop Nature", "OM",
    "轉出 SKU 數", "轉出總量 (件)", "已分配 (件)", "未分配 (件)", "接收店舖數",
]

_TRANSFER_IN_COLUMNS = [
    "Site Code", "Shop Name", "Shop Nature", "OM",
    "接收 SKU 數", "接收總量 (件)", "接收前庫存 (件)", "接收後庫存 (件)",
    "接收店 Sold Qty (P90D) 合計", "涉及轉出店舖數",
]

_TRANSFER_SKU_COLUMNS = [
    "Brand", "SKU", "SKU Desc", "Buyer", "Adj. MCH1", "轉出店舖數", "轉出總量 (件)",
    "可接收店舖數 (有銷售)", "實際接收店舖數", "接收店 Sold Qty (P90D) 合計",
    "已分配總量 (件)", "未分配量 (件)", "狀態",
]

_TRANSFER_PCT_COLUMNS = frozenset({"接收店銷售佔比 (同 SKU)"})
_TRANSFER_STATUS_ALLOCATED = "已全數分配"
_TRANSFER_STATUS_NO_RECEIVER = "無接收店舖（保留原店）"
_TRANSFER_STATUS_CAPPED = "接收店 14 日需求已滿（保留原店）"

# 接收店集中模式：None＝不限（原有行為）；2／4＝每個 SKU 只取銷量最高的 1–2／3–4 間接收店
_TRANSFER_RECEIVER_LIMIT_NONE = "不限（原有：全數有銷售的接收店舖）"
_TRANSFER_RECEIVER_LIMIT_LABELS = {
    2: "集中 1–2 間（同 SKU 銷量最高的接收店）",
    4: "集中 3–4 間（同 SKU 銷量最高的接收店）",
}


def receiver_limit_label(receiver_limit: int | None = None) -> str:
    """接收店集中模式的顯示文字（None＝不限）。"""
    if receiver_limit is None:
        return _TRANSFER_RECEIVER_LIMIT_NONE
    limit = int(receiver_limit)
    return _TRANSFER_RECEIVER_LIMIT_LABELS.get(
        limit, f"集中最多 {limit} 間（同 SKU 銷量最高的接收店）"
    )


def _validate_receiver_limit(receiver_limit: int | None) -> int | None:
    if receiver_limit is None:
        return None
    try:
        limit = int(receiver_limit)
    except (TypeError, ValueError):
        raise ValueError("接收店舖集中上限必須為正整數或「不限」。") from None
    if limit <= 0:
        raise ValueError("接收店舖集中上限必須為正整數或「不限」。")
    return limit


# 「轉貨單」欄位分組著色（表頭底色、資料底色）：黃＝轉出店與轉貨數量、淺藍＝接收店相關欄位。
# 只影響轉貨單的顯示格式，不改任何欄位名稱、欄序或數值。
_TRANSFER_PLAN_HEAD_OUT = "FFD966"
_TRANSFER_PLAN_TINT_OUT = "FFF2CC"
_TRANSFER_PLAN_HEAD_IN = "9DC3E6"
_TRANSFER_PLAN_TINT_IN = "DDEBF7"
_TRANSFER_PLAN_STYLES = {
    "轉出 Site Code": (_TRANSFER_PLAN_HEAD_OUT, _TRANSFER_PLAN_TINT_OUT),
    "轉出 Shop Name": (_TRANSFER_PLAN_HEAD_OUT, _TRANSFER_PLAN_TINT_OUT),
    "轉出店 Sold Qty (P90D)": (_TRANSFER_PLAN_HEAD_OUT, _TRANSFER_PLAN_TINT_OUT),
    "轉貨數量 (件)": (_TRANSFER_PLAN_HEAD_OUT, _TRANSFER_PLAN_TINT_OUT),
    "接收 Site Code": (_TRANSFER_PLAN_HEAD_IN, _TRANSFER_PLAN_TINT_IN),
    "接收 Shop Name": (_TRANSFER_PLAN_HEAD_IN, _TRANSFER_PLAN_TINT_IN),
    "接收店 Shop Nature": (_TRANSFER_PLAN_HEAD_IN, _TRANSFER_PLAN_TINT_IN),
    "接收店 OM": (_TRANSFER_PLAN_HEAD_IN, _TRANSFER_PLAN_TINT_IN),
    "接收店 Sold Qty (P90D)": (_TRANSFER_PLAN_HEAD_IN, _TRANSFER_PLAN_TINT_IN),
    "接收店銷售佔比 (同 SKU)": (_TRANSFER_PLAN_HEAD_IN, _TRANSFER_PLAN_TINT_IN),
    "接收店同 SKU 分配總量 (件)": (_TRANSFER_PLAN_HEAD_IN, _TRANSFER_PLAN_TINT_IN),
    "接收店 Shop Stock Qty (CM) 調前": (_TRANSFER_PLAN_HEAD_IN, _TRANSFER_PLAN_TINT_IN),
    "接收店 14 日需求 (件)": (_TRANSFER_PLAN_HEAD_IN, _TRANSFER_PLAN_TINT_IN),
}
_TRANSFER_PLAN_BOLD_DATA_COLUMNS = ("轉貨數量 (件)",)
_TRANSFER_PLAN_HEADER_ROW_HEIGHT = 30


def _first_nonblank(values: pd.Series) -> str:
    vals = values.fillna("").astype(str).str.strip()
    vals = vals[vals != ""]
    return str(vals.iloc[0]) if len(vals) else ""


def _mode_nonblank(values: pd.Series) -> str:
    vals = values.fillna("").astype(str).str.strip()
    vals = vals[vals != ""]
    return str(vals.value_counts().index[0]) if len(vals) else ""


def _text_column(df: pd.DataFrame, name: str) -> pd.Series:
    return df.get(name, _safe_default_series("", df.index)).fillna("").astype(str).str.strip()


def _check_transfer_columns(df_input: pd.DataFrame, require_rp: bool = False) -> None:
    missing = validate_input(df_input)
    if missing:
        raise ValueError(f"缺少必要欄位: {', '.join(missing)}")
    avail = set(df_input.columns)
    required = list(_TRANSFER_REQUIRED)
    if require_rp:
        required += _TRANSFER_ND_REQUIRED
    missing = [c for c in required if _resolve_column_name(c, avail) is None]
    if missing:
        raise ValueError(f"缺少必需欄位: {', '.join(missing)}")


def _prepare_transfer_base(df_input: pd.DataFrame, shop_natures: list[str] | None = None,
                           progress_callback=None, require_rp: bool = False) -> tuple[pd.DataFrame, list[str], dict]:
    """清理輸入並組出轉貨引擎的基礎欄位；回傳 (base, warnings, meta)。

    ``meta`` 記錄選用欄位是否存在（LM／LLM／Launch Date／RP Type），
    供 90 日有貨準則、新品除外與接收店排除 ND 的「欄存在時生效」判定。
    「Site Code × SKU」重複列一律停止計算（與調貨段同一準則）。
    """
    _check_transfer_columns(df_input, require_rp=require_rp)
    warnings: list[str] = []
    df = _clean_dataframe(df_input, shop_natures=shop_natures, progress_callback=progress_callback)
    df = df.reset_index(drop=True)
    avail = set(df.columns)
    # 選用欄位以實際解析出的欄名取用（None＝來源檔沒有該欄，相關準則不生效）
    meta = {
        "lm_col": _resolve_column_name(_TRANSFER_LM_COL, avail),
        "llm_col": _resolve_column_name(_TRANSFER_LLM_COL, avail),
        "launch_col": _resolve_column_name(_TRANSFER_LAUNCH_COL, avail),
        "rp_col": _resolve_column_name(_TRANSFER_RP_COL, avail),
    }
    if df.empty:
        return pd.DataFrame(), warnings, meta
    if "Item_Shop" not in df.columns:
        warnings.append("來源缺少 Item_Shop 欄位，Shop Name 與 SKU Desc 留空。")
    df = _compute_raw_extra_columns(df)
    base = pd.DataFrame({
        "_site": _text_column(df, "Site Code"),
        "_sku": df["SKU"].apply(_safe_sku).fillna("").astype(str).str.strip(),
        "_sales": _safe_numeric(df["Sold Qty (P90D)"]),
        "_stock": np.floor(_safe_numeric(df["Shop Stock Qty (CM)"]) + 1e-9),
        "_nature": _text_column(df, "Shop Nature"),
        "_om": _text_column(df, "OM"),
        "_name": _text_column(df, "Shop Name"),
        "_rp": _text_column(df, _TRANSFER_RP_COL).str.upper(),
        "_lm": _safe_numeric(df[meta["lm_col"]]) if meta["lm_col"] else np.nan,
        "_llm": _safe_numeric(df[meta["llm_col"]]) if meta["llm_col"] else np.nan,
        "_launch": (
            _parse_launch_dates(df[meta["launch_col"]].fillna("").astype(str).str.strip())
            if meta["launch_col"] else pd.Series(pd.NaT, index=df.index)
        ),
        "Brand": _text_column(df, "Brand"),
        "Buyer": _text_column(df, "Buyer"),
        "SKU Desc": _text_column(df, "SKU Desc"),
        "Adj. MCH1": _text_column(df, _TRANSFER_MCH1_COL),
        "Supply Source": _text_column(df, "Supply Source"),
    })
    keep = (base["_site"] != "") & (base["_sku"] != "")
    skipped = int((~keep).sum())
    if skipped:
        warnings.append(f"{skipped} 列因 Site Code 或 SKU 空白已略過。")
    base = base[keep].reset_index(drop=True)
    dup_n = int(base.duplicated(["_site", "_sku"]).sum())
    if dup_n:
        dup_keys = base.loc[base.duplicated(["_site", "_sku"], keep=False), ["_site", "_sku"]].drop_duplicates()
        examples = "、".join(
            f"{site} / {sku}" for site, sku in dup_keys.head(5).itertuples(index=False, name=None)
        )
        raise ValueError(
            f"偵測到「Site Code × SKU」重複列 {dup_n} 列、{len(dup_keys)} 組（例如：{examples}），"
            "可能重複計算庫存與銷售；每間店舖的每個 SKU 必須只有一列，請修正來源資料後再執行。"
        )
    return base, warnings, meta


def _mch_blank_mask(values: pd.Series) -> pd.Series:
    return values.isna() | (values.astype(str).str.strip() == "")


def _mch_values(values: pd.Series) -> list[str]:
    """不重複的 MCH 值（去除前後空格），依字串排序，空白以 __BLANK__ 表示。"""
    out: set[str] = set()
    for v in pd.unique(values):
        if pd.isna(v):
            out.add(_TRANSFER_MCH_BLANK_TOKEN)
            continue
        txt = str(v).strip()
        out.add(txt if txt else _TRANSFER_MCH_BLANK_TOKEN)
    return sorted(out)


def _apply_mch_filter(base: pd.DataFrame, mch1_filter=None) -> pd.DataFrame:
    """排除指定 MCH1 分類以外的列（None 或包含全部值＝不排除）。

    mch1_filter 為 None 或空集合時不作限制；'' 代表空白分類；
    只會收窄母體，不改變已選分類內的分配結果。
    """
    if base.empty or not mch1_filter:
        return base
    vals = {"" if v in ("", _TRANSFER_MCH_BLANK_TOKEN) else v for v in mch1_filter}
    blank = _mch_blank_mask(base["Adj. MCH1"])
    mask = (blank if "" in vals else pd.Series(False, index=base.index)) | base["Adj. MCH1"].isin(vals - {""})
    return base[mask].reset_index(drop=True)


def _mch_excluded_warning(base: pd.DataFrame, filtered: pd.DataFrame) -> str | None:
    """MCH 篩選排除列時的提示（只列數，不含金額）。"""
    removed = len(base) - len(filtered)
    if removed <= 0:
        return None
    sites = int(filtered["_site"].nunique()) if not filtered.empty else 0
    return (
        f"MCH 篩選已排除 {removed:,} 列（未排除前共 {len(base):,} 列），"
        f"剩餘 {len(filtered):,} 列、{sites:,} 間店舖；被排除的列不會轉出，亦不會作為接收店。"
    )


def build_mch_options(df_input: pd.DataFrame, shop_natures: list[str] | None = None) -> list[str]:
    """執行前 MCH 篩選的 Adj. MCH1 選項清單（依 Shop Nature 篩選後的實際資料值）；空白值以 ``__BLANK__`` 表示。"""
    base, _warnings, _meta = _prepare_transfer_base(df_input, shop_natures=shop_natures)
    if base.empty:
        return []
    return _mch_values(base["Adj. MCH1"])


def build_mch_breakdown(df_input: pd.DataFrame, shop_natures: list[str] | None = None,
                        nd_mode: bool = False) -> pd.DataFrame:
    """候選轉出列按 Adj. MCH1 的分佈（供 MCH 篩選前參考；未套用 MCH 篩選）。

    與候選清單、執行轉換共用 ``_prepare_transfer_base``＋``_source_pool`` 同一準則
    （P90D 模式含 90 日有貨準則與新品除外），確保介面各步驟數字一致。
    """
    base, _warnings, meta = _prepare_transfer_base(df_input, shop_natures=shop_natures)
    pool = _source_pool(base, nd_mode, meta)
    if pool.empty:
        return pd.DataFrame()
    out = pd.DataFrame({
        "Adj. MCH1": pool["Adj. MCH1"].replace("", _TRANSFER_MCH_BLANK_TOKEN),
        "SKU": pool["_sku"],
        "Stock": pool["_stock"],
    })
    g = out.groupby("Adj. MCH1", sort=True)
    sku_col = "ND SKU 數" if nd_mode else "0 Sales SKU 數"
    return pd.DataFrame({
        "Adj. MCH1": [_TRANSFER_MCH_BLANK_LABEL if v == _TRANSFER_MCH_BLANK_TOKEN else str(v) for v in g.size().index],
        sku_col: g["SKU"].nunique().astype(int).values,
        "可轉出庫存 (件)": g["Stock"].sum().astype(int).values,
    })


def _build_mch_summary(out_agg: pd.DataFrame, in_map: dict, plan_df: pd.DataFrame) -> pd.DataFrame:
    """MCH1 分類彙總：按 Adj. MCH1，統計轉出／接收 SKU 數與件數、未分配件數。"""
    if out_agg.empty:
        return pd.DataFrame(columns=_TRANSFER_MCH_SUMMARY_COLUMNS)
    sku_mch = out_agg.groupby("_sku", sort=True).agg(mch1=("Adj. MCH1", _first_nonblank))
    alloc_map: dict[tuple, int] = {}
    if not plan_df.empty:
        for (sku, site), qty in plan_df.groupby(["SKU", "接收 Site Code"])["轉貨數量 (件)"].sum().items():
            alloc_map[(str(sku), site)] = int(qty)
    work = out_agg.copy()
    work["_label"] = work["_sku"].map(sku_mch["mch1"]).fillna("")
    rows: list[dict] = []
    for cat, grp in work.groupby("_label", sort=True, dropna=False):
        name = str(cat) if str(cat).strip() else _TRANSFER_MCH_BLANK_LABEL
        skus = set(grp["_sku"])
        in_skus = skus & set(in_map)
        alloc = 0
        for sku in skus:
            for site, _sales_sub, _stock_sub in in_map.get(sku, []):
                alloc += alloc_map.get((sku, site), 0)
        qty = int(grp["qty"].sum())
        rows.append({
            "Adj. MCH1": name,
            "轉出 SKU 數": len(skus),
            "接收 SKU 數": len(in_skus),
            "轉出總量 (件)": qty,
            "接收總量 (件)": int(alloc),
            "未分配量 (件)": qty - int(alloc),
        })
    mch_df = pd.DataFrame(rows, columns=_TRANSFER_MCH_SUMMARY_COLUMNS)
    if mch_df.empty:
        return mch_df
    return _with_total_row(mch_df, {
        "轉出 SKU 數": int(out_agg["_sku"].nunique()),
        "接收 SKU 數": len(set(out_agg["_sku"]) & set(in_map)),
        "轉出總量 (件)": int(out_agg["qty"].sum()),
        "接收總量 (件)": int(plan_df["轉貨數量 (件)"].sum()) if not plan_df.empty else 0,
        "未分配量 (件)": int(out_agg["qty"].sum()) - (int(plan_df["轉貨數量 (件)"].sum()) if not plan_df.empty else 0),
    })


def _site_meta(base: pd.DataFrame) -> dict:
    g = base.groupby("_site", sort=False)
    names = g["_name"].agg(_first_nonblank)
    natures = g["_nature"].agg(_mode_nonblank)
    oms = g["_om"].agg(_mode_nonblank)
    return {
        site: {"name": names[site], "nature": natures[site], "om": oms[site]}
        for site in names.index
    }


def _allocate_proportional(total: int, weights: list[float]) -> list[int]:
    """按權重分配整數件（最大餘數法）；權重總和 > 0 時分配結果加總恆等於 total。"""
    n = len(weights)
    if n == 0 or total <= 0:
        return [0] * n
    wsum = float(sum(weights))
    if wsum <= 0:
        weights = [1.0] * n
        wsum = float(n)
    raw = [total * w / wsum for w in weights]
    alloc = [math.floor(x) for x in raw]
    rest = max(total - sum(alloc), 0)
    order = sorted(range(n), key=lambda i: (-(raw[i] - alloc[i]), -weights[i], i))
    for i in order[:rest]:
        alloc[i] += 1
    return alloc


def _receiver_cap(sales: float, stock: float) -> int:
    """接收店 14 日需求上限：max(0, ceil(Sold Qty (P90D) ÷ 90 × 14) − 調前庫存)。

    與調貨段（Safety Stock x NetStock Review）的「調入需求量 (補到14日)」同一公式。
    """
    need = math.ceil(float(sales) / 90.0 * _TRANSFER_CAP_DAYS) - int(round(float(stock)))
    return max(need, 0)


def _allocate_with_caps(total: int, weights: list[float], caps: list[int]) -> tuple[list[int], int]:
    """按權重比例分配，惟每個接收地不得超出上限（caps）；超出上限的餘量在尚有需求的
    接收地之間按權重重新分配，直至派完或全部到頂。回傳 (alloc, unalloc)。

    ``alloc[i] ≤ caps[i]`` 恆成立；``sum(alloc) + unalloc == total``。
    """
    n = len(weights)
    alloc = [0] * n
    if n == 0 or total <= 0:
        return alloc, max(int(total), 0)
    remaining = int(total)
    active = [i for i in range(n) if caps[i] > 0]
    while remaining > 0 and active:
        share = _allocate_proportional(remaining, [weights[i] for i in active])
        moved = 0
        nxt: list[int] = []
        for i, s in zip(active, share):
            room = caps[i] - alloc[i]
            take = min(s, room)
            if take > 0:
                alloc[i] += take
                moved += take
            if caps[i] > alloc[i]:
                nxt.append(i)
        remaining -= moved
        if moved == 0:
            break
        active = nxt
    return alloc, remaining


def _pair_transfer(sources: list[tuple[str, int]], receivers: list[tuple[str, int]]) -> list[tuple[str, str, int]]:
    """依序配對轉出店與接收店，產生「轉出店 → 接收店 → 數量」行；兩側總量相同時全數配對。"""
    src = [[s, q] for s, q in sources if q > 0]
    rcv = [[s, q] for s, q in receivers if q > 0]
    pairs: list[tuple[str, str, int]] = []
    i = j = 0
    while i < len(src) and j < len(rcv):
        take = min(src[i][1], rcv[j][1])
        pairs.append((src[i][0], rcv[j][0], take))
        src[i][1] -= take
        rcv[j][1] -= take
        if src[i][1] == 0:
            i += 1
        if rcv[j][1] == 0:
            j += 1
    return pairs


def _normalize_sites(values) -> set[str]:
    out: set[str] = set()
    for v in values or []:
        txt = str(v).strip() if v is not None else ""
        if txt:
            out.add(txt)
    return out


def _empty_transfer(warnings: list[str], receiver_limit: int | None = None, receiver_cap: bool = False):
    stats = {
        "out_sites": 0, "out_skus": 0, "out_qty": 0, "in_sites": 0,
        "plan_lines": 0, "alloc_qty": 0, "unalloc_qty": 0, "unalloc_skus": 0,
        "receiver_limit": receiver_limit, "receiver_cap": bool(receiver_cap),
    }
    return (
        pd.DataFrame(columns=_TRANSFER_PLAN_COLUMNS),
        pd.DataFrame(columns=_TRANSFER_OUT_COLUMNS),
        pd.DataFrame(columns=_TRANSFER_IN_COLUMNS),
        pd.DataFrame(columns=_TRANSFER_SKU_COLUMNS),
        pd.DataFrame(columns=_TRANSFER_MCH_SUMMARY_COLUMNS),
        stats,
        warnings,
    )


def _with_total_row(df: pd.DataFrame, total: dict) -> pd.DataFrame:
    row = {c: total.get(c) for c in df.columns}
    row[df.columns[0]] = "總計"
    out = pd.concat([df, pd.DataFrame([row], columns=df.columns)], ignore_index=True)
    for c in df.columns:
        if pd.api.types.is_integer_dtype(df[c]):
            out[c] = out[c].astype("Int64")
    return out


def _source_pool(base: pd.DataFrame, nd_mode: bool, meta: dict | None = None) -> pd.DataFrame:
    """可轉出母體（未套用已選店舖）。

    ND 模式：RP Type＝ND 且有庫存（不看銷售記錄，不受 90 日有貨準則約束）。
    P90D 模式：0 Sales（Sold Qty (P90D) ≤ 0）且有庫存；``Shop Stock Qty (LM)``／``(LLM)``
    欄存在時另須 > 0（90 日有貨準則，v1.93.0：M-2／M-1 無貨＝期間剛調入／派貨，
    非全程有貨，不算真坐貨）；``Launch Date`` 欄存在時排除 90 天內新品。
    """
    if nd_mode:
        return base[(base["_rp"] == _TRANSFER_RP_ND) & (base["_stock"] > 0)]
    meta = meta or {}
    pool = base[(base["_sales"] <= 0) & (base["_stock"] > 0)]
    if meta.get("lm_col"):
        pool = pool[pool["_lm"] > 0]
    if meta.get("llm_col"):
        pool = pool[pool["_llm"] > 0]
    if meta.get("launch_col"):
        ref = pd.Timestamp(_normalize_ref_date(None))
        is_new = pool["_launch"].notna() & (pool["_launch"] >= ref - pd.Timedelta(days=_TRANSFER_NEW_PRODUCT_DAYS))
        pool = pool[~is_new]
    return pool


def _zero_sales_criteria_warnings(base: pd.DataFrame, meta: dict | None = None) -> list[str]:
    """P90D 模式轉出列被 90 日有貨準則／新品除外排除的數量提示（欄存在且有排除時才提示）。

    依次序歸因（LM → LLM → 新品），與 ``_source_pool`` 的排除次序一致，
    各項只計入前面準則未排除的列，合計等於實際排除列數。
    """
    meta = meta or {}
    warnings: list[str] = []
    pool = base[(base["_sales"] <= 0) & (base["_stock"] > 0)]
    if pool.empty:
        return warnings
    if meta.get("lm_col"):
        excluded = pool[~(pool["_lm"] > 0)]
        if len(excluded):
            warnings.append(
                f"有 {len(excluded)} 列 0 Sales 且有庫存，但 Shop Stock Qty (LM) 為 0"
                "（上月無貨，90 日窗口內非全程有貨），已按 90 日有貨準則排除轉出。"
            )
            pool = pool[pool["_lm"] > 0]
    if meta.get("llm_col"):
        excluded = pool[~(pool["_llm"] > 0)]
        if len(excluded):
            warnings.append(
                f"有 {len(excluded)} 列 0 Sales 且有庫存，但 Shop Stock Qty (LLM) 為 0"
                "（M-2 無貨，90 日窗口內非全程有貨），已按 90 日有貨準則排除轉出。"
            )
            pool = pool[pool["_llm"] > 0]
    if meta.get("launch_col"):
        ref = pd.Timestamp(_normalize_ref_date(None))
        is_new = pool["_launch"].notna() & (pool["_launch"] >= ref - pd.Timedelta(days=_TRANSFER_NEW_PRODUCT_DAYS))
        excluded = pool[is_new]
        if len(excluded):
            warnings.append(
                f"有 {len(excluded)} 列 0 Sales 且有庫存屬 {_TRANSFER_NEW_PRODUCT_DAYS} 天內新品"
                f"（Launch Date 距 {ref.date().isoformat()} 不足 {_TRANSFER_NEW_PRODUCT_DAYS} 天），"
                "銷售尚未起步，已排除轉出。"
            )
    return warnings


def build_transfer_candidates(df_input: pd.DataFrame, shop_natures: list[str] | None = None,
                              mch1_filter=None, nd_mode: bool = False) -> pd.DataFrame:
    """候選轉出店舖一覽，依店舖彙總。

    預設（P90D 模式）：Sold Qty (P90D) ≤ 0 且 Shop Stock Qty (CM) > 0 的列；
    ``Shop Stock Qty (LM)``／``(LLM)`` 欄存在時另須 > 0（90 日有貨準則）、
    ``Launch Date`` 欄存在時排除 90 天內新品；
    ``nd_mode=True``（ND 模式）：RP Type＝ND 且 Shop Stock Qty (CM) > 0 的列
    （另附「有銷售 ND SKU 數」＝其中 Sold Qty (P90D) > 0 的 SKU 數，供核對仍錄得銷售的轉出列）。

    ``mch1_filter`` 為執行前的 MCH1 篩選（None 或包含全部值＝不排除）。
    """
    base, _warnings, meta = _prepare_transfer_base(df_input, shop_natures=shop_natures, require_rp=nd_mode)
    base = _apply_mch_filter(base, mch1_filter=mch1_filter)
    columns = _TRANSFER_ND_CANDIDATE_COLUMNS if nd_mode else _TRANSFER_CANDIDATE_COLUMNS
    sku_col = "ND SKU 數" if nd_mode else "0 Sales SKU 數"
    if base.empty:
        return pd.DataFrame(columns=columns)
    pool = _source_pool(base, nd_mode, meta)
    if pool.empty:
        return pd.DataFrame(columns=columns)
    g = pool.groupby("_site", sort=True)
    out = pd.DataFrame({
        "Site Code": g.size().index,
        "Shop Name": g["_name"].agg(_first_nonblank).values,
        "Shop Nature": g["_nature"].agg(_mode_nonblank).values,
        "OM": g["_om"].agg(_mode_nonblank).values,
        sku_col: g["_sku"].nunique().astype(int).values,
        "可轉出庫存 (件)": g["_stock"].sum().round().astype(int).values,
    })
    if nd_mode:
        sales_sites = pool[pool["_sales"] > 0].groupby("_site")["_sku"].nunique()
        out["有銷售 ND SKU 數"] = out["Site Code"].map(sales_sites).fillna(0).astype(int)
    return out[columns].reset_index(drop=True)


def transform_p90d_transfer(
    df_input: pd.DataFrame,
    transfer_out_sites: list[str],
    progress_callback=None,
    shop_natures: list[str] | None = None,
    mch1_filter=None,
    nd_mode: bool = False,
    receiver_limit: int | None = None,
    receiver_cap: bool = False,
):
    """依指定轉出店舖產生轉貨計畫。

    ``mch1_filter`` 為執行前的 MCH1 篩選（None 或包含全部值＝不排除），
    對轉出與接收兩側同時生效。

    ``receiver_limit`` 為每個 SKU 的接收店舖上限（None＝不限，原有行為）：
    指定 2／4 時，按 Sold Qty (P90D) 由高至低只取最多 2／4 間接收店
    （同銷量以 Site Code 排序），再按這批店舖的銷售比例分配，轉出總量仍力求全數接收。

    ``receiver_cap=True`` 時啟用接收量「補到 14 日」封頂：每間接收店最多接收其
    14 日需求（Sold Qty (P90D) ÷ 90 × 14 無條件進位 − 調前庫存，最低 0），
    超出的餘量在尚有需求的接收店之間重新分配，仍未派完者列未分配（保留原店）。

    ``nd_mode=True`` 時改用 ND SKU with Stock Review(轉貨模式) 的判定：
    轉出＝已選店舖中 RP Type＝ND 且有庫存的列（不看銷售記錄）、
    接收＝同 SKU、非轉出店舖、RP Type＝RF 且 Sold Qty (P90D) > 0 的列；
    ``RP Type`` 為必要欄位，非 RF／ND 的值一律不納入轉出與接收。

    P90D 模式（``nd_mode=False``）的轉出另受 90 日有貨準則與新品除外約束
    （LM／LLM 欄存在時須 > 0、Launch Date 90 天內新品排除，欄不存在時不生效）；
    接收店於 ``RP Type`` 欄存在時排除 RP Type＝ND 的店舖。

    回傳 (plan_df, out_shop_df, in_shop_df, sku_df, mch_df, stats, warnings)。
    """
    if progress_callback:
        progress_callback(0.0, "驗證輸入資料...")
    sites = _normalize_sites(transfer_out_sites)
    if not sites:
        raise ValueError("請至少選擇一間轉出店舖。")
    receiver_limit = _validate_receiver_limit(receiver_limit)
    receiver_cap = bool(receiver_cap)
    base, warnings, _colmeta = _prepare_transfer_base(df_input, shop_natures=shop_natures,
                                                      progress_callback=progress_callback,
                                                      require_rp=nd_mode)
    if base.empty:
        warnings.append("過濾後無有效資料列，沒有可轉貨的記錄。")
        return _empty_transfer(warnings, receiver_limit=receiver_limit, receiver_cap=receiver_cap)
    _filtered = _apply_mch_filter(base, mch1_filter=mch1_filter)
    _mch_note = _mch_excluded_warning(base, _filtered)
    if _mch_note:
        warnings.append(_mch_note)
    base = _filtered
    if base.empty:
        warnings.append("MCH 篩選後無有效資料列，本次沒有轉出。")
        return _empty_transfer(warnings, receiver_limit=receiver_limit, receiver_cap=receiver_cap)

    if progress_callback:
        progress_callback(0.25, "辨識轉出與接收店舖...")
    meta = _site_meta(base)
    if nd_mode:
        _nonstd = int((~base["_rp"].isin([_TRANSFER_RP_RF, _TRANSFER_RP_ND])).sum())
        if _nonstd:
            warnings.append(
                f"有 {_nonstd:,} 列 RP Type 非 RF／ND（空白或其他值），一律不納入轉出與接收。"
            )
        _pool = _source_pool(base, True, _colmeta)
        out_rows = _pool[_pool["_site"].isin(sites)].copy()
        in_rows = base[(~base["_site"].isin(sites)) & (base["_rp"] == _TRANSFER_RP_RF) & (base["_sales"] > 0)]
        _out_missing_txt = "以下已選轉出店舖沒有 RP Type＝ND 且有庫存的列，未納入轉出"
        _out_empty_txt = "已選轉出店舖沒有符合「RP Type＝ND 且有庫存（Shop Stock Qty (CM) > 0）」的資料，本次沒有轉出。"
        _in_unalloc_txt = "但同 SKU 沒有其他 RP Type＝RF 且有銷售記錄（Sold Qty (P90D) > 0）的店舖可接收"
    else:
        warnings.extend(_zero_sales_criteria_warnings(base, _colmeta))
        _pool = _source_pool(base, False, _colmeta)
        out_rows = _pool[_pool["_site"].isin(sites)].copy()
        in_rows = base[(~base["_site"].isin(sites)) & (base["_sales"] > 0)]
        if _colmeta.get("rp_col"):
            _nd_recv = in_rows["_rp"] == _TRANSFER_RP_ND
            _n_nd_recv = int(_nd_recv.sum())
            if _n_nd_recv:
                warnings.append(
                    f"有 {_n_nd_recv:,} 列接收候選為 RP Type＝ND（不可補貨店舖，調入只會形成新的積壓），"
                    "已排除接收。"
                )
            in_rows = in_rows[~_nd_recv]
        _out_missing_txt = "以下已選轉出店舖沒有 0 Sales 且有庫存的列，未納入轉出"
        _out_empty_txt = "已選轉出店舖沒有符合「0 Sales（Sold Qty (P90D) ≤ 0）且有庫存（Shop Stock Qty (CM) > 0）」的資料，本次沒有轉出。"
        _in_unalloc_txt = "但同 SKU 沒有其他有銷售記錄（Sold Qty (P90D) > 0）的店舖可接收"

    missing_sites = sorted(sites - set(out_rows["_site"]))
    if missing_sites:
        warnings.append(f"{_out_missing_txt}：{'、'.join(missing_sites)}")
    if out_rows.empty:
        warnings.append(_out_empty_txt)
        return _empty_transfer(warnings, receiver_limit=receiver_limit, receiver_cap=receiver_cap)

    out_agg = out_rows.groupby(["_sku", "_site"], sort=True).agg(
        qty=("_stock", "sum"),
        sales=("_sales", "sum"),
        **{"Adj. MCH1": ("Adj. MCH1", _first_nonblank), "Supply Source": ("Supply Source", _first_nonblank)},
    ).reset_index()
    in_agg = in_rows.groupby(["_sku", "_site"], sort=True).agg(
        sales=("_sales", "sum"), stock=("_stock", "sum")
    ).reset_index()
    sku_info = out_rows.groupby("_sku", sort=True).agg(
        Brand=("Brand", _first_nonblank),
        Buyer=("Buyer", _first_nonblank),
        **{
            "SKU Desc": ("SKU Desc", _first_nonblank),
            "Adj. MCH1": ("Adj. MCH1", _first_nonblank),
            "Supply Source": ("Supply Source", _first_nonblank),
        },
    ).to_dict("index")

    in_map: dict[str, list[tuple[str, float, float]]] = {}
    for sku, site, sales, stock in in_agg[["_sku", "_site", "sales", "stock"]].itertuples(index=False, name=None):
        in_map.setdefault(sku, []).append((site, float(sales), float(stock)))

    if progress_callback:
        progress_callback(0.4, "分配同 SKU 轉出量...")
    plan_rows: list[dict] = []
    sku_rows: list[dict] = []
    pool_rows: list[dict] = []
    capped_out_pairs = 0
    no_receiver_skus = 0
    no_receiver_qty = 0
    cap_short_skus = 0
    cap_short_qty = 0
    for sku, sku_out in out_agg.groupby("_sku", sort=True):
        sources = [
            (site, int(round(q)), float(s))
            for site, q, s in zip(sku_out["_site"], sku_out["qty"], sku_out["sales"])
        ]
        total = sum(q for _, q, _ in sources)
        all_pool = sorted(in_map.get(sku, []))
        if receiver_limit is None:
            pool = all_pool
        else:
            # 集中模式：只取 Sold Qty (P90D) 最高的 N 間接收店（同銷量以 Site Code 排序）
            pool = sorted(all_pool, key=lambda r: (-r[1], r[0]))[:receiver_limit]
            capped_out_pairs += max(len(all_pool) - len(pool), 0)
        weights = [sales for _, sales, _ in pool]
        # 接收店 14 日需求（恆計算，供轉貨單稽核；只在封頂模式下約束分配）
        caps = [_receiver_cap(sales, stock) for _, sales, stock in pool]
        if receiver_cap and pool:
            alloc, _unalloc_cap = _allocate_with_caps(total, weights, caps)
        else:
            alloc = _allocate_proportional(total, weights)
        recv = [(site, sales, stock, a, c) for (site, sales, stock), a, c in zip(pool, alloc, caps)]
        recv_map = {site: (sales, stock, a, c) for site, sales, stock, a, c in recv}
        sales_sum = sum(sales for _, sales, _ in pool)
        src_sales = {site: s for site, _, s in sources}
        info = sku_info.get(sku, {})
        order = sorted(recv, key=lambda r: (-r[3], -r[1], r[0]))
        pairs = _pair_transfer([(site, q) for site, q, _ in sources], [(site, a) for site, _, _, a, _ in order])
        for from_site, to_site, qty in pairs:
            sales_to, stock_to, alloc_to, cap_to = recv_map[to_site]
            plan_rows.append({
                "轉出 Site Code": from_site,
                "轉出 Shop Name": meta.get(from_site, {}).get("name", ""),
                "轉出店 Sold Qty (P90D)": int(round(src_sales.get(from_site, 0.0))),
                "接收 Site Code": to_site,
                "接收 Shop Name": meta.get(to_site, {}).get("name", ""),
                "Brand": info.get("Brand", ""),
                "SKU": sku,
                "SKU Desc": info.get("SKU Desc", ""),
                "Buyer": info.get("Buyer", ""),
                "Adj. MCH1": info.get("Adj. MCH1", ""),
                "Supply Source": info.get("Supply Source", ""),
                "轉貨數量 (件)": int(qty),
                "接收店 Shop Nature": meta.get(to_site, {}).get("nature", ""),
                "接收店 OM": meta.get(to_site, {}).get("om", ""),
                "接收店 Sold Qty (P90D)": int(round(sales_to)),
                "接收店銷售佔比 (同 SKU)": (sales_to / sales_sum) if sales_sum > 0 else 0.0,
                "接收店同 SKU 分配總量 (件)": int(alloc_to),
                "接收店 Shop Stock Qty (CM) 調前": int(round(stock_to)),
                "接收店 14 日需求 (件)": int(cap_to),
            })
        allocated = sum(alloc)
        if total - allocated == 0:
            status = _TRANSFER_STATUS_ALLOCATED
        elif not pool:
            status = _TRANSFER_STATUS_NO_RECEIVER
            no_receiver_skus += 1
            no_receiver_qty += total - allocated
        else:
            status = _TRANSFER_STATUS_CAPPED
            cap_short_skus += 1
            cap_short_qty += total - allocated
        sku_rows.append({
            "Brand": info.get("Brand", ""),
            "SKU": sku,
            "SKU Desc": info.get("SKU Desc", ""),
            "Buyer": info.get("Buyer", ""),
            "Adj. MCH1": info.get("Adj. MCH1", ""),
            "轉出店舖數": len(sources),
            "轉出總量 (件)": int(total),
            "可接收店舖數 (有銷售)": len(pool),
            "實際接收店舖數": sum(1 for a in alloc if a > 0),
            "接收店 Sold Qty (P90D) 合計": int(round(sales_sum)),
            "已分配總量 (件)": int(allocated),
            "未分配量 (件)": int(total - allocated),
            "狀態": status,
        })
        for site, sales, stock, a, _cap in recv:
            pool_rows.append({"接收 Site Code": site, "SKU": sku, "sales": sales, "stock": stock, "alloc": a})

    if progress_callback:
        progress_callback(0.7, "建立轉貨單與彙總表...")
    plan_df = pd.DataFrame(plan_rows, columns=_TRANSFER_PLAN_COLUMNS)
    if not plan_df.empty:
        plan_df = plan_df.sort_values(["轉出 Site Code", "Brand", "SKU", "接收 Site Code"], kind="stable").reset_index(drop=True)

    sku_df = pd.DataFrame(sku_rows, columns=_TRANSFER_SKU_COLUMNS)
    if not sku_df.empty:
        sku_df = sku_df.sort_values(["未分配量 (件)", "轉出總量 (件)", "SKU"], ascending=[False, False, True], kind="stable").reset_index(drop=True)
        sku_df = _with_total_row(sku_df, {
            "轉出總量 (件)": int(sku_df["轉出總量 (件)"].sum()),
            "已分配總量 (件)": int(sku_df["已分配總量 (件)"].sum()),
            "未分配量 (件)": int(sku_df["未分配量 (件)"].sum()),
        })

    out_qty_total = int(out_agg["qty"].sum())
    alloc_qty = int(plan_df["轉貨數量 (件)"].sum()) if not plan_df.empty else 0
    alloc_by_site = plan_df.groupby("轉出 Site Code")["轉貨數量 (件)"].sum() if not plan_df.empty else pd.Series(dtype=int)
    to_count_by_site = plan_df.groupby("轉出 Site Code")["接收 Site Code"].nunique() if not plan_df.empty else pd.Series(dtype=int)
    out_site_df = out_agg.groupby("_site", sort=True).agg(
        sku_n=("_sku", "nunique"), qty=("qty", "sum")
    ).reset_index().rename(columns={"_site": "Site Code"})
    out_site_df["qty"] = out_site_df["qty"].round().astype(int)
    out_site_df["Shop Name"] = out_site_df["Site Code"].map(lambda s: meta.get(s, {}).get("name", ""))
    out_site_df["Shop Nature"] = out_site_df["Site Code"].map(lambda s: meta.get(s, {}).get("nature", ""))
    out_site_df["OM"] = out_site_df["Site Code"].map(lambda s: meta.get(s, {}).get("om", ""))
    out_site_df["已分配"] = out_site_df["Site Code"].map(alloc_by_site).fillna(0).astype(int)
    out_site_df["未分配"] = out_site_df["qty"] - out_site_df["已分配"]
    out_site_df["接收店舖數"] = out_site_df["Site Code"].map(to_count_by_site).fillna(0).astype(int)
    out_shop_df = out_site_df.rename(columns={
        "sku_n": "轉出 SKU 數", "qty": "轉出總量 (件)", "已分配": "已分配 (件)", "未分配": "未分配 (件)",
    })[_TRANSFER_OUT_COLUMNS]
    out_shop_df = _with_total_row(out_shop_df, {
        "轉出 SKU 數": int(out_agg["_sku"].nunique()),
        "轉出總量 (件)": out_qty_total,
        "已分配 (件)": alloc_qty,
        "未分配 (件)": out_qty_total - alloc_qty,
        "接收店舖數": int(plan_df["接收 Site Code"].nunique()) if not plan_df.empty else 0,
    })

    pool_df = pd.DataFrame(pool_rows)
    pool_df = pool_df[pool_df["alloc"] > 0] if not pool_df.empty else pool_df
    if not pool_df.empty:
        in_site_df = pool_df.groupby("接收 Site Code", sort=True).agg(
            sku_n=("SKU", "nunique"), alloc=("alloc", "sum"),
            before=("stock", "sum"), sales=("sales", "sum"),
        ).reset_index()
        from_count = plan_df.groupby("接收 Site Code")["轉出 Site Code"].nunique()
        in_site_df["Shop Name"] = in_site_df["接收 Site Code"].map(lambda s: meta.get(s, {}).get("name", ""))
        in_site_df["Shop Nature"] = in_site_df["接收 Site Code"].map(lambda s: meta.get(s, {}).get("nature", ""))
        in_site_df["OM"] = in_site_df["接收 Site Code"].map(lambda s: meta.get(s, {}).get("om", ""))
        in_site_df["涉及轉出店舖數"] = in_site_df["接收 Site Code"].map(from_count).fillna(0).astype(int)
        in_site_df["接收後庫存 (件)"] = (in_site_df["before"] + in_site_df["alloc"]).round().astype(int)
        in_site_df = in_site_df.rename(columns={
            "接收 Site Code": "Site Code", "sku_n": "接收 SKU 數", "alloc": "接收總量 (件)",
            "before": "接收前庫存 (件)", "sales": "接收店 Sold Qty (P90D) 合計",
        })
        in_site_df["接收 SKU 數"] = in_site_df["接收 SKU 數"].astype(int)
        in_site_df["接收總量 (件)"] = in_site_df["接收總量 (件)"].astype(int)
        in_site_df["接收前庫存 (件)"] = in_site_df["接收前庫存 (件)"].round().astype(int)
        in_site_df["接收店 Sold Qty (P90D) 合計"] = in_site_df["接收店 Sold Qty (P90D) 合計"].round().astype(int)
        in_shop_df = in_site_df[_TRANSFER_IN_COLUMNS]
        in_shop_df = _with_total_row(in_shop_df, {
            "接收 SKU 數": int(pool_df["SKU"].nunique()),
            "接收總量 (件)": int(pool_df["alloc"].sum()),
            "接收前庫存 (件)": int(in_site_df["接收前庫存 (件)"].sum()),
            "接收後庫存 (件)": int(in_site_df["接收後庫存 (件)"].sum()),
            "接收店 Sold Qty (P90D) 合計": int(in_site_df["接收店 Sold Qty (P90D) 合計"].sum()),
            "涉及轉出店舖數": int(plan_df["轉出 Site Code"].nunique()),
        })
    else:
        in_shop_df = pd.DataFrame(columns=_TRANSFER_IN_COLUMNS)

    if progress_callback:
        progress_callback(0.8, "建立 MCH1 分類彙總...")
    mch_df = _build_mch_summary(out_agg, in_map, plan_df)

    unalloc_skus = sum(1 for r in sku_rows if r["未分配量 (件)"] > 0)
    unalloc_qty = out_qty_total - alloc_qty
    stats = {
        "out_sites": int(len(out_site_df)),
        "out_skus": int(out_agg["_sku"].nunique()),
        "out_qty": out_qty_total,
        "in_sites": int(plan_df["接收 Site Code"].nunique()) if not plan_df.empty else 0,
        "plan_lines": int(len(plan_df)),
        "alloc_qty": alloc_qty,
        "unalloc_qty": unalloc_qty,
        "unalloc_skus": unalloc_skus,
        "receiver_limit": receiver_limit,
        "receiver_cap": receiver_cap,
    }
    if no_receiver_skus:
        warnings.append(
            f"有 {no_receiver_skus:,} 個 SKU（共 {no_receiver_qty:,} 件）在已選轉出店舖有可轉出庫存，"
            f"{_in_unalloc_txt}，本次未能轉出，保留於原店舖；詳見「SKU 轉貨彙總」。"
        )
    if cap_short_skus:
        warnings.append(
            f"已啟用接收量「補到 14 日」封頂：有 {cap_short_skus:,} 個 SKU（共 {cap_short_qty:,} 件）"
            "超出接收店合計 14 日需求（Sold Qty (P90D) ÷ 90 × 14 − 調前庫存），"
            "未能全數轉出，保留於原店舖；詳見「SKU 轉貨彙總」。"
        )
    if receiver_limit is not None and capped_out_pairs:
        warnings.append(
            f"已啟用接收店集中模式（{receiver_limit_label(receiver_limit)}）："
            f"共 {capped_out_pairs:,} 個「SKU × 有銷售接收店」組合超出上限，未納入分配；"
            "如需全數有銷售店舖接收，請將接收店舖數量改為「不限」。"
        )

    if progress_callback:
        progress_callback(0.9, "轉貨計畫完成...")
    return plan_df, out_shop_df, in_shop_df, sku_df, mch_df, stats, warnings


def _write_transfer_table(ws, df: pd.DataFrame, auto_filter: bool = False,
                          column_styles: dict | None = None,
                          bold_data_columns=()) -> None:
    """寫入一張轉貨表（表頭＋資料列）。

    ``column_styles``：{欄名: (表頭底色, 資料底色)}，只影響該欄的顯示底色與表頭字色
    （著色欄的表頭改用黑字；未列出的欄位沿用預設深藍表頭）；欄名、欄序與數值一律不變。
    ``bold_data_columns``：資料列需加粗的欄名。
    """
    styles = dict(column_styles or {})
    bold_cols = set(bold_data_columns or ())
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    tinted_header_font = Font(bold=True, color="000000", size=11)
    bold_font = Font(bold=True, size=11)
    center_align = Alignment(horizontal="center", vertical="center")
    headers = list(df.columns)
    tinted_header_fills = {
        h: PatternFill(start_color=styles[h][0], end_color=styles[h][0], fill_type="solid")
        for h in headers if h in styles
    }
    tinted_data_fills = {
        i: PatternFill(start_color=styles[h][1], end_color=styles[h][1], fill_type="solid")
        for i, h in enumerate(headers, 1) if h in styles and styles[h][1]
    }
    bold_data_cols = {i for i, h in enumerate(headers, 1) if h in bold_cols}
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        if h in tinted_header_fills:
            cell.font = tinted_header_font
            cell.fill = tinted_header_fills[h]
        else:
            cell.font = header_font
            cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    if tinted_header_fills:
        ws.row_dimensions[1].height = _TRANSFER_PLAN_HEADER_ROW_HEIGHT
    if df.empty:
        ws.cell(row=2, column=1, value="無符合資料").font = Font(size=12)
    else:
        for row_idx, vals in enumerate(df.itertuples(index=False, name=None)):
            excel_row = row_idx + 2
            is_total = str(vals[0]) == "總計"
            for col_idx, h in enumerate(headers, 1):
                val = vals[col_idx - 1]
                if isinstance(val, np.generic):
                    val = val.item()
                if val is not None and not isinstance(val, str) and pd.isna(val):
                    val = None
                cell = ws.cell(row=excel_row, column=col_idx, value=val)
                cell.border = thin_border
                cell.alignment = center_align
                if col_idx in tinted_data_fills:
                    cell.fill = tinted_data_fills[col_idx]
                if is_total or col_idx in bold_data_cols:
                    cell.font = bold_font
                if isinstance(val, (int, float)) and not isinstance(val, bool):
                    cell.number_format = "0.0%" if h in _TRANSFER_PCT_COLUMNS else "#,##0"
    ws.freeze_panes = "A2"
    if auto_filter and ws.max_row > 1 and ws.max_column > 1:
        ws.auto_filter.ref = ws.dimensions
    _auto_width(ws)


def generate_excel_p90d_transfer(
    plan_df: pd.DataFrame,
    out_shop_df: pd.DataFrame,
    in_shop_df: pd.DataFrame,
    sku_df: pd.DataFrame,
    mch_df: pd.DataFrame | None = None,
    stats: dict | None = None,
    nature_label: str = "全部",
    mch_filter_label: str = "全部（未排除）",
    progress_callback=None,
    nd_mode: bool = False,
    receiver_limit: int | None = None,
    receiver_cap: bool = False,
) -> BytesIO:
    """產生六個工作表：說明、轉貨單、轉出店舖彙總、接收店舖彙總、SKU 轉貨彙總、MCH1 分類彙總（無 RAW）。

    ``nd_mode=True`` 時說明頁改用 ND SKU with Stock Review(轉貨模式) 的措辭
    （轉出＝RP Type＝ND 且有庫存、接收＝RP Type＝RF 且有銷售記錄）。

    ``receiver_limit`` 只影響 ND 模式說明頁的「接收店舖集中模式」參數
    （None＝不限、2／4＝每個 SKU 最多 2／4 間接收店）。

    ``receiver_cap`` 只影響說明頁的「接收量上限」參數（False＝不限、True＝補到 14 日封頂）。
    """
    stats = dict(stats or {})
    _cap_label = (
        "接收店按「補到 14 日」需求封頂"
        if receiver_cap else "不限（按銷售比例全數分配）"
    )
    if progress_callback:
        progress_callback(0.0, "建立 Excel 結構...")
    wb = Workbook()

    ws_guide = wb.active
    ws_guide.title = "說明"
    if nd_mode:
        _title = "ND SKU with Stock Review（轉貨模式）說明"
        _subtitle = "由指定轉出店舖（RP Type＝ND 且有庫存）轉貨至有銷售記錄的 RF 店舖；不分 OM 群組，轉貨數量不設上限"
        _params = [
            ("轉出店舖", f"已選店舖中共 {int(stats.get('out_sites', 0)):,} 間有 ND 且有庫存的列（未納入者見警告）"),
            ("Shop Nature 篩選", nature_label),
            ("MCH 篩選（執行前）", mch_filter_label),
            ("ND 定義", "RP Type＝ND（去除前後空格、大小寫不敏感）"),
            ("轉出數量", "Shop Stock Qty (CM) 全數，不設上限（只計 RP Type＝ND 且有庫存的列，不看銷售記錄）"),
            ("接收店定義", "同 SKU、非轉出店舖、RP Type＝RF 且 Sold Qty (P90D) > 0（有銷售記錄），不論有無庫存"),
            ("接收店舖集中模式", receiver_limit_label(receiver_limit)),
            ("接收量上限", _cap_label),
        ]
        _logic = [
            "同 SKU 的轉出總量＝所有已選轉出店舖的 RP Type＝ND 有庫存列 Shop Stock Qty (CM) 合計，全數轉出，不設上限（即使該列仍有銷售記錄亦照轉，轉出店銷量已列於轉貨單「轉出店 Sold Qty (P90D)」欄供核對）。",
            "只有一間接收店時，全數分配給該店；有多間接收店時，按各店 Sold Qty (P90D) 的銷售比例分配（最大餘數法取整數件），力求轉出總量全數被接收。",
            "接收店舖集中模式（可選）：「不限」時全數有銷售記錄的 RF 接收店舖按銷售比例分配；「集中 1–2 間」或「集中 3–4 間」時，每個 SKU 只取 Sold Qty (P90D) 最高的 2 間或 4 間接收店（同銷量以 Site Code 排序），再按這批店舖的銷售比例分配，轉出總量仍力求全數接收；未納入的店舖不會收到該 SKU 的轉貨。",
            "接收量上限（可選）：「補到 14 日」封頂時，每間接收店最多接收其 14 日需求（Sold Qty (P90D) ÷ 90 × 14 無條件進位 − 調前庫存，最低 0，見轉貨單「接收店 14 日需求 (件)」欄）；按銷售比例分配後超出上限的餘量，會在尚有需求的接收店之間重新分配，仍未派完者列未分配（保留原店）。",
            "轉貨單配對：同 SKU 內，轉出店依 Site Code 排序，接收店依分配量由大至小排序，依序配對成「轉出店 → 接收店」行，各行數量合計與分配一致。",
            "接收店不受庫存限制（封頂模式開啟時受 14 日需求限制）；接收後庫存＝調前庫存＋分配總量。",
            "同 SKU 沒有任何接收店（即沒有其他 RP Type＝RF 且有銷售記錄的店舖）時不轉出，保留於原店舖，並列於「SKU 轉貨彙總」的未分配量。",
            "RP Type 為必要欄位，來源檔缺少即停止計算；空白或其他非 RF／ND 的值一律不納入轉出與接收（列警告）。",
            "「Site Code × SKU」重複列一律停止計算（重複列會重複計算庫存與銷售），請先修正來源資料。",
            "MCH 篩選（執行前）：按 Adj. MCH1 排除分類；被排除的分類不會轉出，亦不會作為接收店；篩選只收窄母體，不改變已選分類內的分配結果。",
            "本模式不做 Safety、周轉天數或調貨觸發判斷，只按上述規則計算。",
        ]
        _tips = [
            "轉出店舖須於網頁先勾選；只有勾選店舖中 RP Type＝ND 且有庫存的列會轉出，其他列不會被轉出或作為接收店。",
            "候選清單「有銷售 ND SKU 數」大於 0 表示該店仍有錄得銷售的 ND 列會照轉，請於轉貨單「轉出店 Sold Qty (P90D)」欄逐 SKU 核對。",
            "MCH 篩選於執行前設定（預設全部 Adj. MCH1 都包含）；被排除分類的列既不轉出，也不接收，排除列數見警告。",
            "來源檔如含同店同 SKU 重複列，執行會停止並提示修正（避免重複計算庫存）。",
            "分配以銷售比例為準；如需調整某店接收量，可於轉貨單直接修改數量，修改後請重新核對各 SKU 的轉出與接收總量。",
        ]
        if receiver_limit is not None:
            _tips.append(
                "已啟用接收店集中模式：集中後個別接收店的數量會相應增加，請按店舖陳列能力於轉貨單核對；"
                "如需全數有銷售店舖接收，請將接收店舖數量改為「不限」。"
            )
        if receiver_cap:
            _tips.append(
                "已啟用接收量「補到 14 日」封頂：接收店最多接收其 14 日需求，超出部分保留於原店舖；"
                "如需全數接收，請關閉封頂。"
            )
    else:
        _title = "P90D 0 Sales Review（轉貨模式）說明"
        _subtitle = "由指定轉出店舖（0 Sales 且有庫存）轉貨至有銷售記錄的店舖；不分 OM 群組，轉貨數量不設上限"
        _params = [
            ("轉出店舖", f"已選店舖中共 {int(stats.get('out_sites', 0)):,} 間有 0 Sales 且有庫存的列（未納入者見警告）"),
            ("Shop Nature 篩選", nature_label),
            ("MCH 篩選（執行前）", mch_filter_label),
            ("0 Sales 定義", "Sold Qty (P90D) ≤ 0；Shop Stock Qty (LM)／(LLM) 欄存在時另須 > 0（90 日有貨準則）；Launch Date 欄存在時排除 90 天內新品"),
            ("轉出數量", "Shop Stock Qty (CM) 全數，不設上限（只計符合上述 0 Sales 準則且有庫存的列）"),
            ("接收店定義", "同 SKU、非轉出店舖、Sold Qty (P90D) > 0（有銷售記錄），不論有無庫存；RP Type 欄存在時排除 RP Type＝ND 的店舖"),
            ("接收量上限", _cap_label),
        ]
        _logic = [
            "同 SKU 的轉出總量＝所有已選轉出店舖的 0 Sales 有庫存列 Shop Stock Qty (CM) 合計，全數轉出，不設上限。",
            "90 日有貨準則（Shop Stock Qty (LM)／(LLM) 欄存在時）：M-2／M-1 無貨的列 90 日窗口內非全程有貨（期間剛調入／派貨），不算真坐貨，排除轉出；90 天內新品（Launch Date）銷售尚未起步，亦排除轉出；欄不存在時該準則不生效。",
            "接收店於 RP Type 欄存在時排除 RP Type＝ND 的店舖（ND 不再補貨，調入只會形成新的積壓）；欄不存在時不檢查。",
            "只有一間接收店時，全數分配給該店；有多間接收店時，按各店 Sold Qty (P90D) 的銷售比例分配（最大餘數法取整數件），力求轉出總量全數被接收。",
            "接收量上限（可選）：「補到 14 日」封頂時，每間接收店最多接收其 14 日需求（Sold Qty (P90D) ÷ 90 × 14 無條件進位 − 調前庫存，最低 0，見轉貨單「接收店 14 日需求 (件)」欄）；按銷售比例分配後超出上限的餘量，會在尚有需求的接收店之間重新分配，仍未派完者列未分配（保留原店）。",
            "轉貨單配對：同 SKU 內，轉出店依 Site Code 排序，接收店依分配量由大至小排序，依序配對成「轉出店 → 接收店」行，各行數量合計與分配一致。",
            "接收店不受庫存限制（封頂模式開啟時受 14 日需求限制）；接收後庫存＝調前庫存＋分配總量。",
            "同 SKU 沒有任何接收店（即沒有其他有銷售記錄的店舖）時不轉出，保留於原店舖，並列於「SKU 轉貨彙總」的未分配量。",
            "「Site Code × SKU」重複列一律停止計算（重複列會重複計算庫存與銷售），請先修正來源資料。",
            "MCH 篩選（執行前）：按 Adj. MCH1 排除分類；被排除的分類不會轉出，亦不會作為接收店；篩選只收窄母體，不改變已選分類內的分配結果。",
            "本模式不做 Safety、周轉天數或調貨觸發判斷，只按上述規則計算。",
        ]
        _tips = [
            "轉出店舖須於網頁先勾選；只有勾選店舖中符合 0 Sales 準則且有庫存的列會轉出，其他列不會被轉出或作為接收店。",
            "MCH 篩選於執行前設定（預設全部 Adj. MCH1 都包含）；被排除分類的列既不轉出，也不接收，排除列數見警告。",
            "來源檔如含同店同 SKU 重複列，執行會停止並提示修正（避免重複計算庫存）。",
            "分配以銷售比例為準；如需調整某店接收量，可於轉貨單直接修改數量，修改後請重新核對各 SKU 的轉出與接收總量。",
        ]
        if receiver_cap:
            _tips.append(
                "已啟用接收量「補到 14 日」封頂：接收店最多接收其 14 日需求，超出部分保留於原店舖；"
                "如需全數接收，請關閉封頂。"
            )
    _write_guide_sheet(
        ws_guide,
        title=_title,
        subtitle=_subtitle,
        params=_params + [
            ("本次結果", (
                f"轉出 {int(stats.get('out_qty', 0)):,} 件｜已分配 {int(stats.get('alloc_qty', 0)):,} 件｜"
                f"未分配 {int(stats.get('unalloc_qty', 0)):,} 件（{int(stats.get('unalloc_skus', 0)):,} 個 SKU）"
            )),
        ],
        logic=_logic,
        sheet_rows=[
            ("轉貨單", "每行一筆：轉出店（含 Sold Qty (P90D)）→ 接收店、SKU（含 Adj. MCH1 與 Supply Source）、轉貨數量、接收店銷售佔比、分配總量與 14 日需求（黃色＝轉出店與轉貨數量、淺藍色＝接收店相關欄位）", "執行轉貨的依據；同一接收店若有多行，代表多間轉出店供貨"),
            ("轉出店舖彙總", "每間轉出店的轉出 SKU 數、轉出總量、已分配與未分配", "未分配 > 0 表示同 SKU 無接收店（或接收店 14 日需求已滿），需人工確認"),
            ("接收店舖彙總", "每間接收店的接收 SKU 數、接收總量與接收前後庫存", "接收後庫存明顯偏高者，可考慮改派"),
            ("SKU 轉貨彙總", "每個 SKU 的轉出店數、轉出總量、接收店數與分配結果", f"狀態為「{_TRANSFER_STATUS_NO_RECEIVER}」或「{_TRANSFER_STATUS_CAPPED}」者不會全數轉出"),
            ("MCH1 分類彙總", "按 Adj. MCH1 的分類：轉出／接收 SKU 數與件數、未分配件數", "未分配 > 0 的分類表示同 SKU 無接收店舖；「(空白)」表示來源檔該欄空白"),
        ],
        tips=_tips + [
            "轉貨單已按欄位分組著色：黃色＝轉出店與轉貨數量（轉貨數量加粗）、淺藍色＝接收店相關欄位；配合首列篩選，可快速檢視個別店舖或 SKU 的轉貨安排。",
        ],
    )

    if progress_callback:
        progress_callback(0.3, "建立轉貨單...")
    ws_plan = wb.create_sheet(title="轉貨單")
    _write_transfer_table(
        ws_plan, plan_df, auto_filter=True,
        column_styles=_TRANSFER_PLAN_STYLES,
        bold_data_columns=_TRANSFER_PLAN_BOLD_DATA_COLUMNS,
    )

    if progress_callback:
        progress_callback(0.5, "建立轉出店舖彙總...")
    ws_out = wb.create_sheet(title="轉出店舖彙總")
    _write_transfer_table(ws_out, out_shop_df)

    if progress_callback:
        progress_callback(0.7, "建立接收店舖彙總...")
    ws_in = wb.create_sheet(title="接收店舖彙總")
    _write_transfer_table(ws_in, in_shop_df)

    if progress_callback:
        progress_callback(0.85, "建立 SKU 轉貨彙總...")
    ws_sku = wb.create_sheet(title="SKU 轉貨彙總")
    _write_transfer_table(ws_sku, sku_df, auto_filter=True)

    if progress_callback:
        progress_callback(0.9, f"建立 {_TRANSFER_MCH_SHEET}...")
    ws_mch = wb.create_sheet(title=_TRANSFER_MCH_SHEET)
    _write_transfer_table(
        ws_mch,
        mch_df if mch_df is not None else pd.DataFrame(columns=_TRANSFER_MCH_SUMMARY_COLUMNS),
        auto_filter=True,
    )

    if progress_callback:
        progress_callback(0.95, "正在儲存 Excel 檔案...")
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    if progress_callback:
        progress_callback(1.0, "完成!")
    return buf
