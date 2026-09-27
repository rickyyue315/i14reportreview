# -*- coding: utf-8 -*-
"""積壓資金 Review：彙總佔用資金最多的滯銷存貨，排列去庫存優先次序。

只引用 modes.common，不引用其他模式模組，避免循環引用。
"""
import gc
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


from modes.common import (
    _auto_width,
    _clean_dataframe,
    _compute_raw_extra_columns,
    _extract_shop_name,
    _reorder_raw_columns,
    _resolve_column_name,
    _safe_default_series,
    _safe_numeric,
    _safe_sku,
    _write_raw_placeholder,
    _write_raw_sheet,
    validate_input,
)

__all__ = [
    "transform_excess_capital",
    "generate_excel_excess_capital",
    "_build_excess_by_sku",
    "_build_excess_by_shop",
]

_EXCESS_SIGNAL_COLUMNS = [
    "Excess Stock Qty",
    "Excess Stock Value",
]

_EXCESS_RAW_OUTPUT_COLUMNS = [
    "Site Code",
    "Shop Name",
    "OM",
    "Shop Nature",
    "Buyer",
    "Brand",
    "SKU",
    "SKU Desc",
    "RP Type",
    "Shop Stock Status",
    "Excess Stock Qty",
    "Excess Stock Value",
    "Shop Stock Qty (CM)",
    "Sold Qty (M-2)",
    "Sold Qty (M-1)",
    "Sold Qty (CM - MTD)",
    "Sold Qty (P90D)",
    "Shop Monthly Sold Qty (P90D)",
    "Turnday (P90D Qty)",
    "Net Sales (M-2)",
    "Net Sales (M-1)",
    "Net Sales (CM - MTD)",
    "Adj. Safety Stock",
]

_EXCESS_BY_SKU_COLUMNS = [
    "Buyer",
    "Brand",
    "SKU",
    "SKU Desc",
    "影響店舖數",
    "Excess Stock Qty 合計",
    "Excess Stock Value 合計",
]

_EXCESS_BY_SHOP_COLUMNS = [
    "Site Code",
    "Shop Name",
    "OM",
    "影響 SKU 數",
    "Excess Stock Qty 總數",
    "Excess Stock Value 總值",
]


def _excess_raw_columns(raw: pd.DataFrame) -> list[str]:
    r = _reorder_raw_columns(_compute_raw_extra_columns(raw))
    available = set(r.columns)
    return [c for c in _EXCESS_RAW_OUTPUT_COLUMNS if _resolve_column_name(c, available) is not None]


def _build_excess_by_sku(raw: pd.DataFrame) -> pd.DataFrame:
    cols = ["Buyer", "Brand", "SKU", "SKU Desc"]
    if raw is None or raw.empty:
        return pd.DataFrame(columns=_EXCESS_BY_SKU_COLUMNS)
    df = raw.copy()
    df["SKU"] = df["SKU"].apply(_safe_sku)
    df["_excess_qty"] = _safe_numeric(df.get("Excess Stock Qty", _safe_default_series(0, df.index)))
    df["_excess_val"] = _safe_numeric(df.get("Excess Stock Value", _safe_default_series(0, df.index)))
    grouped = (
        df.groupby(cols, dropna=False, observed=True)
        .agg(
            影響店舖數=("Site Code", "nunique"),
            Excess_Qty=("_excess_qty", "sum"),
            Excess_Val=("_excess_val", "sum"),
        )
        .reset_index()
    )
    grouped = grouped.rename(columns={
        "Excess_Qty": "Excess Stock Qty 合計",
        "Excess_Val": "Excess Stock Value 合計",
    })
    for c in ["影響店舖數", "Excess Stock Qty 合計", "Excess Stock Value 合計"]:
        grouped[c] = pd.to_numeric(grouped[c], errors="coerce").fillna(0).round().astype(int)
    grouped = grouped.sort_values(
        ["Excess Stock Value 合計", "Excess Stock Qty 合計", "Brand", "SKU"],
        ascending=[False, False, True, True],
        kind="stable",
    ).reset_index(drop=True)
    if not grouped.empty:
        total_row = pd.DataFrame([{
            "Buyer": "總計",
            "Brand": None,
            "SKU": None,
            "SKU Desc": None,
            "影響店舖數": int(df["Site Code"].dropna().nunique()),
            "Excess Stock Qty 合計": int(grouped["Excess Stock Qty 合計"].sum()),
            "Excess Stock Value 合計": int(grouped["Excess Stock Value 合計"].sum()),
        }])
        grouped = pd.concat([grouped, total_row], ignore_index=True)
    return grouped[_EXCESS_BY_SKU_COLUMNS]


def _build_excess_by_shop(raw: pd.DataFrame) -> pd.DataFrame:
    if raw is None or raw.empty:
        return pd.DataFrame(columns=_EXCESS_BY_SHOP_COLUMNS)
    df = raw.copy()
    site_code = df.get("Site Code", _safe_default_series("", df.index)).fillna("").astype(str).str.strip()
    item_shop = df.get("Item_Shop", _safe_default_series("", df.index))
    df["Site Code"] = site_code
    df["Shop Name"] = [_extract_shop_name(item_shop.iloc[i], site_code.iloc[i]) for i in range(len(df))]
    df["OM"] = df.get("OM", _safe_default_series("", df.index)).fillna("").astype(str).str.strip()
    df["SKU"] = df["SKU"].apply(_safe_sku)
    df["_excess_qty"] = _safe_numeric(df.get("Excess Stock Qty", _safe_default_series(0, df.index)))
    df["_excess_val"] = _safe_numeric(df.get("Excess Stock Value", _safe_default_series(0, df.index)))
    grouped = (
        df.groupby(["Site Code", "Shop Name", "OM"], dropna=False, observed=True)
        .agg(
            影響SKU數=("SKU", "nunique"),
            Excess_Qty=("_excess_qty", "sum"),
            Excess_Val=("_excess_val", "sum"),
        )
        .reset_index()
    )
    grouped = grouped.rename(columns={
        "影響SKU數": "影響 SKU 數",
        "Excess_Qty": "Excess Stock Qty 總數",
        "Excess_Val": "Excess Stock Value 總值",
    })
    grouped["影響 SKU 數"] = pd.to_numeric(grouped["影響 SKU 數"], errors="coerce").fillna(0).astype(int)
    for c in ["Excess Stock Qty 總數", "Excess Stock Value 總值"]:
        grouped[c] = pd.to_numeric(grouped[c], errors="coerce").fillna(0).round().astype(int)
    grouped = grouped.sort_values(
        ["Excess Stock Value 總值", "Excess Stock Qty 總數", "Site Code"],
        ascending=[False, False, True],
        kind="stable",
    ).reset_index(drop=True)
    if not grouped.empty:
        total_row = pd.DataFrame([{
            "Site Code": "總計",
            "Shop Name": "",
            "OM": "",
            "影響 SKU 數": int(df["SKU"].dropna().nunique()),
            "Excess Stock Qty 總數": int(grouped["Excess Stock Qty 總數"].sum()),
            "Excess Stock Value 總值": int(grouped["Excess Stock Value 總值"].sum()),
        }])
        grouped = pd.concat([grouped, total_row], ignore_index=True)
    return grouped[_EXCESS_BY_SHOP_COLUMNS]


def transform_excess_capital(
    df_input: pd.DataFrame,
    progress_callback=None,
    shop_natures: list[str] | None = None,
):
    """積壓資金轉換：篩選有超額存貨列，輸出 RAW、By SKU、By Shop 與提示訊息。"""
    if progress_callback:
        progress_callback(0.0, "驗證輸入資料...")
    missing = validate_input(df_input)
    if missing:
        raise ValueError(f"缺少必要欄位: {', '.join(missing)}")
    avail = set(df_input.columns)
    if _resolve_column_name("Site Code", avail) is None:
        raise ValueError("缺少必需欄位: Site Code")
    if all(_resolve_column_name(c, avail) is None for c in _EXCESS_SIGNAL_COLUMNS):
        raise ValueError(f"缺少積壓訊號欄位（需至少一個）: {'、'.join(_EXCESS_SIGNAL_COLUMNS)}")

    if progress_callback:
        progress_callback(0.1, "清理資料中...")
    df = _clean_dataframe(df_input, shop_natures=shop_natures, progress_callback=progress_callback)
    if df.empty:
        raise ValueError("過濾後無有效資料列，請確認來源檔案中的 Buyer 欄位是否包含有效數據")

    if progress_callback:
        progress_callback(0.2, "解像 SKU Desc / Shop Name...")
    df = _compute_raw_extra_columns(df)
    df["SKU"] = df["SKU"].apply(_safe_sku)

    excess_qty = _safe_numeric(df.get("Excess Stock Qty", _safe_default_series(0, df.index)))
    excess_val = _safe_numeric(df.get("Excess Stock Value", _safe_default_series(0, df.index)))
    mask = ((excess_qty > 0) | (excess_val > 0)).fillna(False)

    raw = df[mask].copy()

    if progress_callback:
        progress_callback(0.35, "建立彙總...")
    by_sku = _build_excess_by_sku(raw)
    by_shop = _build_excess_by_shop(raw)

    warnings = []
    for col in _EXCESS_SIGNAL_COLUMNS:
        if col not in df.columns:
            warnings.append(f"來源檔案缺少「{col}」，相關合計以 0 計。")
    if raw.empty:
        warnings.append("沒有符合「有超額存貨（Excess Stock Qty 或 Excess Stock Value 大於 0）」的資料。")

    if progress_callback:
        progress_callback(0.4, "完成資料轉換...")
    return raw, by_sku, by_shop, warnings


def _write_excess_table(ws, headers, data: pd.DataFrame, progress_callback=None, progress_start=0.0, progress_end=0.1, sheet_name=""):
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    bold_font = Font(bold=True, size=11)
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    if data is not None and not data.empty:
        total = len(data)
        for row_idx, (_, row_data) in enumerate(data.iterrows()):
            excel_row = row_idx + 2
            is_total = row_data.iloc[0] == "總計"
            for col_idx, h in enumerate(headers, 1):
                val = row_data.get(h)
                if pd.isna(val):
                    val = None
                cell = ws.cell(row=excel_row, column=col_idx, value=val)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if is_total:
                    cell.font = bold_font
                if isinstance(val, (int, float)) and col_idx > 4:
                    cell.number_format = "#,##0"
            if progress_callback and (row_idx % 500 == 0 or row_idx == total - 1):
                p = progress_start + (progress_end - progress_start) * ((row_idx + 1) / max(total, 1))
                progress_callback(p, f"寫入 {sheet_name} ({row_idx+1}/{total})...")
    else:
        ws.cell(row=2, column=1, value="無符合資料").font = Font(size=12)
    _auto_width(ws)


def generate_excel_excess_capital(
    raw: pd.DataFrame,
    by_sku: pd.DataFrame,
    by_shop: pd.DataFrame,
    progress_callback=None,
    include_raw: bool = True,
) -> BytesIO:
    wb = Workbook()
    if progress_callback:
        progress_callback(0.0, "建立 Excel 結構...")

    out_cols = _excess_raw_columns(raw if raw is not None else pd.DataFrame())

    ws_raw = wb.active
    ws_raw.title = "RAW"
    if include_raw:
        _write_raw_sheet(ws_raw, raw, progress_callback, progress_start=0.0, progress_end=0.55, columns=out_cols)
    else:
        _write_raw_placeholder(ws_raw)

    if progress_callback:
        progress_callback(0.6, "建立 By SKU...")
    ws_sku = wb.create_sheet(title="By SKU")
    _write_excess_table(ws_sku, _EXCESS_BY_SKU_COLUMNS, by_sku, progress_callback, 0.6, 0.75, "By SKU")

    if progress_callback:
        progress_callback(0.78, "建立 By Shop...")
    ws_shop = wb.create_sheet(title="By Shop")
    _write_excess_table(ws_shop, _EXCESS_BY_SHOP_COLUMNS, by_shop, progress_callback, 0.78, 0.9, "By Shop")

    if progress_callback:
        progress_callback(0.95, "正在儲存 Excel 檔案...")

    gc.collect()
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    if progress_callback:
        progress_callback(1.0, "完成!")
    return buf
