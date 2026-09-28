# -*- coding: utf-8 -*-
"""共用基礎（檢核／清理／Excel 原語）：由 transformer.py 拆分而來的共用核心。
transformer 以墊片 re-export，舊有 import 路徑不變。
"""
import datetime
import gc
import re
import unicodedata
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


__all__ = [
    "RAW_COLUMN_ORDER_NAMES",
    "REQUIRED_COLUMNS",
    "_ADJ_SAFETY_DISPLAY_COL",
    "_ADJ_SAFETY_SOURCE_COL",
    "_SAFETY_DISPLAY_NAMES",
    "_auto_width",
    "_clean_dataframe",
    "_compute_raw_extra_columns",
    "_compute_raw_extra_columns_impl",
    "_display_safety_header",
    "_downcast_numeric_dtypes",
    "_drop_unnamed_empty_columns",
    "_ensure_list",
    "_extract_shop_name",
    "_extract_sku_desc",
    "_normalize_ref_date",
    "_optimize_input_dtypes",
    "_parse_launch_dates",
    "_reorder_raw_columns",
    "_resolve_column_name",
    "_safe_default_series",
    "_safe_numeric",
    "_safe_sku",
    "_standardize_safety_column",
    "_write_metric_cards",
    "_write_raw_placeholder",
    "_write_raw_sheet",
    "validate_input",
    "_nature_rate_arr",
    "_write_guide_sheet",
]


REQUIRED_COLUMNS = ["Buyer", "Brand", "SKU", "Shop Stock Status"]



# Adj. Safety Stock 顯示／來源名稱契約：
# - 來源檔案欄名（內部運算一律沿用）：Adj. Safety Stock
# - 使用者可見輸出（Excel／UI，一律顯示）：Safety Stock (現時)
# - 輸入相容：同時接受上述兩名；Safety Stock (現時) 一律視為來源檔案的 Adj. Safety Stock。
_ADJ_SAFETY_SOURCE_COL = "Adj. Safety Stock"



_ADJ_SAFETY_DISPLAY_COL = "Safety Stock (現時)"



_SAFETY_DISPLAY_NAMES = {_ADJ_SAFETY_SOURCE_COL: _ADJ_SAFETY_DISPLAY_COL}



RAW_COLUMN_ORDER_NAMES = [
    "Item_Shop",
    "SKU",
    "Brand",
    "SKU Desc",
    "Buyer",
    "Adj. SST",
    "Vendor MC",
    "MCH2",
    "Item Status",
    "Launch Date",
    "Item Class",
    "Site Code",
    "Shop Name",
    "OM",
    "Shop Nature",
    "RP Type",
    "建議新RP Type",
    "Buyer 回覆 ND/RF",
    "保留RF原因",
    "Shop Stock Status",
    "Adj. Safety Stock",
    "Shop Monthly Sold Qty (P90D)",
    "Safety VS Monthly Sold Qty",
    "Projected Monthly Sold Qty",
    "Recommendation of new Safety Qty",
    "Safety Qty Increase",
    "Recommendation of new Safety Qty / ND",
    "After new Safety Qty / ND can transfer out of QTY",
    "Overstock Cap (X)",
    "OOS Floor (Y)",
    "Recommend Safety Stock",
    "Safety Qty Change",
    "Buyer reply",
    "D001 Net Stock Qty (CM)",
    "Shop Stock Qty (CM)",
    "Sold Qty (P90D)",
    "Turnday (P90D Qty)",
    "Sold Qty (M-2)",
    "Sold Qty (M-1)",
    "Sold Qty (CM - MTD)",
    "Net Sales (M-2)",
    "Net Sales (M-1)",
    "Net Sales (CM - MTD)",
    "Excess Stock Qty",
    "Excess Stock Value",
    "Shop On Order PO Qty (CM)",
    "Shop Pending Receive",
    "Is OOS (M-2)",
    "Is OOS (M-1)",
    "Is OOS (CM)",
    "OOS Periods Count",
    "OOS Pattern",
    "Max Sales Loss (M-2)",
    "Max Sales Loss (P90D)",
    "Total Sales Loss (M-2)",
    "Total Sales Loss (P90D)",
    "Is Stock Impacted",
    "Turnday (Days)",
    "Mix Turnday Band",
    "Mix Excess Stock",
    "Non Tourist Turnday Band",
]



def _resolve_column_name(name: str, available: set[str]) -> str | None:
    stripped = name.strip()
    for col in available:
        if isinstance(col, str) and col.strip() == stripped:
            return col
    # Adj. Safety Stock 別名相容：內部一律以來源名運算，但同時接受顯示名輸入
    #（Safety Stock (現時) ＝ 來源檔案的 Adj. Safety Stock）。
    if stripped == _ADJ_SAFETY_SOURCE_COL:
        for col in available:
            if isinstance(col, str) and col.strip() == _ADJ_SAFETY_DISPLAY_COL:
                return col
    return None



def _display_safety_header(name: str) -> str:
    """使用者可見表頭：內部來源名轉顯示名（避免誤會為系統建議值）。"""
    return _SAFETY_DISPLAY_NAMES.get(name, name)



def _standardize_safety_column(df: pd.DataFrame) -> pd.DataFrame:
    """輸入別名正規化：Safety Stock (現時) 視為來源檔案的 Adj. Safety Stock。

    內部運算一律以「Adj. Safety Stock」為正規欄名；顯示層才轉為「Safety Stock (現時)」。
    - 只有顯示名時：rename 為來源名，之後的任何讀取或變動皆沿用來源語意；
    - 兩者並存時：保留來源名，移除顯示名避免重複語意。
    """
    cols = list(df.columns)
    src_col = None
    disp_col = None
    for c in cols:
        if not isinstance(c, str):
            continue
        s = c.strip()
        if s == _ADJ_SAFETY_SOURCE_COL and src_col is None:
            src_col = c
        elif s == _ADJ_SAFETY_DISPLAY_COL and disp_col is None:
            disp_col = c
    if disp_col is not None and src_col is None:
        return df.rename(columns={disp_col: _ADJ_SAFETY_SOURCE_COL})
    if disp_col is not None and src_col is not None:
        return df.drop(columns=[disp_col])
    return df



def _drop_unnamed_empty_columns(df: pd.DataFrame) -> pd.DataFrame:
    """P0：移除 Excel 尾端全空的 Unnamed 欄，減少寬表記憶體與後續計算欄數。"""
    drop = [
        c for c in df.columns
        if isinstance(c, str) and c.strip().startswith("Unnamed")
        and pd.to_numeric(df[c], errors="coerce").isna().all()
        and df[c].fillna("").astype(str).str.strip().eq("").all()
    ]
    return df.drop(columns=drop) if drop else df



def _downcast_numeric_dtypes(df: pd.DataFrame) -> pd.DataFrame:
    """P0：整數欄降精度（int64→int8/16/32），不改值、不改欄名。

    刻意「不」降浮點精度：float64→float32 只有約 7 位有效數字，會靜默改值
    （例 43.85 → 43.849998），RAW 為稽核明細不容失真；整數降精度則為完全無損。
    """
    for col in df.columns:
        try:
            kind = df[col].dtype.kind
        except Exception:
            continue
        if kind == "i":
            df[col] = pd.to_numeric(df[col], downcast="integer")
    return df



def _optimize_input_dtypes(df: pd.DataFrame) -> pd.DataFrame:
    """P0：上傳後一次瘦身（去空欄＋降精度），各模式共用，行為不變。"""
    df = _drop_unnamed_empty_columns(df)
    return _downcast_numeric_dtypes(df)



def _reorder_raw_columns(df: pd.DataFrame) -> pd.DataFrame:
    available = set(df.columns)
    ordered: list[str] = []
    for name in RAW_COLUMN_ORDER_NAMES:
        resolved = _resolve_column_name(name, available)
        if resolved:
            ordered.append(resolved)
    trailing = [c for c in df.columns if c not in ordered]
    return df[ordered + trailing]



def validate_input(df: pd.DataFrame) -> list[str]:
    missing = [col for col in REQUIRED_COLUMNS if col not in df.columns]
    return missing



def _safe_sku(x):
    if pd.notna(x):
        try:
            return str(int(float(x)))
        except (ValueError, TypeError):
            return str(x)
    return x



def _safe_numeric(series: pd.Series, default: float = 0.0) -> pd.Series:
    return pd.to_numeric(series, errors="coerce").fillna(default)



def _safe_default_series(value, index):
    if len(index) == 0:
        return pd.Series(dtype='object')
    return pd.Series(value, index=index)



def _normalize_ref_date(ref_date) -> datetime.date:
    """New SKU 參考日期正規化：None／缺值退回今天；datetime 取日期部分；
    字串經 `pd.to_datetime` 解析；其餘型別直接回傳錯誤。"""
    if ref_date is None:
        return datetime.date.today()
    if isinstance(ref_date, datetime.datetime):
        return ref_date.date()
    if isinstance(ref_date, datetime.date):
        return ref_date
    if isinstance(ref_date, pd.Timestamp):
        return ref_date.date()
    if isinstance(ref_date, str):
        parsed = pd.to_datetime(ref_date.strip(), errors="coerce")
        if pd.notna(parsed):
            return parsed.date()
    raise ValueError(f"參考日期「{ref_date}」無法解析為日期，請傳入 date 或 YYYY-MM-DD 字串。")



def _parse_launch_dates(launch_raw: pd.Series) -> pd.Series:
    """Launch Date 原始字串 → datetime（自動推斷 + 常見格式重試，無法解析為 NaT）。

    來源檔 Launch Date 常為 YYYYMMDD 數字（讀入後為 float，如 20260925.0）；
    第一輪解析維持原行為，未解析者再去除 float 尾綴以固定格式重試。
    """
    launch_txt = launch_raw.fillna("").astype(str).str.strip()
    launch_dt = pd.to_datetime(launch_txt, errors="coerce")
    if launch_dt.isna().any():
        still_na = launch_dt.isna()
        # 數字型 YYYYMMDD：去除 float 尾綴（'20260925.0' → '20260925'）後重試
        retry_txt = launch_txt.str.replace(r"^(\d{8})\.0+$", r"\1", regex=True)
        for fmt in ("%Y/%m/%d", "%d/%m/%Y", "%Y%m%d"):
            if not still_na.any():
                break
            retry = pd.to_datetime(retry_txt[still_na], errors="coerce", format=fmt)
            launch_dt[still_na] = retry
            still_na = launch_dt.isna()
    return launch_dt



def _ensure_list(val):
    if not isinstance(val, list):
        return [val]
    return val



def _extract_sku_desc(item_shop, site_code) -> str:
    if pd.isna(item_shop) or pd.isna(site_code):
        return ""

    item_text = str(item_shop).strip()
    site_text = str(site_code).strip()
    if not item_text or not site_text:
        return ""

    item_text = re.sub(r'^\([^)]+\)\s*', '', item_text)
    item_text = re.sub(r'^\d+\s*-\s*', '', item_text)

    marker = f"{site_text}-"
    idx = item_text.rfind(marker)
    if idx == -1:
        return ""

    return item_text[:idx].strip().strip(",，;；:： ")



def _compute_raw_extra_columns_impl(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    site_code = df.get("Site Code", _safe_default_series("", df.index)).fillna("").astype(str).str.strip()
    item_shop = df.get("Item_Shop", _safe_default_series("", df.index))
    shop_name = [
        _extract_shop_name(item_shop.iloc[i], site_code.iloc[i])
        for i in range(len(df))
    ]
    sku_desc = [
        _extract_sku_desc(item_shop.iloc[i], site_code.iloc[i])
        for i in range(len(df))
    ]
    df["Shop Name"] = shop_name
    df["SKU Desc"] = sku_desc
    return df



def _compute_raw_extra_columns(df: pd.DataFrame) -> pd.DataFrame:
    return _compute_raw_extra_columns_impl(df)



def _extract_shop_name(item_shop, site_code) -> str:
    if pd.isna(item_shop) or pd.isna(site_code):
        return ""

    item_text = str(item_shop).strip()
    site_text = str(site_code).strip()
    if not item_text or not site_text:
        return ""

    marker = f"{site_text}-"
    idx = item_text.rfind(marker)
    if idx == -1:
        pattern = rf"{re.escape(site_text)}-(.+)$"
        match = re.search(pattern, item_text)
        if not match:
            return ""
        shop_text = match.group(1)
    else:
        shop_text = item_text[idx + len(marker):]

    return shop_text.strip().strip(",，;；:：")



def _clean_dataframe(df_input: pd.DataFrame, shop_natures: list[str] | None = None, progress_callback=None) -> pd.DataFrame:
    df = _standardize_safety_column(df_input).dropna(subset=["Buyer"])
    if progress_callback:
        progress_callback(0.1, "清理資料中...")
    _footer_mask = df["Buyer"].astype(str).str.strip().str.lower().str.startswith(("total", "applied filters"))
    if "Item_Shop" in df.columns:
        _item_shop_mask = df["Item_Shop"].astype(str).str.strip().str.lower().str.startswith("total")
        _footer_mask = _footer_mask | _item_shop_mask
    df = df[~_footer_mask]
    df = df[df["Buyer"].astype(str).str.strip() != "117 - Marketing"]
    if shop_natures is not None and "Shop Nature" in df.columns:
        sn_col = df["Shop Nature"].fillna("").astype(str).str.strip().str.lower()
        mask = _safe_default_series(False, df.index)
        if "Tourist" in shop_natures:
            # 精確比對：'Non Tourist' 含 'tourist' 子字串，contains 會把 Non Tourist 誤當 Tourist
            mask |= sn_col == "tourist"
        if "Mix" in shop_natures:
            mask |= sn_col == "mix"
        if "Non Tourist" in shop_natures:
            # 非 Tourist/Mix 的其他值（如 'Other'）隨 Non Tourist 一併保留，
            # 與倍率邏輯一致（非 tourist/mix 一律套 default 倍率）
            is_tourist_or_mix = (sn_col == "tourist") | (sn_col == "mix")
            mask |= (sn_col != "") & ~is_tourist_or_mix
        df = df[mask]
    return df



def _write_raw_sheet(ws, raw_df, progress_callback=None, progress_start=0.0, progress_end=0.3, columns=None):
    raw_df = _compute_raw_extra_columns(raw_df)
    raw_df = _reorder_raw_columns(raw_df)

    # 選擇性只輸出指定欄位（節省記憶體）
    if columns is not None:
        avail = [c for c in columns if c in raw_df.columns]
        raw_df = raw_df[avail]

    raw_cols = list(raw_df.columns)
    header_row = [_display_safety_header(c) for c in raw_cols]
    ws.append(header_row)

    total_raw = len(raw_df)
    if total_raw == 0:
        ws.auto_filter.ref = ws.dimensions
        _auto_width(ws)
        return
    # P0：數值欄預先四捨五入以減少逐格分支。
    # 註：必須用 Python 內建 round，不可改用 np.round／Series.round ——
    # numpy 採「縮放後 rint」法，於 .x5 進位邊界與 Python 的十進位正確捨入結果不同
    # （例 43.85、403.05），會使 RAW 輸出與拆分前不一致。
    _round_cols = [
        c for c in raw_cols
        if pd.api.types.is_numeric_dtype(raw_df[c])
        and not pd.api.types.is_bool_dtype(raw_df[c])
    ]
    if _round_cols:
        raw_df = raw_df.copy()
        for c in _round_cols:
            try:
                raw_df[c] = pd.Series(
                    [None if pd.isna(v) else round(float(v), 1) for v in raw_df[c]],
                    index=raw_df.index, dtype="object",
                )
            except Exception:
                pass
    # P0：直接逐列 append，免除批次暫存串列的二次迴圈與峰值；
    # 仍保留非數值欄的逐格 isinstance 捨入與 pd.isna 判定，確保與原行為一致
    for i, row in enumerate(raw_df.itertuples(index=False, name=None)):
        ws.append([
            None if pd.isna(v)
            else (v if isinstance(v, str) or isinstance(v, bool)
                  else (round(float(v), 1) if isinstance(v, (int, float)) else v))
            for v in row
        ])
        if progress_callback and (i % 2000 == 0 or i == total_raw - 1):
            p = progress_start + (progress_end - progress_start) * ((i + 1) / total_raw)
            progress_callback(p, f"寫入 RAW ({i+1}/{total_raw})...")

    header_fill_raw = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    ncols = len(raw_cols)
    for col_idx in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        cell.fill = header_fill_raw
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    ws.auto_filter.ref = ws.dimensions
    _auto_width(ws)



def _write_raw_placeholder(ws, reason="已略過 RAW 明細（匯出選項：只匯出主表與摘要）"):
    """P0：略過 RAW 時仍保留同名工作表，維持表數契約，僅寫入一行說明。"""
    ws.append(["RAW"])
    ws.append([reason])
    ws.column_dimensions["A"].width = 60



def _write_metric_cards(ws, title, row1_metrics, row2_metrics, thin_border, desc_font, progress_callback=None, progress_pct=0.95, row3_metrics=None, row4_metrics=None, row5_metrics=None, row6_metrics=None, row7_metrics=None):
    ws.merge_cells("A1:D1")
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font = Font(bold=True, size=14, color="1E40AF")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    row_offsets = [(3, row1_metrics), (7, row2_metrics)]
    if row3_metrics:
        row_offsets.append((11, row3_metrics))
    if row4_metrics:
        row_offsets.append((15, row4_metrics))
    if row5_metrics:
        row_offsets.append((19, row5_metrics))
    if row6_metrics:
        row_offsets.append((23, row6_metrics))
    if row7_metrics:
        row_offsets.append((27, row7_metrics))

    for row_offset, metrics in row_offsets:
        for i, (name, value, desc, color) in enumerate(metrics):
            col = i + 1
            tc = ws.cell(row=row_offset, column=col, value=name)
            tc.font = Font(bold=True, size=11, color="FFFFFF")
            tc.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
            tc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            tc.border = thin_border
            vc = ws.cell(row=row_offset + 1, column=col, value=value)
            vc.font = Font(bold=True, size=20, color=color)
            vc.alignment = Alignment(horizontal="center", vertical="center")
            vc.number_format = "#,##0"
            vc.border = thin_border
            dc = ws.cell(row=row_offset + 2, column=col, value=desc)
            dc.font = desc_font
            dc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            dc.border = thin_border

    if progress_callback:
        progress_callback(progress_pct, "調整格式中...")
    _auto_width(ws)



def _auto_width(ws, min_width=8, max_width=50, sample_rows=200):
    for col_idx in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0.0
        scan_end = min(ws.max_row, 1 + sample_rows)
        for row_idx in range(1, scan_end + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value is not None:
                s = str(cell.value)
                val_len = sum(2.0 if unicodedata.east_asian_width(c) in ("W", "F") else 1.0 for c in s)
                if val_len > max_len:
                    max_len = val_len
        adjusted = min(max(max_len + 2, min_width), max_width)
        ws.column_dimensions[col_letter].width = adjusted



try:
    import streamlit as st
    # 必須限制緩存大小與壽命 — 無 TTL/無 max_entries 時每次轉換都累積一份 DataFrame 拷貝，
    # 長時間運作的容器記憶體會持續增長且「釋放記憶體」按鈕前的執行間隔內無法自動回收
    # P0（單人使用）：由 4 份降為 1 份，TTL 縮短，減少常駐記憶體
    _compute_raw_extra_columns_impl = st.cache_data(_compute_raw_extra_columns_impl, max_entries=1, ttl=1800)
except ImportError:
    pass



def _nature_rate_arr(shop_nature, default_rate, tourist_rate, mix_rate):
    # Shop Nature → 每列 rate。必須精確比對：'Non Tourist' 含 'tourist' 子字串，
    # 用 contains 會把 Non Tourist 誤套 Tourist 倍率；其餘值（如 'Other'）歸 default。
    arr = np.full(len(shop_nature), default_rate)
    arr[shop_nature == "tourist"] = tourist_rate
    arr[shop_nature == "mix"] = mix_rate
    return arr



def _write_guide_sheet(ws, *, title, subtitle, params, logic, sheet_rows, tips):
    """通用「說明」導讀表 writer：版面與 _write_hybrid_guide_sheet 相同。
    params: [(標籤, 內容)]；logic: [段落]；sheet_rows: [(工作表, 內容, 建議留意)]；tips: [段落]。"""
    def _dw(text):
        # 顯示寬度估算（全形字元約 2、半形約 1），供換行高度估算
        return sum(2 if ord(ch) > 0x2E7F else 1 for ch in str(text))

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 56
    ws.column_dimensions["C"].width = 66
    title_font = Font(size=15, bold=True, color="1A1A2E")
    sub_font = Font(size=9, color="6B7280", italic=True)
    section_fill = PatternFill("solid", start_color="0077A3")
    section_font = Font(size=11, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", start_color="E4EEF5")
    label_font = Font(size=10, bold=True)
    body_font = Font(size=10)
    wrap = Alignment(vertical="top", wrap_text=True)
    r = 1

    def _merged(text, font, height=None, fill=None):
        nonlocal r
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        c = ws.cell(row=r, column=1, value=text)
        c.font = font
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if fill:
            c.fill = fill
        ws.row_dimensions[r].height = height or 18
        r += 1

    def _sec(text):
        _merged(text, section_font, height=20, fill=section_fill)

    def _kv(label, value):
        nonlocal r
        ws.cell(row=r, column=1, value=label).font = label_font
        ws.cell(row=r, column=1).alignment = wrap
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        c = ws.cell(row=r, column=2, value=value)
        c.font = body_font
        c.alignment = wrap
        lines = max(1, -(-_dw(value) // 110))
        ws.row_dimensions[r].height = 14 * lines + 6
        r += 1

    def _bullet(text):
        _kv("•", text)

    def _row3(a, b, c):
        nonlocal r
        for col, val, f in ((1, a, label_font), (2, b, body_font), (3, c, body_font)):
            cell = ws.cell(row=r, column=col, value=val)
            cell.font = f
            cell.alignment = wrap
        lines = max(1, -(-_dw(a) // 25), -(-_dw(b) // 54), -(-_dw(c) // 64))
        ws.row_dimensions[r].height = 14 * lines + 6
        r += 1

    _merged(title, title_font, height=26)
    _merged(subtitle, sub_font, height=16)

    _sec("① 本次執行參數（供核對）")
    for label, value in params:
        _kv(label, value)

    _sec("② 核心邏輯（這份報告怎樣算出來）")
    for text in logic:
        _bullet(text)

    _sec("③ 工作表導讀（先看哪一頁、應留意什麼）")
    for col, val in ((1, "工作表"), (2, "內容"), (3, "建議留意")):
        c = ws.cell(row=r, column=col, value=val)
        c.font = label_font
        c.fill = head_fill
        c.alignment = Alignment(vertical="center")
    ws.row_dimensions[r].height = 16
    r += 1
    for row in sheet_rows:
        _row3(*row)

    _sec("④ 實用提示")
    for text in tips:
        _bullet(text)
