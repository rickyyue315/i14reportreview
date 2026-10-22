# -*- coding: utf-8 -*-
"""店舖精簡SKU模式：把 By SKUxShop 明細轉化為店舖精簡 SKU 回覆表。

每列一店×一 SKU，保留來源檔全部欄位，並在最後附加五個固定欄位：

- 過去90日有無賣過：Sold Qty (P90D) > 0 → 「有」，否則（等於 0、負數或空白）→「無賣」；
- 過去60日有無存貨：Shop Stock Qty (LM) 或 Shop Stock Qty (CM) > 0 →「有存貨」，
  兩者同時為 0（或空白／負數）→「無存貨」；
- SCT建議、Shop 回覆新的NDRF、如果RF的Safety Qty：一律留空，供 SCT／店舖填寫。

本模式不作調貨、Safety 或退倉建議，亦不篩選列（全店全 SKU 輸出）；清理規則沿用
`_clean_dataframe`（剔除合計列與 117 - Marketing），Shop Nature 篩選為選填。
另由 Item_Shop 以 Site Code 解析 SKU Desc 與 Shop Name（SKU Desc 排於 Brand 之後、
Shop Name 排於 Site Code 之後），其餘來源欄序不變；欄名含「Value」的欄位一律刪除
（大小寫敏感，與「RAW 明細轉化 (不含 Value)」同一判定，故 Valuation Class 不受影響）。
輸出經 Brand／SKU（文字）排序，並產生 Summary（Adj. MCH1 × Item Status 計數）與
Export（回覆明細）兩個工作表；Export 套用店舖精簡回覆表範本的欄寬、隱藏欄、
數字格式、凍結窗格與顏色標示。
"""
import datetime
import gc
import re
from io import BytesIO

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from modes.common import (
    _cached_safe_sku_series,
    _clean_dataframe,
    _compute_raw_extra_columns,
    _resolve_column_name,
    _safe_numeric,
    validate_input,
)
from modes.raw_convert import _is_raw_convert_value_column

__all__ = [
    "_SHOP_SIMPLIFY_SOLD_COLUMN",
    "_SHOP_SIMPLIFY_STOCK_COLUMN",
    "_SHOP_SIMPLIFY_SCT_COLUMN",
    "_SHOP_SIMPLIFY_NDRF_COLUMN",
    "_SHOP_SIMPLIFY_SAFETY_COLUMN",
    "_SHOP_SIMPLIFY_REPLY_COLUMNS",
    "_SHOP_SIMPLIFY_SOLD_YES",
    "_SHOP_SIMPLIFY_SOLD_NO",
    "_SHOP_SIMPLIFY_STOCK_YES",
    "_SHOP_SIMPLIFY_STOCK_NO",
    "_SHOP_SIMPLIFY_EXPORT_SHEET",
    "_SHOP_SIMPLIFY_SUMMARY_SHEET",
    "shop_simplify_stats",
    "transform_shop_simplify",
    "generate_excel_shop_simplify",
]


# ── 輸出欄位契約（欄名固定，測試與導出檔依賴，未經明確要求不得改動）──
_SHOP_SIMPLIFY_SOLD_COLUMN = "過去90日有無賣過"
_SHOP_SIMPLIFY_STOCK_COLUMN = "過去60日有無存貨"
_SHOP_SIMPLIFY_SCT_COLUMN = "SCT建議"
_SHOP_SIMPLIFY_NDRF_COLUMN = "Shop 回覆新的NDRF"
_SHOP_SIMPLIFY_SAFETY_COLUMN = "如果RF的Safety Qty"

# 需要由 SCT／店舖填寫的三欄（一律留空；不沿用來源檔舊值）
_SHOP_SIMPLIFY_REPLY_COLUMNS = [
    _SHOP_SIMPLIFY_SCT_COLUMN,
    _SHOP_SIMPLIFY_NDRF_COLUMN,
    _SHOP_SIMPLIFY_SAFETY_COLUMN,
]

# 兩個判定欄的固定取值
_SHOP_SIMPLIFY_SOLD_YES = "有"
_SHOP_SIMPLIFY_SOLD_NO = "無賣"
_SHOP_SIMPLIFY_STOCK_YES = "有存貨"
_SHOP_SIMPLIFY_STOCK_NO = "無存貨"


# ── Excel 輸出格式（對齊店舖精簡回覆表範本）──
_SHOP_SIMPLIFY_EXPORT_SHEET = "Export"
_SHOP_SIMPLIFY_SUMMARY_SHEET = "Summary"

# 固定表頭顏色（使用者範本）
_SHOP_SIMPLIFY_HEADER_BLUE = "D9E2F3"
_SHOP_SIMPLIFY_GREEN = "CCFF99"
_SHOP_SIMPLIFY_PURPLE = "E4DFEC"
_SHOP_SIMPLIFY_NO_SALE_PINK = "F2DCDB"

# 綠色表頭（另加綠色資料格：RP Type／Shop Stock Qty (CM)）
_SHOP_SIMPLIFY_GREEN_HEADERS = {
    "RP Type", "Shop Stock Qty (CM)", _SHOP_SIMPLIFY_NDRF_COLUMN, _SHOP_SIMPLIFY_SAFETY_COLUMN,
}
_SHOP_SIMPLIFY_GREEN_DATA_COLUMNS = {"RP Type", "Shop Stock Qty (CM)"}
# 紫色表頭與資料格（Sold Qty (P90D)）
_SHOP_SIMPLIFY_PURPLE_COLUMNS = {"Sold Qty (P90D)"}

# 隱藏欄（使用者範本；按欄名對應，來源缺欄時自動略過）
_SHOP_SIMPLIFY_HIDDEN_COLUMNS = {
    "Item_Shop", "Valuation Class", "MCH2", "MC", "Shop Nature", "Shop Type",
    "D001 Stock Qty (LM)", "D001 Blocked Qty (CM)", "D001 QI Qty (CM)",
    "Turnday (P90D Qty)", "Net Sales (M-2)", "Excess Stock Qty",
    "D001 On Order PO Qty (CM)", "Shop On Order PO Qty (CM)", "Shop Pending Receive (±7 days)",
}

# 欄寬（使用者範本；未列出的欄沿用 Excel 預設寬度）
_SHOP_SIMPLIFY_COLUMN_WIDTHS = {
    "Item_Shop": 79.7109375,
    "SKU": 13.140625,
    "Valuation Class": 25.0,
    "Brand": 8.7109375,
    "SKU Desc": 46.42578125,
    "Buyer": 18.85546875,
    "Adj. SST": 29.42578125,
    "Vendor MC": 23.140625,
    "Adj. MCH1": 17.28515625,
    "MCH2": 48.28515625,
    "MC": 56.140625,
    "Item Status": 12.7109375,
    "Launch Date": 10.5703125,
    "Item Class": 10.0,
    "Supply Source": 11.140625,
    "Site Code": 10.140625,
    "Shop Name": 10.85546875,
    "OM": 8.85546875,
    "Shop Nature": 11.28515625,
    "Shop Type": 10.0,
    "RP Type": 9.85546875,
    "Adj. Safety Stock": 10.28515625,
    "Shop Monthly Sold Qty (P90D)": 13.140625,
    "Safety VS Monthly Sold Qty": 10.0,
    "D001 Stock Qty (LM)": 10.28515625,
    "D001 Net Stock Qty (CM)": 9.7109375,
    "D001 Blocked Qty (CM)": 11.42578125,
    "D001 QI Qty (CM)": 10.0,
    "Shop Stock Qty (LM)": 9.0,
    "Shop Stock Qty (CM)": 9.0,
    "Sold Qty (P90D)": 10.7109375,
    "Turnday (P90D Qty)": 12.0,
    "Sold Qty (M-2)": 10.5703125,
    "Sold Qty (M-1)": 10.5703125,
    "Sold Qty (CM - MTD)": 10.42578125,
    "Net Sales (M-2)": 11.0,
    "Net Sales (CM - MTD)": 10.42578125,
    "Shop Stock Status": 56.85546875,
    "Excess Stock Qty": 10.42578125,
    "D001 On Order PO Qty (CM)": 12.85546875,
    "Shop On Order PO Qty (CM)": 13.0,
    "Shop Pending Receive (±7 days)": 13.42578125,
    _SHOP_SIMPLIFY_SOLD_COLUMN: 10.85546875,
    _SHOP_SIMPLIFY_STOCK_COLUMN: 10.85546875,
    _SHOP_SIMPLIFY_NDRF_COLUMN: 8.42578125,
    _SHOP_SIMPLIFY_SAFETY_COLUMN: 9.7109375,
}

# 數字格式（使用者範本）
_SHOP_SIMPLIFY_NUMBER_FORMATS = {
    "Shop Monthly Sold Qty (P90D)": '_(* #,##0_);_(* \\(#,##0\\);_(* "-"??_);_(@_)',
    "Safety VS Monthly Sold Qty": "0%",
    "Turnday (P90D Qty)": "#,##0_);[Red](#,##0)",
    "Sold Qty (M-2)": "#,##0_);[Red](#,##0)",
    "Sold Qty (M-1)": "#,##0_);[Red](#,##0)",
    "Sold Qty (CM - MTD)": "#,##0_);[Red](#,##0)",
    "Net Sales (M-2)": "#,##0_);[Red](#,##0)",
    "Net Sales (M-1)": "#,##0_);[Red](#,##0)",
    "Net Sales (CM - MTD)": "#,##0_);[Red](#,##0)",
}


def shop_simplify_stats(df: pd.DataFrame) -> dict:
    """結果頁指標卡用的統計（不影響輸出檔）：列數、無賣、無存貨、同時無賣且無存貨、店舖數、SKU 數。"""
    stats = {
        "rows": 0, "sold_no": 0, "stock_no": 0, "both_no": 0, "sites": 0, "skus": 0,
    }
    if df is None or len(df) == 0:
        return stats
    sold_no = df.get(_SHOP_SIMPLIFY_SOLD_COLUMN, pd.Series("", index=df.index)).astype(str) == _SHOP_SIMPLIFY_SOLD_NO
    stock_no = df.get(_SHOP_SIMPLIFY_STOCK_COLUMN, pd.Series("", index=df.index)).astype(str) == _SHOP_SIMPLIFY_STOCK_NO
    stats["rows"] = int(len(df))
    stats["sold_no"] = int(sold_no.sum())
    stats["stock_no"] = int(stock_no.sum())
    stats["both_no"] = int((sold_no & stock_no).sum())
    if "Site Code" in df.columns:
        sites = df["Site Code"].dropna().astype(str).str.strip()
        stats["sites"] = int(sites[sites != ""].nunique())
    if "SKU" in df.columns:
        skus = df["SKU"].dropna().astype(str).str.strip()
        stats["skus"] = int(skus[skus != ""].nunique())
    return stats


def _shop_simplify_place_derived_columns(df: pd.DataFrame) -> pd.DataFrame:
    """把解析出的 SKU Desc／Shop Name 移到慣用位置，其餘來源欄序不變。

    SKU Desc 排於 Brand 之後（缺 Brand 時依序退回 SKU／Item_Shop）；
    Shop Name 排於 Site Code 之後（缺 Site Code 時依序退回 OM／Shop Nature／SKU Desc）。
    """
    def _move(frame: pd.DataFrame, column: str, anchors: tuple[str, ...]) -> pd.DataFrame:
        if column not in frame.columns:
            return frame
        cols = [c for c in frame.columns if c != column]
        for anchor in anchors:
            if anchor in cols:
                cols.insert(cols.index(anchor) + 1, column)
                break
        else:
            cols.append(column)
        return frame[cols]

    out = _move(df, "SKU Desc", ("Brand", "SKU", "Item_Shop"))
    out = _move(out, "Shop Name", ("Site Code", "OM", "Shop Nature", "SKU Desc"))
    return out


def _shop_simplify_launch_int(value):
    """Launch Date 統一為 YYYYMMDD 整數（來源多為 float 如 20120520.0）；
    日期物件、YYYY-MM-DD、YYYY/MM/DD 皆可解析；無法解析或空白回傳 NaN。"""
    if value is None:
        return np.nan
    try:
        if pd.isna(value):
            return np.nan
    except (TypeError, ValueError):
        pass
    if isinstance(value, (pd.Timestamp, datetime.datetime, datetime.date)):
        return int(value.strftime("%Y%m%d"))
    text = re.sub(r"\.0+$", "", str(value).strip())
    text = text.replace("-", "").replace("/", "")
    if len(text) == 8 and text.isdigit():
        return int(text)
    return np.nan


def _shop_simplify_drop_existing_columns(df: pd.DataFrame) -> pd.DataFrame:
    """移除來源檔已有的五個輸出欄（含前後空格變體），避免重複欄位。

    回傳移除後的 df；是否有欄被移除由呼叫方比對欄數自行判斷。
    """
    avail = set(df.columns)
    drop = []
    for name in [_SHOP_SIMPLIFY_SOLD_COLUMN, _SHOP_SIMPLIFY_STOCK_COLUMN, *_SHOP_SIMPLIFY_REPLY_COLUMNS]:
        resolved = _resolve_column_name(name, avail)
        if resolved is not None:
            drop.append(resolved)
    return df.drop(columns=drop) if drop else df


def transform_shop_simplify(
    df_input: pd.DataFrame,
    progress_callback=None,
    shop_natures: list[str] | None = None,
):
    """清理來源資料並附加「過去90日有無賣過／過去60日有無存貨」與三個空白回覆欄。

    回傳 (raw_df, warnings)。缺少 Sold Qty (P90D)，或 Shop Stock Qty (LM) 與
    Shop Stock Qty (CM) 同時缺失時拋出 ValueError（無法判定，不應靜默輸出）。
    """
    if progress_callback:
        progress_callback(0.0, "驗證輸入資料...")
    missing = validate_input(df_input)
    if missing:
        raise ValueError(f"缺少必需欄位: {', '.join(missing)}")

    avail = set(df_input.columns)
    sold_col = _resolve_column_name("Sold Qty (P90D)", avail)
    lm_col = _resolve_column_name("Shop Stock Qty (LM)", avail)
    cm_col = _resolve_column_name("Shop Stock Qty (CM)", avail)
    if sold_col is None:
        raise ValueError("來源檔案缺少「Sold Qty (P90D)」，無法判定過去90日有無賣過。")
    if lm_col is None and cm_col is None:
        raise ValueError(
            "來源檔案同時缺少「Shop Stock Qty (LM)」與「Shop Stock Qty (CM)」，無法判定過去60日有無存貨。"
        )

    if progress_callback:
        progress_callback(0.15, "清理資料中...")
    df = _clean_dataframe(df_input, shop_natures=shop_natures, progress_callback=progress_callback).copy()

    warnings: list[str] = []
    if lm_col is None:
        warnings.append("來源檔案缺少「Shop Stock Qty (LM)」，過去60日有無存貨只以 Shop Stock Qty (CM) 判定。")
    if cm_col is None:
        warnings.append("來源檔案缺少「Shop Stock Qty (CM)」，過去60日有無存貨只以 Shop Stock Qty (LM) 判定。")

    had_output_columns = any(
        _resolve_column_name(name, avail) is not None
        for name in [_SHOP_SIMPLIFY_SOLD_COLUMN, _SHOP_SIMPLIFY_STOCK_COLUMN, *_SHOP_SIMPLIFY_REPLY_COLUMNS]
    )
    if had_output_columns:
        warnings.append(
            "來源檔已含本模式的輸出欄（過去90日有無賣過／過去60日有無存貨／SCT建議／"
            "Shop 回覆新的NDRF／如果RF的Safety Qty），將以本次計算結果覆蓋，三個回覆欄一律留空。"
        )

    if progress_callback:
        progress_callback(0.35, "判定過去90日有無賣過...")
    df = _shop_simplify_drop_existing_columns(df)

    if "Item_Shop" not in df.columns:
        warnings.append("來源檔案缺少「Item_Shop」，SKU Desc 與 Shop Name 留空。")
    df = _compute_raw_extra_columns(df)
    df = _shop_simplify_place_derived_columns(df)

    value_columns = [c for c in df.columns if _is_raw_convert_value_column(c)]
    if value_columns:
        df = df.drop(columns=value_columns)

    sold = _safe_numeric(df[sold_col]) if sold_col in df.columns else pd.Series(0.0, index=df.index)
    sold_status = np.where(
        sold.to_numpy(dtype=float) > 0,
        _SHOP_SIMPLIFY_SOLD_YES,
        _SHOP_SIMPLIFY_SOLD_NO,
    )

    if progress_callback:
        progress_callback(0.55, "判定過去60日有無存貨...")
    lm_stock = (
        _safe_numeric(df[lm_col]) if lm_col is not None and lm_col in df.columns else pd.Series(0.0, index=df.index)
    )
    cm_stock = (
        _safe_numeric(df[cm_col]) if cm_col is not None and cm_col in df.columns else pd.Series(0.0, index=df.index)
    )
    has_stock = (lm_stock.to_numpy(dtype=float) > 0) | (cm_stock.to_numpy(dtype=float) > 0)
    stock_status = np.where(has_stock, _SHOP_SIMPLIFY_STOCK_YES, _SHOP_SIMPLIFY_STOCK_NO)

    df[_SHOP_SIMPLIFY_SOLD_COLUMN] = sold_status
    df[_SHOP_SIMPLIFY_STOCK_COLUMN] = stock_status
    for name in _SHOP_SIMPLIFY_REPLY_COLUMNS:
        df[name] = ""

    # SKU 統一為文字（與其他模式一致，避免 Excel 以科學記號顯示長編號）
    if "SKU" in df.columns:
        df["SKU"] = _cached_safe_sku_series(df["SKU"])

    # Launch Date 統一為 YYYYMMDD 整數（Excel 顯示為數字，與範本一致）
    if "Launch Date" in df.columns:
        df["Launch Date"] = pd.Series(
            [_shop_simplify_launch_int(v) for v in df["Launch Date"]],
            index=df.index,
            dtype=object,
        )

    # 非有限數值（如 0 銷售時 Turnday 的 inf）一律留空，避免 Excel 出現 inf
    for col in df.columns:
        if pd.api.types.is_float_dtype(df[col]):
            arr = df[col].to_numpy(dtype=float)
            if np.isinf(arr).any():
                df[col] = df[col].where(np.isfinite(arr))

    # 排序：Brand、SKU（文字升序），與店舖精簡回覆表範本一致
    if "Brand" in df.columns and "SKU" in df.columns:
        df = df.sort_values(
            ["Brand", "SKU"],
            key=lambda s: s.astype(str).str.strip(),
            kind="stable",
        ).reset_index(drop=True)

    if df.empty:
        warnings.append("清理後沒有可輸出的資料列，請確認來源檔案與 Shop Nature 篩選。")

    if progress_callback:
        progress_callback(0.7, "完成資料轉換...")
    return df, warnings


def _shop_simplify_cell(value):
    """Excel 儲存格值正規化：空白字串與各類缺值一律寫成 None（真正空白）；
    浮點數保留完整精度（不四捨五入），inf／-inf 視為空白。"""
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, str):
        return value if value != "" else None
    if isinstance(value, np.bool_):
        return bool(value)
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, (float, np.floating)):
        number = float(value)
        return number if np.isfinite(number) else None
    return value


def _write_shop_simplify_summary(ws, df: pd.DataFrame) -> None:
    """Summary 工作表：Adj. MCH1 × Item Status 的 SKU 列數樞紐（含總計；零值留空）。"""
    ws["A3"] = "計數 - SKU"
    ws["B3"] = "欄標籤"

    index_col = "Adj. MCH1" if "Adj. MCH1" in df.columns else (
        "Item Class" if "Item Class" in df.columns else None
    )
    status_col = "Item Status" if "Item Status" in df.columns else None
    if df is None or df.empty or index_col is None or status_col is None or "SKU" not in df.columns:
        return

    summary_df = df.copy()
    index_values = summary_df[index_col].fillna("").astype(str).str.strip()
    status_values = summary_df[status_col].fillna("").astype(str).str.strip()
    summary_df[index_col] = index_values.mask(index_values == "", "(空白)")
    summary_df[status_col] = status_values.mask(status_values == "", "(空白)")

    table = summary_df.pivot_table(
        index=index_col,
        columns=status_col,
        values="SKU",
        aggfunc="count",
        margins=True,
        margins_name="總計",
        fill_value=0,
    )
    # 空白組排在最後（Excel 樞紐慣例），總計列／欄置末
    statuses = [c for c in table.columns if c != "總計"]
    statuses = [c for c in statuses if c != "(空白)"] + [c for c in statuses if c == "(空白)"]
    total_col = 2 + len(statuses)
    row_order = [v for v in table.index if v not in ("總計", "(空白)")]
    if "(空白)" in table.index:
        row_order.append("(空白)")
    if "總計" in table.index:
        row_order.append("總計")
    table = table.reindex(
        index=row_order,
        columns=statuses + (["總計"] if "總計" in table.columns else []),
    )

    ws.cell(row=4, column=1, value="列標籤")
    for j, status in enumerate(statuses, start=2):
        ws.cell(row=4, column=j, value=str(status))
    ws.cell(row=4, column=total_col, value="總計")

    for i, (idx, row) in enumerate(table.iterrows(), start=5):
        cell = ws.cell(row=i, column=1, value=str(idx))
        cell.alignment = Alignment(horizontal="left")
        for j, status in enumerate(statuses, start=2):
            count = int(row[status])
            if count:
                ws.cell(row=i, column=j, value=count)
        total = int(row["總計"])
        if total:
            ws.cell(row=i, column=total_col, value=total)

    # 欄寬（對齊範本；狀態欄多於 4 個時其餘沿用預設）
    widths = [35.28515625, 9.7109375, 12.7109375, 8.28515625, 10.0, 5.28515625]
    for j, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = width


def generate_excel_shop_simplify(raw: pd.DataFrame, progress_callback=None) -> BytesIO:
    """產生兩個工作表：Summary（Adj. MCH1 × Item Status 計數）與 Export（回覆明細）。

    Export 沿用店舖精簡回覆表範本的格式：Brand／SKU 排序、凍結 F2、指定欄寬與隱藏欄、
    數字格式、綠色／紫色重點欄，以及「無賣」列的淡紅標示。
    """
    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = _SHOP_SIMPLIFY_SUMMARY_SHEET
    ws = wb.create_sheet(title=_SHOP_SIMPLIFY_EXPORT_SHEET)

    if progress_callback:
        progress_callback(0.0, "建立 Excel 結構...")

    df = raw if raw is not None else pd.DataFrame()
    headers = list(df.columns)
    ws.append([_shop_simplify_cell(c) for c in headers])

    total = len(df)
    for i, row in enumerate(df.itertuples(index=False, name=None)):
        ws.append([_shop_simplify_cell(v) for v in row])
        if progress_callback and (i % 2000 == 0 or i == total - 1):
            progress_callback(0.05 + 0.75 * ((i + 1) / total), f"寫入 Export ({i+1}/{total})...")

    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    fill_blue = PatternFill(start_color=_SHOP_SIMPLIFY_HEADER_BLUE, end_color=_SHOP_SIMPLIFY_HEADER_BLUE, fill_type="solid")
    fill_green = PatternFill(start_color=_SHOP_SIMPLIFY_GREEN, end_color=_SHOP_SIMPLIFY_GREEN, fill_type="solid")
    fill_purple = PatternFill(start_color=_SHOP_SIMPLIFY_PURPLE, end_color=_SHOP_SIMPLIFY_PURPLE, fill_type="solid")
    fill_pink = PatternFill(start_color=_SHOP_SIMPLIFY_NO_SALE_PINK, end_color=_SHOP_SIMPLIFY_NO_SALE_PINK, fill_type="solid")

    # 表頭：粗體、指定底色、細框、置中置頂、自動換行；列高 70.5
    ws.row_dimensions[1].height = 70.5
    for col_idx, name in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        if name in _SHOP_SIMPLIFY_GREEN_HEADERS:
            cell.fill = fill_green
        elif name in _SHOP_SIMPLIFY_PURPLE_COLUMNS:
            cell.fill = fill_purple
        else:
            cell.fill = fill_blue
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)

        letter = get_column_letter(col_idx)
        width = _SHOP_SIMPLIFY_COLUMN_WIDTHS.get(name)
        if width is not None:
            ws.column_dimensions[letter].width = width
        if name in _SHOP_SIMPLIFY_HIDDEN_COLUMNS:
            ws.column_dimensions[letter].hidden = True

    # 資料格：綠色重點欄、紫色 Sold Qty (P90D)、「無賣」淡紅標示
    green_cols = [i for i, name in enumerate(headers, start=1) if name in _SHOP_SIMPLIFY_GREEN_DATA_COLUMNS]
    purple_cols = [i for i, name in enumerate(headers, start=1) if name in _SHOP_SIMPLIFY_PURPLE_COLUMNS]
    sold_col_idx = headers.index(_SHOP_SIMPLIFY_SOLD_COLUMN) + 1 if _SHOP_SIMPLIFY_SOLD_COLUMN in headers else None
    number_formats = [
        (i, _SHOP_SIMPLIFY_NUMBER_FORMATS[name])
        for i, name in enumerate(headers, start=1)
        if name in _SHOP_SIMPLIFY_NUMBER_FORMATS
    ]
    for row_idx in range(2, total + 2):
        for col_idx in green_cols:
            ws.cell(row=row_idx, column=col_idx).fill = fill_green
        for col_idx in purple_cols:
            ws.cell(row=row_idx, column=col_idx).fill = fill_purple
        if sold_col_idx is not None:
            sold_cell = ws.cell(row=row_idx, column=sold_col_idx)
            if sold_cell.value == _SHOP_SIMPLIFY_SOLD_NO:
                sold_cell.fill = fill_pink
        for col_idx, fmt in number_formats:
            ws.cell(row=row_idx, column=col_idx).number_format = fmt

    # 凍結於 SKU Desc 之後（範本為 F2）；縮放 85%
    if "SKU Desc" in headers:
        ws.freeze_panes = f"{get_column_letter(headers.index('SKU Desc') + 2)}2"
    elif headers:
        ws.freeze_panes = "A2"
    ws.sheet_view.zoomScale = 85
    ws.auto_filter.ref = ws.dimensions

    if progress_callback:
        progress_callback(0.85, "建立 Summary 工作表...")
    _write_shop_simplify_summary(ws_summary, df)

    # 開啟檔案時預設顯示 Export（範本行為）
    wb.active = wb.sheetnames.index(_SHOP_SIMPLIFY_EXPORT_SHEET)

    if progress_callback:
        progress_callback(0.95, "正在儲存 Excel 檔案...")
    gc.collect()
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    if progress_callback:
        progress_callback(1.0, "完成!")
    return buf
